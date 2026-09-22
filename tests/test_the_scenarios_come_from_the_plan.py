"""THE SCENARIOS HAVE ONE OWNER: THE IMPLEMENTATION PLAN SHEET.

The product owner keeps each feature row's Given/When/Then scenarios in the column
"Executable scenarios (Gherkin)" of `docs/Nyaymalaw_Implementation_Plan.xlsx`. The
`.feature` files pytest-bdd runs are generated from that column by
`assurance/control_plane/plan_scenarios.py`. A file edited by hand, a row with no
file, or a file with no row would each let the tests say something the sheet does not.
"""

from __future__ import annotations

import shutil

import pytest
from openpyxl import load_workbook

from assurance.control_plane import plan_scenarios

pytestmark = pytest.mark.class_a


def test_every_feature_file_is_exactly_what_the_sheet_says():
    expected = plan_scenarios.expected_files(plan_scenarios.sheet_rows())
    assert expected, "no feature row carries scenarios, so this check asserts nothing"
    found = plan_scenarios.problems(expected)
    assert not found, (
        "the feature files and the plan sheet disagree. Edit the sheet, then run "
        "`python assurance/control_plane/plan_scenarios.py --write`:\n  " + "\n  ".join(found)
    )


def test_requirement_rows_reconcile_both_sheets():
    assert plan_scenarios.requirement_problems() == []


def test_active_sheet_does_not_choose_the_schema():
    rows = plan_scenarios.sheet_rows()
    assert len(rows) == 91
    assert sum(row[plan_scenarios.STATE_COLUMN] == "Executable" for row in rows) == 19
    assert sum(row[plan_scenarios.STATE_COLUMN] == "Planned" for row in rows) == 72


@pytest.mark.parametrize(
    "mutation",
    [
        "missing_header",
        "duplicate_header",
        "bad_state",
        "empty_body",
        "planned_body",
        "missing_notes",
        "duplicate_id",
        "bad_id",
        "empty_features",
        "empty_executable",
        "invalid_gherkin",
    ],
)
def test_scenario_schema_rejects_real_mutations(tmp_path, mutation):
    path = tmp_path / "mutated.xlsx"
    workbook = load_workbook(plan_scenarios.SHEET)
    sheet = workbook[plan_scenarios.WORKSHEET]
    features = [r for r in range(2, sheet.max_row + 1) if sheet.cell(r, 3).value == "Feature"]
    executable = [r for r in features if sheet.cell(r, 58).value == "Executable"]
    assert features and executable
    r = executable[0]
    edits = {
        "missing_header": [(1, 57, "renamed")],
        "duplicate_header": [(1, 56, plan_scenarios.COLUMN)],
        "bad_state": [(r, 58, "PASS")],
        "empty_body": [(r, 57, None)],
        "planned_body": [(r, 58, "Planned")],
        "missing_notes": [(r, 59, None)],
        "duplicate_id": [(features[1], 1, sheet.cell(features[0], 1).value)],
        "bad_id": [(features[-1], 1, "no identity")],
        "empty_features": [(n, 3, "Removed") for n in features],
        "empty_executable": [
            (n, c, v) for n in executable for c, v in ((57, None), (58, "Planned"))
        ],
        "invalid_gherkin": [(r, 57, "This prose is not an executable scenario")],
    }[mutation]
    for row, col, value in edits:
        assert sheet.cell(row, col).value != value
        sheet.cell(row, col).value = value
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError):
        plan_scenarios.expected_files(plan_scenarios.sheet_rows(path))


@pytest.mark.parametrize("mutation", ["removed", "duplicate", "drift"])
def test_requirement_reconciliation_can_reject_drift(tmp_path, mutation):
    workbook = load_workbook(plan_scenarios.SHEET)
    sheet = workbook[plan_scenarios.WORKSHEET]
    row = next(r for r in range(2, sheet.max_row + 1) if sheet.cell(r, 1).value == "LB-01")
    if mutation == "removed":
        sheet.cell(row, 1).value = "removed"
    elif mutation == "duplicate":
        sheet.append([cell.value for cell in sheet[row]])
    else:
        sheet.cell(row, 9).value = "Changed without reconciling the owner"
    path = tmp_path / "mutated.xlsx"
    workbook.save(path)
    workbook.close()
    assert plan_scenarios.requirement_problems(path)


def test_the_check_sees_an_edited_a_missing_and_an_orphan_file(tmp_path):
    """POSITIVE CONTROL: each way the files can drift from the sheet is reported."""
    expected = plan_scenarios.expected_files(plan_scenarios.sheet_rows())
    features = tmp_path / "features"
    shutil.copytree(plan_scenarios.FEATURES, features)
    assert plan_scenarios.problems(expected, features) == []

    edited, missing = sorted(expected)[:2]
    (features / edited).write_text(
        expected[edited] + "\n  Scenario: added by hand\n", encoding="utf8"
    )
    (features / missing).unlink()
    (features / "arrive" / "F-A-99.feature").write_text("Feature: orphan\n", encoding="utf8")

    found = plan_scenarios.problems(expected, features)
    assert any(f"{edited}: differs from the sheet" in line for line in found), found
    assert any(f"{missing}: in the sheet and not generated" in line for line in found), found
    assert any("F-A-99.feature: a feature file with no row" in line for line in found), found
