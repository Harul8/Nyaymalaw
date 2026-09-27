"""EVERY REQUIREMENT CLAUSE RESOLVES TO A REGISTERED OWNER, OR SAYS IT DOES NOT. LB-43-AC4.

THE MEASURED GAP, 27 September 2026. LB-43-AC4 asked for "all 92 LB
requirements and 324 clauses" to resolve to an owner. The sheet held 170 LB
and 23 OM requirements and 497 clauses, nothing computed either number, and
the first run of the check found a delivery-owner line (LB-163) still pairing
BK-101-AC1 with the packet it had left an hour earlier.

THE RULES, each asserted below:

* the population is read from the sheet when the check runs -- never counted
  in the check;
* no delivery-owner line names a criterion, item or packet the registries do
  not carry, or pairs a criterion with a packet that does not carry it;
* a slice may not start with a selected row that is not owned, a selected row
  with no acceptance clause, or a packet criterion no requirement claims.
"""
from __future__ import annotations

import re

import pytest
from openpyxl import load_workbook

from assurance.control_plane import requirement_owners as owners

pytestmark = pytest.mark.class_a


@pytest.fixture(scope="module")
def registry():
    return owners.Registry.load()


@pytest.fixture(scope="module")
def population(registry):
    return owners.requirements(registry=registry)


def test_the_population_is_the_sheets_own(population):
    """A POSITIVE CONTROL: the check sees exactly the requirement rows the sheet holds."""
    book = load_workbook(owners.SHEET, read_only=True)
    try:
        ids = [str(r[0] or "").strip()
               for r in book[owners.WORKSHEET].iter_rows(min_row=2, values_only=True)]
    finally:
        book.close()
    held = [i for i in ids if owners.REQUIREMENT.match(i)]
    assert [r.ident for r in population] == held and len(held) > 100
    assert sum(len(r.clauses) for r in population) > len(held)
    assert any(r.state == "owned" for r in population), (
        "no row resolves, so the resolver is not running")


def test_no_delivery_owner_line_names_what_the_registries_do_not_carry(population):
    problems = [f"{r.ident}: {'; '.join(r.problems)}" for r in population if r.problems]
    assert problems == [], "\n".join(problems)


def _registry() -> owners.Registry:
    return owners.Registry(
        criteria=frozenset({"BK-1-AC1", "BK-1-AC2", "BK-2-AC1"}), items=frozenset({"BK-1", "BK-2"}),
        carried_by={"BK-1-AC1": {"P01"}, "BK-1-AC2": {"P01"}, "BK-2-AC1": {"P02"}},
        carries={"P01": {"BK-1-AC1", "BK-1-AC2"}, "P02": {"BK-2-AC1"}})


def _row(ident, line, clauses=("X-AC1",)):
    return owners.resolve(owners.Requirement(ident, tuple(clauses), line), _registry())


@pytest.mark.parametrize("line,state", [
    ("Delivery owner (registered): BK-1-AC1 and AC2, packet P01.", "owned"),
    # moved, line left behind
    ("Delivery owner (registered): BK-1-AC1, packet P02.", "problem"),
    ("Delivery owner (registered): BK-9-AC1 (P01).", "problem"),                # no such criterion
    ("Delivery owner (registered): BK-2-AC1 (P02); reviewed under BK-1-AC2.", "owned"),
    ("Delivery owner (registered): No delivery owner yet: awaiting the owner.", "declared"),
    ("", "unowned"),
])
def test_each_line_resolves_to_its_state(line, state):
    assert _row("LB-900", line).state == state


def test_a_slice_with_an_unowned_or_clauseless_row_or_an_unclaimed_criterion_is_blocked():
    registry = _registry()
    owned = _row("LB-901", "Delivery owner (registered): BK-1-AC1, packet P01.")
    clauseless = _row("LB-902", "Delivery owner (registered): BK-1-AC2, packet P01.", clauses=())
    assert owners.slice_blockers("P01", [owned], registry) == [
        "BK-1-AC2 (P01) is claimed by no requirement row -- work nobody specified"]
    blockers = owners.slice_blockers("P01", [owned, clauseless], registry)
    assert blockers == ["LB-902 names no acceptance clause, so nothing could show it done"]
    stale = _row("LB-903", "Delivery owner (registered): BK-2-AC1, packet P01.")
    assert any(b.startswith("LB-903 is problem")
               for b in owners.slice_blockers("P01", [owned, stale], registry))
    assert owners.slice_blockers("P01", [owned, _row("LB-904", "BK-1-AC2 and nothing else (P01)")],
                                 registry) == []


def test_a_packaged_bare_item_claims_its_criteria_in_the_selected_slice():
    registry = _registry()
    row = _row("LB-901", "Delivery owner (registered): BK-1, packet P01.")
    assert row.state == "owned" and row.criteria == () and row.items == ("BK-1",)
    assert owners.registered_claims(row.line, registry) == {
        "P01": frozenset({"BK-1-AC1", "BK-1-AC2"})}
    assert owners.slice_blockers("P01", [row], registry) == []


def test_a_bare_item_claims_only_its_own_registered_criteria_in_its_paired_packet():
    registry = owners.Registry(
        criteria=frozenset({"BK-1-AC1", "BK-1-AC2", "BK-2-AC1"}),
        items=frozenset({"BK-1", "BK-2"}),
        carried_by={"BK-1-AC1": {"P01"}, "BK-1-AC2": {"P02"}, "BK-2-AC1": {"P01"}},
        carries={"P01": {"BK-1-AC1", "BK-2-AC1", "BK-1-AC99"}, "P02": {"BK-1-AC2"}})
    line = "Delivery owner (registered): BK-1, packet P01."
    row = owners.resolve(owners.Requirement("LB-901", ("LB-901-AC1",), line), registry)
    assert row.state == "owned"
    assert owners.registered_claims(line, registry) == {"P01": frozenset({"BK-1-AC1"})}
    blockers = owners.slice_blockers("P01", [row], registry)
    assert len(blockers) == 2
    assert all("is claimed by no requirement row" in blocker for blocker in blockers)
    assert any(blocker.startswith("BK-2-AC1 ") for blocker in blockers)
    assert any(blocker.startswith("BK-1-AC99 ") for blocker in blockers)
    assert owners.slice_blockers("P02", [row], registry)


def test_mixed_explicit_and_bare_groups_keep_their_local_packet_pairings():
    registry = _registry()
    line = "Delivery owner (registered): BK-1, packet P01; BK-2-AC1 (P02)."
    row = _row("LB-901", line)
    assert row.state == "owned"
    assert owners.registered_claims(line, registry) == {
        "P01": frozenset({"BK-1-AC1", "BK-1-AC2"}), "P02": frozenset({"BK-2-AC1"})}
    assert owners.slice_blockers("P01", [row], registry) == []
    assert owners.slice_blockers("P02", [row], registry) == []
    unpaired = "Delivery owner (registered): BK-1; BK-2-AC1 (P02)."
    assert owners.registered_claims(unpaired, registry) == {"P02": frozenset({"BK-2-AC1"})}


@pytest.mark.parametrize("line", [
    "Delivery owner (registered): BK-1, packet P02.",
    "Delivery owner (registered): BK-9, packet P01.",
    "Delivery owner (registered): BK-1-AC99, packet P01.",
])
def test_mismatched_or_unregistered_references_claim_no_slice(line):
    row = _row("LB-901", line)
    assert row.state == "problem"
    assert owners.registered_claims(line, _registry()) == {}
    assert owners.slice_blockers("P01", [row], _registry())


@pytest.mark.parametrize("reference", ["BK-3", "BK-3-AC1"])
def test_a_registered_but_unpackaged_reference_does_not_claim_a_slice(reference):
    base = _registry()
    registry = owners.Registry(
        criteria=base.criteria | {"BK-3-AC1"}, items=base.items | {"BK-3"},
        carried_by=base.carried_by, carries=base.carries)
    line = f"Delivery owner (registered): {reference}, packet P01."
    row = owners.resolve(owners.Requirement("LB-901", ("LB-901-AC1",), line), registry)
    assert row.state == "problem"
    assert owners.registered_claims(line, registry) == {}
    assert owners.slice_blockers("P01", [row], registry)


@pytest.mark.parametrize("line,state", [
    ("Delivery owner (registered): No delivery owner: BK-1, packet P01.", "declared"),
    ("Delivery owner (registered): BK-1, packet P01; BK-9-AC1 (P01).", "problem"),
])
def test_an_invalid_owner_cannot_cover_a_slices_unclaimed_criteria(line, state):
    registry = _registry()
    row = _row("LB-901", line)
    assert row.state == state
    blockers = owners.slice_blockers("P01", [row], registry)
    assert any(blocker.startswith(f"LB-901 is {state}") for blocker in blockers)
    for criterion in registry.carries["P01"]:
        assert any(blocker.startswith(f"{criterion} (P01) is claimed by no requirement row")
                   for blocker in blockers)


def test_cli_row_ownership_pass_does_not_claim_full_permission_to_start(monkeypatch, capsys):
    registry = _registry()
    row = _row("LB-901", "Delivery owner (registered): BK-1, packet P01.")
    monkeypatch.setattr(owners.Registry, "load", staticmethod(lambda: registry))
    monkeypatch.setattr(owners, "requirements", lambda **_kwargs: [row])
    assert owners.main(["--slice", "P01"]) == 0
    output = capsys.readouterr().out
    assert "P01: row ownership passes" in output
    assert "may start" not in output
    assert ("does not assess clause/evidence coverage, prerequisite closure "
            "or owner decisions") in output


def test_the_check_holds_no_count_of_its_own():
    """The number the criterion carried is what went stale; the check may not carry one."""
    source = owners.__file__
    code = open(source, encoding="utf-8").read().split('"""', 2)[2]
    assert not re.search(r"\b(92|324|170|193|497)\b", code), (
        "a population count is written into the check")
