"""The scratch pad, the guided method, the per-dispute answer and opposing counsel in three passes.

OWNER DIRECTION AND DECISIONS, 26 September 2026, given in conversation:

  * A scratch pad, closed by default, shows the loop's work live: the
    understanding of the message, the disputes, and for each dispute the Acts
    and sections considered, kept and set aside, the judgments and passages
    found relevant, and what the dispute needs.
  * The final answer summarises it per dispute: Act passages, case-law
    passages, evidence to collect, the case to prepare, the arguments, what the
    other side will say and how to strengthen.
  * The matter board asks, per dispute, for the details the passages require.
  * DECIDED: the order is guided by the principles with the final sections
    required, not enforced step by step.
  * DECIDED: each turn's scratch pad is saved with the matter.
  * DECIDED: opposing counsel runs in three passes -- anticipated defences
    early per dispute, a full attack per dispute once its details are in, and
    one pass across the matter at the end.
  * Everything is entered into the legal brain.

WHAT CHANGES. LB-139 (live progress, drafted in the interface section) becomes
the scratch pad and moves into L.2 beside the loop it shows. LB-140 (opposing
counsel) is rewritten as the three passes. Four rows are added: LB-163 the
guided method and LB-164 the step events in L.2; LB-165 the per-dispute answer
and LB-166 the matter-board questions in L.7. Both are owner-directed drafts;
every owner-authored row is proved byte-identical.
"""
import copy
import hashlib
import importlib.util
import re
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

_REPO = Path(__file__).resolve().parents[2]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))
from assurance.common._console import utf8_console  # noqa: E402

utf8_console()

_spec = importlib.util.spec_from_file_location(
    "regroup", Path(__file__).with_name("legal_brain_regroup_20260926.py"))
regroup = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(regroup)

SOURCE, TODAY, FIRST = regroup.SOURCE, regroup.TODAY, regroup.FIRST
HEADER = re.compile(r"^(L|L\.\d|U|U\.\d|P|X)  ")
SECTION = {"L.2": "Legal brain › L.2 The reasoning loop",
           "L.7": "Legal brain › L.7 Legal reasoning and advice"}

GUIDED = (" Owner decision, 26 September 2026: the order is guided by the principles and the final "
          "sections are required; it is not enforced step by step.")
SAVED = " Owner decision, 26 September 2026: each turn's scratch pad is saved with the matter."
PASSES = (" Owner decision, 26 September 2026: opposing counsel runs in three passes -- anticipated "
          "defences early per dispute, a full attack per dispute once its details are in, and one pass "
          "across the matter at the end.")

#: Rows written in full: the two rewritten and the four new. group, then ten columns.
ROWS = {
    "LB-139": ("L.2", [
        "The scratch pad: the loop's work, shown live and kept with the matter",
        "A panel, closed by default, shows every step of the loop as it happens, in plain words, grouped by "
        "dispute; it is saved with each turn so it can be reviewed later.",
        "See how NM reached its answer -- what it understood, what it read, what it kept and what it set "
        "aside -- and correct it while it works.",
        "Any turn; the advocate opens the panel whenever they choose.",
        "A collapsible panel beside the conversation, closed by default, on desktop and phone. It renders the "
        "step events (LB-164) live as the loop streams them, grouped under each dispute: what NM understood "
        "the message to say; the disputes identified; for each dispute, the Acts retrieved, the sections "
        "considered, those kept and those set aside with the reason; the judgments retrieved and the passages "
        "found relevant with the reason; what the dispute needs; and the other side's anticipated defences, "
        "attacks and cross-matter risks, each labelled with its pass (LB-140). Everything in the panel is "
        "marked as working, not yet checked -- only the final answer has passed the harness. A reason for "
        "keeping or setting aside a section is labelled as NM's judgment, so the advocate can challenge it. "
        "The advocate can stop the turn or correct a fact mid-turn; the correction is used by the next step "
        "and recorded (LB-142). Each turn's scratch pad is saved with the matter and can be reopened. Plain "
        "words throughout, never internal identifiers.",
        "The advocate can see, live or later, every step behind an answer.",
        "A lost connection keeps the saved record; reopening shows where the turn reached. If streaming is "
        "unavailable, the saved record is shown when the turn ends.",
        "Never let an interim line read as advice. Never show internal identifiers. Never show a step the "
        "loop did not take, or hide one it did.",
        "LB-139-AC1: a two-dispute message shows its understanding, both disputes, and each one's statutes, "
        "authorities and needs, in plain words, under the right dispute.\n"
        "LB-139-AC2: a section set aside shows its reason, labelled as NM's judgment.\n"
        "LB-139-AC3: a scratch pad reopened the next day shows the same steps it showed live.\n"
        "LB-139-AC4: the panel is closed by default and opens on phone and desktop.",
        "LB-128, LB-140, LB-142, LB-163, LB-164; LB-70, LB-82, LB-86; frontend/matter-workspace.js (the "
        "existing matter board); test_no_enum_value_reaches_the_advocate. Streaming does not exist today: "
        "/api/turn returns one finished answer.",
    ], SAVED),
    "LB-140": (None, [
        "Opposing counsel, in three passes",
        "The other side is argued three times, each when it is most useful: anticipated defences early per "
        "dispute, a full attack per dispute once its details are in, and one pass across the whole matter at "
        "the end.",
        "Hear what the other side will say early enough to gather the facts that answer it, and receive advice "
        "that has already survived the strongest attack.",
        "PASS 1: a dispute has been identified and its law retrieved. PASS 2: that dispute's key details are "
        "on the file, or the advocate asks for advice on it. PASS 3: every dispute has been through pass 2, or "
        "the advocate asks for the overall position.",
        "Each pass is a nested loop with a fresh context, the same retrieval tools and an adversarial brief; "
        "it sees the claims and their sources, never the main loop's reasoning. PASS 1 -- ANTICIPATED "
        "DEFENCES, light: the defences the law allows against this cause, each drawn from a retrieved "
        "provision or authority; they shape the matter-board questions (LB-166), each question naming the "
        "defence it prepares for. PASS 2 -- FULL ATTACK, full budget: objections to our facts, law, "
        "authorities and procedure on that dispute, each with its span and locator; the main loop must "
        "revise, refute with a source, or disclose each one; the result is the dispute's 'what the other side "
        "will say' and 'how to strengthen' (LB-165). PASS 3 -- ACROSS THE MATTER, once: a position in one "
        "dispute that damages another -- an inconsistent account, an admission usable elsewhere (G-EXPOSURE). "
        "A pass re-runs only when something it rested on changes -- a corrected fact, a new authority -- "
        "through the dependency ledger, not on every turn. This subsumes the adverse, exposure and attacks "
        "reads.",
        "Each dispute carries its anticipated defences, its answered objections and its strengthening steps; "
        "the matter carries its cross-dispute risks.",
        "A pass that cannot run is disclosed -- 'not tested against the other side' -- never presented as "
        "passed.",
        "Never let an anticipated defence or an objection reach the advocate without the passage it rests on. "
        "Never ask a pass-1 question whose answer would not change our position (OM-P05). Never let the "
        "opposing loop see the main loop's reasoning. Never drop an objection silently.",
        "LB-140-AC1: a cheque-dishonour dispute's pass 1 names the enforceable-debt and notice-receipt "
        "defences, and the board asks for the loan evidence and the proof of delivery.\n"
        "LB-140-AC2 (planted): a draft relying on an overruled authority meets that objection in pass 2 and is "
        "revised.\n"
        "LB-140-AC3: pass 3 flags a sum described as a loan in one dispute and as an investment in another.\n"
        "LB-140-AC4 (planted): an objection without a span is refused by the harness.",
        "G-ADVERSE, G-EXPOSURE; the adverse, exposure and attacks reads; LB-18, LB-23, LB-65, LB-115, LB-122, "
        "LB-165, LB-166. OPEN: what counts as a dispute's key details for pass 2 -- proposed: the elements "
        "table's required facts are answered or marked unobtainable.",
    ], PASSES),
    "LB-163": ("L.2", [
        "How each message is worked: the guided method",
        "Each message is worked through a method the principles set out, and the final answer must carry "
        "every section that applies; the order guides the model rather than binding it.",
        "Get the same thorough treatment on every message, without NM being forced through steps that do not "
        "apply.",
        "The advocate sends a message on a matter.",
        "THE METHOD, set out in the principles (LB-126): (1) state the understanding of the message; (2) "
        "identify the disputes it concerns, new or existing; (3) for each dispute, retrieve the governing "
        "Acts, consider the candidate sections and finalise those that apply; (4) retrieve the authorities and "
        "identify the passages relevant to that dispute; (5) from the passages, work out what the dispute "
        "needs -- facts, evidence, procedure; (6) run the other side's passes (LB-140); (7) submit the answer "
        "in the structure of LB-165. The model may go back -- re-reading a statute after a judgment changes "
        "the picture -- and skips what does not apply: a message that only corrects a date re-runs only what "
        "rested on that date. The harness requires the final answer to carry each section of LB-165 that "
        "applies, or to say why it does not.",
        "The answer covers every step the message called for, and the scratch pad shows how.",
        "A step that could not be completed is stated in the answer as a limit, with what would complete it.",
        "Never skip identifying the disputes. Never finalise a section that was not read. Never impose steps a "
        "message does not call for.",
        "LB-163-AC1: a first brief with two disputes produces an understanding, both disputes, and each one's "
        "statutes, authorities and needs.\n"
        "LB-163-AC2: a message that only corrects a date re-runs only what rested on that date.\n"
        "LB-163-AC3 (planted): an answer missing a required section with no stated reason is refused.",
        "LB-01, LB-03, LB-109, LB-110, LB-126, LB-128.",
    ], GUIDED),
    "LB-164": ("L.2", [
        "Step events: every step recorded, streamed and saved",
        "Each step of the loop is a typed event, appended to the turn's log and streamed to the scratch pad, "
        "so the display and the record are one thing.",
        "Trust that what the scratch pad shows is exactly what NM did.",
        "Every step of every loop.",
        "Event types: understanding; dispute identified (new or existing); tool called (the tool and its "
        "arguments, in plain words); section considered; section kept or set aside (with the reason); "
        "authority retrieved; passage relevant (with the reason); dispute needs; anticipated defence (pass "
        "1); objection and how it was dealt with (pass 2); cross-matter risk (pass 3); question asked; budget "
        "stop; answer ready. Each carries the dispute it belongs to, its sources and locators, and a sequence "
        "number. Events are appended to the turn's step log and streamed from a server-sent-events endpoint "
        "beside /api/turn; the saved log is the source when a turn is reopened. The same log is what the "
        "harness reads as process evidence (LB-143) and what record-and-replay captures (LB-146).",
        "A turn's step log is complete, ordered and saved, and the stream and the saved log agree.",
        "If the stream drops, the client resumes from the last sequence number it received.",
        "Never let the stream show an event the log does not hold. Never lose the log when a turn fails.",
        "LB-164-AC1: the events streamed for a turn equal, in order, the events saved with it.\n"
        "LB-164-AC2: a client that reconnects mid-turn receives every later event exactly once.\n"
        "LB-164-AC3: a failed turn keeps its log up to the failure.",
        "LB-128, LB-139, LB-143, LB-146; backend/nm/edge/api.py (/api/turn is request-response today).",
    ], SAVED),
    "LB-165": ("L.7", [
        "The answer, dispute by dispute: passages, evidence, case, arguments, the other side, how to strengthen",
        "The final answer summarises the work per dispute in a fixed structure, every item resting on "
        "retrieved passages or the advocate's own words.",
        "Know, for each dispute, what the law says, what I must gather, what case to prepare, what to argue, "
        "what the other side will say and how to meet it.",
        "The loop submits its answer (LB-161).",
        "After a short statement of the understanding and the disputes identified, each dispute carries: "
        "RELEVANT ACT PASSAGES (retrieved text, with locators); RELEVANT CASE-LAW PASSAGES (retrieved "
        "paragraphs, with court, binding status and treatment); EVIDENCE TO COLLECT (each item tied to the "
        "passage or element that requires it); THE CASE TO PREPARE (cause, forum and procedural steps -- "
        "pre-institution steps, limitation, periods); ARGUMENTS TO MAKE (each resting on cited passages); "
        "WHAT THE OTHER SIDE WILL SAY (from pass 2, with its passages); HOW TO STRENGTHEN (what meets each "
        "objection -- evidence, authority or another route). Then the cross-matter risks (pass 3) and the "
        "questions outstanding. Each item is a claim in submit_answer, rendered as readable prose; a section "
        "that does not apply says why.",
        "The advocate has a complete, checkable working brief for each dispute.",
        "A section that cannot be supported is stated as a limit, never filled with an unsupported item.",
        "Never include an item without its support. Never merge two disputes' analysis. Never present an "
        "anticipated defence as the other side's actual case.",
        "LB-165-AC1: on a golden two-dispute matter each dispute carries all seven sections, or a reason for "
        "each one absent.\n"
        "LB-165-AC2 (planted): an evidence item tied to no passage or element is refused.\n"
        "LB-165-AC3: every passage quoted in the answer matches its retrieved span.",
        "LB-19, LB-48, LB-54, LB-65, LB-130, LB-140, LB-161, LB-163.",
    ], GUIDED),
    "LB-166": ("L.7", [
        "What each dispute needs, asked on the matter board",
        "For each dispute the matter board lists the details it still needs -- from the retrieved passages, "
        "the elements and the anticipated defences -- and asks for them there, each saying why it matters.",
        "Be asked, dispute by dispute, only for what the law and the other side's likely case make necessary.",
        "A dispute is identified or re-analysed, or pass 1 anticipates a defence.",
        "The left-hand matter board lists every dispute with its status. Under each, NM lists the details "
        "still needed: facts an element requires, dates a period runs from, documents that prove a step, "
        "evidence that answers an anticipated defence. Each item names the passage or defence behind it "
        "('s.138(b) runs from receipt of the notice'; 'the drawer may argue the cheque was only security'). "
        "The advocate answers on the board or in the conversation; the answer is recorded through the checked "
        "writes (LB-142) and reopens that dispute's analysis. Items are ordered by what they unblock (G-GAP). "
        "A detail already on the file is never asked for again, and one the advocate says cannot be obtained "
        "is marked so and the analysis proceeds with that limit.",
        "Each dispute shows what is known, what is still needed, and why.",
        "A question whose basis is later withdrawn -- a passage no longer relied on -- is removed with a note.",
        "Never ask for a detail whose answer would not change the position (OM-P05). Never ask twice. Never "
        "let a board answer bypass the checked writes.",
        "LB-166-AC1: a cheque-dishonour dispute's board asks for the notice date and the proof of debt, each "
        "with its reason.\n"
        "LB-166-AC2: answering on the board updates the dispute and removes the item.\n"
        "LB-166-AC3: a detail marked unobtainable is not asked again, and the analysis states the limit.",
        "LB-05, LB-111, LB-117, LB-140, LB-142; G-GAP; frontend/matter-workspace.js (the existing board, "
        "with dispute rows and a need per dispute).",
    ], PASSES),
}
NEW = ("LB-163", "LB-164", "LB-165", "LB-166")
REWRITTEN = ("LB-139", "LB-140")

#: Where the rows land inside their groups, in this order.
L2_AFTER_LB140 = ("LB-163", "LB-164", "LB-139")
L7_FIRST = ("LB-165", "LB-166")

HEADER_ADDITIONS = {
    "L.2": (" Each message is worked by a guided method shown live in the scratch pad (LB-163, LB-164, "
            "LB-139), and the other side is argued in three passes (LB-140)."),
    "L.7": " Each dispute's answer follows a set structure, and its questions are asked on the matter board (LB-165, LB-166).",
}


def values_for(ident, group):
    title, lead, *rest = ROWS[ident][1]
    section = SECTION[group]
    column_a = f"{ident}\n{section}\n{title}\n\nOWNER-DIRECTED REQUIREMENT, DRAFTED FOR REVIEW:\n{lead}"
    return [column_a, *rest, regroup.READINESS + ROWS[ident][2]]


def main() -> int:
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    book = load_workbook(SOURCE)
    sheet, plan = book["Before Build"], book["Implementation Plan"]
    features = {s.title: regroup.sheet_features(s) for s in book}
    before = {(s.title, c.coordinate): (c.value, copy.copy(c._style))
              for s in book for row in s for c in row}
    heights = {r: sheet.row_dimensions[r].height for r in range(1, sheet.max_row + 1)}

    order = []
    for r in range(FIRST, sheet.max_row + 1):
        text = str(sheet.cell(r, 1).value or "")
        header = HEADER.match(text)
        key = f"HEADER:{header[1]}" if header else regroup.key_of(text)
        order.append((key, {"values": [sheet.cell(r, c).value for c in range(1, 11)],
                            "styles": [copy.copy(sheet.cell(r, c)._style) for c in range(1, 11)],
                            "height": sheet.row_dimensions[r].height, "row": r}))
    source = dict(order)
    for ident in NEW:
        assert ident not in source, ident
    lb140_group = None
    group = None
    for key, _ in order:
        if key.startswith("HEADER:"):
            group = key.split(":", 1)[1]
        if key == "LB-140":
            lb140_group = group
    assert lb140_group == "L.2", lb140_group
    template = source["LB-153"]

    # ---- the new body ----------------------------------------------------------
    body = []
    group = None
    for key, row in order:
        if key.startswith("HEADER:"):
            group = key.split(":", 1)[1]
        if key == "LB-139":
            continue                               # moves into L.2
        body.append((key, row))
        if key == "LB-140":
            for ident in L2_AFTER_LB140:
                body.append((ident, {"row": source.get(ident, {}).get("row"),
                                     "styles": [copy.copy(s) for s in
                                                (source.get(ident) or template)["styles"]],
                                     "height": 280, "values": None}))
        if key == "HEADER:L.7":
            for ident in L7_FIRST:
                body.append((ident, {"row": None, "styles": [copy.copy(s) for s in template["styles"]],
                                     "height": 280, "values": None}))
    assert len(body) == len(order) + len(NEW)

    changed = set()
    layout = []
    group, index = None, 0
    reshade = {"L.2", "L.7", "U.1"}
    for offset, (key, row) in enumerate(body):
        r = FIRST + offset
        is_header = key.startswith("HEADER:")
        if is_header:
            group, index = key.split(":", 1)[1], 0
        if key in ROWS:
            values = values_for(key, ROWS[key][0] or group)
        else:
            values = list(row["values"])
            if is_header and group in HEADER_ADDITIONS:
                values[0] += HEADER_ADDITIONS[group]
        for c in range(1, 11):
            cell = sheet.cell(r, c)
            cell.value = values[c - 1]
            cell._style = copy.copy(row["styles"][c - 1])
            if not is_header and group in reshade:
                cell.fill = PatternFill("solid", fgColor=regroup.ZEBRA[index % 2])
            changed.add(("Before Build", cell.coordinate))
        if not is_header:
            index += 1
        sheet.row_dimensions[r].height = row["height"]
        layout.append((key, r, values, group))
    last = FIRST + len(body) - 1
    sheet.auto_filter.ref = f"A4:J{last}"

    # ---- the row-2 note ----------------------------------------------------------
    population = sorted({regroup.LBOM.match(str(sheet.cell(x, 1).value or ""))[1]
                         for x in range(5, last + 1)
                         if regroup.LBOM.match(str(sheet.cell(x, 1).value or ""))})
    lb = [p for p in population if p.startswith("LB-")]
    note = str(sheet.cell(2, 1).value)
    assert "LB-126–162" in note
    note = note.replace("LB-126–162", "LB-126–166")
    note, n = re.subn(r"The legal brain holds \d+ LB requirements", f"The legal brain holds {len(lb)} LB requirements", note)
    assert n == 1
    note, n = regroup.NOTE_PATTERN.subn(f"All {len(population)} LB/OM requirements are linked into Implementation Plan", note)
    assert n == 1
    note = note.replace("context management and the tool catalogue",
                        "context management, the tool catalogue, the scratch pad and the per-dispute answer")
    sheet.cell(2, 1).value = note

    # ---- the Implementation Plan -------------------------------------------------
    position = {k: r for k, r, _v, _g in layout}
    plan_changed = set()
    for x in range(2, plan.max_row + 1):
        ident = plan.cell(x, 1).value
        if ident in REWRITTEN:
            updates = {6: ROWS[ident][1][0]}
            if ident == "LB-139":
                updates[5] = "L.2 The reasoning loop"
            for t, origin in regroup.MIRROR.items():
                updates[t] = sheet.cell(position[ident], origin).value
            for c, v in updates.items():
                if plan.cell(x, c).value != v:
                    plan.cell(x, c).value = v
                    plan_changed.add(plan.cell(x, c).coordinate)
    for ident in NEW:
        group = ROWS[ident][0]
        target = plan.max_row + 1
        for c in range(1, plan.max_column + 1):
            plan.cell(target, c)._style = copy.copy(plan.cell(target - 1, c)._style)
        fixed = {1: ident, 3: "Requirement", 4: "Legal brain",
                 5: "L.2 The reasoning loop" if group == "L.2" else "L.7 Legal reasoning and advice",
                 6: ROWS[ident][1][0], 7: "Pilot", 38: "Draft", 39: "Not started", 40: "Not verified",
                 32: f"{TODAY}: owner-directed row mirrored from Before Build. Not acceptance, not an "
                     "approved requirement."}
        for c, v in fixed.items():
            plan.cell(target, c, v)
            plan_changed.add(plan.cell(target, c).coordinate)
        for t, origin in regroup.MIRROR.items():
            plan.cell(target, t, sheet.cell(position[ident], origin).value)
            plan_changed.add(plan.cell(target, t).coordinate)

    out = Path("/tmp/claude-0/-home-user-Nyaymalaw/82a504d3-1374-501f-9fe3-441ac9aa966e/scratchpad") / SOURCE.name
    book.save(out)

    # ================= THE PROOF, on the saved bytes ==========================
    from openpyxl.xml.functions import tostring

    def xml(part):
        return tostring(part.to_tree())

    check = load_workbook(out)
    for tab in check:
        assert regroup.sheet_features(tab) == features[tab.title], tab.title
    bb = check["Before Build"]
    original = load_workbook(SOURCE)["Before Build"]
    for (title, coord), old in before.items():
        if title == "Before Build" and int(re.sub(r"[A-Z]+", "", coord)) < FIRST and coord != "A2":
            assert (bb[coord].value, bb[coord]._style) == old, coord
    for x in range(1, FIRST):
        assert bb.row_dimensions[x].height == heights.get(x)

    for key, r, values, group in layout:
        got = [bb.cell(r, c).value for c in range(1, 11)]
        assert got == values, key
        if key in ROWS:
            continue
        want = source[key]
        expected = list(want["values"])
        if key.startswith("HEADER:") and key.split(":", 1)[1] in HEADER_ADDITIONS:
            expected[0] += HEADER_ADDITIONS[key.split(":", 1)[1]]
        assert got == expected, f"{key}: content changed"
        for c in range(1, 11):
            a, o = bb.cell(r, c), original.cell(want["row"], c)
            assert xml(a.border) == xml(o.border) and xml(a.alignment) == xml(o.alignment), key
            assert (a.font.name, a.font.sz, a.font.color, a.font.b) == \
                (o.font.name, o.font.sz, o.font.color, o.font.b), key

    keys_after = sorted(k for k, *_ in layout)
    assert keys_after == sorted(list(source) + list(NEW)), "row population changed"
    groups = {k: g for k, _r, _v, g in layout}
    for ident in ("LB-139", "LB-140", "LB-163", "LB-164"):
        assert groups[ident] == "L.2", (ident, groups[ident])
    for ident in ("LB-165", "LB-166"):
        assert groups[ident] == "L.7", (ident, groups[ident])

    ip = check["Implementation Plan"]
    for (title, coord), old in before.items():
        if title == "Implementation Plan" and coord not in plan_changed:
            assert (ip[coord].value, ip[coord]._style) == old, f"plan {coord} moved"
    from assurance.control_plane import plan_scenarios
    problems = plan_scenarios.requirement_problems(out)
    assert not problems, problems
    assert len(plan_scenarios.sheet_rows(out)) == 93

    shutil.copy2(out, SOURCE)
    print(f"source digest before: {digest}")
    print(f"source digest after:  {hashlib.sha256(SOURCE.read_bytes()).hexdigest()}")
    print(f"last row: {last}; LB/OM: {len(population)} (LB {len(lb)})")
    for ident in ("LB-140", "LB-163", "LB-164", "LB-139", "LB-165", "LB-166"):
        print(f"  {ident}  row {position[ident]:>3}  {groups[ident]}  {ROWS[ident][1][0]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
