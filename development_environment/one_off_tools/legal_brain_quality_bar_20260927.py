"""Record the owner's approval of the LB-40 quality bar, and state its thresholds once.

OWNER DIRECTION, 27 September 2026: "go with your recommendations" -- the
proposed severity scale, populations and thresholds of the absolute bar are
approved as drafted. The time targets per kind of turn stay the owner's to set
before the first switch.

ONE OWNER FOR EACH PART. The thresholds are now rows RG-30 to RG-38 of
assurance/specification/release.yaml, where release.yaml says a threshold
lives, scored by pipeline/quality/releasegate.py (NOT MEASURED until the
held-out populations are graded). Writing them into the sheet as well would be
the rule-in-two-places shape CLAUDE.md refuses, so LB-40's 'Pass thresholds'
cell points at the rows; the severity scale and the populations stay in LB-40's
'Quality targets', which release.yaml points back to.

Same discipline as the other workbook tools: snapshot, write to a temporary
file, reload, prove nothing else moved, check the reconciler and the ownership
check, then replace. A second run refuses.
"""
import copy
import hashlib
import importlib.util
import re
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

_REPO = Path(__file__).resolve().parents[2]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))
from assurance.common._console import utf8_console  # noqa: E402

utf8_console()

_spec = importlib.util.spec_from_file_location(
    "regroup", Path(__file__).with_name("legal_brain_regroup_20260926.py"))
regroup = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(regroup)

SOURCE, FIRST = regroup.SOURCE, regroup.FIRST
HEADER = re.compile(r"^(L|L\.\d|U|U\.\d|P|X)  ")
TODAY = "27 September 2026"
ROWS = "RG-30 to RG-38 of assurance/specification/release.yaml"

PLAN_REPLACE = {
    ("LB-40", "Quality targets"): (
        f"PROPOSED {TODAY}, FOR OWNER APPROVAL -- the severity scale and the populations the bar is measured on. "
        "Nothing here binds until the owner approves it.",
        f"APPROVED by the owner, {TODAY} -- the severity scale and the populations the bar is measured on. The "
        f"thresholds are {ROWS}."),
}
PLAN_SET = {
    ("LB-40", "Pass thresholds"): (
        f"APPROVED by the owner, {TODAY}. The thresholds are stated once, as {ROWS} -- critical findings, major "
        "findings, unsupported release, excessive refusal, independent judgment under pressure, question burden, "
        "comprehension, cost and latency per satisfactorily completed task, and every required capability "
        "assessed. Each is blocking and scored by pipeline/quality/releasegate.py as PASS, FAIL or NOT MEASURED; "
        "NOT MEASURED fails like FAIL, and a row the scorer has no measurement for is reported NOT MEASURED, never "
        "dropped. Relative improvement over the pipeline offsets none of them (LB-138). The time targets per kind "
        "of turn remain the owner's to set before the first switch."),
    ("LB-39", "Pass thresholds"): (
        f"The bar is stated once: the severity scale and populations in LB-40, the thresholds as {ROWS} "
        f"(approved {TODAY})."),
    ("LB-72", "Pass thresholds"): (
        f"The bar is stated once: the severity scale and populations in LB-40, the thresholds as {ROWS} "
        f"(approved {TODAY})."),
}
DECISION = (f"{TODAY}, owner decision: the proposed severity scale, populations and thresholds are approved as "
            f"drafted; the thresholds are {ROWS}; the time targets per kind of turn remain to be set before the "
            f"first switch.")
CHANGE = f"{TODAY}: the quality bar approved by the owner; thresholds recorded as {ROWS}."


def main() -> int:
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    book = load_workbook(SOURCE)
    sheet, plan = book["Before Build"], book["Implementation Plan"]
    features = {s.title: regroup.sheet_features(s) for s in book}
    before = {(s.title, c.coordinate): (c.value, copy.copy(c._style)) for s in book for row in s for c in row}

    rows = {}
    for r in range(FIRST, sheet.max_row + 1):
        text = str(sheet.cell(r, 1).value or "")
        h = HEADER.match(text)
        rows[f"HEADER:{h[1]}" if h else regroup.key_of(text)] = r
    header = [str(c.value or "").strip() for c in plan[1]]
    col = {name: header.index(name) + 1 for name in
           {n for _k, n in (*PLAN_REPLACE, *PLAN_SET)} | {"Requirement version / change summary"}}
    at = {plan.cell(x, 1).value: x for x in range(2, plan.max_row + 1)}

    decisions = sheet.cell(rows["LB-40"], 10)
    assert DECISION not in str(decisions.value), "already applied"
    decisions.value = f"{decisions.value}\n\n{DECISION}" if decisions.value else DECISION

    written = set()
    for (key, name), (old, new) in PLAN_REPLACE.items():
        cell = plan.cell(at[key], col[name])
        assert str(cell.value).count(old) == 1, (key, name)
        cell.value = cell.value.replace(old, new)
        written.add(cell.coordinate)
    for (key, name), value in PLAN_SET.items():
        cell = plan.cell(at[key], col[name])
        cell.value = value
        written.add(cell.coordinate)
    for target, origin in regroup.MIRROR.items():
        cell = plan.cell(at["LB-40"], target)
        if cell.value != sheet.cell(rows["LB-40"], origin).value:
            cell.value = sheet.cell(rows["LB-40"], origin).value
            written.add(cell.coordinate)
    change = plan.cell(at["LB-40"], col["Requirement version / change summary"])
    change.value = f"{change.value}\n\n{CHANGE}" if change.value else CHANGE
    written.add(change.coordinate)

    out = Path(tempfile.gettempdir()) / "nm_legal_brain_quality_bar_20260927" / SOURCE.name
    out.parent.mkdir(exist_ok=True)
    book.save(out)

    check = load_workbook(out)
    for tab in check:
        assert regroup.sheet_features(tab) == features[tab.title], tab.title
    for (title, coord), old in before.items():
        cell = check[title][coord]
        if (title == "Implementation Plan" and coord in written) or \
                (title == "Before Build" and coord == decisions.coordinate):
            assert cell._style == old[1], coord
            continue
        assert (cell.value, cell._style) == old, f"{title} {coord} moved"
    from assurance.control_plane import plan_scenarios, requirement_owners
    assert plan_scenarios.requirement_problems(out) == []
    registry = requirement_owners.Registry.load()
    assert [r for r in requirement_owners.requirements(out, registry) if r.problems] == []

    shutil.copy2(out, SOURCE)
    print(f"source digest before: {digest}")
    print(f"source digest after:  {hashlib.sha256(SOURCE.read_bytes()).hexdigest()}")
    print(f"plan cells written: {len(written)}; Before Build cells written: 1")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
