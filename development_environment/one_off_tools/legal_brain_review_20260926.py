"""Record the 26 September 2026 review of the legal-brain rows, each finding in the row it concerns.

OWNER DIRECTION, 26 September 2026: "record this into the Before Build sheet,
place them appropriately in their place instead of just adding them at the end".

The review read every legal-brain row against the code, this week's
measurements and the corpus. Nothing here is appended to the end of the sheet:

  * DRAFT rows (LB-126..167, written by Claude from owner direction) are
    amended in place -- exact substring replacements, each asserted to occur
    once, and dated additions.
  * OWNER-AGREED and diagnostic rows keep their wording. A finding that
    concerns one is added as a dated "review (proposed)" paragraph -- the
    shape the 22 September amendments already use -- in the column it belongs
    to. Nothing is marked superseded: each carries a proposal the owner decides.
  * Three new practice-layer rows (LB-168..170) are inserted INSIDE L.8, after
    LB-125, and mirrored into the Implementation Plan as the reconciler needs.
  * Two Implementation Plan "Remaining gaps" cells (LB-121, LB-125) stated an
    absence that the corpus measurement below refutes; they are corrected.

CORPUS, MEASURED 26 September 2026 (raw_data/, chunks.db, legal.db and
pipeline/manifest.yaml): the Commercial Courts Act 2015, the Telangana
Court-fees and Suits Valuation Act 1956 and the Telangana Civil Courts Act 1972
are HELD in all three layers and absent only from the manifest; the court-fee
schedules are in the raw text but no schedule atoms reached chunks.db; the
Andhra Pradesh Civil Rules of Practice 1980 are in raw_data/ only.

The discipline is the other workbook tools': snapshot every cell and native
sheet feature, write, reload the SAVED file, prove nothing moved except what
was meant to, check the reconciler, and only then replace the source.
"""
import copy
import hashlib
import importlib.util
import re
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

_REPO = Path(__file__).resolve().parents[2]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))
from assurance.common._console import utf8_console  # noqa: E402

utf8_console()

_spec = importlib.util.spec_from_file_location(
    "regroup", Path(__file__).with_name("legal_brain_regroup_20260926.py"))
regroup = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(regroup)

SOURCE, FIRST = regroup.SOURCE, regroup.FIRST
HEADER = re.compile(r"^(L|L\.\d|U|U\.\d|P|X)  ")
R = "26 September 2026 review"
P = f"{R} (proposed, for owner approval)"
MAP = "Ownership map: LB-43."

# =============================================================================
# 1. EXACT REPLACEMENTS -- draft rows and the L.4 header only.
# =============================================================================
REPLACE = [
    ("HEADER:L.4", 1,
     "so a measured gap -- today the Commercial Courts Act and the Telangana court-fee and suits-valuation "
     "schedules, measured against the manifest -- is closed by acquiring the source, never by the model "
     "filling it.",
     "so a gap is closed by acquiring or indexing the source, never by the model filling it. A gap is measured "
     "against raw_data/ and every derived store before it is called absent. Measured 26 September 2026: the "
     "Commercial Courts Act 2015, the Telangana Court-fees and Suits Valuation Act 1956 and the Telangana Civil "
     "Courts Act 1972 are HELD -- in raw_data/, chunks.db and legal.db -- and missing only from "
     "pipeline/manifest.yaml, so exact identification cannot reach them; the court-fee schedules are in the raw "
     "text but no schedule atoms reached chunks.db. What is missing is the manifest entry and the schedule "
     "atoms, not the source (LB-121, LB-125, LB-170)."),

    ("LB-139", 4,
     "Everything in the panel is marked as working, not yet checked -- only the final answer has passed the "
     "harness.",
     "Everything in the panel is marked as working, not yet checked -- only the final answer has passed the "
     "harness -- and every model-written line in it has first passed the one citation detector every channel "
     "uses (LB-67), so an authority the loop never retrieved cannot appear even as working."),

    ("LB-166", 4,
     "A detail already on the file is never asked for again, and one the advocate says cannot be obtained is "
     "marked so and the analysis proceeds with that limit.",
     "A detail already on the file is never asked for again. Each item carries LB-117's four states -- held, "
     "outstanding, promised, unavailable -- and the advocate's replies move it through LB-118's reply read: one "
     "the advocate says cannot be obtained is marked unavailable and the analysis proceeds with that limit, and "
     "a promised document comes back when it is due. Every item's text, including the passage it quotes, passes "
     "LB-67's detector before it is shown."),

    ("LB-128", 5,
     "The turn ends with an answer that passed the harness, a question to the advocate, or a budget stop that "
     "says what was not finished.",
     "The turn ends in one of four distinct, truthful states (carried from LB-64): an answer that passed the "
     "harness (sufficient for the task), a question to the advocate (awaiting information or authority), a "
     "no-progress stop, or a budget stop -- each saying what was not finished."),

    ("LB-128", 8,
     "LB-128-AC3: the order of work differs across two different matters -- no hidden fixed sequence.",
     "LB-128-AC3 (replay): a scripted model chooses two different orders of work on the same matter and the "
     "runner follows both to the same checked result -- the runner imposes no sequence. (Replaces 'the order of "
     "work differs across two matters', which could pass by chance and could not fail.)"),

    ("LB-130", 4,
     "Each claim carries its kind -- fact from the file, law from a retrieved provision, authority from a "
     "retrieved judgment, inference drawn from named claims, question, or limit --",
     "Each claim carries its kind -- fact from the file, kept as what it is (the advocate's instruction, the "
     "client's account, a third party's allegation, or what a document records; never promoted to proven: "
     "LB-07, OM-P06), law from a retrieved provision, authority from a retrieved judgment, inference drawn from "
     "named claims, a research hypothesis labelled as unverified (LB-57), question, or limit --"),

    ("LB-141", 4,
     "A separate model instance, on the cheaper tier, sees one claim",
     "A separate model instance -- starting on the cheaper tier, and moved to the hard tier only on the "
     "LB-141-AC2 measurement (PRD 7.4.1; qualified for the role under LB-71) -- sees one claim"),

    ("LB-141", 4,
     "and answers supports, does not support, or supports in part, quoting the words it relies on.",
     "and gives its reason first -- quoting the words it relies on -- then answers supports, does not support, "
     "or supports in part (the reason-first pattern recorded in LB-60)."),

    ("LB-126", 6,
     "A change takes effect on the next turn, never mid-turn.",
     "A change takes effect at the next turn boundary by forcing a compaction (LB-150), because the principles "
     "sit in the stable prefix (LB-147), which is never edited mid-conversation; never mid-turn."),

    ("LB-150", 3,
     "A conversation approaches its context budget; never in the middle of a tool round.",
     "A conversation approaches its context budget, or the principles (LB-126) have changed since it began; "
     "never in the middle of a tool round."),

    ("LB-145", 1,
     "Each turn's context is assembled from the matter file in layers, and every item in it says where it came "
     "from and how it is known.",
     "The model's context is assembled from the matter file in layers when a conversation begins and at each "
     "compaction, then extended by appending (LB-147); every item in it says where it came from and how it is "
     "known."),

    ("LB-145", 3,
     "Every turn, when the loop starts; and when the turn's working material outgrows its budget.",
     "When a conversation begins and at each compaction (LB-147, LB-150) -- within a conversation, changes are "
     "appended rather than rebuilt; and when the turn's working material outgrows its budget."),

    ("LB-138", 4,
     "(2) the result envelope, the tool registry and the core tools wrapping existing code;",
     "(2) the result envelope, the tool registry and the core tools wrapping existing code, WITH TEXT EXTRACTION "
     "FROM BORN-DIGITAL PDFs so read_document and search_matter have text to read -- eleven rows wait on "
     "extraction (OM-Q03, OM-I01, OM-I02, F-B-05, F-B-16, F-C-06, LB-08, LB-11, LB-12, LB-33, LB-37) and a "
     "Telangana brief usually arrives as a document; the governing_code tool here also makes LB-120 reachable "
     "without a criminal cause;"),

    ("LB-138", 4,
     "(8) document extraction, which unblocks the rows waiting on uploaded material;",
     "(8) the rest of document extraction -- scanned pages (OCR), photographs and audio;"),

    ("LB-138", 4,
     "The switch moves only when the loop matches or beats the pipeline;",
     "The switch moves one kind of turn at a time -- a statute question with no matter first (GS-02), then the "
     "frame, dates, proof, theory and duty suites -- each only when the loop matches or beats the pipeline on "
     "that kind's golden subset;"),
]

# =============================================================================
# 2. DATED ADDITIONS, each in the column it belongs to.
# =============================================================================
APPEND: dict[str, list[tuple[int, str]]] = {}


def add(key: str, column: int, text: str) -> None:
    APPEND.setdefault(key, []).append((column, text))


def golden(key: str, text: str) -> None:
    add(key, 8, f"Golden binding ({R}, proposed): {text}")


# ---- Finding 1: the scratch pad and every other channel pass ONE detector ----
add("LB-67", 4,
    f"{P}: ONE DETECTOR ON EVERY CHANNEL. Under the loop the channels include the scratch pad and its step "
    "events (LB-139, LB-164), the matter-board items (LB-166), the model's questions (LB-161 ask_advocate) and "
    "the compaction handover (LB-150). Every model-written line on any of them passes the same citation "
    "detector G-GROUND uses (grounding.unretrieved_authorities) before it leaves the server; a line naming an "
    "authority or provision not retrieved on this matter is left out, the omission is shown, and the dropped "
    "text is kept only in encrypted diagnostics. This row is the single owner of that rule; the rows it names "
    "point here rather than restating it.")
add("LB-67", 8,
    f"{P}: LB-67-AC5 (planted): an opposing-counsel step event naming a case that was never retrieved never "
    "reaches the stream, the saved scratch pad or the board, and the omission is shown.\n"
    "LB-67-AC6: every channel calls G-GROUND's own detector -- a repository check fails on a second notion of "
    "'names an authority' (the one-detector rule, tests/test_one_remembered_authority_does_not_cost_the_turn.py).")
add("LB-67", 10,
    f"{R}: proposed amendment recorded -- one detector on every channel, the scratch pad included. The measured "
    "defect behind it: on 23 September a model read wrote a case recalled from memory into an opposing "
    "argument; the scratch pad as drafted in LB-139 would have streamed that name live before any check ran. "
    "This row and LB-70 already forbade unreviewed model text as progress; LB-139 is reconciled to them here. "
    "Not approved until the owner reviews it.")

add("LB-139", 7, "Never stream a line that names an authority or provision not retrieved on this matter (LB-67).")
add("LB-139", 8,
    "LB-139-AC5 (planted): a pass-2 event naming an unretrieved case is left out of the live panel and the saved "
    "scratch pad, and the omission is shown (LB-67-AC5).")
golden("LB-139", "LB-139-AC1 on GS-12 (two causes in one brief).")
add("LB-139", 9,
    f"{R}: reconciled with LB-70 ('no ... unreviewed model text masquerading as safe progress') and LB-67 ('show "
    "safe progress before release, not unchecked legal tokens'): the panel's lines pass LB-67's detector before "
    "they stream; the fuller release checks remain the answer's.")
add("LB-139", 10, f"{R}: amended so the scratch pad cannot carry an invented authority (LB-67). Still a draft for "
                  "owner review.")

add("LB-164", 4,
    f"{R}: an event carrying model-written text is released to the stream and the saved log only after LB-67's "
    "detector; what the detector removes is kept in encrypted diagnostics, never in the stream.")
add("LB-164", 7, "Never stream model-written text the detector has not passed (LB-67).")
add("LB-164", 10, f"{R}: amended -- step events pass LB-67's detector. Still a draft for owner review.")

add("LB-161", 4, f"{R}: the question's text passes LB-67's detector like every other channel.")
add("LB-161", 8,
    "LB-161-AC4 (carried from LB-62-AC4): an external act whose acknowledgement was lost is reconciled, never "
    "re-executed blindly.\n"
    "LB-161-AC5 (planted): a question submitted through ask_advocate that names an unretrieved authority is "
    "stopped by LB-67's detector.")
golden("LB-161", "the refusals on GS-05 (a backdated acknowledgment) and GS-18 (an original the client does not "
                 "hold).")
add("LB-161", 10, f"{R}: questions pass LB-67's detector; LB-62-AC4 carried. Still a draft for owner review.")

add("LB-70", 10,
    f"{P}: the scratch pad (LB-139) is this row's live view of the loop. The two are reconciled through LB-67: "
    "every model-written line in the scratch pad passes LB-67's detector before it streams, so it is checked "
    f"progress, not unreviewed model text. {MAP}")

# ---- Finding 2: one owner per rule -- the map, and the clauses carried ----
add("LB-43", 4,
    f"{P}: ONE OWNER PER RULE. The loop rows (LB-126 to LB-167) were drafted beside rows that already state "
    "parts of the same rules, and the L header's 'the L.2 rows govern' marks no clause as replaced. Proposed "
    "map, each owner carrying the others' surviving clauses as acceptance criteria, named in each row: the "
    "loop's run and stop -- LB-128 (from LB-01, LB-36, LB-61, LB-64, LB-69); the answer's claims -- LB-130 (from "
    "LB-57, LB-59); the adversarial mechanism -- LB-140 (from LB-18, LB-65, LB-115; also F-D-11 to F-D-13); "
    "which checks run on the loop's output -- LB-131 (from LB-38); research -- LB-136 (from LB-105, at the "
    "switch); the board's questions -- LB-166, building on LB-117 and LB-118 and carrying LB-05; context -- "
    "LB-145, LB-147 and LB-150, building on LB-58 and carrying LB-10; the scratch pad -- LB-139, reconciled with "
    "LB-70 through LB-67. OLDER ROWS THAT STAY THE OWNERS, the new rows implementing them on the loop path: "
    "LB-62 (admission; LB-143), LB-63 (one writer; LB-142), LB-66 (verification records; LB-131, LB-141), LB-67 "
    "(every channel), LB-71 (model qualification; LB-141, LB-153), LB-103 (support states; LB-141), LB-60 (task "
    "contracts; LB-126), LB-39 and LB-40 (evaluation of OM-P01-14). The practice layer restates later-phase "
    "rows: LB-121 and F-D-06; LB-125 and F-D-04, F-D-05, F-D-07. Nothing is marked superseded until the owner "
    "approves the owning row. MEASURED TODAY: 14 L.0 rows have columns B-J empty (F-B-02 to F-B-04, F-B-06 to "
    "F-B-15, F-C-03) and six sets repeat another row's B-J (F-B-01 = OM-Q01; F-C-04 = LB-01; F-C-05 = LB-38; "
    "F-C-07 and F-C-09 = LB-10; F-C-11 = LB-30; F-C-08 and F-C-12 = OM-Q05).")
add("LB-43", 8,
    f"{P}: LB-43-AC5: requirement_problems() fails on a requirement row whose columns B-J are empty, and on one "
    "whose B-J repeat another row's without a declared alias.\n"
    "LB-43-AC6: a row that restates another names its owner; a rule stated by two rows with no named owner "
    "fails the check.")
add("LB-43", 10, f"{R}: ownership map and two proposed checks recorded. Not approved until the owner reviews it.")

add("LB-128", 6,
    f"{R}: a no-progress stop fires when repeated equivalent tool calls or observations, oscillating positions "
    "or revisions that change nothing recur -- a rephrased or relabelled query is not new work (LB-105) -- and it "
    "keeps the unresolved needs. One allowance covers the turn, its retries and every nested loop; a resumed turn "
    "never resets it without a recorded decision (carried from LB-69).")
add("LB-128", 7, "Never iterate until the answer favours the client; never count rising confidence as progress "
                 "(LB-64).")
add("LB-128", 8,
    "LB-128-AC4 (planted, carried from LB-64-AC2): repeated equivalent observations end in a no-progress stop "
    "with the unresolved needs kept.\n"
    "LB-128-AC5 (carried from LB-64-AC4): an answerable, an externally blocked and a budget-limited task end in "
    "three distinct, truthful stop states.\n"
    "LB-128-AC6 (planted, carried from LB-69-AC1): nested loops and retries draw on one allowance, which holds "
    "under restart.\n"
    "LB-128-AC7 (carried from LB-36-AC1 and LB-01-AC3): a narrow question is answered without a compulsory full "
    "intake, and the matter is not closed by it.\n"
    "LB-128-AC8 (carried from LB-61-AC1): rephrasing a message without changing its meaning does not change "
    "what the loop does in substance.")
golden("LB-128", "LB-128-AC1 on GS-02 (one statute question, no matter) first, then the frame suite, GS-06 to "
                 "GS-11.")
add("LB-128", 9,
    f"{R}: SINGLE OWNER of the loop's run-and-stop behaviour. Carried here on the owner's approval of this row: "
    f"LB-64's four stop states and no-progress rule, LB-69's single allowance, LB-36's sufficiency, LB-61's "
    f"rephrasing invariance and LB-01's narrow-task rule; those rows then point here. {MAP}")
add("LB-128", 10, f"{R}: four stop states, carried criteria, AC3 replaced by a replay test, golden binding "
                  "proposed. Still a draft for owner review.")

for ident, carried in {
    "LB-01": "LB-01-AC3 (a narrow task finishes without closing the matter) as LB-128-AC7; LB-01-AC1 and AC2 "
             "(a relevant next action on unseen tasks; switching between briefing, research and advice) by "
             "LB-128 and LB-163",
    "LB-36": "LB-36-AC1 as LB-128-AC7; the distinct stop states as LB-128's four endings",
    "LB-61": "LB-61-AC1 as LB-128-AC8; LB-61-AC4 (an unsupported action refused at admission) by LB-127-AC2 "
             "and LB-143",
    "LB-64": "the four stop states and the no-progress rule as LB-128's exit and failure clauses, LB-128-AC4 and "
             "AC5; LB-64-AC3 (no drift under pressure) stays with LB-17 and LB-39, which own independence",
    "LB-69": "one allowance across retries and nested loops, never reset on resume, as LB-128-AC6; LB-69-AC4 (a "
             "cached result revalidated after a change) by LB-149 and LB-152",
}.items():
    add(ident, 10,
        f"{P}: the loop rows restate this row. Proposed single owner: LB-128. Carried there: {carried}. On the "
        f"owner's approval of LB-128 this row is marked 'carried into LB-128'; until then it stands unchanged. "
        f"{MAP}")

add("LB-130", 8,
    "LB-130-AC4 (planted, carried from LB-57-AC1): an invented fact, an invented provision and a valid-looking "
    "source identifier from another matter are each refused as support.\n"
    "LB-130-AC5 (carried from LB-57-AC4): with a store blocked, the answer gives honest limited help and neither "
    "supplies law from memory nor says the law is absent.\n"
    "LB-130-AC6 (planted, carried from LB-59-AC2): a circular dependency or a missing necessary premise refuses "
    "the dependent inference while independent claims are served.")
golden("LB-130", "LB-130-AC2 on the grounding suite -- GS-02, GS-12, GS-14, GS-17, GS-25.")
add("LB-130", 9,
    f"{R}: SINGLE OWNER of the answer's claim structure. LB-57 keeps the grounding rule as stated and "
    "owner-agreed; this row is its mechanism on the loop path. LB-59's legal-application record -- the test, "
    "its exceptions, jurisdiction, date and stage -- is what a law or inference claim carries, and claim "
    f"identities are the dependency ledger's, not a parallel claim store (LB-59). {MAP}")
add("LB-130", 10, f"{R}: claim kinds widened to keep allegation, document and hypothesis apart; LB-57 and LB-59 "
                  "criteria carried; golden binding proposed. Still a draft for owner review.")
add("LB-57", 10,
    f"{P}: LB-130 is this row's mechanism on the loop path; LB-57-AC1 and AC4 are carried there as LB-130-AC4 and "
    f"AC5, and the research-hypothesis label as a claim kind. This row stays the statement of the rule. {MAP}")
add("LB-59", 10,
    f"{P}: the claim structure is owned by LB-130; this row's legal-application record and LB-59-AC2 are carried "
    f"there (LB-130-AC6). On the owner's approval of LB-130 this row is marked 'carried into LB-130'; until then "
    f"it stands. {MAP}")

add("LB-140", 4,
    f"{R}: each argument is labelled raised (in a pleading or submission on the file), reasonably anticipated, "
    "contingent or speculative, and only the first three reach the answer (carried from LB-18 and LB-65). An "
    "objection with no supporting passage does not change the assessment (LB-65); a sound position survives a "
    "pass unchanged -- agreement is not required, and the loop does not iterate until the answer favours the "
    "client (LB-64). Before the extra passes are enabled, a single-pass and a challenged run are compared on "
    "held-out matters, reporting material gains, regressions, added latency and cost (LB-65-AC4).")
add("LB-140", 8,
    "LB-140-AC5 (planted, carried from LB-65-AC1): a plausible objection with no passage leaves the assessment "
    "unchanged and is recorded as unsupported.\n"
    "LB-140-AC6 (carried from LB-65-AC3 and LB-18): actual submissions on the file, anticipated defences and "
    "contingent ones are labelled apart in the answer and the scratch pad.\n"
    "LB-140-AC7 (carried from LB-65-AC4): the three passes are enabled only after a recorded single-pass versus "
    "three-pass comparison, with its cost.")
golden("LB-140", "LB-140-AC1 on GS-13 (cheque dishonour); pass 2 on GS-23 (the opponent at their strongest); "
                 "LB-140-AC3 on GS-22 (cross-thread exposure) and GS-21 (two arguments that cannot both be true).")
add("LB-140", 9,
    f"{R}: SINGLE OWNER of the adversarial mechanism. Carried from LB-18, LB-65 and LB-115 on the owner's "
    f"approval of this row. Also restated by F-D-11, F-D-12 and F-D-13 in the later phases, which point here. "
    f"{MAP}")
add("LB-140", 10, f"{R}: argument labels, the unsupported-objection rule and the single-versus-three-pass "
                  "comparison carried; golden binding proposed. Still a draft for owner review.")
add("LB-18", 10,
    f"{P}: the adversarial mechanism is owned by LB-140; this row's four argument categories (raised, "
    "anticipated, contingent, speculative) are carried there. LB-18-AC4 (reliable disconfirming material "
    "updates the view; pressure alone does not) stays with LB-17 and LB-39. On the owner's approval of LB-140 "
    f"this row is marked 'carried into LB-140'; until then it stands. {MAP}")
add("LB-65", 10,
    f"{P}: owned by LB-140 on the loop path; LB-65-AC1, AC3 and AC4 are carried there as LB-140-AC5, AC6 and "
    f"AC7. Until the owner approves LB-140 this row stands. {MAP}")
add("LB-115", 10,
    f"{P}: the live defect this row records stays OPEN here. Its served-path checks on the loop are LB-140-AC4 "
    "(an objection without a span is refused) and LB-130-AC1 (a provision never retrieved is refused); it closes "
    f"only on a live run, not on those planted tests. {MAP}")

add("LB-166", 8,
    "LB-166-AC4 (carried from LB-118): an item the advocate promises returns when due and is not re-asked "
    "before then; an item marked unavailable is not re-asked without a new reason.\n"
    "LB-166-AC5 (carried from LB-05-AC2): a detail already held in an uploaded document is found by "
    "search_matter and not asked for.")
golden("LB-166", "LB-166-AC1 on GS-13 (cheque dishonour).")
add("LB-166", 9,
    f"{R}: BUILDS ON, does not replace: LB-117 (the source-backed checklist, its four states and its cache by "
    "content identity) and LB-118 (replies update the checklist), both partly built and tested. This row adds "
    "the per-dispute board, the passage or defence behind each item, and pass-1 defences as a source of items; "
    f"LB-05's questioning principles are carried. {MAP}")
add("LB-166", 10, f"{R}: items carry LB-117's four states and move through LB-118; criteria carried; golden "
                  "binding proposed. Still a draft for owner review.")
add("LB-117", 10, f"{P}: LB-166 builds on this row -- the board's items carry these four states and this cache "
                  f"-- and does not replace it. {MAP}")
add("LB-118", 10, f"{P}: LB-166's board items are moved by this row's reply read; its promise and unavailable "
                  f"rules are carried as LB-166-AC4. {MAP}")
add("LB-05", 10, f"{P}: LB-166 carries this row's questioning principles on the board; LB-05-AC2 (retrieve what "
                 f"is held instead of re-asking) is carried as LB-166-AC5. This row stays a guiding principle "
                 f"(LB-126). {MAP}")

add("LB-136", 6,
    f"{R}: a research loop does not repeat work -- rephrasing or relabelling a query does not make it new work, "
    "a proposal against an out-of-date matter snapshot is refused, and its stop names which happened: provider "
    "failure, a bad proposal, retrieval failure, repetition, no progress, or budget (carried from LB-105).")
add("LB-136", 8, "LB-136-AC3 (planted, carried from LB-105): a relabelled repeat of an earlier query is refused as "
                 "repeated work, and a stale-snapshot proposal is refused.")
add("LB-136", 9, f"{R}: LB-105 is today's bounded research loop, built and tested; this row replaces its fixed "
                 f"rounds when the loop is switched on (LB-138), carrying its rules above. {MAP}")
add("LB-136", 10, f"{R}: LB-105's rules carried. Still a draft for owner review.")
add("LB-105", 10, f"{P}: when the loop is switched on, LB-136 replaces this loop's fixed rounds and carries its "
                  f"repeated-query, stale-snapshot and distinct-stop rules (LB-136-AC3). Until the switch this row "
                  f"is the live behaviour. {MAP}")
add("LB-137", 7, "Never count a rephrased or relabelled query as new work (LB-105).")
add("LB-137", 10, f"{R}: LB-105's repeated-query rule added. Still a draft for owner review.")

add("LB-143", 8, "LB-143-AC4 (planted, carried from LB-62-AC3): access revoked mid-turn refuses every later tool "
                 "call on that matter.")
add("LB-143", 9,
    f"{R}: IMPLEMENTS, does not replace: LB-62 (capability and permission checks before action) and LB-63 (one "
    "controlled writer) stay the owners; this row is them on the loop's tool calls. Carried: LB-62-AC3 as "
    "LB-143-AC4; LB-62-AC4 into LB-161 (LB-161-AC4); LB-63-AC1 and AC2 into LB-142 (LB-142-AC4, AC5). "
    f"{MAP}")
add("LB-143", 10, f"{R}: owners named (LB-62, LB-63); revocation criterion carried. Still a draft for owner "
                  "review.")
add("LB-62", 10, f"{P}: stays the owner of admission. LB-143 applies it to every tool call in the loop; LB-62-AC3 "
                 f"is carried as LB-143-AC4 and LB-62-AC4 as LB-161-AC4. {MAP}")
add("LB-63", 10, f"{P}: stays the owner of the single controlled writer. LB-142's checked writes are this writer "
                 f"on the loop path; LB-63-AC1 and AC2 are carried as LB-142-AC4 and AC5. {MAP}")
add("LB-142", 8,
    "LB-142-AC4 (carried from LB-63-AC1): duplicate and out-of-order writes have one accepted effect, and a newer "
    "value is kept.\n"
    "LB-142-AC5 (carried from LB-63-AC2): a crash between a write's check and its commit resumes with no lost "
    "and no duplicated effect.")
golden("LB-142", "LB-142-AC2 on GS-15 (a date corrected mid-conversation).")
add("LB-142", 9, f"{R}: LB-63 stays the owner of the single writer; this row is it on the loop path. {MAP}")
add("LB-142", 10, f"{R}: LB-63's criteria carried; golden binding proposed. Still a draft for owner review.")

add("LB-131", 4,
    f"{R}: when one check fails and repair cannot fix it, only the affected claim and the conclusions resting on "
    "it are withheld; the rest of the answer is re-checked for coherence and served with the limit stated "
    "(carried from LB-38 and LB-67). Each check outcome is written as an LB-66 verification record. During the "
    "overlap with the pipeline, the loop and the pipeline call the same gate functions -- one implementation, "
    "two callers.")
add("LB-131", 8,
    "LB-131-AC3 (planted, carried from LB-38-AC3 and LB-67-AC3): one failed dependency withholds that conclusion "
    "and what rests on it, and the independent help is still served.\n"
    "LB-131-AC4: a repository check fails the build if the loop path calls a gate function the pipeline does "
    "not, or a copy of one (LB-138).")
add("LB-131", 9,
    f"{R}: SINGLE OWNER of which checks run on the loop's output. LB-67 owns every channel through one release "
    "service; LB-66 owns the verification record; LB-38's release checks are carried here. F-C-05 repeats "
    f"LB-38's columns B-J. {MAP}")
add("LB-131", 10, f"{R}: scoped withholding and one-implementation rule added. Still a draft for owner review.")
add("LB-38", 10,
    f"{P}: on the loop path this row's checks are LB-131's, and LB-38-AC3 is carried there as LB-131-AC3; LB-67 "
    "owns every channel. On the owner's approval of LB-131 this row is marked 'carried into LB-131'; until then "
    f"it stands. F-C-05 repeats this row's B-J. {MAP}")
add("LB-66", 10, f"{P}: stays the owner of the verification record. On the loop path, each of LB-131's check "
                 f"outcomes and each LB-141 verdict is written as one of these records. {MAP}")
add("LB-133", 4, f"{R}: after the last attempt, what is still failing is withheld with the conclusions that rest "
                 "on it, and the rest of the answer is re-evaluated for coherence before it is served (LB-67).")
add("LB-133", 10, f"{R}: LB-67's dependent-withholding rule added. Still a draft for owner review.")

add("LB-58", 10, f"{P}: stays the owner of the durable, versioned matter snapshot. LB-145, LB-147 and LB-150 are "
                 f"how the loop builds its context from it; LB-58-AC2 is exercised on the loop path by LB-152-AC1. "
                 f"{MAP}")
add("LB-10", 10, f"{P}: the loop's context rows (LB-135, LB-145, LB-147, LB-150) carry this row's composition of "
                 f"context; LB-10-AC3 (context limits disclosed, never false completeness) is carried as "
                 f"LB-135-AC4. F-C-07 and F-C-09 repeat this row's B-J. {MAP}")
add("LB-135", 8, "LB-135-AC4 (carried from LB-10-AC3): when the context budget forces material out, the answer "
                 "discloses the coverage limit rather than implying it read everything.")
golden("LB-135", "none yet -- no golden scenario runs longer than six turns (see LB-150).")
add("LB-135", 10, f"{R}: LB-10-AC3 carried. Still a draft for owner review.")

add("LB-153", 9,
    f"{R}: OPEN -- the tolerance in the exit condition is not set; it is set from the first recorded provider "
    "comparison and names which differences count (a grounding failure counts at any rate; wording does not). "
    "LB-71 stays the owner of qualifying each model for each role on held-out behaviour; this row owns the "
    f"context policy. LB-71-AC4 and LB-153-AC1 are one replay test, run once. {MAP}")
golden("LB-153", "LB-153-AC3 on the smoke and grounding suites.")
add("LB-153", 10, f"{R}: tolerance recorded as OPEN; LB-71 named as owner of qualification. Still a draft for "
                  "owner review.")
add("LB-71", 10, f"{P}: stays the owner of per-role model qualification -- which is how LB-141's verifier tier is "
                 f"chosen. LB-153 owns the context policy across providers; LB-71-AC4 and LB-153-AC1 are one "
                 f"replay test. {MAP}")

# ---- Finding 3: the two contradictions between draft rows ----
add("LB-126", 4,
    f"{R}: OM-P01-14 become the numbered sections of this document rather than rows with their own build "
    "status; their evaluation is owned by LB-39 and LB-40.")
add("LB-126", 8,
    "LB-126-AC4 (carried from LB-60-AC1): the served model request carries the current principles version, read "
    "from the request itself, not from the file.\n"
    "LB-126-AC5: a principles change made mid-conversation forces a compaction at the next turn boundary; the "
    "prefix of the running conversation is never edited (LB-147-AC1 holds).")
add("LB-126", 10, f"{R}: a principles change now forces a compaction (it contradicted LB-147's frozen prefix); "
                  "OM-P01-14 proposed as sections. Still a draft for owner review.")
add("LB-147", 4, f"{R}: the principles (LB-126) are part of the prefix, so a change to them forces a compaction "
                 "at the next turn boundary rather than an edit to the running conversation.")
add("LB-147", 10, f"{R}: reconciled with LB-126. Still a draft for owner review.")
add("LB-150", 4, f"{R}: the handover's model-written lines also pass LB-67's detector.")
golden("LB-150", "none of the 25 golden scenarios runs longer than six turns, so LB-150-AC1 and LB-135-AC1 need a "
                 "long-matter scenario added to docs/GOLDEN_SET.md -- OPEN.")
add("LB-150", 10, f"{R}: a principles change also triggers compaction; the handover passes LB-67. Still a draft "
                  "for owner review.")
add("LB-145", 4, f"{R}: each run names the matter snapshot its brief was built from (LB-58, which stays the owner "
                 "of the durable snapshot).")
golden("LB-145", "LB-145-AC2 on GS-04 (an injected instruction in an upload).")
add("LB-145", 10, f"{R}: aligned with LB-147, which had refined this row without its text being changed; LB-58 "
                  "named as the snapshot's owner. Still a draft for owner review.")

# ---- Finding 4: criteria bound to named golden conversations ----
golden("LB-163", "LB-163-AC1 on GS-12 (two causes in one brief); LB-163-AC2 on GS-15 (a date corrected "
                 "mid-conversation).")
add("LB-163", 10, f"{R}: golden binding proposed. Still a draft for owner review.")
golden("LB-165", "LB-165-AC1 on GS-12 and GS-08 (three threads).")
add("LB-165", 10, f"{R}: golden binding proposed. Still a draft for owner review.")
golden("LB-144", "LB-144-AC2 on the dates suite -- GS-13, GS-14, GS-15 -- and GS-07 (a custody clock).")
add("LB-144", 10, f"{R}: golden binding proposed. Still a draft for owner review.")
golden("LB-146", "LB-146-AC1 first on the smoke suite, GS-01 to GS-05, which runs on every commit.")
add("LB-146", 10, f"{R}: golden binding proposed. Still a draft for owner review.")
golden("LB-156", "LB-156-AC3 on GS-16 (the era rule, straddled) and GS-06.")
add("LB-156", 10, f"{R}: golden binding proposed. Still a draft for owner review.")
golden("LB-157", "LB-157-AC3 on GS-25 (a judgment against the exact point).")
add("LB-157", 10, f"{R}: golden binding proposed. Still a draft for owner review.")
golden("LB-132", "LB-132-AC2 on GS-04 -- which needs text extraction (LB-138 slice 2) to run at all.")
add("LB-132", 10, f"{R}: golden binding proposed. Still a draft for owner review.")
golden("LB-123", "GS-24 (salvage by coordinate, Order 39).")

# ---- Finding 5: this week's measurements ----
add("LB-141", 8,
    "LB-141-AC4: the agreement measured under AC2 counts both error directions apart -- a supported claim marked "
    "'does not support' and an unsupported claim marked 'supports' -- and the tier is recorded with the "
    "measurement that chose it.")
add("LB-141", 9,
    f"{R}: the verdict is carried in LB-103's three states (assessed true, assessed false, not assessed) and "
    "recorded as an LB-66 verification record; those rows stay the owners of the carrier and the record. "
    "MEASURED THIS WEEK: the one comparable harness read measured -- the consistency check, G-CONSISTENT -- was "
    "moved to the hard tier on its measurement (commit f90df19); the step-independence read answered "
    "'dependent' for every step until it was given the reason first and two observable halves "
    "(tests/test_the_independence_read_reasons_before_it_rules.py). A verifier's tier is a measurement, not a "
    f"design choice. {MAP}")
add("LB-141", 10, f"{R}: tier set by measurement, reason before verdict, both error directions counted; owners "
                  "of the carrier and record named. Still a draft for owner review.")
add("LB-103", 10, f"{P}: LB-141's independent verifier is what fills 'assessed true / assessed false' here; this "
                  f"row stays the carrier of the three states. {MAP}")

add("LB-134", 4,
    f"{R}: MEASUREMENT BEFORE CONVERSION. The 'addressed' check a principle relies on is itself a model read, and "
    "the one gate read measured this week -- G-LIMITATION's step-independence read -- answered 'dependent' for "
    "every step as shipped (10 of 18 correct) and reached 16 of 18 only once it gave its reason first and "
    "answered two observable halves from which code derived the verdict. So a gate converts to a principle only "
    "when its 'addressed' check has been measured on a labelled set, both error directions counted apart, and "
    "built on the LB-60 patterns; until then it stays a floor, whatever the decision in principle.")
add("LB-134", 8, "LB-134-AC2: no gate is recorded as converted without a labelled measurement of its 'addressed' "
                 "check, both error directions counted apart, cited by path.")
add("LB-134", 9, f"{R}: evidence -- tests/test_the_independence_read_reasons_before_it_rules.py; "
                 "development_environment/one_off_tools/independence_classifier_20260922.py; "
                 "development_environment/reviews/LIVE_MATTER_LOOP_20260922.md.")
add("LB-134", 10, f"{R}: a measured precondition added to every conversion. Still a draft for owner review.")

add("LB-60", 4,
    f"{P}: TWO MEASURED PATTERNS FOR THE HARNESS'S OWN MODEL READS -- the verifier (LB-141), the 'addressed' "
    "checks (LB-134), consistency, and any read that decides a gate. (1) Reason before verdict: the schema emits "
    "the reason before the verdict, because strict structured output generates in schema order and a verdict "
    "emitted first commits before it has reasoned. (2) Observable halves, verdict derived in code: where the "
    "model cannot hold an abstract test, it answers the concrete questions the test is made of and code derives "
    "the verdict; an answer whose halves contradict its own verdict is UNKNOWN, which the gate treats as the "
    "unsafe side. These describe how the harness builds its reads; they are not principles for the reasoning "
    "model (LB-126).")
add("LB-60", 8,
    f"{P}: LB-60-AC5: every harness read that decides a gate emits its reason before its verdict -- a repository "
    "check reads the schema order.\n"
    "LB-60-AC6: where a gate's read is built from halves, a contradiction between the halves and the stated "
    "verdict yields UNKNOWN, never a release (tests/test_the_independence_read_reasons_before_it_rules.py).")
add("LB-60", 10, f"{P}: LB-126 carries LB-60-AC1 as LB-126-AC4 for the principles document. This row stays the "
                 f"owner of the harness's task contracts and now records the two measured patterns. {MAP}")

# ---- Finding 6 and 7: slice order, switch per kind of turn, cost ----
add("LB-138", 4,
    f"{R}: EACH SLICE'S PULL REQUEST lists its acceptance criteria in two groups -- replay (every commit, no API "
    "key, no cost) and live (each run approved by the owner, with its ledger estimate). Before slice 3 serves a "
    "live turn, the cost per turn is measured from replay call counts times price and scored as a release row "
    "(LB-35). The golden set's slice tags (S1 to S10 in docs/GOLDEN_SET.md) belong to the earlier slice plan: "
    "each scenario is re-tagged with the slice above at which it can first run, or 'slice-N' selects the wrong "
    "set.")
add("LB-138", 8,
    "LB-138-AC3: during the overlap, a repository check fails the build if the loop and the pipeline call "
    "different implementations of a gate (LB-131-AC4).\n"
    "LB-138-AC4: every golden scenario carries an LB-138 slice tag, and each slice's pull request lists its "
    "replay and live criteria with the ledger estimate for the live ones.")
add("LB-138", 9, f"{R}: OPEN -- the revised slice order (extraction of born-digital PDFs moved into slice 2), the "
                 "per-kind switch order, and the cost-per-turn threshold (LB-35, LB-44).")
add("LB-138", 10, f"{R}: born-digital PDF extraction proposed for slice 2; switch per kind of turn; replay/live "
                  "split and cost gate added; golden tags to be re-mapped. Still a draft for owner review.")
add("LB-155", 9, f"{R}: read_document and search_matter depend on text extraction, proposed for slice 2 "
                 "(LB-138).")
add("LB-155", 10, f"{R}: extraction dependency recorded. Still a draft for owner review.")
add("LB-35", 4,
    f"{P}: COST PER TURN BECOMES A RELEASE ROW. The loop multiplies calls -- a full attack per dispute, a verifier "
    "call per claim, research sub-loops, a cross-matter pass -- and the budgets in this row, LB-69 and LB-128 "
    "are all OPEN. A cost-per-turn row is added to assurance/specification/release.yaml, scored PASS / FAIL / NOT "
    "MEASURED by pipeline/quality/releasegate.py like every other row, estimated from replay call counts times "
    "price before slice 3 serves a live turn (LB-138) and re-measured on each golden comparison. NOT MEASURED "
    "blocks exactly like FAIL.")
add("LB-35", 10, f"{P}: cost per turn proposed as a release row. Not approved until the owner reviews it.")

add("LB-44", 4,
    f"{P}: decisions added -- (1) which golden scenarios bind which acceptance criteria, and the 'ten rows that "
    "define great' raised on 25 September and never answered; (2) the verifier's tier, from its measurement "
    "(LB-141); (3) the cost-per-turn threshold (LB-35); (4) the revised slice order and the per-kind switch "
    "order (LB-138); (5) the one-owner map (LB-43); (6) which new practice-layer rows to build (LB-168, LB-169, "
    "LB-170) and who does counsel review; (7) still open from the 25 September review: practising Telangana "
    "advocates as LB-40's reviewers with a rubric agreed before release; Telugu-language material (F-C-08); "
    "Indian sources, including the Bar Council of India Rules, for the expert-practice research.")
add("LB-44", 10, f"{R}: open decisions listed. Not approved until the owner reviews it.")

# ---- Finding 9: features -- playbook pointers, governing code, the corpus ----
add("LB-148", 4,
    f"{R}: every legal pointer in a playbook is an exact key -- an Act and section, or a reporter citation -- "
    "that resolves through identify_act, read_provision or resolve_citation. A case is never named in a "
    "playbook by its parties alone: a name in context invites a citation that was never retrieved.")
add("LB-148", 8, "LB-148-AC4 (planted): a repository check resolves every pointer in every playbook through the "
                 "exact tools, and fails the build on one that does not resolve or names a case by parties alone.")
add("LB-148", 10, f"{R}: playbook pointers made exact keys. Still a draft for owner review.")
add("LB-120", 10,
    f"{P}: under the loop the governing-code table is reached through the governing_code tool (LB-156), which "
    "needs dates, not a cause of action -- so this row's 'unwired' status ends in LB-138's slice 2. Golden "
    "binding: GS-16 (the era rule, straddled) and GS-06.")
add("LB-121", 9,
    f"MEASURED {R[:17]} against raw_data/ and the derived stores, as this row required: the Commercial Courts "
    "Act 2015 IS HELD -- raw_data/BareActs/Union of India/2015_3_...Commercial Courts...Act, 2015.txt; 371 "
    "chunks in chunks.db under its long title; and in legal.db as 'The Commercial Courts Act, 2015' -- with a "
    "stray byte-order mark in the file name and one act_id. It is absent only from pipeline/manifest.yaml, so "
    "exact identification cannot reach it. s.12A therefore needs a manifest entry and the stray character "
    "removed, then curation -- not an acquisition. This row restates F-D-06 (pre-filing requirements); LB-43 "
    "proposes this row as the owner.")
add("LB-125", 9,
    f"MEASURED {R[:17]} against raw_data/ and the derived stores, as this row required: the Telangana "
    "Court-fees and Suits Valuation Act, 1956 is HELD (raw_data/BareActs/Telangana/2016_10_...; 461 chunks in "
    "chunks.db; 79 sections in legal.db), as are the Andhra Pradesh version (456 chunks) and the Telangana Civil "
    "Courts Act, 1972 (122 chunks; 31 sections in legal.db). The schedules are in the raw text (18 'Schedule' "
    "headings, 41 'Rs.' amounts) but chunks.db holds only section atoms -- no schedule atoms -- and none of these "
    "Acts is in pipeline/manifest.yaml. The gap is the manifest entry and the schedule atoms, not the source; "
    "whether the schedules are current is NOT MEASURED. This row restates F-D-04, F-D-05 and F-D-07; LB-43 "
    "proposes this row as the owner.")

# ---- Finding 10: status per acceptance criterion; OM-P rows become principles ----
add("LB-167", 4,
    f"{R}: STATUS PER ACCEPTANCE CRITERION. 147 of 189 rows read 'In progress', which does not steer. Each "
    "acceptance criterion carries two facts -- is there a runnable check for it (a test path or a golden "
    "scenario), and did it pass as written -- and a row's status is the count. The OM-P rows, whose only gap is "
    "held-out evaluation, become sections of the principles document (LB-126) and stop carrying a build status "
    "of their own.")
add("LB-167", 8, "LB-167-AC4: the status pass reports, per row, how many acceptance criteria have a runnable check "
                 "and how many passed as written.")
add("LB-167", 10, f"{R}: status per acceptance criterion proposed. Still a draft for owner review.")
for n in range(1, 15):
    add(f"OM-P{n:02d}", 10,
        f"{P}: to become a section of the LB-126 principles document; its evaluation is owned by LB-39 and LB-40, "
        f"and it stops carrying a build status of its own. Until the owner approves LB-126 this row stands.")
for ident in ("LB-39", "LB-40"):
    add(ident, 10, f"{P}: OM-P01-14 are evaluated here once they become sections of the principles document "
                   f"(LB-126). {MAP}")

# ---- Finding 8: L.0 rows with no B-J, rows that repeat another's B-J ----
EMPTY = ["F-B-02", "F-B-03", "F-B-04", "F-B-06", "F-B-07", "F-B-08", "F-B-09", "F-B-10", "F-B-11", "F-B-12",
         "F-B-13", "F-B-14", "F-B-15", "F-C-03"]
for ident in EMPTY:
    add(ident, 10,
        f"{R}: columns B-J have not been written for this row; its current requirement is the text under CURRENT "
        "REQUIREMENT in column A. Writing it into B-J -- objective, trigger, interaction, exits, failure, "
        "must/never, acceptance, dependencies -- is OPEN under LB-43.")
for ident, owner in {"F-B-01": "OM-Q01", "F-C-04": "LB-01", "F-C-05": "LB-38", "F-C-07": "LB-10",
                     "F-C-09": "LB-10", "F-C-11": "LB-30", "F-C-12": "OM-Q05"}.items():
    add(ident, 10,
        f"{R}: columns B-J repeat {owner}'s requirement, most columns word for word. One requirement is stated "
        f"twice; LB-43 proposes that this row name {owner} as its owner and keep only what is its own.")
add("F-C-08", 10,
    f"{R}: columns B-J hold OM-Q05's retention requirement, approved 21 September 2026, which superseded this "
    "row's retention wording. The other half of this row -- a voice note or document in Telugu, Hindi or Urdu "
    "transcribed in its own language and translated into English, both kept, and anything resting on a "
    "translated word showing the original word beside it -- was never carried into B-J and exists only in "
    "column A. Recorded here so it is not lost: OPEN for the owner (the 25 September review's 'Telugu-language "
    "material'); not started in code.")
for ident, owner in {"F-D-04": "LB-125", "F-D-05": "LB-125", "F-D-06": "LB-121", "F-D-07": "LB-125",
                     "F-D-11": "LB-140", "F-D-12": "LB-140", "F-D-13": "LB-140"}.items():
    add(ident, 10, f"{R}: restated in the legal brain by {owner}; LB-43 proposes {owner} as the owner.")

# =============================================================================
# 3. NEW ROWS, inside L.8 after LB-125.
# =============================================================================
READINESS = (
    f"26 September 2026: DRAFT prepared for owner review from the {R} of the legal-brain rows; not a verbatim "
    "owner statement, not counsel-reviewed and not an approved requirement. Implementation NOT_ASSESSED; "
    "acceptance NOT_RUN. Delivery mapping OPEN under LB-43; quality and release decisions OPEN under LB-44. Every "
    "legal proposition is a pointer to text that must be retrieved and read back before any table is curated. "
    "No completion or release approval.")

NEW = {
    "LB-168": [
        "Whether each document relied on can be read in evidence",
        "Before a document is relied on, NM must establish whether it can be read in evidence at all -- stamped, "
        "registered, the original or a lawful copy, an electronic record with the certificate it needs -- as a "
        "question separate from whether it exists and what weight it carries.",
        "Never build a case on a document the court will refuse to read, and learn in time what would cure it.",
        "A document on the file, or one the advocate means to rely on, is material to a dispute's proof (LB-165, "
        "evidence to collect).",
        "ONE MECHANISM, SHARED with LB-120 to LB-125: a curated table of admissibility conditions, each row "
        "carrying curated_from, identified on an exact key -- the kind of instrument, the purpose it is relied on "
        "for, its date, and whether an original is held -- never by fuzzy match. First population, each entry "
        "verified against the retrieved text before curation: an instrument chargeable with duty and not duly "
        "stamped, and its cure by impounding and penalty (Indian Stamp Act 1899, s.35, with the Andhra Pradesh "
        "and Telangana amendments); a compulsorily registrable document not registered, and what it may still be "
        "read for -- a collateral purpose; part performance under TPA s.53A as the 2001 amendments left it "
        "(Registration Act 1908, s.17 and s.49); secondary evidence where the original is not produced, and "
        "notice to produce (Evidence Act ss.63 to 66, and their BSA 2023 equivalents by the governing date, "
        "LB-120); an electronic record and the certificate it needs (Evidence Act s.65B; BSA 2023). The model "
        "applies the table to the document as the file describes it and returns, per document: admissible as "
        "held; admissible if a named step is taken, with its owner; not admissible for this purpose; or not "
        "assessed -- with the quoted span. The determination is written to the existing evidence inventory, "
        "which already keeps existence, admissibility and weight apart, and enters the premise record, so a "
        "corrected date or a newly found original reopens it.",
        "Each document relied on reads admissible as held, admissible if a named step is taken, not admissible "
        "for this purpose, or not assessed, with its source.",
        "A document described but not seen reads not assessed, with one question each; it is never assumed "
        "stamped, registered or original. A cure that costs time or money -- impounding and penalty -- is stated "
        "as the advocate's decision, never taken.",
        "Must keep existence, admissibility and weight as three questions. Never treat an unregistered or "
        "unstamped document as dead without asking what else it may be read for. Never advise that a document be "
        "described as lost to let in secondary evidence -- that is a duty refusal with the lawful route. Never "
        "let a table entry stand in for the provision; the provision is read (LB-156).",
        "LB-168-AC1: an unregistered agreement relied on for possession is assessed for what it may still be read "
        "for, with the provisions retrieved (GS-20, GS-17).\n"
        "LB-168-AC2: an original held by a third party yields the notice-to-produce route and a preservation "
        "step with an owner, and a request to call it lost is refused (GS-18).\n"
        "LB-168-AC3 (planted): a document of a kind the table does not hold returns 'no curated table', never a "
        "neighbouring entry.\n"
        "LB-168-AC4: a stamp-duty defect is surfaced with its cure, and correcting the instrument's date reopens "
        "it.",
        "LB-16 (authenticity, admissibility and weight), LB-19, LB-48, LB-159, LB-165; "
        "backend/nm/core/evidence_item.py (the inventory already carries admissibility as its own facet, not "
        "assessed by default, and the model is not asked to decide it). CORPUS, MEASURED 26 September 2026: "
        "the_indian_stamp_act_1899 (475 chunks), two Indian Stamp (Andhra Pradesh Amendment) Acts, "
        "the_registration_act_1908 (320 chunks) with the Andhra Pradesh Rules under it, "
        "the_indian_evidence_act_1872 (462 chunks), the Information Technology Act 2000 and the BSA 2023 are "
        "held in chunks.db. Golden scenarios already expect this row: GS-17, GS-18, GS-19, GS-20. Raised in the "
        f"{R}. OPEN: the owner's selection, and counsel review of every entry (BK-85-AC3).",
    ],
    "LB-169": [
        "What a pleading must contain, and what gets it rejected or returned",
        "Before a plaint or written statement is shown as a draft, NM must check it against the pleading rules "
        "that decide whether it is received at all -- what Order VII requires of a plaint and the grounds on "
        "which it rejects or returns one, what Order VIII requires of a written statement, and Order VI on "
        "material facts, particulars and verification.",
        "Never file a pleading that is returned or rejected before its merits are heard, or a written statement "
        "whose silence admits what we deny.",
        "A draft pleading is prepared (F-F-02, F-F-03), or the advocate asks whether a plaint or written "
        "statement on the file is sound.",
        "ONE MECHANISM, SHARED with LB-120 to LB-125: a curated table of pleading requirements, each row carrying "
        "curated_from, keyed exactly by the pleading and the rule family, and checked by the draft's verification "
        "step (F-F-03) against the drafter brief. First population, each entry verified against the retrieved CPC "
        "text and the Telangana amendments before curation (LB-170): the particulars a plaint must state (Order "
        "VII r.1); the grounds for rejection -- no cause of action disclosed, undervaluation or insufficient "
        "stamp not corrected in time, a suit barred by law on the plaint's own statements (Order VII r.11) -- and "
        "return to the proper court (Order VII r.10); specific denial, and an evasive or absent denial taken as "
        "an admission (Order VIII r.3 to r.5); set-off and counter-claim (Order VIII r.6, r.6A); material facts "
        "and not evidence, particulars where fraud or misrepresentation is pleaded, and verification (Order VI "
        "r.2, r.4, r.15), with the statement of truth a commercial suit adds. The written statement's time stays "
        "with LB-124. The model applies the table to the draft and returns each requirement met, not met (naming "
        "the paragraph), or not assessed.",
        "Each requirement the draft engages reads met, not met with the paragraph named, or not assessed; a draft "
        "with an unmet ground for rejection or return is not shown as ready.",
        "A requirement that depends on a fact not on the file -- the valuation, the date the cause arose -- is "
        "not assessed with one question, never assumed met. A draft that cannot meet a requirement shows it as a "
        "visible marked blank (F-F-03), never smoothed over.",
        "Must check the draft actually produced, not the brief it came from. Never plead evidence as material "
        "facts, or omit particulars where fraud is pleaded. Never show a draft as ready while an Order VII r.11 "
        "ground is unmet. Never let a table entry stand in for the rule; the rule is read (LB-156).",
        "LB-169-AC1: a draft plaint that does not state when the cause of action arose is not shown as ready, and "
        "the paragraph is named.\n"
        "LB-169-AC2: a draft written statement that does not deal with an allegation specifically names that "
        "allegation as at risk of being taken as admitted.\n"
        "LB-169-AC3 (planted): a draft whose valuation does not support the relief claimed is flagged under the "
        "rejection ground, with the valuation workings (LB-125).\n"
        "LB-169-AC4: every requirement result names its curated_from source.",
        "F-F-02 (the drafter brief) and F-F-03 (draft, and verify the draft), inside which these checks run; "
        "LB-23, LB-124, LB-125; backend/nm/domain/drafting.py and backend/nm/core/drafting.py (the drafting "
        "package, which keeps drafting readiness apart from filing readiness). CORPUS, MEASURED 26 September "
        "2026: the_code_of_civil_procedure_1908 (2,901 chunks in chunks.db) and the Code of Civil Procedure "
        "(Telangana) Second Amendment Act, 1953 (7 chunks) are held; state amendments to the Orders are checked "
        "before any entry is curated (LB-170). Raised in the 25 September review ('Order VII / VIII pleadings'), "
        "never answered. OPEN: the owner's selection, and counsel review of every entry (BK-85-AC3).",
    ],
    "LB-170": [
        "Local procedure: the Telangana rules of practice, the Civil Courts Act and the CPC's state amendments",
        "The rules that decide how a Telangana court receives and hears a matter -- the Civil and Criminal Rules "
        "of Practice, the Telangana Civil Courts Act and the state amendments to the CPC -- are identified "
        "exactly and read wherever a period, a forum, a form or a fee turns on them.",
        "Have NM follow the procedure my court actually applies, not only the central text.",
        "A period, forum, form, filing step or fee in a Telangana matter is governed or modified by a local rule "
        "or a state amendment.",
        "Each local source is given an exact identity in pipeline/manifest.yaml, so it is identified by exact "
        "match like every other Act (CLAUDE.md section 5), and is read through read_provision with its locator. "
        "MEASURED 26 September 2026 -- held but not identifiable: the Andhra Pradesh Civil Rules of Practice and "
        "Circular Orders, 1980 (raw_data/ only; no chunks in chunks.db); the Criminal Rules of Practice and "
        "Circular Orders, 1990 (1,139 chunks); the Telangana Civil Courts Act, 1972 (122 chunks); the Code of "
        "Civil Procedure (Telangana) Second Amendment Act, 1953 (7 chunks); the Telangana Court-fees and Suits "
        "Valuation Act, 1956 (461 chunks, no schedule atoms) -- none of them in the manifest. Whether each "
        "applies in Telangana as adapted is verified before it is relied on. Where a local rule or state "
        "amendment modifies a central provision, the answer reads both and says which governs. The curated "
        "tables of LB-121, LB-124, LB-125 and LB-169 name the local source for each entry that depends on one.",
        "Every local source a Telangana answer relies on is identified exactly, read with its locator, and named "
        "beside the central provision it modifies.",
        "A local source not in the manifest, or not indexed, returns 'held, not identifiable' or 'not held', "
        "naming the store -- never a silent fall-back to the central text as though no local rule existed.",
        "Never state a Telangana period, form or fee from the central text alone where a local rule is engaged "
        "and not read. Never identify a local rule by a shared word -- the Andhra Pradesh and Telangana versions "
        "share titles. Never report a source absent from one store as absent from the corpus.",
        "LB-170-AC1: a question on the pecuniary jurisdiction of a civil court in Telangana reads the Telangana "
        "Civil Courts Act, identified exactly.\n"
        "LB-170-AC2 (planted): the Court-fees and Suits Valuation Act named without its state, on a Telangana "
        "matter, is not silently resolved to the Andhra Pradesh version -- it is inferred as the Telangana Act "
        "with the basis stated, and the advocate can correct it (ActBasis.INFERRED).\n"
        "LB-170-AC3: a zero result for a local source names the store it came from and tells 'held, not "
        "identifiable' from 'not held'.\n"
        "LB-170-AC4: every local source in the manifest carries a check that its text is the adaptation it is "
        "named as.",
        "LB-121, LB-124, LB-125, LB-169; pipeline/manifest.yaml; backend/nm/knowledge/manifest.py (exact "
        "resolution); the L.4 header; docs/BASELINE.md, to be updated with these holdings. Raised in the 25 "
        "September review ('CPC state amendments and the Telangana High Court rules, 0 mentions'), never "
        "answered. OPEN: whether each source applies in Telangana as adapted; which further High Court rules are "
        "needed -- a file-name search of raw_data/ found no original-side or writ rules; and who maintains them.",
    ],
}
AFTER = "LB-125"
STATUS = {
    "LB-168": ("In progress", "backend/nm/core/evidence_item.py; tests/test_inventory_on_a_served_turn.py",
               "The evidence inventory keeps existence, admissibility and weight apart, admissibility not assessed "
               "by default; no curated admissibility conditions exist, so nothing assesses it."),
    "LB-169": ("Not started", "backend/nm/domain/drafting.py; backend/nm/core/drafting.py",
               "The drafting package exists; no pleading requirement from Order VI, VII or VIII is checked "
               "against a draft."),
    "LB-170": ("Not started", "pipeline/manifest.yaml",
               "The local sources are held (measured 26 September 2026) but none is in the manifest, and the Civil "
               "Rules of Practice are not in the derived store; nothing identifies or reads them exactly."),
}
PLAN_CELLS = {
    ("LB-121", 44): "Engaged conditions are named on the served turn; whether the file shows them done is not read; "
                    "s.12A is not curated -- the Commercial Courts Act is held (raw_data/, chunks.db and legal.db, "
                    "measured 26 September 2026) but absent from the manifest, so it cannot be identified exactly.",
    ("LB-125", 44): "No forum, valuation or fee is computed. The Telangana Court-fees and Suits Valuation Act and "
                    "the Civil Courts Act are held (raw_data/, chunks.db and legal.db, measured 26 September 2026) "
                    "but not in the manifest, and the fee schedules were never atomised.",
}


def main() -> int:
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    book = load_workbook(SOURCE)
    sheet, plan = book["Before Build"], book["Implementation Plan"]
    features = {s.title: regroup.sheet_features(s) for s in book}
    before = {(s.title, c.coordinate): (c.value, copy.copy(c._style)) for s in book for row in s for c in row}
    heights = {r: sheet.row_dimensions[r].height for r in range(1, sheet.max_row + 1)}

    order = []
    for r in range(FIRST, sheet.max_row + 1):
        text = str(sheet.cell(r, 1).value or "")
        h = HEADER.match(text)
        key = f"HEADER:{h[1]}" if h else regroup.key_of(text)
        order.append((key, {"values": [sheet.cell(r, c).value for c in range(1, 11)],
                            "styles": [copy.copy(sheet.cell(r, c)._style) for c in range(1, 11)],
                            "height": sheet.row_dimensions[r].height, "row": r}))
    source = dict(order)
    assert len(source) == len(order), "duplicate row key"
    for ident in NEW:
        assert ident not in source, ident
    for key, *_ in REPLACE:
        assert key in source, key
    for key in APPEND:
        assert key in source, key

    # ---- the expected values of every row that changes -----------------------
    expected = {}
    for key, row in order:
        values = list(row["values"])
        for ident, col, old, new in REPLACE:
            if key == ident:
                assert str(values[col - 1]).count(old) == 1, (ident, col, old[:60])
                values[col - 1] = values[col - 1].replace(old, new)
        for col, text in APPEND.get(key, []):
            values[col - 1] = f"{values[col - 1]}\n\n{text}" if values[col - 1] else text
        expected[key] = values
    for ident, cols in NEW.items():
        assert len(cols) == 10, ident  # title, summary, then columns 2 to 9
        expected[ident] = [f"{ident}\nIndian practice layer\n{cols[0]}\n\nDRAFT REQUIREMENT FOR OWNER REVIEW:\n"
                           f"{cols[1]}", *cols[2:], READINESS]
        assert len(expected[ident]) == 10

    # ---- the new rows sit after LB-125, with the zebra parity of their place --
    body, group_rows = [], []
    for key, row in order:
        body.append((key, row))
        if key == AFTER:
            for i, ident in enumerate(NEW):
                body.append((ident, {"values": None, "styles": None, "height": row["height"], "row": None}))
    assert len(body) == len(order) + len(NEW)

    group, layout = None, []
    for offset, (key, row) in enumerate(body):
        if key.startswith("HEADER:"):
            group = key[7:]
        layout.append((key, FIRST + offset, group))
    l8 = [k for k, _r, g in layout if g == "L.8" and not k.startswith("HEADER:")]
    for ident in NEW:
        index = l8.index(ident)
        same_parity = next(k for k in reversed(l8[:index]) if k not in NEW and (l8.index(k) - index) % 2 == 0)
        body[[k for k, _ in body].index(ident)][1]["styles"] = [copy.copy(s) for s in source[same_parity]["styles"]]

    for (key, r, _group), (_k, row) in zip(layout, body):
        for c in range(1, 11):
            cell = sheet.cell(r, c)
            cell.value = expected[key][c - 1]
            cell._style = copy.copy(row["styles"][c - 1])
        sheet.row_dimensions[r].height = row["height"]
    last = FIRST + len(body) - 1
    sheet.auto_filter.ref = f"A4:J{last}"

    population = sorted({regroup.LBOM.match(str(sheet.cell(x, 1).value or ""))[1]
                         for x in range(5, last + 1) if regroup.LBOM.match(str(sheet.cell(x, 1).value or ""))})
    lb = [p for p in population if p.startswith("LB-")]
    note = str(sheet.cell(2, 1).value)
    anchor = "and the per-dispute answer and are drafted for owner review."
    assert note.count(anchor) == 1
    note = note.replace(anchor, anchor + " LB-168–170 (L.8: admissibility of documents, pleadings, local procedure) "
                                         "and the 26 September 2026 review's amendments across the legal-brain "
                                         "rows are likewise drafted for owner review.")
    note, n = re.subn(r"The legal brain holds \d+ LB requirements", f"The legal brain holds {len(lb)} LB requirements",
                      note)
    assert n == 1
    note, n = regroup.NOTE_PATTERN.subn(f"All {len(population)} LB/OM requirements are linked into Implementation Plan",
                                        note)
    assert n == 1
    sheet.cell(2, 1).value = note

    # ---- the Implementation Plan ---------------------------------------------
    position = {k: r for k, r, _g in layout}
    touched = {k for k, *_ in REPLACE} | set(APPEND) | set(NEW)
    plan_changed = set()
    for x in range(2, plan.max_row + 1):
        ident = plan.cell(x, 1).value
        if ident in touched and regroup.LBOM.fullmatch(str(ident)):
            for t, origin in regroup.MIRROR.items():
                v = sheet.cell(position[ident], origin).value
                if plan.cell(x, t).value != v:
                    plan.cell(x, t).value = v
                    plan_changed.add(plan.cell(x, t).coordinate)
        for (pid, col), v in PLAN_CELLS.items():
            if ident == pid:
                plan.cell(x, col).value = v
                plan_changed.add(plan.cell(x, col).coordinate)
    for ident, cols in NEW.items():
        target = plan.max_row + 1
        for c in range(1, plan.max_column + 1):
            plan.cell(target, c)._style = copy.copy(plan.cell(target - 1, c)._style)
        build, evidence, gaps = STATUS[ident]
        fixed = {1: ident, 3: "Requirement", 4: "Legal brain", 5: "L.8 Indian practice layer", 6: cols[0],
                 7: "Pilot", 38: "Draft", 39: build, 40: "Not verified", 41: evidence,
                 43: "26 September 2026: code inspection at commit 71426a9; no test run for this row. Acceptance "
                     "criteria not run as written.",
                 44: gaps,
                 32: "26 September 2026: draft practice-layer row from the legal-brain review, mirrored from Before "
                     "Build. Not acceptance, not an approved requirement."}
        for c, v in fixed.items():
            plan.cell(target, c, v)
            plan_changed.add(plan.cell(target, c).coordinate)
        for t, origin in regroup.MIRROR.items():
            plan.cell(target, t, sheet.cell(position[ident], origin).value)
            plan_changed.add(plan.cell(target, t).coordinate)

    out_dir = Path(tempfile.gettempdir()) / "nm_legal_brain_review_20260926"
    out_dir.mkdir(exist_ok=True)
    out = out_dir / SOURCE.name
    book.save(out)

    # ================= THE PROOF, on the saved bytes ==========================
    from openpyxl.xml.functions import tostring

    def xml(part):
        return tostring(part.to_tree())

    check = load_workbook(out)
    for tab in check:
        assert regroup.sheet_features(tab) == features[tab.title], tab.title
    bb = check["Before Build"]
    original = load_workbook(SOURCE)["Before Build"]
    for (title, coord), old in before.items():
        if title == "Before Build" and int(re.sub(r"[A-Z]+", "", coord)) < FIRST and coord != "A2":
            assert (bb[coord].value, bb[coord]._style) == old, coord
    for x in range(1, FIRST):
        assert bb.row_dimensions[x].height == heights.get(x)
    changed = 0
    for key, r, group in layout:
        got = [bb.cell(r, c).value for c in range(1, 11)]
        assert got == expected[key], key
        if key in NEW:
            assert group == "L.8", key
            continue
        if key not in touched:
            assert got == source[key]["values"], f"{key}: changed but not meant to"
        else:
            changed += 1
            for c in range(1, 11):
                old = source[key]["values"][c - 1]
                if old and not any(k == key and col == c for k, col, *_ in REPLACE):
                    assert str(got[c - 1]).startswith(str(old)), f"{key} col {c}: earlier text not kept"
        for c in range(1, 11):
            a, o = bb.cell(r, c), original.cell(source[key]["row"], c)
            assert xml(a.border) == xml(o.border) and xml(a.alignment) == xml(o.alignment), key
            assert xml(a.fill) == xml(o.fill), key
            assert (a.font.name, a.font.sz, a.font.color, a.font.b) == (o.font.name, o.font.sz, o.font.color,
                                                                         o.font.b), key
    ip = check["Implementation Plan"]
    for (title, coord), old in before.items():
        if title == "Implementation Plan" and coord not in plan_changed:
            assert (ip[coord].value, ip[coord]._style) == old, f"plan {coord} moved"
    from assurance.control_plane import plan_scenarios
    problems = plan_scenarios.requirement_problems(out)
    assert problems == [], problems
    assert len(plan_scenarios.sheet_rows(out)) == 93

    shutil.copy2(out, SOURCE)
    print(f"source digest before: {digest}")
    print(f"source digest after:  {hashlib.sha256(SOURCE.read_bytes()).hexdigest()}")
    print(f"rows amended in place: {changed}; replacements: {len(REPLACE)}; additions: "
          f"{sum(len(v) for v in APPEND.values())}; new rows: {', '.join(f'{k} at row {position[k]}' for k in NEW)}; "
          f"plan cells written: {len(plan_changed)}; LB/OM {len(population)} (LB {len(lb)})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
