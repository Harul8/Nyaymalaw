"""Register today's legal-brain work in the backlog, reconciled to the owners it already has.

OWNER DIRECTION, 26 September 2026: register the day's legal-brain work so the
generated plan shows it -- and, choosing among three options, "Reconcile":
new items only for work nobody owns yet; each LB-138 slice an execution packet;
existing items keep the LB rows they already cover. No rule stated twice.

WHY RECONCILE. An item per slice would have restated about fifteen criteria
the backlog already owns -- the loop's choices (BK-91-AC1/AC5), the golden
comparison (BK-91-AC4), delegation (BK-92), claims tied to sources (BK-64-AC1),
opposing counsel and exposure (BK-95-AC5/AC6), the answer's shape (BK-37),
board questions (BK-54-AC2/AC3), document extraction (BK-54-AC1), progress
(BK-41-AC1), notice, forum and deadlines (BK-95-AC1/AC2), governing law
(BK-65-AC2) and interim protection (BK-70-AC1). Those stay where they are.

WHAT IS NEW, and so registered here:
  BK-98  the curated Indian practice tables -- exact, sourced, counsel-reviewed
  BK-99  the provider-neutral tool loop -- tool calling, envelope and registry,
         step log, replay, scratch-pad stream, and the switch
  BK-100 the loop's answer checked, repaired and independently verified on
         every channel
  BK-101 what the loop's model reads -- principles, stable prefix, compaction
         from the file, advocate memory and playbooks
  P48    the practice tables; P49 to P54 six of the nine LB-138 slices. A
         packet contributes only to criteria it finally owns: the blueprint
         requires contributions to precede the final packet, so slices 5, 6
         and 8 -- which deliver only behaviour already owned by P23 to P26 --
         are recorded as re-sequencings for the owner to decide at the
         switch, not as packets.

Each of today's LB rows (LB-120 to LB-170) names its delivery owner in its
Dependencies column, mirrored to the Implementation Plan.

NOT RE-RUNNABLE BY DESIGN: it refuses to start if BK-98 is already registered.
The workbook part keeps the other tools' discipline -- prove on the saved file
that nothing else moved, check the reconciler, then replace the source.
"""
import copy
import hashlib
import importlib.util
import json
import re
import shutil
import sys
import tempfile
from pathlib import Path

import yaml
from openpyxl import load_workbook

_REPO = Path(__file__).resolve().parents[2]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))
from assurance.common._console import utf8_console  # noqa: E402

utf8_console()

_spec = importlib.util.spec_from_file_location(
    "regroup", Path(__file__).with_name("legal_brain_regroup_20260926.py"))
regroup = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(regroup)

STATUS = _REPO / "docs/backlog/status.yaml"
PLAN = _REPO / "docs/backlog/plan.json"
MODULES = _REPO / "docs/blueprint/modules.json"
PACKETS = _REPO / "docs/blueprint/packets.json"
BACKLOG = _REPO / "docs/BACKLOG.md"
WORKBOOK = regroup.SOURCE
TODAY = "26 September 2026"
ALL_PHASES = ["B", "C", "D", "E", "F", "G", "H"]


def slug(heading: str) -> str:
    """The anchor convention BACKLOG.md records already use."""
    text = re.sub(r"[^a-z0-9 -]", "", heading.lower())
    return re.sub(r"-{2,}", "-", text.replace(" ", "-")).strip("-")


def criterion(ident, requirement, evidence, mutation, failure):
    return {"id": ident, "requirement": requirement, "required_evidence": evidence, "evidence": {},
            "negative_control": {"mutation": mutation, "expected_failure": failure}}


# =============================================================================
# THE FOUR ITEMS
# =============================================================================
ITEMS = [
    {
        "id": "BK-98", "title": "curated Indian practice tables: exact, sourced and counsel-reviewed",
        "kind": "substrate", "priority": "P1", "delivery_status": "in_progress", "implementation": "partial",
        "verification": "partial", "affects_phases": ["D", "E"], "started": True,
        "next_action": "put every curated entry through counsel review (none is reviewed yet); add manifest "
                       "entries and schedule atoms for the held Telangana court-fee and Civil Courts Acts and the "
                       "Commercial Courts Act with BK-84; the behaviours the tables serve stay owned by BK-95-AC1, "
                       "BK-95-AC2, BK-65-AC2 and BK-70-AC1",
        "acceptance": [
            criterion("BK-98-AC1",
                      "each curated practice table -- governing code by date, pre-institution conditions, binding "
                      "authority, interim-relief tests, procedural periods and filing requirements -- identifies its "
                      "entry on an exact key and never by fuzzy match, carries curated_from on every entry, and "
                      "returns applies, does not apply, not assessed or no curated table",
                      ["domain_test", "adversarial_test"],
                      "look an entry up by a neighbouring key, or add an entry without curated_from",
                      "the lookup returns no curated table and the type refuses the entry"),
            criterion("BK-98-AC2",
                      "a table entry never stands in for the provision it points to: the provision is retrieved and "
                      "read back before the answer relies on it, and no section is mapped across the old and new "
                      "codes by equal number",
                      ["domain_test", "adversarial_test"],
                      "rely on an entry without retrieving its provision, or map an IPC section to the BNS by number",
                      "the reliance is withheld by G-GROUND and correspondence comes only from the curated table"),
            criterion("BK-98-AC3",
                      "every curated entry has a recorded counsel review against the actual section or judgment "
                      "before any criterion relying on it is signed off",
                      ["counsel_review"],
                      "sign off a criterion that relies on an unreviewed table entry",
                      "sign-off is refused and the unreviewed entry is named"),
            criterion("BK-98-AC4",
                      "a table determination enters the premise record, so correcting the date, court, relief or "
                      "amount it rests on reopens it and everything resting on it through the dependency ledger",
                      ["domain_test", "integration_test"],
                      "correct a service date or a claimed amount across a boundary",
                      "the determination and its dependants reopen, with the prior value reported"),
        ],
    },
    {
        "id": "BK-99", "title": "the provider-neutral tool loop: tool calling, the tool envelope, the step log, "
                                "replay and the switch",
        "kind": "substrate", "priority": "P1", "delivery_status": "planned", "implementation": "none",
        "verification": "none", "affects_phases": ALL_PHASES, "started": False,
        "next_action": "owner to approve LB-127, LB-129, LB-138, LB-146, LB-154, LB-164 and LB-139 and to confirm "
                       "the LB-138 slice order; the loop's reasoning stays owned by BK-91-AC1, AC4 and AC5, and "
                       "delegation by BK-92",
        "acceptance": [
            criterion("BK-99-AC1",
                      "tool calling is one model-port method translated by each provider's adapter: one scripted "
                      "conversation produces identical tool invocations through the OpenAI, Anthropic and scripted "
                      "adapters, and a call to an undeclared tool or with malformed arguments is returned to the "
                      "model as an error and never executed",
                      ["domain_test", "integration_test", "adversarial_test"],
                      "execute an undeclared tool, act on malformed arguments, or let a provider tool format reach "
                      "the core",
                      "the call is refused and returned to the model, and the core check fails on a provider format"),
            criterion("BK-99-AC2",
                      "every tool is declared once in a registry and returns one envelope -- found, not held, not "
                      "found, not assessed or error, with its source, locator, index consulted and exclusions -- so a "
                      "zero names its index; a closed vocabulary is required only as a tool's input key, and a tool "
                      "absent from the registry cannot be called",
                      ["domain_test", "adversarial_test"],
                      "return an empty list with no status, or call a tool the registry does not declare",
                      "the result is refused before the model sees it and the repository check fails the build"),
            criterion("BK-99-AC3",
                      "every step of the loop is a typed event appended in order to the turn's saved log, which "
                      "survives a failed turn and is what the harness reads as process evidence",
                      ["domain_test", "integration_test"],
                      "fail a turn mid-loop, or append a step out of order",
                      "the log keeps every step up to the failure, in order"),
            criterion("BK-99-AC4",
                      "a recorded conversation replays through the loop, tools and harness with no network or API "
                      "key and reaches the same answer; a recording made under older principles or tools is marked "
                      "stale, never replayed as current",
                      ["domain_test", "integration_test"],
                      "change the harness so an unsupported claim passes, or replay a recording across a principles "
                      "change",
                      "the replay turns red, and the stale recording is refused"),
            criterion("BK-99-AC5",
                      "the step events stream to a scratch pad that is closed by default on desktop and phone, groups "
                      "steps by dispute in plain words and marks them as working; the streamed events equal the saved "
                      "ones in order, and a client that reconnects receives each later event exactly once",
                      ["integration_test", "browser_journey"],
                      "drop the stream mid-turn, or stream an event the log does not hold",
                      "the client resumes from its last sequence number and the stream and the log still agree"),
            criterion("BK-99-AC6",
                      "the loop is built beside the pipeline and switched one kind of turn at a time: during the "
                      "overlap both call one implementation of each gate, after a kind is switched no code path "
                      "reaches its retired stage, and the cost per turn is a release row measured before a live turn "
                      "is served",
                      ["domain_test", "integration_test", "production_measure"],
                      "call a copied gate on the loop path, reach a retired stage after the switch, or serve a live "
                      "turn with its cost not measured",
                      "the repository check fails the build and the release row reads NOT MEASURED"),
        ],
    },
    {
        "id": "BK-100", "title": "the loop's answer is checked, repaired and independently verified on every channel",
        "kind": "substrate", "priority": "P0", "delivery_status": "planned", "implementation": "none",
        "verification": "none", "affects_phases": ALL_PHASES, "depends_on": ["BK-99"], "started": False,
        "next_action": "owner to approve LB-133, LB-141, LB-144 and the LB-67 amendment; the eighteen existing "
                       "output checks and the professional boundaries keep their owners and run on the loop path "
                       "through BK-99-AC6",
        "acceptance": [
            criterion("BK-100-AC1",
                      "a failed check goes back to the model with its reason and is revised within a bound; what "
                      "still fails is withheld with the conclusions resting on it, every attempt is in the step log, "
                      "and the served bytes are the answer the harness checked",
                      ["domain_test", "adversarial_test"],
                      "submit an answer with one fabricated quotation, or edit the answer after its check",
                      "it is repaired or withheld with its reason, and the edited answer is refused"),
            criterion("BK-100-AC2",
                      "an independent verifier -- never the author and never shown the conversation -- gives its "
                      "reason first and judges each claim against its cited passage, filling LB-103's assessed "
                      "states; its tier and agreement are set by a labelled measurement counting both error "
                      "directions",
                      ["domain_test", "adversarial_test", "model_eval"],
                      "cite a real passage for a proposition it does not state, or let the author grade its own claim",
                      "the claim is marked does not support and repaired, and the self-grade is refused"),
            criterion("BK-100-AC3",
                      "every number and date in an answer maps to a computation tool's result or to the advocate's "
                      "or a document's own words; an unmatched figure goes to repair, and a figure resting on an "
                      "inferred premise is labelled conditional",
                      ["domain_test", "adversarial_test"],
                      "state a limitation expiry the tool did not compute",
                      "the figure is refused before release"),
            criterion("BK-100-AC4",
                      "every model-written line on every channel the advocate sees -- the answer, the step events "
                      "and scratch pad, board items, questions and the compaction handover -- passes G-GROUND's own "
                      "citation detector before it leaves, proved by one test that plants an unretrieved authority "
                      "on each channel",
                      ["domain_test", "integration_test", "adversarial_test"],
                      "plant a case that was never retrieved in a step event, a board item or a question",
                      "it reaches the advocate on none of them and the omission is shown"),
        ],
    },
    {
        "id": "BK-101", "title": "what the loop's model reads: the principles, a stable prefix, compaction from the "
                                 "file and the advocate's memory",
        "kind": "substrate", "priority": "P1", "delivery_status": "planned", "implementation": "none",
        "verification": "none", "affects_phases": ALL_PHASES, "depends_on": ["BK-99"], "started": False,
        "next_action": "owner to approve LB-126 and LB-147 to LB-151 and to choose the first playbooks; matter-state "
                       "freshness and durability stay owned by BK-36-AC2 and BK-65-AC1",
        "acceptance": [
            criterion("BK-101-AC1",
                      "one owner-edited principles document is read at the start of every turn and its version is "
                      "recorded on the turn; a missing document stops the loop with a visible reason, and a change "
                      "takes effect at the next turn boundary through compaction",
                      ["domain_test", "integration_test"],
                      "copy a principle into the turn path as a hard-coded rule, or run a turn without the document",
                      "the repository check refuses the copy and the turn is refused with its reason"),
            criterion("BK-101-AC2",
                      "within a conversation the principles and tools form a byte-stable prefix and the history is "
                      "append-only on every provider; a corrected fact, a change notice or the remaining budget is "
                      "appended, never edited into an earlier message",
                      ["domain_test", "integration_test"],
                      "inject a timestamp into the prefix or edit an earlier message",
                      "the prefix comparison reports the defect"),
            criterion("BK-101-AC3",
                      "spent results are cleared to a stub with their locator and re-fetched when needed; past the "
                      "context budget the conversation restarts from a brief regenerated from the checked file with "
                      "every line tagged by source and status, a handover that alters a fact is rejected, and an "
                      "earlier turn stays readable word for word",
                      ["domain_test", "integration_test", "adversarial_test"],
                      "alter a date in the compaction handover, or clear a span a pending claim cites",
                      "the handover is rejected against the file and the clearing is refused"),
            criterion("BK-101-AC4",
                      "the advocate's memory holds working preferences, written only with approval and never a "
                      "client's facts, and practice-area playbooks load on demand with every legal pointer an exact "
                      "key that resolves through the exact tools",
                      ["domain_test", "adversarial_test"],
                      "write a client's name to the advocate memory, or give a playbook a period with no source or "
                      "a case named by its parties alone",
                      "the write is refused and the repository check fails on the playbook"),
        ],
    },
]
WAVES = {"BK-98": "W3", "BK-99": "W3", "BK-100": "W3", "BK-101": "W3"}
MODULE_OF = {"BK-98": "M07", "BK-99": "M07", "BK-100": "M07", "BK-101": "M04"}


# =============================================================================
# THE PACKETS -- P48 the tables; P49 to P54 six of the nine LB-138 slices
# =============================================================================
def packet(pid, module, title, criteria, final, prerequisites, decisions, commands, inputs, outputs, steps,
           proof, expected, rollback, existing, planned=()):
    return {"id": pid, "module": module, "title": title, "kind": "implementation", "criteria": criteria,
            "prerequisites": prerequisites, "decisions": decisions, "commands": commands,
            "boundaries": [{"path": p, "kind": "existing"} for p in existing]
                          + [{"path": p, "kind": "planned"} for p in planned],
            "inputs": inputs, "outputs": outputs, "steps": steps,
            "proof": {"positive": [proof[0]], "negative": [proof[1]], "live": [proof[2]]},
            "expected": expected, "rollback": rollback, "final_criteria": final, "requires_completed_items": []}


PACKET_ROWS = [
    packet("P48", "M07", "Curate the Indian practice tables and put each through counsel review",
           ["BK-98-AC1", "BK-98-AC2", "BK-98-AC3", "BK-98-AC4"],
           ["BK-98-AC1", "BK-98-AC2", "BK-98-AC3", "BK-98-AC4"],
           ["P21", "P22"], ["CHOICE-01", "CHOICE-07"], ["submit-turn", "get-advice"],
           ["The retrieved and read-back text of each provision or judgment an entry points to.",
            "The owner's selection of rows, and a practising Telangana advocate's counsel review of every entry."],
           ["Curated tables in backend/nm/knowledge with curated_from on every entry, each behind a port and "
            "adapter.", "A counsel-review record per table before any dependent criterion is signed off."],
           ["Keep one curated table per rule family, identified on an exact key and never by fuzzy match.",
            "Retrieve and read back the provision or judgment behind every entry; the table points to text and "
            "never replaces it.",
            "Write each determination into the premise record so a correction reopens it through the dependency "
            "ledger.",
            "With BK-84, add manifest entries and schedule atoms for the held Telangana court-fee and Civil Courts "
            "Acts and the Commercial Courts Act before any fee or s.12A condition is computed."],
           ("EVAL-015", "EVAL-011", "EVAL-015"),
           ["A table entry never stands in for the provision it points to.",
            "An uncurated key returns 'no curated table', never a neighbouring entry.",
            "No criterion relying on a table is signed off before the table's counsel review."],
           "Withdraw an unreviewed or disputed entry to 'no curated table' and keep the retrieval path; never serve "
           "an entry whose source cannot be read back.",
           ["backend/nm/knowledge/governing_law.py", "backend/nm/knowledge/institution.py",
            "backend/nm/knowledge/interim_relief.py", "backend/nm/knowledge/procedural_period.py",
            "backend/nm/knowledge/filing_requirement.py", "backend/nm/knowledge/identity.py"]),
    packet("P49", "M07", "Slice 1 -- the loop foundation: tool calling on every provider, the budgeted runner, the "
                         "step log and replay",
           ["BK-99-AC1", "BK-99-AC3", "BK-99-AC4"],
           ["BK-99-AC1", "BK-99-AC3", "BK-99-AC4"],
           [], ["CHOICE-02", "CHOICE-05", "CHOICE-06"], ["submit-turn"],
           ["The model port (complete, structured, embed) and its OpenAI and scripted adapters.",
            "The owner's approval of LB-127, LB-128, LB-146, LB-147 and LB-164 and of the LB-138 slice order."],
           ["A tool-calling method on the model port with OpenAI, Anthropic and scripted adapters.",
            "A loop runner with declared step, token, cost and time budgets, beside the TurnEngine behind a switch.",
            "A typed, ordered, saved step log per turn, and a replay adapter from recordings."],
           ["Add one tool-calling method to the model port; provider formats stay inside the adapters.",
            "Run the loop -- decide, call, read, decide again -- to an answer, a question, a no-progress stop or a "
            "budget stop, each disclosed; build it beside the TurnEngine behind a switch.",
            "Append every step as a typed event to the turn's saved log.",
            "Record model responses with the principles and tool versions in force, and replay them with no "
            "network on every commit."],
           ("EVAL-031", "EVAL-031", "EVAL-031"),
           ["No provider format reaches the core.",
            "No loop runs without a budget, and a budget stop is never presented as a finished answer.",
            "A replay proves the harness and the tools, never live model quality."],
           "Keep the switch off: the TurnEngine serves every turn, and recordings and step logs are kept as evidence.",
           ["backend/nm/ports/model.py", "backend/nm/adapters/model", "backend/nm/core/turn.py"],
           ["backend/nm/core/loop.py", "backend/nm/adapters/model/anthropic_adapter.py",
            "backend/nm/adapters/model/replay.py"]),
    packet("P50", "M06", "Slice 2 -- tools: the envelope, the registry and the core tools over existing code, with "
                         "born-digital document text",
           ["BK-99-AC2"], ["BK-99-AC2"],
           ["P49"], ["CHOICE-02", "CHOICE-07"], ["submit-turn", "get-source", "get-casefile"],
           ["The search, evidence and matter ports, the curated tables, and the limitation and deadline arithmetic.",
            "Born-digital PDF uploads already admitted and stored."],
           ["A registry declaring each tool once with its description, schema, parallel safety, kind and tests.",
            "Core tools wrapping existing code, each returning the envelope.",
            "Page-located text extracted from born-digital PDFs for read_document and search_matter."],
           ["Declare the envelope and the registry, with a planted-failure test for every tool.",
            "Wrap existing code in the core tools: matter reads, identify_act and read_provision, authority search "
            "and resolve_citation, compute_limitation, ask_advocate and submit_answer.",
            "Extract text with page locators from born-digital PDFs; a page with no text layer is reported as not "
            "extracted, never as empty."],
           ("EVAL-013", "EVAL-014", "EVAL-013"),
           ["A zero result names its index.", "Tools that identify never rank, and tools that rank never identify.",
            "Nothing unextracted is read as searched and empty."],
           "Remove a tool from the registry and return 'not available' to the model; extracted text is a derivative "
           "and the original upload is untouched.",
           ["backend/nm/ports", "backend/nm/knowledge/manifest.py", "backend/nm/edge/uploads.py"],
           ["backend/nm/core/tools.py"]),
    packet("P51", "M07", "Slice 3 -- the harness after the loop: checks, boundaries on tool calls, repair and the "
                         "independent verifier",
           ["BK-100-AC1", "BK-100-AC2", "BK-100-AC3", "BK-100-AC4"],
           ["BK-100-AC1", "BK-100-AC2", "BK-100-AC3", "BK-100-AC4"],
           ["P50"], ["CHOICE-05", "CHOICE-07"], ["submit-turn", "get-advice"],
           ["The gate matrix and its existing checks.",
            "A labelled set of claim-passage pairs for the verifier, chosen by the owner."],
           ["The existing output checks and professional boundaries run on the loop path, with checks before and "
            "after every tool call.", "A bounded repair step and an independent verifier.",
            "One citation detector applied on every advocate-visible channel."],
           ["Run the eighteen output checks after the loop and the six boundaries around every tool call, calling "
            "the gate functions the pipeline calls.",
            "Return a failed check to the model with its reason; withhold what still fails after the bound.",
            "Add the independent verifier, reason first, its tier set by the labelled measurement.",
            "Pass every model-written line on every channel through G-GROUND's detector."],
           ("EVAL-016", "EVAL-013", "EVAL-016"),
           ["The model never skips, disables or waives a check.", "A check that cannot run is not assessed, never "
            "passed.", "The author never grades its own claim."],
           "Keep the switch off for any kind of turn whose checks do not all run on the loop path.",
           ["backend/nm/domain/gates.py", "backend/nm/core/grounding.py", "backend/nm/core/consistency.py"],
           ["backend/nm/core/verifier.py"]),
    packet("P52", "M09", "Slice 4 -- stream the loop's steps to the scratch pad",
           ["BK-99-AC5"], ["BK-99-AC5"],
           ["P51"], ["CHOICE-02", "CHOICE-06"], [],
           ["The saved step log (P49) and the citation detector on every channel (P51).",
            "The matter board in frontend/matter-workspace.js."],
           ["A server-sent-events endpoint beside /api/turn, and a collapsible scratch-pad panel, closed by "
            "default, saved with each turn."],
           ["Stream each step event with its sequence number; resume a reconnecting client from its last number.",
            "Render the panel grouped by dispute, in plain words, marked as working, with NM's judgment labelled.",
            "Save each turn's scratch pad with the matter and show it again on reopening."],
           ("EVAL-031", "EVAL-034", "EVAL-031"),
           ["The stream never shows an event the log does not hold.",
            "No interim line reads as advice or shows an internal identifier."],
           "Hide the panel and show the saved record when the turn ends; the step log is unaffected.",
           ["backend/nm/edge/api.py", "frontend/matter-workspace.js"]),
    packet("P53", "M04", "Slice 7 -- context for long matters: principles, stable prefix, clearing, compaction from "
                         "the file, the research loop and memory",
           ["BK-101-AC1", "BK-101-AC2", "BK-101-AC3", "BK-101-AC4"],
           ["BK-101-AC1", "BK-101-AC2", "BK-101-AC3", "BK-101-AC4"],
           ["P49"], ["CHOICE-02", "CHOICE-08"], ["submit-turn", "get-casefile"],
           ["The matter file and its summary.", "Owner-edited principles and the first playbooks, once approved."],
           ["A principles document read every turn; a stable prefix with append-only history; clearing by locator; "
            "compaction to a brief regenerated from the file; a fresh-context research loop; the advocate memory."],
           ["Read the versioned principles at the start of every turn and freeze the prefix for the conversation.",
            "Clear spent results to locator stubs, and compact to a brief regenerated from the checked file.",
            "Run long reading as a fresh-context research loop that returns only findings with spans.",
            "Keep the advocate's memory apart from every matter, written only with approval."],
           ("EVAL-010", "EVAL-006", "EVAL-010"),
           ["A summary of the chat never replaces the file as the thing worked from.",
            "No client fact enters the advocate's memory."],
           "Build each request from the file as today; keep the full transcript and step log on the record.",
           ["backend/nm/domain/summary.py", "backend/nm/domain/matter.py"], ["backend/nm/core/context.py"]),
    packet("P54", "M09", "Slice 9 -- compare on the golden set and switch, one kind of turn at a time",
           ["BK-99-AC6"], ["BK-99-AC6"],
           ["P51", "P52", "P53"], ["CHOICE-06", "CHOICE-10"], [],
           ["docs/GOLDEN_SET.md re-tagged with these slices, and the owner's approval of each live run.",
            "The ledger for every live comparison."],
           ["A recorded comparison per kind of turn with its date, versions, grounding, correctness and cost.",
            "The switch per kind of turn, and the retired pipeline stages removed."],
           ["Run each kind of turn's golden subset through the pipeline and the loop, with the owner's approval of "
            "each run.",
            "Switch a kind of turn only when the loop matches or beats the pipeline, with the cost per turn "
            "measured as a release row.",
            "Remove the pipeline stages a switched kind no longer uses, so no second live path survives."],
           ("EVAL-034", "EVAL-034", "EVAL-028"),
           ["No live golden run happens without the owner's approval of that run.",
            "No second live path survives a switch."],
           "Turn the switch back for that kind of turn; pipeline stages are removed only after the switch has held.",
           ["assurance/journeys/run_goldens.py", "docs/GOLDEN_SET.md", "backend/nm/core/turn.py"]),
]


# =============================================================================
# EACH OF TODAY'S LB ROWS NAMES ITS DELIVERY OWNER (Before Build column 9)
# =============================================================================
OWNER_PREFIX = f"Delivery owner (registered {TODAY}, LB-43): "
LB_OWNERS = {
    "LB-120": "BK-98 (the curated table), packet P48; the behaviour it serves stays BK-65-AC2 (final packet P22).",
    "LB-121": "BK-98 (the curated table), packet P48; the behaviour it serves stays BK-95-AC1 (P22).",
    "LB-122": "BK-98 (the bench and court ranking), packet P48; reviewed under BK-67-AC3.",
    "LB-123": "BK-98 (the curated table), packet P48; the behaviour it serves stays BK-70-AC1 (P23).",
    "LB-124": "BK-98 (the curated table), packet P48; the behaviour it serves stays BK-95-AC2 (P22).",
    "LB-125": "BK-98 (the curated table), packet P48; the behaviour it serves stays BK-95-AC1 (P22); the corpus gap "
              "is closed with BK-84.",
    "LB-126": "BK-101-AC1, packet P53 (slice 7).",
    "LB-127": "BK-99-AC1, packet P49 (slice 1).",
    "LB-128": "BK-91-AC1 and AC5, already owned (final packet P46); the shared budget across nested loops is BK-92-AC1 "
              "(P47). The loop is built beside them in slice 1 (P49).",
    "LB-129": "BK-99-AC2, packet P50 (slice 2).",
    "LB-130": "BK-64-AC1 (P17) and BK-91-AC2 (P46), already owned.",
    "LB-131": "The checks keep their current owners and run on the loop path through BK-99-AC6 (P54).",
    "LB-132": "BK-63-AC1 (P32) and BK-92-AC1 (P47), already owned.",
    "LB-133": "BK-100-AC1, packet P51 (slice 3).",
    "LB-134": "No delivery owner: it is an owner decision, recorded under LB-44.",
    "LB-135": "BK-101-AC3, packet P53 (slice 7).",
    "LB-136": "BK-92-AC1 and AC2 (P47), with BK-95-AC3 (P21), already owned.",
    "LB-137": "BK-38 and BK-97 (P21), already owned.",
    "LB-138": "BK-99-AC6, packet P54 (slice 9); the comparison itself is BK-91-AC4 (P35), already owned.",
    "LB-139": "BK-99-AC5, packet P52 (slice 4); what the advocate can see stays BK-41-AC1 (P37) and BK-66-AC1 (P34).",
    "LB-140": "BK-95-AC5 and AC6 (P23), enabled only on BK-92-AC4's evidence (P37), already owned. Slice 6 is a "
              "re-sequencing of P23 onto the loop, the owner's decision at the switch.",
    "LB-141": "BK-100-AC2, packet P51 (slice 3).",
    "LB-142": "BK-36-AC2 (P10), BK-65-AC1 (P18) and BK-92-AC2 (P47), already owned.",
    "LB-143": "BK-63-AC1 (P32) and BK-92-AC1 (P47), already owned.",
    "LB-144": "BK-100-AC3, packet P51 (slice 3).",
    "LB-145": "BK-101-AC3, packet P53 (slice 7).",
    "LB-146": "BK-99-AC4, packet P49 (slice 1).",
    "LB-147": "BK-101-AC2, packet P53 (slice 7).",
    "LB-148": "BK-101-AC4, packet P53 (slice 7).",
    "LB-149": "BK-101-AC3, packet P53 (slice 7).",
    "LB-150": "BK-101-AC3, packet P53 (slice 7).",
    "LB-151": "BK-101-AC4, packet P53 (slice 7).",
    "LB-152": "BK-36-AC2 (P10) and BK-65-AC1 (P18), already owned.",
    "LB-153": "BK-101-AC2, packet P53 (slice 7), with BK-29-AC1 for budgets, already owned.",
    "LB-154": "BK-99-AC2, packet P50 (slice 2).",
    "LB-155": "BK-99-AC2 for the tools, packet P50; reading uploads stays BK-54-AC1 (P25).",
    "LB-156": "BK-99-AC2 for the tools, packet P50; exact Act identity stays BK-97 (P21).",
    "LB-157": "BK-99-AC2 for the tools, packet P50; authority research stays BK-38 and BK-95-AC3 (P21).",
    "LB-158": "BK-99-AC2 for the tools, packet P50; the computations stay BK-35 and BK-95-AC2 (P22).",
    "LB-159": "BK-99-AC2 for the tools, packet P50, over BK-98's tables (P48).",
    "LB-160": "BK-99-AC2 for the tools, packet P50, over BK-100 (P51) and BK-92 (P47).",
    "LB-161": "BK-99-AC2 for the tools, packet P50; authority to act stays BK-63-AC1 (P32) and asking stays BK-54-AC3 "
              "(P24).",
    "LB-162": "BK-99-AC2, packet P50 (slice 2).",
    "LB-163": "BK-101-AC1 (the principles carry the method), packet P53, and BK-91-AC1 (P46), already owned.",
    "LB-164": "BK-99-AC3 (the saved log, packet P49) and BK-99-AC5 (the stream, packet P52).",
    "LB-165": "BK-37-AC1 and AC2 (P26), with BK-95-AC9 (P23), already owned. Slice 5 is a re-sequencing of P26 and P23 "
              "onto the loop, the owner's decision at the switch.",
    "LB-166": "BK-54-AC2 and AC3 (P24), already owned. Slice 5 re-sequences P24 onto the loop at the switch.",
    "LB-167": "BK-74 (the Start, Build, Test and Sign-off records) and BK-98-AC3 (counsel review of the tables, P48).",
    "LB-168": "No delivery owner yet: a candidate for BK-98, awaiting the owner's selection.",
    "LB-169": "No delivery owner yet: a candidate for BK-98, awaiting the owner's selection; drafting itself stays "
              "BK-56 (F-F-02, F-F-03).",
    "LB-170": "No delivery owner yet: a candidate awaiting the owner's selection; corpus identity stays BK-84.",
}


# =============================================================================
# THE BACKLOG RECORD -- why each item exists, and what it does not own
# =============================================================================
def heading(item):
    return f"BK-{item['id'][3:]} — {item['title']}"


RECORDS = {
    "BK-98": """Registered 26 September 2026 under LB-43, reconciling the legal-brain rows to the backlog on
the owner's instruction. The practice-layer rows LB-120 to LB-125 were built on 25 September 2026 as curated
tables in `backend/nm/knowledge/` (governing law, institution, interim relief, procedural periods, filing
requirements) with the bench ranking through `identity.supersedes`, each behind a port and adapter and wired in
the composition root; LB-120 is built but unwired.

**What this item owns is the tables themselves** -- exact identification, `curated_from` on every entry, the
provision read back rather than recited, counsel review, and the premise record. **It does not own the
behaviours the tables serve**, which stay where the backlog already has them: notice, forum and maintainability
thresholds (BK-95-AC1), the deadline register (BK-95-AC2), governing legal premises (BK-65-AC2) and interim
protection (BK-70-AC1). Binding authority is validated by qualified review under BK-67-AC3.

Tests exist for each table (`tests/test_which_code_governs_this_matter.py`,
`tests/test_what_must_happen_before_a_filing.py`, `tests/test_which_authority_this_court_must_follow.py`,
`tests/test_an_interim_application_is_decided_on_its_own_test.py`,
`tests/test_the_clocks_that_run_inside_a_proceeding.py`, `tests/test_whether_a_filing_will_be_accepted.py`);
they passed in the build container's full suite on 25 September and are not bound to a published Class-A
result, so every criterion's evidence is empty. **No curated entry has been counsel-reviewed.**

Candidates awaiting the owner's selection, not registered as criteria: LB-168 (admissibility of documents
relied on), LB-169 (pleading requirements, checked inside BK-56's F-F-03) and LB-170 (Telangana local procedure;
corpus identity with BK-84). Measured 26 September 2026: the Commercial Courts Act, the Telangana Court-fees and
Suits Valuation Act and the Telangana Civil Courts Act are held in every store and absent only from
`pipeline/manifest.yaml`.""",
    "BK-99": """Registered 26 September 2026 under LB-43 from the owner-directed loop rows LB-126 to LB-167,
which are drafts awaiting the owner's approval; the item is `planned` and its Start Record is BLOCKED on that
approval and on the LB-138 slice order.

**What is new here**: tool calling on the model port across providers (LB-127), the tool envelope and registry
(LB-129, LB-154, LB-162), the typed and saved step log (LB-164), record-and-replay (LB-146), the scratch-pad
stream (LB-139) and building beside the pipeline with a switch per kind of turn (LB-138).

**What it does not own**: the loop's choices of next action, completion and escalation are BK-91-AC1 and AC5;
the golden comparison is BK-91-AC4; delegated work, its shared budget and its enabling evidence are BK-92. Six of
the nine LB-138 slices are execution packets (P49 to P54); slices 5, 6 and 8 deliver only behaviour the backlog
already owns and are re-sequencings of its existing packets, decided at the switch.""",
    "BK-100": """Registered 26 September 2026 under LB-43 from LB-133 (repair), LB-141 (the independent
verifier), LB-144 (every figure from a tool) and the review's LB-67 amendment (one citation detector on every
channel, the scratch pad included). Planned; depends on BK-99.

**What it does not own**: the eighteen output checks and six professional boundaries already exist and keep
their current owners; they run on the loop path through BK-99-AC6's one-implementation rule. Claims tied to
sources are BK-64-AC1 and BK-91-AC2; authority to act is BK-63-AC1; the single acceptance path for findings is
BK-92-AC2. The verifier's verdict fills the assessed states LB-103 already carries.""",
    "BK-101": """Registered 26 September 2026 under LB-43 from LB-126 (the principles document), LB-147 (stable
prefix and append-only history), LB-135, LB-145, LB-149 and LB-150 (clearing, compaction from the file, tagged
context), LB-148 (playbooks) and LB-151 (the advocate's memory). Planned; depends on BK-99.

**What it does not own**: freshness and durable, version-checked writes are BK-36-AC2 and BK-65-AC1; the
research sub-loop is delegated work under BK-92-AC1 and AC2 with BK-95-AC3; model output budgets are BK-29-AC1.""",
}

RECONCILIATION = """## 26 September 2026 — the legal-brain rows reconciled to the backlog (LB-43)

The owner chose to reconcile rather than register an item per slice: an item per slice would have restated
about fifteen criteria the backlog already owns. New items exist only for work nobody owned (BK-98 to BK-101),
and every one of today's LB rows names its delivery owner in its Dependencies column in
`docs/Nyaymalaw_Implementation_Plan.xlsx`.

**A packet contributes only to criteria it finally owns.** The blueprint requires every contribution to come
before its criterion's final packet, so the loop's packets cannot contribute to criteria that P10 to P47 already
finalise without making those earlier packets wait for the loop. Whether the loop re-delivers them is a
re-sequencing the owner decides at the switch (LB-138); until then they are named, not rewired:

| Slice | Packet | Final owner of | Related criteria already owned (their final packet) |
|---|---|---|---|
| — the practice tables | P48 | BK-98-AC1 to AC4 | BK-95-AC1, AC2 (P22), BK-65-AC2 (P22), BK-70-AC1 (P23) |
| 1 loop foundation | P49 | BK-99-AC1, AC3, AC4 | BK-91-AC1, AC5 (P46), BK-92-AC1 (P47) |
| 2 tools and born-digital text | P50 | BK-99-AC2 | BK-54-AC1 (P25), BK-97-AC1 (P21) |
| 3 the harness after the loop | P51 | BK-100-AC1 to AC4 | BK-64-AC1 (P17), BK-63-AC1 (P32), BK-92-AC2 (P47), BK-91-AC2 (P46) |
| 4 the scratch pad | P52 | BK-99-AC5 | BK-41-AC1 (P37), BK-66-AC1 (P34) |
| 5 per-dispute answer and board | none | — | BK-37-AC1, AC2 (P26), BK-95-AC9 (P23), BK-54-AC2, AC3 (P24) |
| 6 opposing counsel | none | — | BK-95-AC5, AC6 (P23), BK-92-AC4 (P37) |
| 7 context | P53 | BK-101-AC1 to AC4 | BK-36-AC2 (P10), BK-65-AC1 (P18), BK-92-AC1 (P47), BK-95-AC3 (P21) |
| 8 scans, photographs, audio | none | — | BK-54-AC1, BK-79-AC1 (P25) |
| 9 comparison and switch | P54 | BK-99-AC6 | BK-91-AC4, BK-67-AC1 (P35) |

The slice order is the review's revised proposal (born-digital PDF text in slice 2) and remains the owner's to
confirm. Each packet's proof cites existing evaluation cases; dedicated cases for the loop are not yet
specified."""


# =============================================================================
def yaml_item(item: dict, nl: str) -> str:
    q = lambda s: json.dumps(s, ensure_ascii=False)  # noqa: E731 -- a JSON string is a YAML double-quoted scalar
    anchor = "docs/BACKLOG.md#" + slug(heading(item))
    started = item["started"]
    lines = [f"- id: {item['id']}", f"  title: {q(item['title'])}", f"  kind: {item['kind']}",
             f"  priority: {item['priority']}", f"  delivery_status: {item['delivery_status']}",
             f"  implementation: {item['implementation']}", f"  verification: {item['verification']}",
             f"  affects_phases: [{', '.join(item['affects_phases'])}]"]
    if item.get("depends_on"):
        lines.append(f"  depends_on: [{', '.join(item['depends_on'])}]")
    lines += [f"  record: {anchor}", f"  next_action: {q(item['next_action'])}", "  stage_records:",
              f"    start: {{result: {'READY' if started else 'BLOCKED'}, ref: {anchor}}}",
              f"    build: {{result: {'OPEN' if started else 'NOT_STARTED'}, ref: {anchor}}}",
              f"    test: {{result: {'OPEN' if started else 'NOT_RUN'}, ref: {anchor}}}",
              f"    signoff: {{result: NOT_RUN, ref: {anchor}}}", "  acceptance:"]
    for ac in item["acceptance"]:
        lines += [f"  - id: {ac['id']}", f"    requirement: {q(ac['requirement'])}",
                  f"    required_evidence: [{', '.join(ac['required_evidence'])}]", "    evidence: {}",
                  "    negative_control:", f"      mutation: {q(ac['negative_control']['mutation'])}",
                  f"      expected_failure: {q(ac['negative_control']['expected_failure'])}"]
    return nl.join(lines) + nl


def expected_item(item: dict) -> dict:
    """What the registry must read back for this item."""
    anchor = "docs/BACKLOG.md#" + slug(heading(item))
    started = item["started"]
    out = {k: item[k] for k in ("id", "title", "kind", "priority", "delivery_status", "implementation",
                                "verification", "affects_phases")}
    if item.get("depends_on"):
        out["depends_on"] = item["depends_on"]
    out.update({"record": anchor, "next_action": item["next_action"], "stage_records": {
        "start": {"result": "READY" if started else "BLOCKED", "ref": anchor},
        "build": {"result": "OPEN" if started else "NOT_STARTED", "ref": anchor},
        "test": {"result": "OPEN" if started else "NOT_RUN", "ref": anchor},
        "signoff": {"result": "NOT_RUN", "ref": anchor}}, "acceptance": item["acceptance"]})
    return out


def json_write(path: Path, data) -> None:
    raw = path.read_text(encoding="utf-8")
    crlf = "\r\n" in path.read_bytes().decode("utf-8")
    out = json.dumps(data, indent=2, ensure_ascii=False) + "\n"
    path.write_bytes((out.replace("\n", "\r\n") if crlf else out).encode("utf-8"))
    assert raw  # the file existed and was read


def json_roundtrips(path: Path) -> bool:
    raw = path.read_bytes().decode("utf-8")
    out = json.dumps(json.loads(raw), indent=2, ensure_ascii=False) + "\n"
    return (out.replace("\n", "\r\n") if "\r\n" in raw else out) == raw


def main() -> int:
    status_raw = STATUS.read_bytes().decode("utf-8")
    registry = yaml.safe_load(status_raw)
    have = {i["id"] for i in registry["items"]}
    if "BK-98" in have:
        print("BK-98 is already registered; this tool registers once and is not re-run.")
        return 1
    for item in ITEMS:
        assert item["id"] not in have, item["id"]
        for ac in item["acceptance"]:
            assert ac["id"].startswith(item["id"] + "-AC")
    catalogue = json.loads(PACKETS.read_text(encoding="utf-8"))
    pids = {p["id"] for p in catalogue["packets"]}
    for row in PACKET_ROWS:
        assert row["id"] not in pids, row["id"]
        for b in row["boundaries"]:
            exists = (_REPO / b["path"]).exists()
            assert exists == (b["kind"] == "existing"), (row["id"], b)
    # every criterion a packet names exists (existing ones in the registry, new ones here)
    known = {ac["id"] for i in registry["items"] for ac in (i.get("acceptance") or [])} | {
        ac["id"] for i in ITEMS for ac in i["acceptance"]}
    for row in PACKET_ROWS:
        missing = set(row["criteria"]) - known
        assert not missing, (row["id"], missing)
    finals = [c for row in PACKET_ROWS for c in row["final_criteria"]]
    new_criteria = {ac["id"] for i in ITEMS for ac in i["acceptance"]}
    assert sorted(finals) == sorted(new_criteria), "each new criterion must have exactly one final owner here"
    assert json_roundtrips(MODULES) and json_roundtrips(PACKETS)

    # ---- status.yaml: the four items, before the events log ----------------------
    nl = "\r\n" if "\r\n" in status_raw else "\n"
    anchor = f"{nl}events:{nl}"
    assert status_raw.count(anchor) == 1
    block = "".join(yaml_item(i, nl) for i in ITEMS)
    status_new = status_raw.replace(anchor, nl + block.rstrip(nl) + anchor, 1)
    reread = yaml.safe_load(status_new)
    by_id = {i["id"]: i for i in reread["items"]}
    for item in ITEMS:
        assert by_id[item["id"]] == expected_item(item), f"{item['id']} does not read back as written"
    assert [i["id"] for i in reread["items"]][:len(registry["items"])] == [i["id"] for i in registry["items"]]
    assert reread["events"] == registry["events"]

    # EVERYTHING IS COMPUTED AND CHECKED BEFORE ANYTHING IS WRITTEN. A first run
    # wrote status.yaml and then failed on plan.json, leaving the registry half
    # changed; the writes now happen together at the end, after every check.

    # ---- plan.json: one wave row each, in the file's own formatting --------------
    plan_raw = PLAN.read_bytes().decode("utf-8")
    pnl = "\r\n" if "\r\n" in plan_raw else "\n"
    last = f'   "id": "BK-97",{pnl}   "wave": "W3"{pnl}  }}'
    assert plan_raw.count(last) == 1
    rows = "".join(f',{pnl}  {{{pnl}   "id": "{k}",{pnl}   "wave": "{v}"{pnl}  }}' for k, v in WAVES.items())
    plan_new = plan_raw.replace(last, last + rows, 1)
    waves = {r["id"]: r["wave"] for r in json.loads(plan_new)["item_waves"]}
    assert all(waves[k] == v for k, v in WAVES.items())

    # ---- modules.json and packets.json ------------------------------------------
    modules = json.loads(MODULES.read_text(encoding="utf-8"))
    by_module = {m["id"]: m for m in modules["modules"]}
    for iid, mid in MODULE_OF.items():
        assert iid not in by_module[mid]["items"]
        by_module[mid]["items"].append(iid)
    catalogue["packets"].extend(PACKET_ROWS)

    # ---- BACKLOG.md: one record per item, then the reconciliation ---------------
    backlog_raw = BACKLOG.read_bytes().decode("utf-8")
    bnl = "\r\n" if "\r\n" in backlog_raw else "\n"
    sections = [f"## {heading(i)}{bnl}{bnl}" + RECORDS[i["id"]].replace("\n", bnl) for i in ITEMS]
    addition = bnl + (bnl + bnl).join(sections + [RECONCILIATION.replace("\n", bnl)]) + bnl
    backlog_new = backlog_raw.rstrip("\r\n") + bnl + addition

    # ---- the workbook: each LB row's delivery owner, in column 9 ----------------
    digest = hashlib.sha256(WORKBOOK.read_bytes()).hexdigest()
    book = load_workbook(WORKBOOK)
    sheet, plan = book["Before Build"], book["Implementation Plan"]
    features = {s.title: regroup.sheet_features(s) for s in book}
    before = {(s.title, c.coordinate): (c.value, copy.copy(c._style)) for s in book for row in s for c in row}
    rows_at = {}
    for r in range(regroup.FIRST, sheet.max_row + 1):
        m = regroup.LBOM.match(str(sheet.cell(r, 1).value or ""))
        if m:
            rows_at[m[1]] = r
    changed = {}
    for ident, owner in LB_OWNERS.items():
        r = rows_at[ident]
        value = sheet.cell(r, 9).value
        assert OWNER_PREFIX not in str(value or ""), ident
        new = f"{value}\n\n{OWNER_PREFIX}{owner}" if value else OWNER_PREFIX + owner
        sheet.cell(r, 9).value = new
        changed[(r, 9)] = new
    plan_changed = set()
    for x in range(2, plan.max_row + 1):
        ident = plan.cell(x, 1).value
        if ident in LB_OWNERS:
            for t, origin in regroup.MIRROR.items():
                v = sheet.cell(rows_at[ident], origin).value
                if plan.cell(x, t).value != v:
                    plan.cell(x, t).value = v
                    plan_changed.add(plan.cell(x, t).coordinate)
    out = Path(tempfile.gettempdir()) / "nm_register_legal_brain_20260926" / WORKBOOK.name
    out.parent.mkdir(exist_ok=True)
    book.save(out)
    check = load_workbook(out)
    for tab in check:
        assert regroup.sheet_features(tab) == features[tab.title], tab.title
    bb, ip = check["Before Build"], check["Implementation Plan"]
    for (title, coord), old in before.items():
        cell = bb[coord] if title == "Before Build" else ip[coord]
        if title == "Before Build" and (cell.row, cell.column) in changed:
            assert cell.value == changed[(cell.row, cell.column)] and cell._style == old[1], coord
            assert str(cell.value).startswith(str(old[0] or "")), f"{coord}: earlier text not kept"
        elif title == "Implementation Plan" and coord in plan_changed:
            continue
        else:
            assert (cell.value, cell._style) == old, f"{title} {coord} moved"
    from assurance.control_plane import plan_scenarios
    problems = plan_scenarios.requirement_problems(out)
    assert problems == [], problems

    # ---- every check has passed: write all six files together --------------------
    STATUS.write_bytes(status_new.encode("utf-8"))
    PLAN.write_bytes(plan_new.encode("utf-8"))
    json_write(MODULES, modules)
    json_write(PACKETS, catalogue)
    BACKLOG.write_bytes(backlog_new.encode("utf-8"))
    shutil.copy2(out, WORKBOOK)

    print(f"items: {', '.join(i['id'] for i in ITEMS)} "
          f"({sum(len(i['acceptance']) for i in ITEMS)} criteria); packets: "
          f"{', '.join(p['id'] for p in PACKET_ROWS)}; LB rows given an owner: {len(LB_OWNERS)}; "
          f"workbook {digest[:12]} -> {hashlib.sha256(WORKBOOK.read_bytes()).hexdigest()[:12]}, "
          f"plan cells {len(plan_changed)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
