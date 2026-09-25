"""Add the Indian practice layer (LB-120..125) to the curated workbook.

DRAFT ROWS. Column 10 of every row says so: not owner-agreed, not
counsel-reviewed, implementation NOT_ASSESSED, acceptance NOT_RUN. The rows
record what must be verified before anything is built, and every legal
proposition in them is a POINTER TO TEXT THAT MUST BE RETRIEVED AND READ BACK,
never a statement of law to rely on.

Same discipline as the other workbook tools here: snapshot every cell's value
and style, write only the intended cells, then reload the saved file and prove
that every other cell and every native sheet feature is unchanged. The source
is only replaced once that proof passes.
"""
import copy
import hashlib
import re
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

_REPO = Path(__file__).resolve().parents[2]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))
from assurance.common._console import utf8_console  # noqa: E402

utf8_console()

ROOT = _REPO
SOURCE = ROOT / "docs/Nyaymalaw_Implementation_Plan.xlsx"
SECTION = "Indian practice layer"

#: The one mechanism every row below shares, stated once here and referenced by
#: each row rather than restated six times.
MECHANISM = (
    "ONE MECHANISM, SHARED: a curated table per rule family in "
    "backend/nm/knowledge/, built the way nm.knowledge.resolution.Edge already "
    "is for Limitation Articles, every row carrying curated_from and the type "
    "refusing a row without it. The table IDENTIFIES on an exact key (cause, "
    "date, court, relief) and never by fuzzy match; the model APPLIES it to "
    "the advocate's words and returns applies / does not apply / not assessed "
    "with the quoted span. The provision is retrieved and read back, never "
    "recited from model memory. The determination enters the premise record so "
    "a corrected date, court or relief reopens it through the existing "
    "dependency ledger."
)

HEADING = (
    f"{SECTION}\nDRAFT rows for owner review, 25 September 2026\n\n"
    "Counted across LB-01..119 and both legal-brain blueprints on 23 September "
    "2026: BNS/BNSS/BSA 0 mentions, pre-institution conditions 0, per incuriam "
    "and larger bench 0, court fee and valuation 0, interim-injunction tests 2. "
    "LB-20 names notice, valuation and fees as a category and leaves the rules "
    "to be invented. These six rows are the procedural craft that decides "
    "Indian matters before their merits are reached.\n\n"
    + MECHANISM + "\n\n"
    "NOT OWNER-AGREED AND NOT COUNSEL-REVIEWED. No table is built until the "
    "owner selects the rows and a practising Telangana advocate signs off its "
    "entries as counsel review, on the BK-85-AC3 pattern. Every provision and "
    "judgment named below is to be retrieved and verified first; none is "
    "asserted as current law by this sheet."
)

READINESS = (
    "25 September 2026: DRAFT prepared for owner review; not a verbatim owner "
    "statement, not counsel-reviewed and not an approved requirement. "
    "Implementation NOT_ASSESSED; acceptance NOT_RUN. Delivery mapping OPEN "
    "under LB-43; quality and release decisions OPEN under LB-44. Every legal "
    "proposition is a pointer to text that must be retrieved and read back "
    "before any table is curated. No completion or release approval."
)

ROWS = [
    [
        "LB-120\n" + SECTION + "\nApply the law in force on the date that governs it\n\n"
        "DRAFT REQUIREMENT FOR OWNER REVIEW:\n"
        "Where a statute has been repealed and replaced, NM must establish which "
        "one governs this matter from its dates before reading either. First "
        "population: IPC to BNS, CrPC to BNSS and the Evidence Act to BSA, in "
        "force from 1 July 2024.",

        "Never receive analysis under a code that does not govern the offence, "
        "the proceeding or the evidence in issue.",

        "The matter involves an offence, a criminal proceeding or a question of "
        "evidence, or cites a section of a repealed or a replacing code.",

        "Establish the dates that decide which law applies: the offence date for "
        "substantive law, and whether an investigation, inquiry, trial or appeal "
        "was pending at commencement for procedure. Read substantive and "
        "procedural law separately; one matter can need the old substantive code "
        "and the new procedural one together. State the governing code, the date "
        "it rests on and the saving provision that decides it. Map a section "
        "cited under one code to its counterpart only through a curated "
        "correspondence table, never by matching section numbers. " + MECHANISM,

        "Every cited provision belongs to the code that governs it for this "
        "matter, and the date that decided it is on the premise record.",

        "An unknown or disputed governing date yields both readings, labelled "
        "conditional, with a question for the date. A corrected date reopens the "
        "choice and everything resting on it through the dependency ledger.",

        "Never apply a replacing penal provision to conduct before its "
        "commencement. Never treat equal section numbers across two codes as "
        "corresponding. Never read a pending proceeding's procedure from the new "
        "code without the saving provision that permits it. Never recite either "
        "code's text from model memory.",

        "LB-120-AC1: an offence on 10 June 2024 charged in August 2024 reaches "
        "the old substantive code for the offence and the new procedural code "
        "for the proceeding, each with its date and saving provision.\n"
        "LB-120-AC2 (planted): an answer citing a replacing penal provision for "
        "a 2023 offence is refused before release.\n"
        "LB-120-AC3 (planted): a section mapped across codes by equal number is "
        "refused; correspondence comes only from the curated table.\n"
        "LB-120-AC4: correcting the offence date across 1 July 2024 reopens the "
        "governing-code premise and every conclusion resting on it.",

        "LB-13/14/19/20/29; premise and dependency owners (P18, P22). VERIFY "
        "BEFORE BUILDING: the commencement notifications; BNSS s.531 (repeal and "
        "savings); BSA s.170 (repeal and savings); Constitution Article 20(1). "
        "CORPUS, MEASURED (BASELINE section 2.1): IPC 574, BNS 358, CrPC 509 and "
        "BNSS 531 sections held; Evidence Act 175. THE BHARATIYA SAKSHYA "
        "ADHINIYAM IS NOT AMONG THE MEASURED PRINCIPAL ACTS -- measure it "
        "against raw_data/ before the evidence half is built. OPEN: whether an "
        "official correspondence table is held, and who curates it.",

        READINESS + " The generalisation is deliberate: the rule is about repeal "
        "with savings, and the three criminal codes are its first population, "
        "not the rule itself.",
    ],
    [
        "LB-121\n" + SECTION + "\nMandatory steps before a proceeding can be instituted\n\n"
        "DRAFT REQUIREMENT FOR OWNER REVIEW:\n"
        "Before recommending that a suit or complaint be filed, NM must "
        "establish whether a statute requires something first -- a notice, a "
        "mediation, a waiting period -- and whether the file shows it done.",

        "Never file something that fails before its merits are reached for want "
        "of a step that had to come first.",

        "NM is about to recommend instituting a proceeding, or the advocate asks "
        "whether they can file.",

        "Identify from the curated table every pre-institution condition the "
        "cause, forum and parties engage. For each, show its source, what the "
        "file establishes about it (done, when, by whom), what is outstanding, "
        "and the date it is satisfied. First population, each entry subject to "
        "retrieval and verification: Commercial Courts Act s.12A pre-institution "
        "mediation where no urgent interim relief is contemplated; CPC s.80 "
        "notice before suing the Government or a public officer, with its leave "
        "route for urgent relief; NI Act s.138 provisos and s.142(1)(b) -- "
        "presentation within validity, demand notice within the statutory period "
        "of the information, the payment window, and the complaint window; TPA "
        "s.106 notice terminating a lease, for the tenancy type. " + MECHANISM,

        "Each engaged condition reads satisfied, outstanding with its date, or "
        "not assessed. A recommendation to file waits on the outstanding ones.",

        "Missing facts about a condition give not assessed and one question "
        "each, never an assumption that it was done. An urgent-relief claim that "
        "would lift a condition is surfaced as the advocate's decision with its "
        "test, never taken by NM.",

        "Never recommend filing while an engaged condition is outstanding "
        "without saying so first. Never read no condition found as none applies "
        "where the cause or the forum is unestablished; that is the third state. "
        "Never treat the advocate saying a notice went as proof that it did.",

        "LB-121-AC1: a dishonoured cheque with the demand notice sent outside "
        "the statutory window is identified as out of time, and ordinary filing "
        "is not recommended.\n"
        "LB-121-AC2: a commercial recovery with no urgency pleaded names s.12A "
        "mediation before filing.\n"
        "LB-121-AC3 (planted): an answer recommending suit against a State "
        "department with no mention of s.80 is refused.\n"
        "LB-121-AC4: an unestablished forum yields not assessed, never no "
        "conditions apply.",

        "LB-20/22/29/48; deadline register and premise owners (P22). VERIFY "
        "BEFORE BUILDING: the current text of each provision named, and Patil "
        "Automation v Rakheja Engineers (2022) on s.12A. CORPUS, MEASURED "
        "(BASELINE section 2.1): CPC 826 and NI Act 261 sections held; TPA 131. "
        "THE COMMERCIAL COURTS ACT IS NOT AMONG THE MEASURED PRINCIPAL ACTS -- "
        "measure it against raw_data/ first. OPEN: the specified-value threshold "
        "and its source; which Telangana courts are designated Commercial "
        "Courts. Share that determination with LB-124 rather than recomputing "
        "it.",

        READINESS + " Extends LB-20 from a named category to a curated, "
        "checkable population.",
    ],
    [
        "LB-122\n" + SECTION + "\nWeigh a precedent by what binds this court\n\n"
        "DRAFT REQUIREMENT FOR OWNER REVIEW:\n"
        "Whether a judgment binds depends on the court that decided it, the "
        "strength of its bench and what happened to it afterwards -- not on how "
        "relevant it reads. NM must weigh authority the way the forum will.",

        "Rely on authority that binds this forum, and be warned about authority "
        "that has been overruled, referred, doubted or decided per incuriam.",

        "NM relies on a judgment, or the advocate cites one.",

        "For each relied-on judgment record: its court, bench strength and date; "
        "its relationship to this forum -- binding, persuasive or not binding, "
        "through the existing nm.knowledge.jurisdiction relationship extended by "
        "bench strength; the proposition it is relied on for, marked ratio or "
        "obiter, with its paragraph; and its subsequent treatment where the "
        "corpus establishes one. Where two co-equal benches conflict, say so and "
        "give the rule for which prevails rather than choosing silently. " + MECHANISM,

        "Each authority carries its binding status for this forum, the basis of "
        "the proposition taken from it, and its treatment state.",

        "Unknown bench strength or unknown treatment leaves the authority usable "
        "with the limit disclosed; it never carries a proposition alone.",

        "Never present a smaller bench as prevailing over a larger one on the "
        "same point. Never present obiter as holding; this extends G-ATTRIB. "
        "Never read a silent citator as clearance -- treatment the corpus cannot "
        "establish is not assessed, never good law.",

        "LB-122-AC1: a division-bench and a single-judge decision of the same "
        "High Court conflict; the division bench is identified as binding on the "
        "single judge and on the trial court.\n"
        "LB-122-AC2 (planted): a judgment whose point stands referred to a "
        "larger bench, presented as settled, is refused or disclosed.\n"
        "LB-122-AC3 (planted): an obiter passage quoted as the holding is "
        "refused.",

        "LB-14/16/18/25/38; G-GROUND(b), G-ATTRIB and jurisdiction owners. "
        "VERIFY BEFORE BUILDING: Constitution Article 141; Central Board of "
        "Dawoodi Bohra Community v State of Maharashtra (2005) on bench "
        "strength; State of UP v Synthetics and Chemicals (1991) on per "
        "incuriam. CORPUS, MEASURED (BASELINE): Bench: header on 30,710 of "
        "34,037 raw judgment files (90.2%) and Equivalent citations: on 82.2% -- "
        "IN raw_data/ ONLY; the derived store dropped both, so bench strength is "
        "read there, by rglob and never by find. THE 0.7% AP-HIGH-COURT BENCH "
        "FIGURE IN CIRCULATION IS ONE OF THE THREE CLAIMS BASELINE RECORDS AS "
        "MEASURED FROM THE DERIVED LAYER AND WRONG; this row would have been "
        "declared unbuildable on it. Re-measure and name the store. OPEN: a "
        "treatment source, since none is measured as held. Honours the standing "
        "decision that Andhra Pradesh High Court judgements bind in Telangana "
        "(BASELINE section 1.1).",

        READINESS,
    ],
    [
        "LB-123\n" + SECTION + "\nAssess interim relief on its own test\n\n"
        "DRAFT REQUIREMENT FOR OWNER REVIEW:\n"
        "An application for interim relief is decided on a different test from "
        "the final merits. NM must assess it separately, on the test for the "
        "relief actually sought.",

        "Know whether interim relief is realistically available now, and what "
        "the application must show, apart from whether the suit ultimately "
        "succeeds.",

        "The objective or the urgency points to interim relief -- injunction, "
        "stay, attachment, receiver -- or the advocate asks for it.",

        "Identify the relief and its source, whether CPC Order XXXIX rules 1 and "
        "2 or a special statute. Assess prima facie case, balance of convenience "
        "and irreparable injury separately, each against the record. Apply the "
        "higher threshold where the relief sought is mandatory rather than "
        "prohibitory, and check the statutory bars on injunctions. For ex parte "
        "relief, show what the court must record and what the applicant must "
        "then do. State what evidence at this stage would strengthen each limb. "
        + MECHANISM,

        "Each limb reads supported, weak or not assessed with its basis, and the "
        "interim position is kept distinct from the final-merits assessment "
        "under LB-21.",

        "A missing limb gives a conditional assessment and a focused question. "
        "Urgency is routed to LB-06 protective handling and never converted into "
        "a verified deadline.",

        "Never infer final success from interim prospects, or interim prospects "
        "from final merits. Never apply the prohibitory threshold to a mandatory "
        "injunction. Never omit an engaged statutory bar.",

        "LB-123-AC1: a wall under construction with a strong title document has "
        "each limb assessed separately, with the status-quo reasoning shown.\n"
        "LB-123-AC2 (planted): a mandatory injunction assessed on the ordinary "
        "triple test alone is refused.\n"
        "LB-123-AC3: an engaged statutory bar is named before the limbs are "
        "reached.",

        "LB-06/20/21/22/50; relief owner (P23). VERIFY BEFORE BUILDING: CPC "
        "Order XXXIX rules 1, 2, 3 and 3A; Dalpat Kumar v Prahlad Singh (1992); "
        "Dorab Cawasji Warden v Coomi Sorab Warden (1990); Wander v Antox (1990) "
        "on appellate review; Specific Relief Act s.41. CORPUS, MEASURED "
        "(BASELINE section 2.1): CPC 826 sections held; Specific Relief Act 44 "
        "-- IN THE UPPERCASE STORE ONLY, so both stores are searched (CLAUDE.md, "
        "three stores three answers).",

        READINESS,
    ],
    [
        "LB-124\n" + SECTION + "\nTime limits that run inside a proceeding\n\n"
        "DRAFT REQUIREMENT FOR OWNER REVIEW:\n"
        "Beyond limitation for filing, a proceeding has its own clocks: the "
        "written statement, leave to defend, a caveat, condonation of delay. NM "
        "must know which run for this matter and whether they bind.",

        "Not lose a defence or a right of response to a period nobody computed, "
        "and know when a period can be extended.",

        "Service, a filing, an order or a listing is recorded, or the advocate "
        "asks what is due.",

        "Compute each engaged period from its trigger date through the existing "
        "deadline register, keeping asserted, computed and court-listed dates "
        "apart as LB-20 requires. State whether the period is mandatory or "
        "directory and whether it can be extended, from its curated source. "
        "First population, each subject to verification: CPC Order VIII rule 1 "
        "written statement, with the commercial-suit outer limit and the "
        "ordinary-suit reading distinguished; Order XXXVII appearance and leave "
        "to defend; CPC s.148A caveat and how long it stays in force; Limitation "
        "Act s.5 condonation and where it does not apply. " + MECHANISM,

        "Each engaged period is on the register with its trigger, its source and "
        "whether it is mandatory or directory.",

        "An unknown trigger date gives a conditional period and a question, "
        "never a guessed date. A corrected trigger moves the period and is "
        "reported with its prior through the cascade.",

        "Never apply the ordinary-suit reading to a commercial suit or the "
        "reverse. Never present a mandatory period as extendable. Never compute "
        "a period from a date the advocate has not given.",

        "LB-124-AC1: a commercial suit served on a stated date computes the "
        "written-statement outer limit and labels it mandatory.\n"
        "LB-124-AC2 (planted): an extension suggested past the commercial outer "
        "limit is refused.\n"
        "LB-124-AC3: correcting the service date moves the period and G-CASCADE "
        "reports the move with its prior value.",

        "LB-20/29/121; deadline register, premise and cascade owners (P18, P22). "
        "VERIFY BEFORE BUILDING: CPC Order VIII rule 1, Order XXXVII and s.148A; "
        "Limitation Act s.5; SCG Contracts v K.S. Chamankar (2019) and Kailash v "
        "Nanhku (2005) on the two readings. CORPUS, MEASURED (BASELINE section "
        "2.1): CPC 826 sections and Limitation Act 169 sections plus 137 "
        "Schedule Articles held. The commercial-or-ordinary determination "
        "belongs to LB-121; share it rather than recomputing it.",

        READINESS + " Reuses the deadline register and the dependency ledger; no "
        "new timing mechanism is introduced.",
    ],
    [
        "LB-125\n" + SECTION + "\nWhether a filing will be accepted: forum, valuation and court fee\n\n"
        "DRAFT REQUIREMENT FOR OWNER REVIEW:\n"
        "Before a filing is recommended, NM must establish the court it goes to "
        "-- subject matter, pecuniary and territorial -- how the suit is valued, "
        "and the court fee payable under the law in force in Telangana.",

        "File in the right court with the right valuation and fee, and know the "
        "cost before the step is recommended.",

        "NM is about to recommend a filing, or the advocate asks where to file "
        "or what it will cost.",

        "Derive the subject-matter forum, then the pecuniary tier from the "
        "valuation, then territorial jurisdiction from the cause of action and "
        "the parties' residence or business, in the CPC ss.15 to 20 order. Value "
        "the suit under the curated valuation rule for the relief sought and "
        "compute the fee from the current schedule, showing the workings. State "
        "the source and the version of every figure used. " + MECHANISM,

        "Forum, valuation and fee are each computed with their workings, or "
        "marked not assessed with the missing input named.",

        "A schedule version that is not held means the fee is not computed and "
        "the gap is disclosed. It is never estimated from a neighbouring "
        "schedule or from model memory.",

        "Never compute a fee from a schedule whose version is unverified. Never "
        "infer the pecuniary tier without a valuation. Never state territorial "
        "jurisdiction without the cause-of-action facts that ground it.",

        "LB-125-AC1: a recovery suit for a stated sum at Hyderabad yields forum, "
        "tier and fee with workings.\n"
        "LB-125-AC2 (planted): a fee computed from an unversioned schedule is "
        "refused.\n"
        "LB-125-AC3: correcting the claimed amount across a tier boundary "
        "reopens both the forum and the fee.",

        "LB-20/121; threshold owners -- nm.core.thresholds already carries "
        "VALUATION, COURT_FEES and JURISDICTION as not assessed, and this row is "
        "what would assess them. CORPUS: NEITHER THE COURT-FEES AND SUITS "
        "VALUATION ACT AS APPLIED IN TELANGANA NOR THE TELANGANA CIVIL COURTS "
        "ACT IS AMONG THE MEASURED PRINCIPAL ACTS (BASELINE section 2.1). Verify "
        "the adapted titles, the current schedules and the pecuniary limits, and "
        "measure against raw_data/, before anything is built. OPEN: who supplies "
        "and maintains the schedule, and its version identity.",

        READINESS + " The most corpus-dependent of the six. If the schedule "
        "cannot be held and versioned, this row ships as an honest not assessed "
        "rather than an estimate.",
    ],
]

#: The row-2 note states a count of LB/OM requirements. MEASURED, not assumed:
#: it read 115 while the sheet already carried 142, so it was stale before this
#: change. Both the old and the new figure are computed from the sheet below.
NOTE_PATTERN = re.compile(r"All \d+ LB/OM requirements are linked into Implementation Plan")

#: Plan-sheet column <- Before Build column. The same fixed contract
#: `assurance/control_plane/plan_scenarios.requirement_problems` enforces.
MIRROR = {8: 1, 9: 4, 10: 2, 12: 3, 15: 4, 16: 5, 17: 6, 19: 7, 28: 9, 30: 10, 33: 8}


def sheet_features(sheet):
    """Preserve native sheet controls, not only visible cell values."""
    from openpyxl.xml.functions import tostring
    return {
        "merged": tuple(str(area) for area in sheet.merged_cells.ranges),
        "freeze": sheet.freeze_panes,
        "state": sheet.sheet_state,
        "validations": tostring(sheet.data_validations.to_tree()),
        "conditional": repr(list(sheet.conditional_formatting)),
        "views": tostring(sheet.views.to_tree()),
        "protection": tostring(sheet.protection.to_tree()),
        "margins": tostring(sheet.page_margins.to_tree()),
        "print_options": tostring(sheet.print_options.to_tree()),
        "columns": {key: (dim.width, dim.hidden, dim.outlineLevel)
                    for key, dim in sheet.column_dimensions.items()},
        "hyperlinks": {cell.coordinate: (cell.hyperlink.target, cell.hyperlink.location)
                       for row in sheet for cell in row if cell.hyperlink},
    }


def main() -> int:
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    book = load_workbook(SOURCE)
    before = {(s.title, c.coordinate): (c.value, copy.copy(c._style))
              for s in book for row in s for c in row}
    features = {s.title: sheet_features(s) for s in book}
    changed: set = set()

    def write(tab, row, col, value):
        tab.cell(row, col, value)
        changed.add((tab.title, tab.cell(row, col).coordinate))

    sheet = book["Before Build"]

    def population(tab):
        return {match[1]
                for r in range(5, tab.max_row + 1)
                for match in [re.match(r"(LB-\d+|OM-[PIQ]\d+)\b",
                                       str(tab.cell(r, 1).value or ""))]
                if match}

    was = len(population(sheet))

    # THE SECTION HEADING, once, if it is not already there.
    if not any(str(sheet.cell(r, 1).value or "").startswith(SECTION + "\n")
               for r in range(5, sheet.max_row + 1)):
        row = sheet.max_row + 1
        sheet.cell(row, 1)._style = copy.copy(sheet.cell(row - 1, 1)._style)
        write(sheet, row, 1, HEADING)
        sheet.row_dimensions[row].height = 280

    indices = {}
    for values in ROWS:
        ident = values[0].splitlines()[0]
        matches = [r for r in range(5, sheet.max_row + 1)
                   if str(sheet.cell(r, 1).value or "").startswith(ident + "\n")]
        assert len(matches) <= 1, ident
        row = matches[0] if matches else sheet.max_row + 1
        indices[ident] = row
        for col, value in enumerate(values, 1):
            if not matches:
                sheet.cell(row, col)._style = copy.copy(sheet.cell(row - 1, col)._style)
            write(sheet, row, col, value)
        sheet.row_dimensions[row].height = 280

    # THE MIRROR, so both sheets carry every requirement. This is the contract
    # `plan_scenarios.requirement_problems` checks in both directions.
    plan = book["Implementation Plan"]
    for ident, source in indices.items():
        matches = [r for r in range(2, plan.max_row + 1) if plan.cell(r, 1).value == ident]
        assert len(matches) <= 1, ident
        row = matches[0] if matches else plan.max_row + 1
        if not matches:
            for col in range(1, plan.max_column + 1):
                plan.cell(row, col)._style = copy.copy(plan.cell(row - 1, col)._style)
            for col, value in {
                    1: ident, 3: "Requirement", 4: "Legal brain",
                    6: sheet.cell(source, 1).value.splitlines()[2],
                    7: "Pilot", 38: "Draft", 39: "Not started",
                    40: "Not verified"}.items():
                write(plan, row, col, value)
        for target, origin in MIRROR.items():
            write(plan, row, target, sheet.cell(source, origin).value)
        write(plan, row, 32,
              "25 September 2026: draft practice-layer row mirrored from Before "
              "Build. Not acceptance, not an approved requirement.")

    # THE NOTE'S COUNT IS A CLAIM ABOUT THIS SHEET, so it is measured from it.
    now = len(population(sheet))
    note = str(sheet.cell(2, 1).value or "")
    replaced, count = NOTE_PATTERN.subn(
        f"All {now} LB/OM requirements are linked into Implementation Plan", note)
    assert count == 1, "the row-2 note no longer states the linked-requirement count"
    write(sheet, 2, 1, replaced)

    scratch = Path("/tmp/claude-0/-home-user-Nyaymalaw/"
                   "82a504d3-1374-501f-9fe3-441ac9aa966e/scratchpad")
    scratch.mkdir(parents=True, exist_ok=True)
    target = scratch / SOURCE.name
    book.save(target)

    # THE PROOF, on the saved bytes rather than on the objects in memory.
    check = load_workbook(target)
    assert check.sheetnames == book.sheetnames
    for tab in check:
        assert sheet_features(tab) == features[tab.title], tab.title
    for key, old in before.items():
        if key not in changed:
            cell = check[key[0]][key[1]]
            assert (cell.value, cell._style) == old, key
    from assurance.control_plane import plan_scenarios
    problems = plan_scenarios.requirement_problems(target)
    assert not problems, problems
    assert len(plan_scenarios.sheet_rows(target)) == 93, "the feature population moved"

    shutil.copy2(target, SOURCE)
    print(f"source digest before: {digest}")
    print(f"source digest after:  {hashlib.sha256(SOURCE.read_bytes()).hexdigest()}")
    print(f"LB/OM requirements: {was} -> {now}; rows added: {sorted(indices)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
