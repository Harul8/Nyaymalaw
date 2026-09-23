"""Narrow Before Build amendment. Artifact-tool unavailable in bundled runtime.

Use bundled Python for this authorised fallback; preserve every existing cell.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import shutil
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.xml.functions import tostring

from reconcile_implementation_20260922 import preview
import sys

_REPO = Path(__file__).resolve().parents[2]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))
from assurance.common._console import utf8_console  # noqa: E402

utf8_console()

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs/Nyaymalaw_Implementation_Plan.xlsx"
OUT = ROOT / "outputs/01a07b76-6b21-71f3-bb09-261f64617594"
RECORD = "docs/backlog/evidence/legal-brain-20260922/chat-references.md"
CONTRACTS = [
    ("LB-93", "One calm conversation surface",
     "Read NM naturally while keeping the matter board available.",
     "New, restored or historical conversation; desktop, mobile, zoom or theme change.",
     "Keep one matter board on the left, compact navigation above and profile below. Use unboxed assistant paragraphs, understated user bubbles, readable spacing and one aligned bounded composer. Reuse existing send/clear/retry and draft protections. No second board or permanent source column.",
     "The transcript occupies the available right-hand area; all navigation and input remain reachable.",
     "Keep persistent error and unsaved-work notices. Source inspection and repaint must preserve text, cursor and reading position.",
     "Do not transplant reference React architecture or CSS wholesale. Do not make text tiny, hide the board on mobile, or turn every sentence into a diagnostic card.",
     "Test at phone, tablet and desktop widths, zoom and dark mode. Send multiline text and observe immediate clear/shrink, safe retry and preservation of newer edits. Verify one board, no overflow and keyboard reachability.",
     "Builds on LB-75/83/84/90 and existing single-board journey tests. UI polish, not a new legal conclusion."),
    ("LB-94", "Prose and supporting material retain their meaning",
     "Read the assessment and inspect its basis without internal process labels.",
     "Render a validated answer containing explanation, qualifications, quotations or references.",
     "Preserve accepted reading order, natural paragraphs and all material qualifications. Display supplied supporting passages by default; ordinary support may be collapsed deliberately. Use restrained citation and quotation styling. Never require one rigid response template.",
     "The advocate can see the assessment, quoted basis and material limits without opening technical details.",
     "Unknown element types, signals and unresolved provenance remain visible. No cosmetic rewriting or deletion of substantive wording.",
     "No chain-of-thought display, section taxonomy as conversation, regex censorship, or blanket confidence/correctness percentages.",
     "Assert visible prose, warnings outside collapsible support, retained quotation words and citation identities. Equal text with different refs or qualifications must not deduplicate into lost evidence.",
     "BK-37; LB-76/77/78/87. Existing language and grounding controls remain authoritative."),
    ("LB-95", "Exact saved citation identity",
     "Open the source actually supplied for this response, not a similarly named source.",
     "A retrieved passage is assembled into a released answer and persisted, replayed or restored.",
     "Capture readable label, locator, source namespace, exact retrieved text and a content identity with the answer. Link only the bound element. Keep multiple same-title sources and same-number sections distinct. Older answers without this binding keep their original reference text and explicitly state that source inspection is unavailable.",
     "The selected reference resolves through the authorised saved answer. The digest identifies saved content, not an enacted-law version or legal-currentness guarantee.",
     "Refuse malformed identity, changed text/digest, missing reference binding and unreleased answers. Do not fetch a current replacement for a historical source.",
     "No section-number matching, fuzzy identity, model URL execution, guessing court/date/version, or invented legacy source reconstruction.",
     "Test duplicate labels across Acts and versions, exact text/digest, persisted reload and replay, old receipts, unknown references, withheld answers and injected markup. Saved snapshot metadata must never bypass grounding.",
     "LB-91; BK-37/BK-38. Full original documents are not retained in existing turn receipts; do not claim they are."),
    ("LB-96", "Contextual source reader",
     "Inspect saved source wording without losing the answer or draft.",
     "Activate a bound citation using pointer or keyboard.",
     "Open a right-hand modal reader, full-width on small screens. Show source label, saved-content identity and actual coverage. Start at the cited saved passage. Page long text, offer return-to-passage, search explicitly within loaded text and copy the loaded passage with its attribution, coverage and qualifications.",
     "Closing or Escape restores the originating focus and conversation position. Plain extracted text is not described as an original facsimile or whole Act/judgment.",
     "Show missing/unavailable/full-document-not-retained states explicitly. Retry only the selected reference. Never highlight an approximate match as the cited anchor.",
     "No automatic new tab, external fetch, model call, unbounded prefetch, raw HTML or silently substituted current law. Whole-document/original controls appear only when exact authorised content exists.",
     "Test accurate scope labels, large-text paging with no dropped characters, loaded-text search, safe attributed copy, keyboard Escape/focus, phone layout, zoom and no model/retrieval calls on inspection.",
     "LB-92. This delivery provides the retained exact passage; larger context and original-document acquisition remain separately dependent on version-bound corpus storage."),
    ("LB-97", "Permission and asynchronous reader isolation",
     "Inspect only material belonging to the current authorised account and matter.",
     "Open, page or copy a source; switch references/matters, close, navigate, expire or sign out during a read.",
     "Authorise every reader request against the matter and validated released turn. Invalidate pending callbacks on close/navigation/session change. Keep no persistent browser source cache. Recheck permission when copying and discard changed pages rather than mixing versions.",
     "A late response never reopens a closed reader or reveals another matter. Sign-out clears protected source content.",
     "Use non-enumerating failures for foreign/missing records, visible retry for unavailable reads and no success notice after refusal. Refused access clears the displayed protected text.",
     "No client-only permission, archive-as-release shortcut, credential-bearing source links or reusable cache that outlives access.",
     "Test another account, expired/revoked session, withheld turn, malformed source, out-of-range page, delayed competing responses, close in flight, matter switch and copy after access is lost.",
     "Existing session, receipt, matter ownership and erasure boundaries remain owners. No change to global OpenAI permissions."),
    ("LB-98", "Accessible and quiet reader controls",
     "Use chat references without relying on colour, hovering or a mouse.",
     "Read or operate references, source dialog, search, copy and recovery controls.",
     "Use readable wrapping underlined citation buttons, named dialog and controls, native modal focus containment, visible focus and polite status announcements. Keep errors persistent. Use existing colour tokens and respect reduced motion.",
     "Long labels and unbroken source tokens fit the viewport; controls remain operable at narrow widths and zoom.",
     "Clipboard denial exposes an actionable persistent message and selectable text. Missing citation focus returns to a safe visible control.",
     "No colour-only status, hover-only identity, invisible focus, flashing notices or successful-copy claim before clipboard success.",
     "Exercise keyboard-only, Escape, labelled search, focus return, mobile, dark mode, reduced motion, long labels and clipboard refusal. Screen-reader professional conformance remains separately reviewed.",
     "LB-88/89/92. Controlled browser checks establish behaviour, not formal accessibility certification."),
    ("LB-99", "Scoped proof and honest completion",
     "Know which presentation changes are built and which broader capabilities still need proof.",
     "Implement, test and record this selected chat-reference slice.",
     "Record the exact changes and tests against this workbook, current source and backlog. Run served persistence/access tests, negative controls and browser journeys, plus a separate interactive browser walkthrough. Preserve parallel edits and all earlier requirements.",
     "Each changed promise has recorded evidence and remaining limitations; no broader expert-quality or release claim follows from UI tests.",
     "Record failed checks and unassessed scope; do not relax tests, invent legal approval or promote old full-gate evidence.",
     "No paid model batch without fresh bounded approval, no changes to the reference project, no automatic commit/push or claim that all LB-90/91/92 obligations are complete.",
     "Compare workbook before/after at cell/style and sheet-feature level. Re-run affected receipt/grounding/history/UI suites. Log interactive browser results separately from scripted-model fixtures.",
     "BK-37/BK-38; this record owns the bounded engineering change. Existing counsel/live-model/full-document obligations remain open."),
]


def sheet_features(sheet):
    """Preserve native sheet controls, not only visible cell values."""
    return {
        "merged": tuple(str(area) for area in sheet.merged_cells.ranges),
        "freeze": sheet.freeze_panes,
        "state": sheet.sheet_state,
        "validations": tostring(sheet.data_validations.to_tree()),
        "conditional": repr(list(sheet.conditional_formatting)),
        "views": tostring(sheet.views.to_tree()),
        "protection": tostring(sheet.protection.to_tree()),
        "margins": tostring(sheet.page_margins.to_tree()),
        "print_options": tostring(sheet.print_options.to_tree()),
        "columns": {key: (dim.width, dim.hidden, dim.outlineLevel)
                    for key, dim in sheet.column_dimensions.items()},
        "hyperlinks": {cell.coordinate: (cell.hyperlink.target, cell.hyperlink.location)
                       for row in sheet for cell in row if cell.hyperlink},
    }


def run(phase):
    OUT.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    w = load_workbook(SOURCE)
    before = {(s.title, c.coordinate): (c.value, copy.copy(c._style))
              for s in w for row in s for c in row}
    features = {sheet.title: sheet_features(sheet) for sheet in w}
    s = w["Before Build"]
    preview(s, [247], [1, 4], OUT / "chat-reference-before.png")
    if phase == "inspect":
        return
    backup = OUT / "implementation-before-chat-references.xlsx"
    if not backup.exists():
        shutil.copy2(SOURCE, backup)
    changed = set()
    for spec in CONTRACTS:
        ident, title, *content = spec
        matches = [row for row in range(5, s.max_row + 1)
                   if str(s.cell(row, 1).value or "").startswith(ident + "\n")]
        row = matches[0] if matches else s.max_row + 1
        values = [f"{ident}\nChat and source presentation\n{title}\n\n"
                  "22 September 2026: owner-authorised consolidation of the selected "
                  "Agentified NM patterns and current-code observations.", *content]
        values.append("22 September 2026: approved build contract. Engineering in progress; "
                      "tests NOT_RUN. Evidence: " + RECORD if phase == "plan" else
                      "22 September 2026: bounded engineering implemented. See " + RECORD
                      + " for counted test/browser outcomes and remaining full-document, "
                      "live-model, legal and accessibility-review limits. No overall release sign-off.")
        assert len(values) == 10
        for col, value in enumerate(values, 1):
            c = s.cell(row, col)
            if not matches:
                c._style = copy.copy(s.cell(248, col)._style)
            c.value = value
            changed.add((s.title, c.coordinate))
        s.row_dimensions[row].height = 240
    for row in (239, 240, 246, 247, 248):
        c = s.cell(row, 10)
        addition = "\n22 September authorised implementation detail: LB-93–99; " + RECORD
        if addition not in str(c.value):
            c.value = str(c.value or "") + addition
        changed.add((s.title, c.coordinate))
    for table in s.tables.values():
        if table.ref.endswith("266") or table.ref.startswith("A4:J"):
            table.ref = f"A4:J{s.max_row}"
            if table.autoFilter:
                table.autoFilter.ref = table.ref
    if s.auto_filter.ref:
        s.auto_filter.ref = f"A4:J{s.max_row}"
    for sheet in w:
        for cells in sheet:
            for cell in cells:
                key = (sheet.title, cell.coordinate)
                if key in before and key not in changed:
                    assert (cell.value, cell._style) == before[key], key
    preview(s, [267, 269], [1, 4], OUT / "chat-reference-after.png")
    result = OUT / SOURCE.name
    w.save(result)
    check = load_workbook(result)
    assert check.sheetnames == w.sheetnames
    for sheet in check:
        assert sheet_features(sheet) == features[sheet.title], sheet.title
        assert sheet.auto_filter == w[sheet.title].auto_filter, sheet.title
        assert list(sheet.tables) == list(w[sheet.title].tables), sheet.title
        for name in sheet.tables:
            assert sheet.tables[name] == w[sheet.title].tables[name], name
    for key, (value, style) in before.items():
        if key not in changed:
            cell = check[key[0]][key[1]]
            assert cell.value == value and cell._style == style, key
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == digest, "parallel workbook edit"
    shutil.copy2(result, SOURCE)
    report = {"phase": phase, "rows": [267, 273], "requirements": 7,
              "changed_cells": len(changed), "preserved_cells": len(before) - len(changed & before.keys()),
              "sheets": check.sheetnames, "unrelated_cell_or_style_changes": 0,
              "native_sheet_features_preserved": True}
    (OUT / "chat-reference-workbook-check.json").write_text(json.dumps(report, indent=2))
    print(json.dumps(report))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("phase", choices=["inspect", "plan", "done"])
    run(parser.parse_args().phase)
