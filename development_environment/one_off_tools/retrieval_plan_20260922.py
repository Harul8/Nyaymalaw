"""Append bounded RAG diagnostics without rewriting the implementation workbook."""
import copy
import hashlib
import json
import shutil
from pathlib import Path

from openpyxl import load_workbook

from chat_reference_plan_20260922 import sheet_features
from reconcile_implementation_20260922 import preview

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'docs/Nyaymalaw_Implementation_Plan.xlsx'
OUT = ROOT / 'outputs/01a07b76-6b21-71f3-bb09-261f64617594'
RECORD = 'docs/backlog/evidence/legal-brain-20260922/retrieval-hardening.md'

# Existing ten-column schema; these are diagnostic requirements, not new phases.
ROWS = [
    ('LB-100', 'Published source withdrawal',
     'Do not receive legal material whose publication has been withdrawn.',
     'Any search, case discovery, expansion, passage, citation resolution, identity or treatment read.',
     'Check the bound publication before reading and immediately before returning. All read surfaces use one guard. Refusal remains distinct from a completed search.',
     'The response contains no withdrawn text, identity or clean-treatment claim.',
     'A withdrawal during the read invalidates the pending result. Existing published-generation and permission policies remain in force.',
     'Never reuse a cached permission or assume metadata reads are exempt. Historical retention is not permission for fresh retrieval.',
     'Exercise all seven reads before withdrawal and during a read. Plant successful identity/treatment results and verify they cannot escape.',
     'BK-84/BK-38. Existing immutable-publication policy is the owner.', 'implemented'),
    ('LB-101', 'Exact case scope',
     'Read the case I selected, not other cases reached by search syntax.',
     'Expand a case with a search term or discover cases in a bounded candidate pool.',
     'Bind case identity as a separate exact database parameter. Search text is literal data. Discovery must examine the pool it reports, not a silently smaller public search limit.',
     'Every returned paragraph belongs to the selected case; the examined discovery population is truthful.',
     'An unknown or punctuation-bearing identifier returns no exact case, without widening the query.',
     'Never interpolate case IDs as query operators, identify Acts by substring, or treat scoping as exact source resolution.',
     'Test ordinary and hostile IDs, positive and empty expansion, 200-candidate discovery and the separate 100-hit public limit.',
     'BK-25/BK-38. Agentified NM provider scope uses substrings and must not be copied.', 'implemented'),
    ('LB-102', 'Search outcomes and limits',
     'Know whether NM searched, failed to search, or searched without finding a candidate.',
     'Authority query, missing/corrupt index, query-term limit, paragraph limit or exhausted turn budget.',
     'Use distinct outcomes for searched-no-match and unavailable work. State the actual index, query terms, examined population and bounds. Keep search reporting separate from legal assumptions and decisions.',
     'A lexical miss does not announce that relevant law is absent from the corpus.',
     'Retain the limit even if every examined candidate is rejected. A budget stop performs no additional search.',
     'Never promote a search note into a decision attributed to the advocate. Never use an old hard-coded corpus population as a present measurement.',
     'Test zero hits, all-rejected bounded hits, unsearchable input, corrupt identity, exhausted budgets and no false decision persistence.',
     'BK-38/BK-91. Informed by Agentified NM missing-stage and exclusion reporting.', 'implemented'),
    ('LB-103', 'Retrieved candidate versus assessed support',
     'Inspect the source without being told keyword overlap proves my legal proposition.',
     'A ranked judgment reaches the reasoning context or response.',
     'Carry support as assessed true, assessed false or not assessed. Search produces the third state. Keep support, binding, treatment and governing date separate in prompts and disclosures.',
     'Unassessed candidates remain readable with qualifications but cannot become verified premises.',
     'Reject truthy strings and numbers as support verdicts. State each consequential unknown, not just a generic treatment warning.',
     'Do not infer support from rank, an exact quotation, or clean treatment. Do not expose hidden model reasoning.',
     'Supply a clean-treatment keyword match and prove support stays unassessed. Test genuine assessed support, assessed mismatch and source text kept outside system instructions.',
     'BK-84/BK-95. Semantic support assessment and qualified evaluation remain open.', 'implemented'),
    ('LB-104', 'Faithful literal query handling',
     'Retrieve using the words and provision numbers I supplied without hidden query corruption.',
     'Tokenise a question and add an already-verified old/new-code subject correspondence.',
     'Preserve punctuation boundaries and numbers. Reserve verified correspondence terms outside the primary-term cut. Count lexical matches with the index tokenizer, not substring containment.',
     'Stemmed words are not rejected by a different scoring rule; substrings do not manufacture extra matches.',
     'Disclose omitted input terms and bounded paragraph examination. Keep original query and corroborating legal correspondence distinct.',
     'No invented aliases, model-memory laws or unverified code mappings. No claim that lexical coverage measures legal relevance.',
     'Test punctuation, section numbers, long queries, retained correspondence terms, inflections and substring false positives on synthetic FTS indexes.',
     'BK-25/BK-38. Does not establish semantic recall or replace adjudicated retrieval evaluation.', 'implemented'),
    ('LB-105', 'Honest bounded research loop',
     'Have NM investigate proportionately without mistaking failure for completed research.',
     'The model proposes an additional judgment search against the current matter snapshot.',
     'Keep source-anchored proposals, closed actions, stale-snapshot rejection, repeated-query checks and shared round budgets. Pass independent source checks to the planner. Stop with retrieval-unavailable when the search cannot complete.',
     'The recorded stop reason describes what actually happened and does not imply full coverage.',
     'Provider failure, bad proposals, retrieval failure, repeated work and zero progress stay distinct.',
     'Changing a purpose label does not make the same query new work. Sources remain untrusted data and cannot grant new authority.',
     'Test failures, repetitions, budgets, source changes and injected proposals. Live judgment quality needs approved model evaluation.',
     'BK-91. Multi-anchor query planning remains pending; a contiguous focus is still the current admitted query contract.', 'implemented'),
    ('LB-106', 'Measured hybrid retrieval and contextual reading',
     'Find relevant and adverse authorities even when they use different wording, and read them in context.',
     'A separately approved retrieval-quality upgrade after integrity repairs.',
     'Evaluate BM25 plus dense retrieval with reciprocal-rank fusion, optional reranking and version-bound neighbouring passages. Carry rhetorical labels without presenting allegations as holdings. Declare unavailable legs, candidate cuts and exclusion reasons.',
     'Promote only after held-out human-vetted measurements show improved relevant and adverse recall within approved latency and cost.',
     'Keep explicit degraded mode. Reject stale or mismatched embedding model, dimensions, IDs or source version. Do not download or build indexes implicitly.',
     'Do not copy Agentified NM wholesale: its served provider uses dense search, sets supports true and can label a scoped similarity hit resolved. No model-predicted role filter may silently remove contrary context.',
     'Compare lexical-only, dense-only, fusion and reranking. Report recall at stated depth, adverse recall, ranking, full-text fidelity, latency, cost and query strata. Require person-vetted gold rather than self-certified cases.',
     'References: https://arxiv.org/html/2608.06828v1 and https://arxiv.org/html/2602.23371v1. Agentified NM index/search/acceptance implementations inform the design.', 'open'),
    ('LB-107', 'Qualified provision identity at grounding',
     'Know that a cited section belongs to the Act actually retrieved.',
     'Ground an assembled response or retrieve a late citation.',
     'Exclude the question/proposition itself from retrieved citation coverage. Complete typed Act-plus-provision-plus-version identity across generation, late retrieval and the final grounding check.',
     'A same-number provision in a different Act cannot satisfy grounding; a reference mentioned only by the user cannot certify itself.',
     'Unresolved or ambiguous qualified references remain unverified and require precise retrieval or clarification.',
     'No fuzzy Act identity, bare-number fallback, source swapping or treating a citation inside another judgment as retrieval of the whole cited provision.',
     'Current negative control rejects a provision appearing only in the question. Pending: cross-Act same-number, schedule-versus-section, aliases and historical-version binding at served output.',
     'BK-84/BK-95. Complete qualified-identity redesign is pending; existing numeric coverage is not sufficient proof.', 'partial'),
    ('LB-108', 'RAG validation and accountable claims',
     'See what retrieval behaviour has actually been proved and what remains open.',
     'Every retrieval change and release review.',
     'Link diagnostic requirements to counted tests, current source, corpus identity and research/browser results. Separate deterministic integrity, retrieval recall, model judgment and qualified professional acceptance.',
     'Recorded results identify their population and limitations; green targeted tests do not become a whole-product assurance claim.',
     'Keep failed or unrun checks visible. Stop repeating browser failures caused by a stale shared journey helper and repair the helper without weakening assertions.',
     'No paid evaluation without bounded approval, no full-gate claim from a subset, no changes to the reference project and no implied commit or push.',
     'Negative controls, served research API and browser flows, bounded real-corpus checks, lint and architecture checks. Record unrelated failures and the exact tested scope.',
     'BK-25/BK-38/BK-84/BK-91/BK-95. Full legal-quality review remains separate.', 'recorded'),
]


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    book = load_workbook(SOURCE)
    before = {(s.title, c.coordinate): (c.value, copy.copy(c._style))
              for s in book for cells in s for c in cells}
    features = {s.title: sheet_features(s) for s in book}
    sheet = book['Before Build']
    original_rows = sheet.max_row
    changed, row_numbers = set(), []
    for ident, title, *values in ROWS:
        state = values.pop()
        existing = [r for r in range(5, sheet.max_row + 1)
                    if str(sheet.cell(r, 1).value or '').startswith(ident + '\n')]
        row_number = existing[0] if existing else sheet.max_row + 1
        row_numbers.append(row_number)
        row_values = [f'{ident}\nRetrieval diagnostic\n{title}', *values,
                      f'22 September 2026: {state}. Evidence: retrieval-hardening.md '
                      'in the backlog evidence folder for 22 September 2026. '
                      'No whole-RAG or release sign-off.']
        assert len(row_values) == 10
        for col, value in enumerate(row_values, 1):
            cell = sheet.cell(row_number, col)
            if not existing:
                cell._style = copy.copy(sheet.cell(original_rows, col)._style)
            cell.value = value
            changed.add((sheet.title, cell.coordinate))
        sheet.row_dimensions[row_number].height = 240
    for table in sheet.tables.values():
        if table.ref.startswith('A4:J'):
            table.ref = f'A4:J{sheet.max_row}'
            if table.autoFilter:
                table.autoFilter.ref = table.ref
    if sheet.auto_filter.ref:
        sheet.auto_filter.ref = f'A4:J{sheet.max_row}'
    target = OUT / SOURCE.name
    book.save(target)
    check = load_workbook(target)
    assert check.sheetnames == book.sheetnames
    for tab in check:
        assert sheet_features(tab) == features[tab.title]
        assert tab.auto_filter == book[tab.title].auto_filter
        for name in tab.tables:
            assert tab.tables[name] == book[tab.title].tables[name]
    for key, value in before.items():
        if key not in changed:
            cell = check[key[0]][key[1]]
            assert (cell.value, cell._style) == value, key
    for row_number in row_numbers:
        assert all(check['Before Build'].cell(row_number, c).value for c in range(1, 11))
    preview(check['Before Build'], row_numbers[:2], [1, 4, 10], OUT / 'retrieval-after.png')
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == digest, 'parallel workbook edit'
    shutil.copy2(target, SOURCE)
    print(json.dumps({'rows': row_numbers, 'requirements': len(ROWS),
                      'preserved_cells': len(before) - len(changed & before.keys()),
                      'unrelated_cell_or_style_changes': 0,
                      'native_sheet_features_preserved': True}))


if __name__ == '__main__':
    main()
