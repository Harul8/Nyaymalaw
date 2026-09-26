"""Add the harness, context and loop rows that make the legal brain stronger.

OWNER DIRECTION, 26 September 2026, given in conversation after reviewing
what loops, tools and context Nyaymalaw needs: add these rows.

  LB-140  L.2  the opposing-counsel loop
  LB-141  L.4  an independent verifier for meaning
  LB-142  L.5  the matter is written only through checked tools
  LB-143  L.4  checks before and after every tool call, and on the step log
  LB-144  L.3  every number and date comes from a tool
  LB-145  L.5  context built fresh from the file, tagged with source and status
  LB-146  L.8  record, replay and compare

Each lands inside its group, directly after the LB-126..139 rows already
there, and is mirrored into the Implementation Plan sheet. Every existing row
keeps all ten columns exactly; only the groups receiving rows are re-shaded,
and the L.2 and L.4 headers gain one sentence naming what they now hold.

The proof is the one the regroup tool used, on the saved file, before the
source is replaced.
"""
import copy
import hashlib
import importlib.util
import re
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

_REPO = Path(__file__).resolve().parents[2]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))
from assurance.common._console import utf8_console  # noqa: E402

utf8_console()

# ONE OWNER of the row shape, the readiness wording and the checks: the regroup
# tool. Loaded by path because one_off_tools is not a package.
_spec = importlib.util.spec_from_file_location(
    "regroup", Path(__file__).with_name("legal_brain_regroup_20260926.py"))
regroup = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(regroup)

SOURCE = regroup.SOURCE
TODAY = regroup.TODAY
FIRST = regroup.FIRST
HEADER = re.compile(r"^(L|L\.\d|U|U\.\d|P|X)  ")

READINESS = regroup.READINESS

NEW_ROWS = {
    "LB-140": ("L.2", [
        "Opposing counsel: a second loop tries to defeat the draft",
        "Before a material answer is served, a separate model instance, seeing "
        "only the draft's claims and their sources, argues the other side.",
        "Receive advice that has already survived the strongest attack the file "
        "and the law allow.",
        "The main loop has a draft that advises, recommends a step or assesses "
        "prospects.",
        "A nested loop with a fresh context and an adversarial brief: find "
        "contrary or distinguishing authority, weak or disputed facts the draft "
        "relies on, procedural bars (limitation, notice, forum, jurisdiction) "
        "and other readings of the provisions. It has the same retrieval tools "
        "and returns objections, each with its span and locator. The main loop "
        "must answer each material objection -- by revising the draft, by "
        "showing with its own source why the objection fails, or by disclosing "
        "it to the advocate as a risk. Objections and how each was dealt with "
        "are in the step log and open when the advocate inspects the basis. "
        "This subsumes the fragments that exist today: the adverse, exposure "
        "and attacks reads.",
        "Every material objection is resolved, refuted with a source, or "
        "disclosed.",
        "If the opposing loop cannot run, the answer is served with 'not tested "
        "against the other side' disclosed, never as though it had been.",
        "Never let the opposing loop see the main loop's reasoning -- only its "
        "claims and sources -- so it attacks the answer rather than agreeing "
        "with it. Never drop an objection silently. Never let an objection "
        "reach the advocate without its own support.",
        "LB-140-AC1 (planted): a draft relying on an authority a larger bench "
        "has overruled meets that objection and is revised.\n"
        "LB-140-AC2: a golden matter with a limitation weakness has it raised "
        "by the opposing loop when the draft omits it.\n"
        "LB-140-AC3 (planted): an objection without a span is refused by the "
        "harness.",
        "G-ADVERSE, G-EXPOSURE; the adverse, exposure and attacks reads; LB-18, "
        "LB-23, LB-65, LB-115; LB-122. OPEN: which answers count as material, "
        "and the loop's budget.",
    ]),
    "LB-141": ("L.4", [
        "An independent verifier: does the passage support the claim?",
        "Deterministic checks prove a quotation exists; an independent verifier "
        "decides whether the passage actually supports the claim made from it.",
        "Rely on every cited source saying what NM says it says.",
        "Every claim of law, authority or document fact in a submitted answer.",
        "A separate model instance, on the cheaper tier, sees one claim and its "
        "cited passage and nothing else, and answers supports, does not "
        "support, or supports in part, quoting the words it relies on. It never "
        "sees the conversation and is never the instance that wrote the claim. "
        "Its verdict is one of the harness's checks, and 'does not support' "
        "goes to repair (LB-133). The deterministic checks run first -- quote "
        "present, citation resolves, in force, binding -- so the verifier "
        "spends its effort only on meaning.",
        "Every served claim has a 'supports' verdict, or its limitation is "
        "disclosed.",
        "A verifier that cannot run leaves the claim 'support not verified', "
        "disclosed, never counted as verified.",
        "Never let the author grade its own claim. Never give the verifier the "
        "conversation, which would let it be persuaded rather than read. Never "
        "treat keyword overlap as support (LB-103).",
        "LB-141-AC1 (planted): a real passage cited for a proposition it does "
        "not state is marked 'does not support' and repaired.\n"
        "LB-141-AC2: on a set of claim-passage pairs labelled by the owner, the "
        "verifier's agreement is measured and reported, never assumed.\n"
        "LB-141-AC3: a verifier that cannot run yields a disclosed 'not "
        "verified'.",
        "G-GROUND, G-ATTRIB; LB-27, LB-59, LB-103, LB-130. OPEN: the labelled "
        "pair set and the agreement threshold.",
    ]),
    "LB-142": ("L.5", [
        "The matter is written only through checked tools",
        "The model keeps its memory by recording facts, premises, issues and "
        "decisions through tools the harness checks, never as free text.",
        "Know that everything NM remembers about my matter came from me, a "
        "document, or a source it can show.",
        "The model establishes, corrects or withdraws anything on the matter "
        "file.",
        "Write tools: record a fact, which must quote the advocate's words or a "
        "document span; record a premise, stated with its quote or inferred "
        "with what it was inferred from; record an issue; record a decision; "
        "correct or withdraw an entry. Each write is checked before it lands: "
        "the quote must exist where it says, a correction replaces rather than "
        "adds (G-CORRECTION), and a changed value marks everything resting on "
        "it for review through the dependency ledger (G-CASCADE). The file is "
        "the memory the next turn is built from (LB-145).",
        "Every entry on the file carries its source and status.",
        "A write that fails its check returns to the model with the reason; "
        "nothing half-written lands.",
        "Never let the model write to the file without a source. Never record "
        "an inference as a stated fact. Never lose the earlier value when a "
        "correction replaces it.",
        "LB-142-AC1 (planted): a fact recorded with a quotation the advocate "
        "never wrote is refused.\n"
        "LB-142-AC2: a corrected date replaces the old one, and every deadline "
        "resting on it is marked for review.\n"
        "LB-142-AC3: an inferred premise is stored as inferred and names its "
        "basis.",
        "Matter, Thread and Fact in backend/nm/domain; the dependency ledger; "
        "LB-07, LB-08, LB-29, LB-49, LB-58.",
    ]),
    "LB-143": ("L.4", [
        "Checks before and after every tool call, and on the step log",
        "The harness checks each tool call before it runs and each result after "
        "it returns, and reads the step log as evidence of how the answer was "
        "reached.",
        "Know that NM cannot act outside my matter and permissions, and that "
        "what it relied on actually came back.",
        "Every tool call in the loop; every submitted answer.",
        "Before a tool runs: is it declared, within the matter's scope and the "
        "advocate's authorisation, clear of conflict and the professional "
        "boundaries (LB-132), and within budget? A refused call returns to the "
        "model with the reason. After a tool returns: does the result carry its "
        "source, its locator and one of the three states? A malformed or empty "
        "result cannot pass as 'nothing found'. Before the answer: the step log "
        "is read as process evidence -- the provision cited was retrieved on "
        "this matter, contrary authority was searched for before advice on "
        "prospects, a figure came from a computation -- without dictating when "
        "in the loop the model did any of it.",
        "No tool acts outside its contract, and no answer rests on a result "
        "that was not returned.",
        "A check that cannot run refuses the call or the answer, and says so.",
        "Never let a tool call execute before its check. Never read a failed or "
        "empty result as a finding. Never impose an order of work through a "
        "process check.",
        "LB-143-AC1 (planted): a tool call reaching another advocate's matter "
        "is refused before it runs.\n"
        "LB-143-AC2 (planted): a result with no locator is refused before the "
        "model sees it.\n"
        "LB-143-AC3 (planted): an answer citing a provision its step log never "
        "retrieved is refused.",
        "LB-62, LB-63, LB-102, LB-128, LB-132; the professional access checks "
        "in backend/nm/core.",
    ]),
    "LB-144": ("L.3", [
        "Every number and date comes from a tool",
        "No figure in an answer -- a date, a period, an amount, a count -- is the "
        "model's own arithmetic; each comes from a computation tool, and the "
        "harness checks it.",
        "Rely on every date and figure without re-computing it myself.",
        "An answer contains a date, a period, an amount or a computed count.",
        "Computation tools -- limitation, date arithmetic, deadlines, interest, "
        "and the court fee once its schedule is held (LB-125) -- return the "
        "figure, the inputs used and the rule applied. The answer's claim for "
        "the figure cites that tool result. The harness matches every number "
        "and date in the answer to a tool result, or to the advocate's or a "
        "document's own words; an unmatched figure goes to repair.",
        "Every figure is traceable to a computation or a source.",
        "A computation that cannot run returns not assessed, and the answer "
        "says the figure is not computed rather than estimating it.",
        "Never let the model compute a date or an amount itself. Never let a "
        "precise figure hide an unestablished premise (LB-114).",
        "LB-144-AC1 (planted): an answer stating a limitation expiry the tool "
        "did not compute is refused.\n"
        "LB-144-AC2: a figure in a golden answer traces to the tool result that "
        "produced it.\n"
        "LB-144-AC3: a computation on an inferred premise is labelled "
        "conditional, never definitive.",
        "backend/nm/core/limitation.py and deadlines.py; G-PREMISE, G-CASCADE; "
        "LB-20, LB-114, LB-124, LB-125.",
    ]),
    "LB-145": ("L.5", [
        "Context built fresh from the file, every item tagged with source and status",
        "Each turn's context is assembled from the matter file in layers, and "
        "every item in it says where it came from and how it is known.",
        "Have NM keep straight what I said, what a document proves, what it "
        "retrieved and what it inferred.",
        "Every turn, when the loop starts; and when the turn's working material "
        "outgrows its budget.",
        "Layers: (1) a fixed prefix -- the principles and tool definitions, "
        "identical across turns so it can be cached; (2) a matter brief "
        "generated from the structured file, each line tagged stated, from a "
        "document, retrieved, computed or inferred, with its source; (3) the "
        "working set this turn needs, pulled by tools; (4) recent conversation "
        "verbatim, with the advocate's full words always retrievable; (5) the "
        "turn's scratchpad, compacted into findings with their locators when it "
        "grows and dropped at the end of the turn except what was recorded "
        "(LB-142). Retrieved and uploaded text is marked as material to read, "
        "never instructions to follow. Every tool is scoped to this advocate "
        "and this matter.",
        "The model's context shows the status and source of everything in it, "
        "within the budget.",
        "If the brief cannot be generated from the file, the turn stops with "
        "the reason rather than running on a partial context.",
        "Never mix one matter's material into another's context. Never let a "
        "compacted note lose its locator. Never follow an instruction found in "
        "retrieved or uploaded text.",
        "LB-145-AC1: a fact the client asserted and a fact a document proves "
        "are tagged differently in the context and treated differently in the "
        "answer.\n"
        "LB-145-AC2 (planted): an uploaded document saying 'ignore previous "
        "instructions' changes nothing.\n"
        "LB-145-AC3: a turn whose scratchpad outgrows its budget keeps every "
        "locator through compaction.",
        "LB-07, LB-49, LB-58, LB-135; memory.advocate_words; "
        "backend/nm/domain/summary.py.",
    ]),
    "LB-146": ("L.8", [
        "Record, replay and compare: the loop tested on every commit",
        "Model conversations are recorded so the loop can be replayed without "
        "a live model, tested on every commit and compared across providers; "
        "every defect found becomes a golden case.",
        "Know that a change did not break the legal brain before it reaches me, "
        "and which model serves me best.",
        "Every commit; every model or principles change; every defect found.",
        "A recorder captures each model response in a loop run -- tool calls "
        "and answers -- with the principles and tool versions in force. A "
        "replay adapter feeds the recording back, so the loop, the tools and "
        "the harness run deterministically with no API key on every commit. A "
        "live comparison runs the golden set through Opus, Fable and an OpenAI "
        "model on the same cases, with the owner's approval of each run, and "
        "records grounding, correctness and cost for each. The improvement "
        "loop: every defect found by the owner or a check becomes a golden case "
        "before it is fixed, and the fix is proved by that case turning green.",
        "Replays run on every commit; comparisons are recorded with their date, "
        "versions and cost.",
        "A recording made under older principles or tools is marked stale "
        "rather than replayed as current.",
        "Never run a live golden or comparison run without the owner's approval "
        "of that run. Never treat a replay as evidence of live model quality: "
        "it proves the harness and tools, not the model.",
        "LB-146-AC1: a recorded golden conversation replays through the loop "
        "with no network access and reaches the same answer.\n"
        "LB-146-AC2 (planted): a harness change that lets an unsupported claim "
        "through turns a replay red.\n"
        "LB-146-AC3: a comparison report names each model's grounding failures "
        "and cost on the same cases.",
        "docs/GOLDEN_SET.md; assurance/journeys/run_goldens.py and judge.py; the "
        "defect register; LB-40, LB-72, LB-138.",
    ]),
}

#: One sentence added to two headers, so each says what it now holds.
HEADER_ADDITIONS = {
    "L.2": " An opposing-counsel loop attacks every material draft before it is served (LB-140).",
    "L.4": (" An independent verifier checks that each cited passage supports its claim (LB-141), and "
            "every tool call is checked before it runs and after it returns (LB-143)."),
}

NEW_NOTE_RANGE = ("LB-126–139", "LB-126–146")


def number(key: str) -> int:
    match = re.match(r"LB-(\d+)$", key)
    return int(match[1]) if match else -1


def main() -> int:
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    book = load_workbook(SOURCE)
    sheet, plan = book["Before Build"], book["Implementation Plan"]
    features = {s.title: regroup.sheet_features(s) for s in book}
    before = {(s.title, c.coordinate): (c.value, copy.copy(c._style))
              for s in book for row in s for c in row}
    heights = {r: sheet.row_dimensions[r].height for r in range(1, sheet.max_row + 1)}

    # ---- the current body, in order, header rows included --------------------
    old_last = sheet.max_row
    order: list[tuple[str, dict]] = []
    for r in range(FIRST, old_last + 1):
        text = str(sheet.cell(r, 1).value or "")
        header = HEADER.match(text)
        key = f"HEADER:{header[1]}" if header else regroup.key_of(text)
        order.append((key, {
            "values": [sheet.cell(r, c).value for c in range(1, 11)],
            "styles": [copy.copy(sheet.cell(r, c)._style) for c in range(1, 11)],
            "height": sheet.row_dimensions[r].height, "row": r}))
    keys = [k for k, _ in order]
    assert len(keys) == len(set(keys)), "duplicate key in the current body"
    for ident in NEW_ROWS:
        assert ident not in keys, f"{ident} already exists"

    template = dict(order)["LB-139"]
    group_title = {}
    for key, row in order:
        if key.startswith("HEADER:"):
            first_line = str(row["values"][0]).split("\n")[0]
            group_title[key.split(":", 1)[1]] = first_line.split("  ", 1)[1].split(regroup.DASH)[0]

    # ---- insert each new row after the last LB-126+ row in its group ----------
    new_by_group: dict[str, list[str]] = {}
    for ident, (group, _cols) in NEW_ROWS.items():
        new_by_group.setdefault(group, []).append(ident)

    body: list[tuple[str, dict]] = []
    current = None
    pending: list[tuple[str, dict]] = []

    def flush():
        # Insert this group's new rows after its last owner-directed row.
        nonlocal pending
        if current in new_by_group:
            anchor = max((i for i, (k, _) in enumerate(pending) if number(k) >= 126), default=-1)
            additions = []
            for ident in new_by_group[current]:
                group, cols = NEW_ROWS[ident]
                title, lead, *rest = cols
                area = "Legal-brain interface" if group.startswith("U") else "Legal brain"
                section = f"{area} › {group} {group_title[group]}"
                column_a = (f"{ident}\n{section}\n{title}\n\n"
                            f"OWNER-DIRECTED REQUIREMENT, DRAFTED FOR REVIEW:\n{lead}")
                values = [column_a, *rest, READINESS]
                assert len(values) == 10, ident
                additions.append((ident, {"values": values,
                                          "styles": [copy.copy(s) for s in template["styles"]],
                                          "height": 280, "row": None}))
            pending = pending[:anchor + 1] + additions + pending[anchor + 1:]
        body.extend(pending)
        pending = []

    for key, row in order:
        if key.startswith("HEADER:"):
            flush()
            current = key.split(":", 1)[1]
            body.append((key, row))
        else:
            pending.append((key, row))
    flush()
    assert len(body) == len(order) + len(NEW_ROWS)

    # ---- write it back ---------------------------------------------------------
    changed: set = set()
    touched_groups = set(new_by_group)
    layout: list[tuple[str, int]] = []
    index_in_group = 0
    group = None
    for offset, (key, row) in enumerate(body):
        r = FIRST + offset
        is_header = key.startswith("HEADER:")
        if is_header:
            group = key.split(":", 1)[1]
            index_in_group = 0
        values = list(row["values"])
        if is_header and group in HEADER_ADDITIONS:
            values[0] = values[0] + HEADER_ADDITIONS[group]
        for c in range(1, 11):
            cell = sheet.cell(r, c)
            if (cell.value, cell._style) != (values[c - 1], row["styles"][c - 1]):
                changed.add(("Before Build", cell.coordinate))
            cell.value = values[c - 1]
            cell._style = copy.copy(row["styles"][c - 1])
            if not is_header and group in touched_groups:
                cell.fill = PatternFill("solid", fgColor=regroup.ZEBRA[index_in_group % 2])
                changed.add(("Before Build", cell.coordinate))
        if not is_header:
            index_in_group += 1
        sheet.row_dimensions[r].height = row["height"]
        layout.append((key, r))
    last = FIRST + len(body) - 1
    sheet.auto_filter.ref = f"A4:J{last}"

    # ---- the row-2 note: range and counts, measured ---------------------------
    population = sorted({regroup.LBOM.match(str(sheet.cell(x, 1).value or ""))[1]
                         for x in range(5, last + 1)
                         if regroup.LBOM.match(str(sheet.cell(x, 1).value or ""))})
    lb = [p for p in population if p.startswith("LB-")]
    note = str(sheet.cell(2, 1).value)
    assert NEW_NOTE_RANGE[0] in note, "the row-2 note no longer names the owner-directed range"
    note = note.replace(*NEW_NOTE_RANGE)
    note, n = re.subn(r"The legal brain holds \d+ LB requirements",
                      f"The legal brain holds {len(lb)} LB requirements", note)
    assert n == 1
    note, n = regroup.NOTE_PATTERN.subn(
        f"All {len(population)} LB/OM requirements are linked into Implementation Plan", note)
    assert n == 1
    note = note.replace("autonomous loop, the harness and context management",
                        "autonomous loop, the harness, the verifier, the opposing-counsel loop and "
                        "context management")
    sheet.cell(2, 1).value = note
    changed.add(("Before Build", "A2"))

    # ---- the mirror ------------------------------------------------------------
    position = dict(layout)
    for ident, (group, cols) in NEW_ROWS.items():
        assert not any(plan.cell(x, 1).value == ident for x in range(2, plan.max_row + 1)), ident
        target = plan.max_row + 1
        for c in range(1, plan.max_column + 1):
            plan.cell(target, c)._style = copy.copy(plan.cell(target - 1, c)._style)
        src = position[ident]
        fixed = {1: ident, 3: "Requirement", 4: "Legal brain", 5: f"{group} {group_title[group]}",
                 6: cols[0], 7: "Pilot", 38: "Draft", 39: "Not started", 40: "Not verified",
                 32: f"{TODAY}: owner-directed harness, context and loop row mirrored from Before Build. "
                     "Not acceptance, not an approved requirement."}
        for c, value in fixed.items():
            plan.cell(target, c, value)
            changed.add(("Implementation Plan", plan.cell(target, c).coordinate))
        for t, origin in regroup.MIRROR.items():
            plan.cell(target, t, sheet.cell(src, origin).value)
            changed.add(("Implementation Plan", plan.cell(target, t).coordinate))

    scratch = Path("/tmp/claude-0/-home-user-Nyaymalaw/82a504d3-1374-501f-9fe3-441ac9aa966e/scratchpad")
    out = scratch / SOURCE.name
    book.save(out)

    # ================= THE PROOF, on the saved bytes ==========================
    from openpyxl.xml.functions import tostring

    def xml(part):
        return tostring(part.to_tree())

    check = load_workbook(out)
    for tab in check:
        assert regroup.sheet_features(tab) == features[tab.title], f"sheet features moved on {tab.title}"
    bb = check["Before Build"]
    original = load_workbook(SOURCE)["Before Build"]

    for (title, coord), old in before.items():
        if title == "Before Build" and int(re.sub(r"[A-Z]+", "", coord)) < FIRST and coord != "A2":
            assert (bb[coord].value, bb[coord]._style) == old, coord
    for x in range(1, FIRST):
        assert bb.row_dimensions[x].height == heights.get(x), f"row {x} height moved"

    moved = dict(order)
    for key, r in layout:
        if key in NEW_ROWS:
            continue
        want = moved[key]
        got = [bb.cell(r, c).value for c in range(1, 11)]
        expected = list(want["values"])
        header_group = key.split(":", 1)[1] if key.startswith("HEADER:") else None
        if header_group in HEADER_ADDITIONS:
            expected[0] += HEADER_ADDITIONS[header_group]
        assert got == expected, f"{key}: content changed"
        for c in range(1, 11):
            a, o = bb.cell(r, c), original.cell(want["row"], c)
            assert xml(a.border) == xml(o.border), f"{key}: border moved"
            assert xml(a.alignment) == xml(o.alignment), f"{key}: alignment moved"
            assert (a.font.name, a.font.sz, a.font.color, a.font.b) == \
                (o.font.name, o.font.sz, o.font.color, o.font.b), f"{key}: font moved"

    after = [k for k, _ in layout]
    assert sorted(after) == sorted(keys + list(NEW_ROWS)), "row population changed"
    for ident, (group, _c) in NEW_ROWS.items():
        header_row = position[f"HEADER:{group}"]
        following = [r for k, r in layout if k.startswith("HEADER:") and r > header_row]
        assert header_row < position[ident] < (min(following) if following else last + 1), \
            f"{ident} landed outside {group}"

    ip = check["Implementation Plan"]
    for (title, coord), old in before.items():
        if title == "Implementation Plan" and (title, coord) not in changed:
            assert (ip[coord].value, ip[coord]._style) == old, f"plan {coord} moved"

    from assurance.control_plane import plan_scenarios
    problems = plan_scenarios.requirement_problems(out)
    assert not problems, problems
    assert len(plan_scenarios.sheet_rows(out)) == 93, "the feature population moved"

    shutil.copy2(out, SOURCE)
    print(f"source digest before: {digest}")
    print(f"source digest after:  {hashlib.sha256(SOURCE.read_bytes()).hexdigest()}")
    print(f"rows added: {len(NEW_ROWS)}; last row: {last}; "
          f"LB/OM requirements: {len(population)} (LB {len(lb)})")
    for ident, (group, cols) in NEW_ROWS.items():
        print(f"  {ident}  row {position[ident]:>3}  {group:<4} {cols[0]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
