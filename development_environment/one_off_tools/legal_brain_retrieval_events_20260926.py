"""Second cross-check: the statute and case-law tools report what they considered, kept and set aside.

OWNER DIRECTION, 26 September 2026: beyond the Claude-derived principles, check
that everything discussed and settled about how NM reads the matter, decomposes
disputes, retrieves, asks follow-up questions, and what the matter board and
scratch pad show, is in the plan.

The session check was extended by 25 items on exactly those topics
(`session_decisions_check_20260926.py`, 85 items in all). 83 were present. The
two gaps were one point: the scratch pad shows the sections considered, kept
and set aside, and the case-law passages found relevant -- but the statute and
case-law tool rows did not say they report their candidates and the model's
decision on each. This adds that to LB-156 and LB-157, and mirrors it.

Only column 4 of two owner-directed draft rows changes, by exact single-match
replacement; the proof checks every other cell in the workbook is unchanged.
"""
import copy
import hashlib
import importlib.util
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

_REPO = Path(__file__).resolve().parents[2]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))
from assurance.common._console import utf8_console  # noqa: E402

utf8_console()

_spec = importlib.util.spec_from_file_location(
    "regroup", Path(__file__).with_name("legal_brain_regroup_20260926.py"))
regroup = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(regroup)
SOURCE = regroup.SOURCE

EDITS = [
    ("LB-156", 4,
     "All wrap existing code: the manifest's exact resolution, the evidence port's fetch, and the governing-law table.",
     "All wrap existing code: the manifest's exact resolution, the evidence port's fetch, and the governing-law table. "
     "The candidate sections each call returns, and the model's decision on each -- kept for this dispute, or set aside "
     "with the reason -- are recorded as step events (LB-164), so the scratch pad shows the sections considered, those "
     "kept and those set aside, and the answer shows only the sections finalised (LB-165)."),
    ("LB-157", 4,
     "the citator, the jurisdiction rule and the identity index.",
     "the citator, the jurisdiction rule and the identity index. The judgments each call returns, and the passages the "
     "model finds relevant to the dispute -- each with the reason, or set aside with the reason -- are recorded as step "
     "events (LB-164), so the scratch pad shows what was retrieved and what was found relevant, and the answer shows "
     "only the relevant passages (LB-165)."),
]


def main() -> int:
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    book = load_workbook(SOURCE)
    sheet, plan = book["Before Build"], book["Implementation Plan"]
    features = {s.title: regroup.sheet_features(s) for s in book}
    before = {(s.title, c.coordinate): (c.value, copy.copy(c._style)) for s in book for row in s for c in row}

    changed = set()
    rows = {}
    for r in range(regroup.FIRST, sheet.max_row + 1):
        head = str(sheet.cell(r, 1).value or "").split("\n")[0]
        rows[head] = r
    for ident, col, old, new in EDITS:
        cell = sheet.cell(rows[ident], col)
        assert cell.value.count(old) == 1, (ident, old[:40])
        cell.value = cell.value.replace(old, new)
        changed.add(("Before Build", cell.coordinate))
        x = next(x for x in range(2, plan.max_row + 1) if plan.cell(x, 1).value == ident)
        for t, origin in regroup.MIRROR.items():
            if origin == col:
                plan.cell(x, t).value = cell.value
                changed.add(("Implementation Plan", plan.cell(x, t).coordinate))

    out = Path("/tmp/claude-0/-home-user-Nyaymalaw/82a504d3-1374-501f-9fe3-441ac9aa966e/scratchpad") / SOURCE.name
    book.save(out)
    check = load_workbook(out)
    for tab in check:
        assert regroup.sheet_features(tab) == features[tab.title], tab.title
    for key, old in before.items():
        if key not in changed:
            cell = check[key[0]][key[1]]
            assert (cell.value, cell._style) == old, key
    from assurance.control_plane import plan_scenarios
    assert plan_scenarios.requirement_problems(out) == []
    assert len(plan_scenarios.sheet_rows(out)) == 93
    shutil.copy2(out, SOURCE)
    print(f"source digest before: {digest}")
    print(f"source digest after:  {hashlib.sha256(SOURCE.read_bytes()).hexdigest()}")
    print(f"cells changed: {len(changed)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
