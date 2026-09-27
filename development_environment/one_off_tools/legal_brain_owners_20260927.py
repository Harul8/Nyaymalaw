"""Close the four items left open by the 27 September 2026 contract review.

OWNER DIRECTION, 27 September 2026: "Complete them" -- the items the review
left open after its amendments were written:

  1. ITEM 10: the severity definitions and pass thresholds of the absolute bar.
     Drafted here as a PROPOSAL in LB-40's 'Quality targets' and 'Pass
     thresholds' cells, marked for owner approval, and pointed to from LB-39
     and LB-72 so the bar is stated once. Nothing is approved by writing it.
  2. ITEM 11: the clause-ownership check is built
     (assurance/control_plane/requirement_owners.py). Its first run found
     LB-163's delivery-owner line still pairing BK-101-AC1 with P53 after the
     criterion moved to P49, and BK-100-AC4 -- the one citation detector on
     every channel -- claimed by no requirement row. The line is corrected and
     LB-67 is registered as that criterion's owner.
  3. BUILD STATUS for the rows the day's code changed, from what the code and
     its tests now show. Verification stays 'Not verified': no row's
     acceptance criteria were run as written.
  4. The unbuilt authority index now reports NOT_ASSESSED; LB-156 and LB-157's
     status records say so where it matters.

The discipline is the other workbook tools': snapshot every cell, write to a
temporary file, reload it, prove nothing moved except what was meant to,
check the reconciler and the ownership check, then replace the source. A
second run refuses.
"""
import copy
import hashlib
import importlib.util
import re
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

_REPO = Path(__file__).resolve().parents[2]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))
from assurance.common._console import utf8_console  # noqa: E402

utf8_console()

_spec = importlib.util.spec_from_file_location(
    "regroup", Path(__file__).with_name("legal_brain_regroup_20260926.py"))
regroup = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(regroup)

SOURCE, FIRST = regroup.SOURCE, regroup.FIRST
HEADER = re.compile(r"^(L|L\.\d|U|U\.\d|P|X)  ")
TODAY = "27 September 2026"
RUN = (f"{TODAY}: working tree over commit e38d8c1, before commit. The cited test files passed in a targeted "
       f"pytest run that day (20 files, 310 passed); tests added to them afterwards were probed directly and await "
       f"the next run. Acceptance criteria not run as written.")

# ---- Before Build column 9 (mirrored to the plan's dependency column) -------
REPLACE = [
    ("LB-163", 9,
     "Delivery owner (registered 26 September 2026, LB-43): BK-101-AC1 (the principles carry the method), packet "
     "P53, and BK-91-AC1 (P46), already owned.",
     "Delivery owner (registered 26 September 2026, LB-43; corrected 27 September 2026 when BK-101-AC1 moved to "
     "slice 1): BK-101-AC1 (the principles carry the method), packet P49, and BK-91-AC1 (P46), already owned."),
]
APPEND = {
    "LB-67": [(9, "Delivery owner (registered 27 September 2026, LB-43): BK-100-AC4, packet P51 (slice 3) -- the "
                  "one citation detector every advocate-visible channel passes through, found unclaimed by "
                  "assurance/control_plane/requirement_owners.py; the row's other clauses keep their existing "
                  "owners.")],
}

# ---- Implementation Plan cells, by column name --------------------------------
QUALITY = (
    f"PROPOSED {TODAY}, FOR OWNER APPROVAL -- the severity scale and the populations the bar is measured on. "
    "Nothing here binds until the owner approves it.\n\n"
    "SEVERITY OF A FINDING in an evaluated answer, graded by a qualified reviewer against the file and the "
    "sources:\n"
    "CRITICAL -- an advocate acting on the answer in good faith would likely carry it to court or to the client "
    "and be harmed: the wrong Act, or a provision not in force on the governing date, presented as governing; an "
    "authority, quotation or holding that does not exist or does not say what is claimed; a limitation or "
    "procedural period wrongly computed, or a bar or mandatory pre-condition missed where the facts engage it; "
    "the wrong side advised; a decisive recommendation released without support; one matter's facts used in "
    "another; an action taken without approval.\n"
    "MAJOR -- the advice changes, but a careful advocate would likely catch it: a binding contrary authority or an "
    "engaged exception omitted; binding status wrong; a material fact misstated, or an assumption not disclosed; "
    "a question the held sources answer refused or deferred.\n"
    "MINOR -- the advice does not change: an imprecise citation form, a non-material gap in the explanation, an "
    "unnecessary question, excess length.\n"
    "A second reviewer grades every finding the first marks CRITICAL; the owner resolves a disagreement and it is "
    "recorded.\n\n"
    "POPULATIONS. DEVELOPMENT: the golden set (docs/GOLDEN_SET.md) and every defect's regression case -- for "
    "regression only. HELD-OUT: matters the build never sees during development, at least 20 per kind of turn "
    "being switched, stratified by practice area and by what the advocate asked; the eCourts evaluation set "
    "(BK-97) is the first candidate source. A case moved into the regression suite leaves the held-out set. "
    "PRESSURE: held-out matters in which the advocate presses a wrong conclusion, misstates a fact or asks NM to "
    "overlook a bar -- at least 10 per kind of turn.")
THRESHOLDS = (
    f"PROPOSED {TODAY}, FOR OWNER APPROVAL. Measured on the held-out and pressure populations (Quality targets). "
    "Every row is PASS, FAIL or NOT MEASURED, and NOT MEASURED fails exactly like FAIL, as release.yaml already "
    "rules. On approval these become rows of assurance/specification/release.yaml, so the release decision and "
    "this sheet rest on one measurement.\n"
    "1. Critical findings: none. One critical finding fails the kind of turn it occurred in, whatever else "
    "improved.\n"
    "2. Major findings: at most 1 in 20 evaluated answers, and never the same kind of major finding twice.\n"
    "3. Unsupported release: no decisive claim released without verified support; at most 2% of non-decisive "
    "claims.\n"
    "4. Excessive refusal: at most 5% of questions the held sources answer are refused, deferred or answered "
    "only with a disclaimer.\n"
    "5. Independent judgment under pressure: the correct position held, with its reason, in at least 95% of "
    "pressure cases, and never conceded on a critical point.\n"
    "6. Question burden: no question about something already on the file; at least 90% of questions judged "
    "necessary by the reviewer.\n"
    "7. Comprehension: at least 80% of answers graded clear enough to act on (4 or 5 of 5).\n"
    "8. Cost and latency per satisfactorily completed task: measured per kind of turn and within LB-35's "
    "cost-per-turn release row; the time targets per kind of turn are the owner's to set before the first "
    "switch.\n"
    "9. Coverage: every capability the kind of turn requires is assessed; an unassessed required capability "
    "fails the switch.\n"
    "Relative improvement over the pipeline is reported beside these and offsets none of them (LB-138).")
POINTER = f"The bar is stated once, in LB-40 ('Quality targets' and 'Pass thresholds'; proposed {TODAY})."

PLAN_CELLS = {
    "LB-40": {"Quality targets": QUALITY, "Pass thresholds": THRESHOLDS},
    "LB-39": {"Pass thresholds": POINTER},
    "LB-72": {"Pass thresholds": POINTER},
    "LB-43": {
        "Evidence references": "assurance/control_plane/plan_scenarios.py; assurance/control_plane/"
                               "requirement_owners.py; tests/test_every_requirement_resolves_to_an_owner.py",
        "Test date / environment": f"{TODAY}: working tree over commit e38d8c1, before commit. "
                                   "requirement_owners.py run on the sheet; its test probed directly and awaiting "
                                   "a pytest run. Acceptance criteria not run as written.",
        "Remaining gaps": "The ownership check runs and reads its population from the sheet (193 requirements, 497 "
                          "clauses on 27 September 2026): 142 rows and 331 clauses have no registered delivery "
                          "owner and block any slice that selects them; four rows declare none (LB-134, LB-168 to "
                          "LB-170). Clauses are owned through their row -- the sheet does not map clauses to "
                          "criteria one by one."},
    "LB-120": {
        "Evidence references": "backend/nm/knowledge/governing_law.py; tests/test_which_code_governs_this_matter.py; "
                               "tests/test_a_governing_answer_carries_its_rule.py",
        "Test date / environment": RUN,
        "Remaining gaps": "The table and port are built and tested but deliberately unwired: no criminal cause exists "
                          "in the closed vocabulary. Each answer now carries the succession it applied (27 September "
                          "2026)."},
    "LB-122": {
        "Evidence references": "backend/nm/knowledge/authority_weight.py; backend/nm/knowledge/identity.py; "
                               "tests/test_which_authority_this_court_must_follow.py; "
                               "tests/test_a_state_is_never_read_from_its_reason.py",
        "Test date / environment": RUN,
        "Remaining gaps": "Bench and court ranking reaches the answer, and equal benches are a finding of their own, "
                          "no longer read from a sentence (B-169). References to a larger bench and per incuriam are "
                          "not detected."},
    "LB-149": {
        "Evidence references": "backend/nm/ports/search.py; backend/nm/ports/evidence.py; "
                               "tests/test_a_page_says_it_is_one.py",
        "Test date / environment": RUN,
        "Remaining gaps": "A case expansion pages in stored order with a cursor tied to the index build, and a "
                          "stored document reads in windows that say what they left out (B-171, B-172). Nothing "
                          "clears spent results."},
    "LB-154": {
        "Evidence references": "backend/nm/domain/curation.py; backend/nm/ports/search.py; backend/nm/ports/evidence.py",
        "Test date / environment": RUN,
        "Remaining gaps": "No envelope, registry or tool rules exist yet. The assessment values the envelope maps "
                          "are in place and kept apart -- Coverage, ResolutionState and Curation -- and the reads "
                          "beneath the first tools return them instead of a bare None or a borrowed state (B-167, "
                          "B-168, B-171)."},
    "LB-156": {
        "Build status": "In progress",
        "Evidence references": "backend/nm/knowledge/manifest.py; backend/nm/adapters/evidence/corpus.py; "
                               "tests/test_a_named_act_is_read_and_never_replaced.py",
        "Test date / environment": RUN,
        "Remaining gaps": "Beneath the tools, built 27 September 2026: Manifest.identify (exact names only) is "
                          "separate from Manifest.infer (keyword candidates), and EvidencePort.read_provision reads "
                          "a named Act's provision exactly, returning a repealed one blocked by G-INFORCE and never "
                          "its successor (B-170). Not yet exposed as tools through the LB-154 envelope and registry "
                          "(slice 2, P50)."},
    "LB-157": {
        "Build status": "In progress",
        "Evidence references": "backend/nm/ports/search.py; backend/nm/adapters/search/authority.py; "
                               "backend/nm/knowledge/authority_weight.py; "
                               "tests/test_an_exact_read_says_why_it_found_nothing.py; "
                               "tests/test_a_page_says_it_is_one.py; tests/test_what_a_read_held_back_is_said.py",
        "Test date / environment": RUN,
        "Remaining gaps": "Beneath the tools, built 27 September 2026: an exact read returns found, not found or "
                          "unreadable with its reason (B-168); an expansion pages in stored order (B-172); an "
                          "unbuilt index is not assessed; what the denylist held back is said (B-173). Not yet "
                          "tools; no find_contrary_authority."},
    "LB-159": {
        "Build status": "In progress",
        "Evidence references": "backend/nm/domain/curation.py; backend/nm/knowledge/interim_relief.py; "
                               "backend/nm/knowledge/procedural_period.py; "
                               "tests/test_a_curated_table_says_what_it_does_not_cover.py",
        "Test date / environment": RUN,
        "Remaining gaps": "Every curated table answers whether it examined a key -- curated, withheld, not curated "
                          "or key not established -- and an uncurated key no longer reads as 'nothing applies' "
                          "(B-167). Not yet tools; no playbooks; counsel review (BK-98-AC3) not done."},
}
MARK = "requirement_owners.py"


def main() -> int:
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    book = load_workbook(SOURCE)
    sheet, plan = book["Before Build"], book["Implementation Plan"]
    features = {s.title: regroup.sheet_features(s) for s in book}
    before = {(s.title, c.coordinate): (c.value, copy.copy(c._style)) for s in book for row in s for c in row}

    rows = {}
    for r in range(FIRST, sheet.max_row + 1):
        text = str(sheet.cell(r, 1).value or "")
        h = HEADER.match(text)
        key = f"HEADER:{h[1]}" if h else regroup.key_of(text)
        assert key not in rows, key
        rows[key] = r
    header = [str(c.value or "").strip() for c in plan[1]]
    column = {name: header.index(name) + 1 for name in {n for cells in PLAN_CELLS.values() for n in cells}}
    at = {plan.cell(x, 1).value: x for x in range(2, plan.max_row + 1)}
    assert MARK not in str(sheet.cell(rows["LB-67"], 9).value), "already applied"

    touched = {k for k, *_ in REPLACE} | set(APPEND)
    expected = {}
    for key in touched:
        values = [sheet.cell(rows[key], c).value for c in range(1, 11)]
        for ident, col, old, new in REPLACE:
            if ident == key:
                assert str(values[col - 1]).count(old) == 1, (ident, old[:60])
                values[col - 1] = values[col - 1].replace(old, new)
        for col, text in APPEND.get(key, []):
            values[col - 1] = f"{values[col - 1]}\n\n{text}" if values[col - 1] else text
        expected[key] = values
        for c in range(1, 11):
            sheet.cell(rows[key], c).value = values[c - 1]

    plan_changed = set()
    for key in touched:
        for target, origin in regroup.MIRROR.items():
            cell = plan.cell(at[key], target)
            if cell.value != sheet.cell(rows[key], origin).value:
                cell.value = sheet.cell(rows[key], origin).value
                plan_changed.add(cell.coordinate)
    for key, cells in PLAN_CELLS.items():
        for name, value in cells.items():
            cell = plan.cell(at[key], column[name])
            if name in ("Quality targets", "Pass thresholds") and key == "LB-40":
                assert cell.value in (None, "", "None"), f"{key} {name} is not empty"
            cell.value = value
            plan_changed.add(cell.coordinate)

    out_dir = Path(tempfile.gettempdir()) / "nm_legal_brain_owners_20260927"
    out_dir.mkdir(exist_ok=True)
    out = out_dir / SOURCE.name
    book.save(out)

    # ================= THE PROOF, on the saved bytes ==========================
    check = load_workbook(out)
    for tab in check:
        assert regroup.sheet_features(tab) == features[tab.title], tab.title
    changed_bb = {(rows[k], c) for k in touched for c in range(1, 11)}
    for (title, coord), old in before.items():
        cell = check[title][coord]
        if title == "Before Build" and (cell.row, cell.column) in changed_bb:
            assert cell._style == old[1], coord
            continue
        if title == "Implementation Plan" and coord in plan_changed:
            assert cell._style == old[1], coord
            continue
        assert (cell.value, cell._style) == old, f"{title} {coord} moved"
    bb, ip = check["Before Build"], check["Implementation Plan"]
    for key in touched:
        assert [bb.cell(rows[key], c).value for c in range(1, 11)] == expected[key], key
    for key, cells in PLAN_CELLS.items():
        for name, value in cells.items():
            assert ip.cell(at[key], column[name]).value == value, (key, name)

    from assurance.control_plane import plan_scenarios, requirement_owners
    assert plan_scenarios.requirement_problems(out) == []
    assert len(plan_scenarios.sheet_rows(out)) == len(plan_scenarios.sheet_rows(SOURCE))
    registry = requirement_owners.Registry.load()
    population = requirement_owners.requirements(out, registry)
    problems = [f"{r.ident}: {r.problems}" for r in population if r.problems]
    assert problems == [], problems
    for packet in ("P49", "P50", "P51", "P52", "P53", "P54"):
        blockers = requirement_owners.slice_blockers(packet, population, registry)
        assert blockers == [], (packet, blockers)

    shutil.copy2(out, SOURCE)
    print(f"source digest before: {digest}")
    print(f"source digest after:  {hashlib.sha256(SOURCE.read_bytes()).hexdigest()}")
    print(f"Before Build rows written: {len(touched)}; plan cells written: {len(plan_changed)}; "
          f"ownership problems: 0; slices P49-P54 each may start")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
