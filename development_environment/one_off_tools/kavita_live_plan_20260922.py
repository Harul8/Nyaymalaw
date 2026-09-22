"""Record measured live diagnostics and repair progress, preserving earlier evidence.

Bundled artifact-tool is unavailable. Use bundled Python/openpyxl fallback.
"""
import copy
import hashlib
import json
import shutil

from openpyxl import load_workbook

from chat_reference_plan_20260922 import sheet_features
from checklist_plan_20260922 import OUT, SOURCE
from reconcile_implementation_20260922 import preview


NOTES = {
    'LB-109': 'OPEN diagnostic: the Kavita live brief describes title, access and shop rent, but only shop rent is recorded. The complete input survives. Raw-reader cause remains unproven.',
    'LB-110': 'OPEN diagnostic: whole-matter iteration cannot be accepted while two described disputes are absent from the inventory. The run stopped before further questions.',
    'LB-112': 'OPEN diagnostic: an explicit no-proceedings/proposed-claims instruction still ends in a posture block without a useful question. Do not invent a filed role or weaken the gate to fix it.',
    'LB-106': 'OPEN diagnostic: rent retrieval shows generic contract Article 55 and suggests Article 113; the opened corpus also contains the rent-specific Article 52, not considered in the response. Legal applicability remains for review.',
    'LB-113': 'OPEN diagnostic: graph-routing prose, raw authority keys and corpus counts reach the response. No independent assessment of the email, title or access is supplied.',
    'LB-117': 'Live checklist acceptance NOT ASSESSED: checklist not established behind the posture block. Controlled-test results do not certify this live journey.',
    'LB-118': 'Live conversational checklist movement NOT ASSESSED: paid messages stopped on the inventory/posture defect before follow-up. The earlier capacity defect is no longer the blocker.',
    'LB-119': 'Observed post-fix: the explicit capacity assessment persisted through restart/sign-in and no longer blocked the live turn. This does not establish legal-quality acceptance. The new failure is inventory/posture.',
    'LB-92': 'Observed live: saved passage and current full document retain separate identities; Full document opens at the highlighted cited passage. Raw full-document title remains a presentation gap.',
}

REPAIR_NOTES = {
    'LB-109': 'IN PROGRESS diagnostic repair: protected trace proved the first model inventory omitted title and access. Later live retries exposed altered quote punctuation, omitted shared instructions and contradictory new-versus-existing verdicts. The reader now uses a complete source-unit allocation contract, restores exact input text, rejects invalid IDs, omissions and contradictory verdicts, and permits one bounded repair. A failed repair cannot admit a subset. Live end-to-end acceptance remains open.',
    'LB-110': 'IN PROGRESS: whole-matter iteration is gated behind the complete inventory; unsuccessful retries remain visible in History. No closure or complete assessment claimed.',
    'LB-112': 'IN PROGRESS: shared representation, proposed claims and authority limits are explicitly allocated with dispute facts. No filed role is invented. Live posture retest remains open.',
    'LB-106': 'IMPLEMENTED, live verification pending: a specific arrears-of-rent cause routes to retrieved Article 52, separately from generic contract damages under Article 55. Each rental instalment has its own due-date question. The route is a candidate, not a finding of applicability. Controlled retrieval fixtures now carry index identity and assert the current search-note contract without weakening the ceiling checks.',
    'LB-113': 'PARTIAL repair: graph-routing boilerplate, raw corpus keys and coverage counts removed from newly assembled prose. Missing judgment names have a human fallback. Historical responses remain unchanged. Full-document statute title and live assessment quality remain open.',
}
NOTES.update(REPAIR_NOTES)
NOTES.update({
    'LB-109': 'DIAGNOSTIC, partial repair verified: the clarification created three distinct disputes, but earlier unallocated brief text did not reach their analyses. Inventory now includes only still-unallocated original accounts, preserves their source-turn provenance, and prevents earlier instructions from setting current focus. Full live revalidation remains OPEN.',
    'LB-110': 'DIAGNOSTIC OPEN: three disputes now persist, but successful whole-matter iteration and saved-session reopen have not yet been demonstrated. Failed attempts remain in History.',
    'LB-112': 'DIAGNOSTIC: ownership incorrectly acquired the shop tenant as opponent. Active-dispute context and an exact-current-words opponent correction contract are implemented; existing resolved roles cannot silently flip. Missing premises no longer render established. Controlled tests pass; live correction pending.',
    'LB-113': 'DIAGNOSTIC: the live assessment invented cross-dispute consequences and exposed internal workup prose. Cross-dispute exposure now requires attributed premise IDs and exact distinct quotes on both disputes; a failed comparison is NOT_RUN, not no-risk. Legal attacks do not run without retrieved law. Assessment purpose now outranks brief length. These repairs are not proof of semantic entailment or expert-quality live prose; live acceptance OPEN.',
})
NOTES.update({
    'LB-109': 'DIAGNOSTIC OPEN: the completed Kavita retry repeated one existing ownership ID and misallocated other disputes to it. Duplicate IDs now refuse; an unavailable first inventory also refuses instead of silently creating one dispute. Structural completeness is not semantic correctness. Original failed history is retained; clean live retest required.',
    'LB-112': 'DIAGNOSTIC OPEN: ownership still carried the shop tenant and used the access obstruction date conditionally. The exact-current opponent correction now selects complete source units, but wrong dispute allocation must not be disguised by a conditional label. Controlled nullable-schema and source-binding repairs passed; live legal correctness not established.',
    'LB-113': 'DIAGNOSTIC OPEN: the live answer contained unsupported cross-dispute/doctrinal assertions and internal workup fragments. The limitation classifier now binds to the whole proposed answer rather than its heading; an explanation is NOT automatically exempt. General source-bound research choices implemented. Model judgment, natural presentation and useful grounded assessment still require live proof.',
    'LB-117': 'DIAGNOSTIC OPEN: unsupported checklist candidates triggered an invalid gate state and HTTP 500. The repaired path preserves existing checklist/cache, records unavailable and refuses unsupported candidates; negative test passes. Dynamic source choices select only retrieved passages. A live established checklist is not yet demonstrated.',
    'LB-110': 'DIAGNOSTIC OPEN: failed Kavita file and histories retained. A fresh complex matter is being run through the normal Chrome journey after repairs. Whole-matter iteration, checklist movement and reopen persistence remain unproven; no end-to-end pass claimed. Ledger checkpoint: 120 real GPT-4o mini calls, USD0.058345 of original USD1.',
})


NOTES.update({
    'LB-109': 'DIAGNOSTIC OPEN: clean Sahana retest initially merged inheritance and access. A correction read identified the missing access dispute but its recovery contract could only retain old targets. Fixed repair now preserves source-supported proposed new targets, refuses invented IDs and refuses silently dropping new work. 44 focused checks passed; semantic allocation still needs live acceptance.',
    'LB-110': 'DIAGNOSTIC OPEN: complete whole-matter review has not yet succeeded. Two failed complex matters and their histories are preserved. Saved-session reopen through My work succeeds; that is persistence evidence, not reasoning acceptance. Original USD1 ledger remains in force, GPT-4o mini only.',
    'LB-112': 'DIAGNOSTIC PARTIAL: explicit prospective claimant correction persisted after reopen. Cause reads now receive active-dispute context. Opponent and chronology allocation still require live review; no conditional deadline may disguise a cross-dispute date error.',
    'LB-113': 'DIAGNOSTIC PARTIAL: consistency now distinguishes discussing an opponent argument from recommending against our client. Refusals no longer quote the rejected model draft into the quotation guard. Theory and posture use bounded population-scaled budgets; a negative control detects unregistered fixed budgets. 167 focused checks passed before the additional inventory repair. Natural grounded live assessment remains OPEN.',
    'LB-119': 'DIAGNOSTIC PARTIAL: withheld response falsely said the brief was unsaved although canonical input had committed. API and UI now distinguish input_only from unknown, unsaved and previously completed receipts; commit-failure negative control passes. Input-only refreshes the current version and avoids blind replay; withheld conclusions remain uncommitted. Live exercise of the new error presentation remains pending.',
})


NOTES.update({
    'LB-109': 'DIAGNOSTIC PARTIAL: Sahana now has three disputes after source-bound allocation repair. Previously misallocated chronology is preserved, not silently rewritten; semantic allocation and correction still need acceptance. Six real browser turns remain in History. No end-to-end pass.',
    'LB-110': 'DIAGNOSTIC OPEN: latest live test stopped on material semantic defects before whole-matter iteration or closing assessment. No claim that access and rent were fully worked. Original USD1 ledger: USD0.134084 measured, plus USD0.030000 reserved/unknown retained; 213 measured GPT-4o mini calls. No stronger-model approval assumed.',
    'LB-106': 'DIAGNOSTIC OPEN: research now selects a stable source ID and code restores exact text plus basis; quotation-containing source text no longer breaks the schema, and inconsistent double-selected bases are impossible in the production contract. Real retrieval runs but returned irrelevant criminal/constitutional judgments for the inheritance question. Source presence is not semantic relevance. Live legal-retrieval quality FAIL.',
    'LB-112': 'DIAGNOSTIC PARTIAL: undated exclusions and current instructions now reach accrual selection. Live ownership date was correctly withdrawn. The model subsequently changed Article 65 to Article 64 without a sound resolution of the ownership instructions. No deadline may be treated as established; professional applicability remains OPEN.',
    'LB-113': 'DIAGNOSTIC OPEN: live model called an ordinary email copy certified, invented a rent-to-inheritance consequence, and treated investigation of an unknown date as contradicting an uncomputed period. Unsupported ancillary prose and long internal workup still reach the user. Structural/source-binding checks are not semantic-support proof. Paid iteration stopped; stronger-model comparison approval remains unanswered.',
})


def main():
    source_digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    book = load_workbook(SOURCE)
    before = {(s.title, c.coordinate): (c.value, copy.copy(c._style))
              for s in book for row in s for c in row}
    features = {s.title: sheet_features(s) for s in book}
    rows = {}
    changed = set()
    sheet, plan = book['Before Build'], book['Implementation Plan']
    for ident, note in NOTES.items():
        matches = [r for r in range(5, sheet.max_row + 1)
                   if str(sheet.cell(r, 1).value or '').startswith(ident + '\n')]
        assert len(matches) == 1, ident
        row = matches[0]
        rows[ident] = row
        addition = ('22 September 2026 live repair progress: ' + note
                    + ' Evidence: kavita-live-followon.md in the 22 September legal-brain evidence folder. No release or professional acceptance claimed.')
        cell = sheet.cell(row, 10)
        if addition not in str(cell.value or ''):
            cell.value = addition + '\n\nEarlier record (retained):\n' + str(cell.value or '')
        changed.add((sheet.title, cell.coordinate))
        mirrors = [r for r in range(2, plan.max_row + 1) if plan.cell(r, 1).value == ident]
        assert len(mirrors) == 1, ident
        mirror = plan.cell(mirrors[0], 30)
        mirror.value = cell.value
        changed.add((plan.title, mirror.coordinate))
    OUT.mkdir(parents=True, exist_ok=True)
    target = OUT / SOURCE.name
    book.save(target)
    check = load_workbook(target)
    assert check.sheetnames == book.sheetnames
    for s in check:
        assert sheet_features(s) == features[s.title], s.title
    for key, old in before.items():
        cell = check[key[0]][key[1]]
        if key not in changed:
            assert (cell.value, cell._style) == old, key
        else:
            assert cell._style == old[1], key
    preview(check['Before Build'], [rows['LB-109'], rows['LB-112']], [1, 10],
            OUT / 'live-kavita-after.png')
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == source_digest, 'parallel workbook edit'
    shutil.copy2(target, SOURCE)
    print(json.dumps({'diagnostic_owners': rows, 'changed_cells': len(changed),
                      'unrelated_value_or_style_changes': 0}))


if __name__ == '__main__':
    main()
