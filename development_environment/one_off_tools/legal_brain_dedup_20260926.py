"""Remove the duplication today's legal-brain updates introduced into the plan.

OWNER DIRECTION, 26 September 2026: "check if there is any duplication from
today's updates".

WHAT WAS FOUND, and the one rule every edit below applies: a requirement,
criterion or measured figure is stated ONCE; every other row that needs it
names it.

  * The review (legal_brain_review_20260926.py) "carried" older criteria into
    the proposed owning rows by COPYING their text while the older rows still
    stand. Until the owner approves an owner, that is two copies. They become
    references: named, not restated; the text moves only on approval.
  * One criterion was written three times (LB-67-AC5, LB-139-AC5, LB-161-AC5)
    and one twice (LB-131-AC4, LB-138-AC3). The detector criterion becomes ONE
    test whose population is every channel, so a channel added later joins it
    rather than getting a copy.
  * LB-161 listed the answer's claim kinds; the review widened LB-130's list
    and LB-161's did not follow -- two copies that had already drifted.
  * Criteria duplicated between the morning's drafts: LB-130-AC1 / LB-143-AC3,
    LB-135-AC1 / LB-150-AC1, LB-137-AC1 / LB-156-AC1, LB-159-AC3 / LB-169-AC4,
    LB-126-AC4 / LB-60-AC1, and LB-140-AC1's board clause / LB-166-AC1.
  * The corpus measurement was written in four cells, and the step-independence
    measurement in two. Each now lives in one row; the others point to it.

Every edit is an exact substring replacement asserted to occur once. The
discipline is the other workbook tools': prove on the SAVED file that nothing
else moved, check the reconciler, then replace the source.
"""
import copy
import hashlib
import importlib.util
import re
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

_REPO = Path(__file__).resolve().parents[2]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))
from assurance.common._console import utf8_console  # noqa: E402

utf8_console()

_spec = importlib.util.spec_from_file_location(
    "regroup", Path(__file__).with_name("legal_brain_regroup_20260926.py"))
regroup = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(regroup)

SOURCE, FIRST = regroup.SOURCE, regroup.FIRST
HEADER = re.compile(r"^(L|L\.\d|U|U\.\d|P|X)  ")
R = "26 September 2026 review"
BY_REF = "ADOPTED BY REFERENCE, not restated, until the owner approves this row as the owner (LB-43): "

#: (row, column, old, new). Each old string must occur exactly once in that cell.
EDITS = [
    # ---- one detector criterion, its population every channel -------------------
    ("LB-67", 8,
     "LB-67-AC5 (planted): an opposing-counsel step event naming a case that was never retrieved never reaches "
     "the stream, the saved scratch pad or the board, and the omission is shown.",
     "LB-67-AC5 (planted): ONE TEST, ITS POPULATION EVERY CHANNEL -- a line naming a case that was never "
     "retrieved is planted in each channel the product has (step events and the scratch pad, board items, the "
     "model's questions, the compaction handover, the answer) and reaches the advocate on none of them, with the "
     "omission shown; a channel added later joins the population rather than getting a test of its own."),
    ("LB-139", 8,
     "LB-139-AC5 (planted): a pass-2 event naming an unretrieved case is left out of the live panel and the saved "
     "scratch pad, and the omission is shown (LB-67-AC5).",
     "The scratch pad is one channel in LB-67-AC5's population; it has no separate test for this."),
    ("LB-139", 7,
     "\n\nNever stream a line that names an authority or provision not retrieved on this matter (LB-67).", ""),
    ("LB-164", 7, "\n\nNever stream model-written text the detector has not passed (LB-67).", ""),
    ("LB-161", 8,
     "LB-161-AC4 (carried from LB-62-AC4): an external act whose acknowledgement was lost is reconciled, never "
     "re-executed blindly.\n"
     "LB-161-AC5 (planted): a question submitted through ask_advocate that names an unretrieved authority is "
     "stopped by LB-67's detector.",
     "LB-62-AC4 (a lost acknowledgement is reconciled, never re-executed blindly) applies to propose_action by "
     "reference. Questions are one channel in LB-67-AC5's population and have no separate test for it."),
    # ---- LB-161's claim kinds had drifted from LB-130's ---------------------------
    ("LB-161", 4,
     "Each claim carries its kind -- fact, law, authority, inference, question, limit -- and its support;",
     "Each claim carries its kind and its support as LB-130 defines them;"),
    ("LB-161", 10, "LB-62-AC4 carried.",
     "LB-62-AC4 applied by reference; the claim kinds now point to LB-130 (the two lists had drifted)."),
    # ---- one repository check, not two --------------------------------------------
    ("LB-138", 8,
     "LB-138-AC3: during the overlap, a repository check fails the build if the loop and the pipeline call "
     "different implementations of a gate (LB-131-AC4).\n"
     "LB-138-AC4: every golden scenario carries an LB-138 slice tag,",
     "LB-138-AC3: every golden scenario carries an LB-138 slice tag,"),
    # A second run of an earlier version of this tool doubled the note below,
    # because that edit's new text contained its old text. The repair comes
    # first; the edit after it no longer contains its own old text.
    ("LB-138", 8,
     "replay and live criteria with the ledger estimate for the live ones. (One gate implementation during the "
     "overlap is LB-131-AC4.) (One gate implementation during the overlap is LB-131-AC4.)",
     "replay and live criteria, with the ledger estimate for the live ones. (One gate implementation during the "
     "overlap is LB-131-AC4.)"),
    ("LB-138", 8, "replay and live criteria with the ledger estimate for the live ones.",
     "replay and live criteria, with the ledger estimate for the live ones. (One gate implementation during the "
     "overlap is LB-131-AC4.)"),
    ("LB-138", 4,
     "Before slice 3 serves a live turn, the cost per turn is measured from replay call counts times price and "
     "scored as a release row (LB-35).",
     "Before slice 3 serves a live turn, LB-35's cost-per-turn release row is measured."),
    # ---- the principles criterion LB-60 already owns --------------------------------
    ("LB-126", 8,
     "LB-126-AC4 (carried from LB-60-AC1): the served model request carries the current principles version, read "
     "from the request itself, not from the file.\n"
     "LB-126-AC5: a principles change made mid-conversation",
     "That the served request carries the current principles is LB-60-AC1's check, which covers this document.\n"
     "LB-126-AC4: a principles change made mid-conversation"),
    ("LB-60", 10, "LB-126 carries LB-60-AC1 as LB-126-AC4 for the principles document.",
     "LB-60-AC1 also covers the principles document (LB-126); it is not restated there."),
    # ---- the loop's carried criteria become references ------------------------------
    ("LB-128", 8,
     "LB-128-AC4 (planted, carried from LB-64-AC2): repeated equivalent observations end in a no-progress stop "
     "with the unresolved needs kept.\n"
     "LB-128-AC5 (carried from LB-64-AC4): an answerable, an externally blocked and a budget-limited task end in "
     "three distinct, truthful stop states.\n"
     "LB-128-AC6 (planted, carried from LB-69-AC1): nested loops and retries draw on one allowance, which holds "
     "under restart.\n"
     "LB-128-AC7 (carried from LB-36-AC1 and LB-01-AC3): a narrow question is answered without a compulsory full "
     "intake, and the matter is not closed by it.\n"
     "LB-128-AC8 (carried from LB-61-AC1): rephrasing a message without changing its meaning does not change "
     "what the loop does in substance.",
     BY_REF + "LB-64-AC2 (no-progress stop), LB-64-AC4 (distinct stop states), LB-69-AC1 (one allowance across "
     "retries and nested loops), LB-36-AC1 and LB-01-AC3 (a narrow task without compulsory intake, the matter not "
     "closed by it), LB-61-AC1 (rephrasing does not change the substance). On approval their text moves here and "
     "those rows are marked carried."),
    ("LB-128", 5,
     "The turn ends in one of four distinct, truthful states (carried from LB-64): an answer that passed the "
     "harness (sufficient for the task), a question to the advocate (awaiting information or authority), a "
     "no-progress stop, or a budget stop -- each saying what was not finished.",
     "The turn ends in one of the four distinct, truthful stop states LB-64 defines -- an answer that passed the "
     "harness, a question to the advocate, a no-progress stop, or a budget stop -- each saying what was not "
     "finished."),
    ("LB-128", 6,
     f"{R}: a no-progress stop fires when repeated equivalent tool calls or observations, oscillating positions "
     "or revisions that change nothing recur -- a rephrased or relabelled query is not new work (LB-105) -- and it "
     "keeps the unresolved needs. One allowance covers the turn, its retries and every nested loop; a resumed turn "
     "never resets it without a recorded decision (carried from LB-69).",
     f"{R}: the no-progress rule is LB-64's and the single allowance across retries and nested loops is LB-69's "
     "(a relabelled query is not new work: LB-105); this row applies both to the loop and does not restate them."),
    ("LB-128", 7,
     "\n\nNever iterate until the answer favours the client; never count rising confidence as progress (LB-64).",
     ""),
    ("LB-128", 10, "four stop states, carried criteria,", "four stop states named, criteria adopted by reference,"),
    ("LB-01", 10,
     "Carried there: LB-01-AC3 (a narrow task finishes without closing the matter) as LB-128-AC7; LB-01-AC1 and "
     "AC2", "Adopted there by reference: LB-01-AC3 (a narrow task finishes without closing the matter); LB-01-AC1 "
     "and AC2"),
    ("LB-36", 10,
     "Carried there: LB-36-AC1 as LB-128-AC7; the distinct stop states as LB-128's four endings.",
     "Adopted there by reference: LB-36-AC1, and the distinct stop states LB-128's four endings name."),
    ("LB-61", 10,
     "Carried there: LB-61-AC1 as LB-128-AC8; LB-61-AC4 (an unsupported action refused at admission) by "
     "LB-127-AC2 and LB-143.",
     "Adopted there by reference: LB-61-AC1; LB-61-AC4 (an unsupported action refused at admission) is already "
     "met by LB-127-AC2 and LB-143."),
    ("LB-64", 10,
     "Carried there: the four stop states and the no-progress rule as LB-128's exit and failure clauses, "
     "LB-128-AC4 and AC5;",
     "Adopted there by reference: the four stop states and the no-progress rule, with LB-64-AC2 and AC4;"),
    ("LB-69", 10,
     "Carried there: one allowance across retries and nested loops, never reset on resume, as LB-128-AC6;",
     "Adopted there by reference: one allowance across retries and nested loops, never reset on resume, with "
     "LB-69-AC1;"),
    # ---- the claim criteria ---------------------------------------------------------
    ("LB-130", 8,
     "LB-130-AC4 (planted, carried from LB-57-AC1): an invented fact, an invented provision and a valid-looking "
     "source identifier from another matter are each refused as support.\n"
     "LB-130-AC5 (carried from LB-57-AC4): with a store blocked, the answer gives honest limited help and neither "
     "supplies law from memory nor says the law is absent.\n"
     "LB-130-AC6 (planted, carried from LB-59-AC2): a circular dependency or a missing necessary premise refuses "
     "the dependent inference while independent claims are served.",
     BY_REF + "LB-57-AC1 (an invented fact or a foreign-matter source identifier never becomes support; an "
     "invented provision is this row's AC1), LB-57-AC4 (a blocked store gives honest limited help, no law from "
     "memory) and LB-59-AC2 (a circular dependency or missing premise refuses only the dependent inference). On "
     "approval their text moves here and those rows are marked carried."),
    ("LB-130", 10, "LB-57 and LB-59 criteria carried;", "LB-57 and LB-59 criteria adopted by reference;"),
    ("LB-57", 10,
     "LB-57-AC1 and AC4 are carried there as LB-130-AC4 and AC5, and the research-hypothesis label as a claim "
     "kind.",
     "LB-130 adopts LB-57-AC1 and AC4 by reference, and the research-hypothesis label as a claim kind."),
    ("LB-59", 10,
     "this row's legal-application record and LB-59-AC2 are carried there (LB-130-AC6).",
     "this row's legal-application record and LB-59-AC2 are adopted there by reference."),
    ("LB-143", 8,
     "LB-143-AC3 (planted): an answer citing a provision its step log never retrieved is refused.",
     "LB-143-AC3: the step log is the evidence LB-130-AC1 reads -- an answer citing a provision the log never "
     "retrieved is refused there, not by a second test here."),
    # ---- the adversarial criteria ---------------------------------------------------
    ("LB-140", 8,
     "LB-140-AC5 (planted, carried from LB-65-AC1): a plausible objection with no passage leaves the assessment "
     "unchanged and is recorded as unsupported.\n"
     "LB-140-AC6 (carried from LB-65-AC3 and LB-18): actual submissions on the file, anticipated defences and "
     "contingent ones are labelled apart in the answer and the scratch pad.\n"
     "LB-140-AC7 (carried from LB-65-AC4): the three passes are enabled only after a recorded single-pass versus "
     "three-pass comparison, with its cost.",
     BY_REF + "LB-65-AC1 (an unsupported objection changes nothing), LB-65-AC3 (actual, anticipated and "
     "conditional arguments labelled apart -- here in the answer and the scratch pad too) and LB-65-AC4 "
     "(single-pass versus challenged comparison, with its cost, before the extra passes are enabled)."),
    ("LB-140", 8,
     "LB-140-AC1: a cheque-dishonour dispute's pass 1 names the enforceable-debt and notice-receipt defences, and "
     "the board asks for the loan evidence and the proof of delivery.",
     "LB-140-AC1: a cheque-dishonour dispute's pass 1 names the enforceable-debt and notice-receipt defences, and "
     "each reaches the board as an item that names it (what the board then asks is LB-166-AC1)."),
    ("LB-140", 4,
     f"{R}: each argument is labelled raised (in a pleading or submission on the file), reasonably anticipated, "
     "contingent or speculative, and only the first three reach the answer (carried from LB-18 and LB-65). An "
     "objection with no supporting passage does not change the assessment (LB-65); a sound position survives a "
     "pass unchanged -- agreement is not required, and the loop does not iterate until the answer favours the "
     "client (LB-64). Before the extra passes are enabled, a single-pass and a challenged run are compared on "
     "held-out matters, reporting material gains, regressions, added latency and cost (LB-65-AC4).",
     f"{R}: this row applies LB-18's and LB-65's argument labels -- raised, reasonably anticipated, contingent, "
     "speculative -- to the passes, and a speculative argument does not reach the answer. That an objection "
     "without support changes nothing (LB-65), that a sound position survives a pass unchanged (LB-64), and that "
     "the extra passes wait for LB-65-AC4's comparison are stated in those rows and not restated here."),
    ("LB-140", 10, "the single-versus-three-pass comparison carried;",
     "the single-versus-three-pass comparison adopted by reference;"),
    ("LB-18", 10,
     "this row's four argument categories (raised, anticipated, contingent, speculative) are carried there.",
     "LB-140 applies this row's four argument categories by reference."),
    ("LB-65", 10, "LB-65-AC1, AC3 and AC4 are carried there as LB-140-AC5, AC6 and AC7.",
     "LB-140 adopts LB-65-AC1, AC3 and AC4 by reference."),
    # ---- writes and admission --------------------------------------------------------
    ("LB-142", 8,
     "LB-142-AC4 (carried from LB-63-AC1): duplicate and out-of-order writes have one accepted effect, and a newer "
     "value is kept.\n"
     "LB-142-AC5 (carried from LB-63-AC2): a crash between a write's check and its commit resumes with no lost "
     "and no duplicated effect.",
     "LB-63-AC1 and AC2 (duplicate, out-of-order and crash-interrupted writes) apply to these writes by reference; "
     "LB-63 stays their owner."),
    ("LB-142", 10, "LB-63's criteria carried;", "LB-63's criteria applied by reference;"),
    ("LB-63", 10, "LB-63-AC1 and AC2 are carried as LB-142-AC4 and AC5.",
     "LB-63-AC1 and AC2 apply to them by reference."),
    ("LB-143", 8,
     "LB-143-AC4 (planted, carried from LB-62-AC3): access revoked mid-turn refuses every later tool call on that "
     "matter.",
     "LB-62-AC3 (access revoked mid-run stops every later call) applies to the loop's tool calls by reference."),
    ("LB-143", 9,
     "Carried: LB-62-AC3 as LB-143-AC4; LB-62-AC4 into LB-161 (LB-161-AC4); LB-63-AC1 and AC2 into LB-142 "
     "(LB-142-AC4, AC5).",
     "By reference, not restated: LB-62-AC3 here; LB-62-AC4 in LB-161; LB-63-AC1 and AC2 in LB-142."),
    ("LB-143", 10, "revocation criterion carried.", "revocation criterion applied by reference."),
    ("LB-62", 10, "LB-62-AC3 is carried as LB-143-AC4 and LB-62-AC4 as LB-161-AC4.",
     "LB-62-AC3 applies there, and LB-62-AC4 to LB-161's external acts, by reference."),
    # ---- research, context, questions -------------------------------------------------
    ("LB-136", 6,
     f"{R}: a research loop does not repeat work -- rephrasing or relabelling a query does not make it new work, "
     "a proposal against an out-of-date matter snapshot is refused, and its stop names which happened: provider "
     "failure, a bad proposal, retrieval failure, repetition, no progress, or budget (carried from LB-105).",
     f"{R}: LB-105's rules -- a relabelled query is not new work, a stale-snapshot proposal is refused, and each "
     "distinct stop is named -- apply to this loop by reference and are not restated here."),
    ("LB-136", 8,
     "LB-136-AC3 (planted, carried from LB-105): a relabelled repeat of an earlier query is refused as repeated "
     "work, and a stale-snapshot proposal is refused.",
     "LB-105's controlled tests apply to this loop by reference until the switch."),
    ("LB-136", 10, "LB-105's rules carried.", "LB-105's rules adopted by reference."),
    ("LB-105", 10, "carries its repeated-query, stale-snapshot and distinct-stop rules (LB-136-AC3).",
     "adopts its repeated-query, stale-snapshot and distinct-stop rules by reference."),
    ("LB-137", 7, "Never count a rephrased or relabelled query as new work (LB-105).",
     "LB-105's repeated-query rule applies."),
    ("LB-137", 8,
     "LB-137-AC1: a question naming s.53A of the Transfer of Property Act retrieves from that Act and no other.",
     "LB-137-AC1: covered by LB-156-AC1 (s.53A of the Transfer of Property Act reads that Act and no other); not a "
     "second test."),
    ("LB-135", 8,
     "LB-135-AC4 (carried from LB-10-AC3): when the context budget forces material out, the answer discloses the "
     "coverage limit rather than implying it read everything.",
     "LB-10-AC3 (context limits disclosed, never false completeness) applies by reference."),
    ("LB-135", 8,
     "LB-135-AC1: on a matter longer than the context budget, a fact stated in the first turn is used correctly "
     "in the last.",
     "LB-135-AC1: covered by LB-150-AC1 (a fact from the first turn used correctly after the budget is passed); "
     "not a second test."),
    ("LB-135", 10, "LB-10-AC3 carried.", "LB-10-AC3 applied by reference."),
    ("LB-10", 10, "is carried as LB-135-AC4.", "applies to LB-135 by reference."),
    ("LB-131", 8,
     "LB-131-AC3 (planted, carried from LB-38-AC3 and LB-67-AC3): one failed dependency withholds that conclusion "
     "and what rests on it, and the independent help is still served.",
     "LB-38-AC3 and LB-67-AC3 (one failed dependency withholds only what rests on it) apply to the loop's output "
     "by reference."),
    ("LB-131", 4,
     "when one check fails and repair cannot fix it, only the affected claim and the conclusions resting on it are "
     "withheld; the rest of the answer is re-checked for coherence and served with the limit stated (carried from "
     "LB-38 and LB-67).",
     "when a check still fails after repair (LB-133), withholding is scoped as LB-67 states it."),
    ("LB-131", 10, "scoped withholding and one-implementation rule added.",
     "scoped withholding (by reference to LB-67) and the one-implementation rule added."),
    ("LB-133", 4,
     "after the last attempt, what is still failing is withheld with the conclusions that rest on it, and the rest "
     "of the answer is re-evaluated for coherence before it is served (LB-67).",
     "after the last attempt, withholding is scoped as LB-67 states it."),
    ("LB-38", 10, "LB-38-AC3 is carried there as LB-131-AC3;", "LB-38-AC3 applies there by reference;"),
    ("LB-166", 8,
     "LB-166-AC4 (carried from LB-118): an item the advocate promises returns when due and is not re-asked before "
     "then; an item marked unavailable is not re-asked without a new reason.\n"
     "LB-166-AC5 (carried from LB-05-AC2): a detail already held in an uploaded document is found by search_matter "
     "and not asked for.",
     "LB-118's controlled tests (a promised item returns when due; an unavailable item is not re-asked without a "
     "new reason) apply to the board by reference. A detail already held in an upload is found rather than asked "
     "for: LB-155-AC1 and LB-161-AC3."),
    ("LB-166", 10, "criteria carried;", "criteria applied by reference;"),
    ("LB-118", 10, "its promise and unavailable rules are carried as LB-166-AC4.",
     "its promise and unavailable rules apply there by reference."),
    ("LB-05", 10, "is carried as LB-166-AC5.", "is covered there by LB-155-AC1 and LB-161-AC3."),
    ("LB-43", 4,
     "each owner carrying the others' surviving clauses as acceptance criteria, named in each row:",
     "each owner adopting the others' surviving clauses by reference -- named in each row, restated nowhere -- "
     "and taking their text only when the owner approves it:"),
    # ---- new practice rows restating owned rules --------------------------------------
    ("LB-169", 8,
     "\nLB-169-AC4: every requirement result names its curated_from source.",
     "\nThat every result names its curated_from source is LB-159-AC3's check, which covers this table."),
    ("LB-168", 7, "Must keep existence, admissibility and weight as three questions.",
     "Must keep existence, admissibility and weight as the three questions LB-16 owns."),
    # ---- each measurement in one row ---------------------------------------------------
    ("HEADER:L.4", 1,
     "Measured 26 September 2026: the Commercial Courts Act 2015, the Telangana Court-fees and Suits Valuation "
     "Act 1956 and the Telangana Civil Courts Act 1972 are HELD -- in raw_data/, chunks.db and legal.db -- and "
     "missing only from pipeline/manifest.yaml, so exact identification cannot reach them; the court-fee "
     "schedules are in the raw text but no schedule atoms reached chunks.db. What is missing is the manifest "
     "entry and the schedule atoms, not the source (LB-121, LB-125, LB-170).",
     "A 26 September 2026 measurement found three Acts the plan had called missing held in every store and "
     "absent only from the manifest; the figures are in LB-121, LB-125 and LB-170, and belong in "
     "docs/BASELINE.md."),
    ("LB-125", 9,
     "as are the Andhra Pradesh version (456 chunks) and the Telangana Civil Courts Act, 1972 (122 chunks; 31 "
     "sections in legal.db).",
     "as is the Andhra Pradesh version (456 chunks); the Telangana Civil Courts Act, which fixes the courts' "
     "pecuniary limits, is measured in LB-170."),
    ("LB-125", 9, "and none of these Acts is in pipeline/manifest.yaml.",
     "and neither version is in pipeline/manifest.yaml."),
    ("LB-170", 4,
     "; the Telangana Court-fees and Suits Valuation Act, 1956 (461 chunks, no schedule atoms) -- none of them in "
     "the manifest.",
     " -- none of them in the manifest; the Court-fees and Suits Valuation Act is measured in LB-125."),
    ("LB-169", 9,
     "CORPUS, MEASURED 26 September 2026: the_code_of_civil_procedure_1908 (2,901 chunks in chunks.db) and the "
     "Code of Civil Procedure (Telangana) Second Amendment Act, 1953 (7 chunks) are held; state amendments to the "
     "Orders are checked before any entry is curated (LB-170).",
     "CORPUS, MEASURED 26 September 2026: the_code_of_civil_procedure_1908 is held (2,901 chunks in chunks.db); "
     "the state amendments, measured in LB-170, are checked before any entry is curated."),
    ("LB-141", 9,
     "the step-independence read answered 'dependent' for every step until it was given the reason first and two "
     "observable halves (tests/test_the_independence_read_reasons_before_it_rules.py).",
     "the step-independence measurement is recorded in LB-134."),
    ("LB-167", 4,
     "The OM-P rows, whose only gap is held-out evaluation, become sections of the principles document (LB-126) "
     "and stop carrying a build status of their own.",
     "The OM-P rows, whose only gap is held-out evaluation, follow LB-126's proposal and stop carrying a build "
     "status of their own."),
    # ---- found by the re-run of the duplicate scan ----------------------------------
    ("LB-150", 8,
     "LB-150-AC3: after compaction the model retrieves an earlier turn verbatim by tool.",
     "LB-150-AC3: covered by LB-155-AC3 (read_turn returns an earlier turn word for word after a compaction); not "
     "a second test."),
]

PLAN_CELLS = {
    ("LB-125", 44): "No forum, valuation or fee is computed. The Telangana Court-fees and Suits Valuation Act is "
                    "held (raw_data/, chunks.db and legal.db, measured 26 September 2026) but not in the manifest, "
                    "and its fee schedules were never atomised.",
}


def main() -> int:
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    book = load_workbook(SOURCE)
    sheet, plan = book["Before Build"], book["Implementation Plan"]
    features = {s.title: regroup.sheet_features(s) for s in book}
    before = {(s.title, c.coordinate): (c.value, copy.copy(c._style)) for s in book for row in s for c in row}

    rows = {}
    for r in range(FIRST, sheet.max_row + 1):
        text = str(sheet.cell(r, 1).value or "")
        h = HEADER.match(text)
        key = f"HEADER:{h[1]}" if h else regroup.key_of(text)
        assert key not in rows, key
        rows[key] = r

    # RE-RUNNABLE, STRUCTURALLY. An edit whose new text contains its old text
    # can match again after it has been applied -- that is how an earlier
    # version doubled a note -- so such an edit is refused before anything is
    # written. An edit whose old text is absent is taken as applied, and that is
    # PROVED below: its new text must be in the final cell.
    for key, col, old, new in EDITS:
        assert old not in new, f"{key} col {col}: new text contains old text, so the edit cannot be re-run"
    expected, skipped = {}, []
    for key, col, old, new in EDITS:
        r = rows[key]
        value = expected.get((r, col), sheet.cell(r, col).value)
        if old not in str(value):
            skipped.append((key, r, col, new))
            continue
        assert str(value).count(old) == 1, (key, col, old[:70])
        expected[(r, col)] = value.replace(old, new)
    for key, r, col, new in skipped:
        final = expected.get((r, col), sheet.cell(r, col).value)
        assert not new or new in str(final), f"{key} col {col}: neither the old nor the new text is present"
    print(f"already applied, skipped: {len(skipped)}")
    touched = {k for k, col, *_ in EDITS if (rows[k], col) in expected}
    for (r, col), value in expected.items():
        sheet.cell(r, col).value = value

    plan_changed = set()
    for x in range(2, plan.max_row + 1):
        ident = plan.cell(x, 1).value
        if ident in touched and regroup.LBOM.fullmatch(str(ident)):
            for t, origin in regroup.MIRROR.items():
                v = sheet.cell(rows[ident], origin).value
                if plan.cell(x, t).value != v:
                    plan.cell(x, t).value = v
                    plan_changed.add(plan.cell(x, t).coordinate)
        for (pid, col), v in PLAN_CELLS.items():
            if ident == pid and plan.cell(x, col).value != v:
                plan.cell(x, col).value = v
                plan_changed.add(plan.cell(x, col).coordinate)
    if not expected and not plan_changed:
        print("nothing to do: every edit is already on the sheet; the source is not rewritten")
        return 0

    out_dir = Path(tempfile.gettempdir()) / "nm_legal_brain_dedup_20260926"
    out_dir.mkdir(exist_ok=True)
    out = out_dir / SOURCE.name
    book.save(out)

    # ================= THE PROOF, on the saved bytes ==========================
    check = load_workbook(out)
    for tab in check:
        assert regroup.sheet_features(tab) == features[tab.title], tab.title
    bb, ip = check["Before Build"], check["Implementation Plan"]
    for (title, coord), old in before.items():
        cell = bb[coord] if title == "Before Build" else ip[coord]
        rc = (cell.row, cell.column)
        if title == "Before Build" and rc in expected:
            assert cell.value == expected[rc] and cell._style == old[1], coord
        elif title == "Implementation Plan" and coord in plan_changed:
            continue
        else:
            assert (cell.value, cell._style) == old, f"{title} {coord} moved"
    from assurance.control_plane import plan_scenarios
    problems = plan_scenarios.requirement_problems(out)
    assert problems == [], problems
    assert len(plan_scenarios.sheet_rows(out)) == 93

    shutil.copy2(out, SOURCE)
    print(f"source digest before: {digest}")
    print(f"source digest after:  {hashlib.sha256(SOURCE.read_bytes()).hexdigest()}")
    print(f"edits: {len(EDITS)} in {len(expected)} cells across {len(touched)} rows; plan cells written: "
          f"{len(plan_changed)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
