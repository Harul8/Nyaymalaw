"""Regroup the Before Build sheet around the legal brain, and add the loop rows.

OWNER DIRECTION, 26 September 2026, given in conversation:

  * Arrive stays the first section.
  * The legal brain comes next and must be stronger. The model works
    autonomously in a loop -- decide, act, verify, decide next -- within
    guiding principles, not fixed rules that work against its capabilities.
  * The harness checks everything the model produces. Zero invention and zero
    hallucination are non-negotiable.
  * The loop, the context management and the scaffolding are all part of the
    legal brain, and retrieval is the ground everything in it stands on.
  * Build beside the existing pipeline and switch on evidence.

WHAT THIS CHANGES AND WHAT IT DOES NOT. Every existing row keeps its own
content -- all ten columns, byte for byte. What moves is WHERE a row sits, and
the header rows that say what each group is. Fourteen new rows (LB-126..139)
record the loop, the harness and the context design; they are mirrored into
the Implementation Plan sheet as the reconciler requires.

The discipline is the one the other workbook tools here keep: snapshot every
cell and every native sheet feature, write, reload the SAVED file and prove
that nothing moved except what was meant to. The source is replaced only after
that proof passes.
"""
import copy
import hashlib
import re
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

_REPO = Path(__file__).resolve().parents[2]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))
from assurance.common._console import utf8_console  # noqa: E402

utf8_console()

SOURCE = _REPO / "docs/Nyaymalaw_Implementation_Plan.xlsx"
TODAY = "26 September 2026"

#: The first row of the body this tool rewrites. Rows 5..70 are Arrive and are
#: left exactly where they are.
FIRST = 71

READINESS = (
    f"{TODAY}: owner direction given in conversation -- the model works "
    "autonomously in a loop within guiding principles, a strong harness checks "
    "everything it produces, zero invention is non-negotiable, and the loop is "
    "built beside the existing pipeline and switched on evidence. The wording "
    "of this row is drafted by Claude from that direction; it is not a "
    "verbatim owner statement and not an approved requirement until the owner "
    "reviews it. Implementation NOT_STARTED; acceptance NOT_RUN. Delivery "
    "mapping OPEN under LB-43; quality and release decisions OPEN under LB-44."
)

#: The ten Before Build columns of each new row. Column 1 is
#: "<ID>\n<group>\n<title>\n\n<description>" -- the shape every existing LB row
#: has, and the Implementation Plan mirror reads the title off line 3.
NEW_ROWS = {
    "LB-126": ("L.1", [
        "Guiding principles the model reads every turn",
        "One principles document tells the model how a careful advocate works. "
        "It guides; it does not script. The owner edits it without a code change.",
        "Have NM work the way a careful advocate works because it has been told "
        "how, not because a pipeline forces each step.",
        "Every turn, before the model decides anything. The owner edits the "
        "principles; the model reads them.",
        "One principles document, owned by the product owner and versioned in "
        "the repository, read by the model at the start of every turn. "
        "Distilled from what is already authored: OM-P01-14, the advocate "
        "tenets in the archived PRD (D16), and the working rules in CLAUDE.md. "
        "Written as how a careful advocate works -- establish the facts, ask "
        "rather than assume, cite only what was read, verify before answering, "
        "say not assessed rather than guess -- not as steps to execute. The "
        "document's version is recorded on every turn, so an answer can be "
        "traced to the principles it was produced under.",
        "The model's context carries the current principles; the turn record "
        "names their version.",
        "A missing or unreadable principles document stops the loop with a "
        "visible reason; the model never runs without them. A change takes "
        "effect on the next turn, never mid-turn.",
        "Must be one document with one owner. Never duplicate a principle in "
        "code as a hidden rule: a rule that must hold whatever the model does "
        "belongs in the harness (L.4), not here. Never let the model edit its "
        "own principles.",
        "LB-126-AC1: the same matter run under two principle versions records "
        "each version on its turns.\n"
        "LB-126-AC2 (planted): a principle copied into the turn path as a "
        "hard-coded rule is refused by a repository check.\n"
        "LB-126-AC3: a missing principles document yields a visible refusal, "
        "not a turn run without them.",
        "OM-P01-14; archived PRD D16 tenets; CLAUDE.md. OPEN: the owner's "
        "review of the first distilled draft.",
    ]),
    "LB-127": ("L.2", [
        "The model calls tools, whichever provider serves it",
        "Tool calling is added to the model layer once, and each provider's "
        "adapter translates it. Nothing in the core names a provider.",
        "Get the same autonomous behaviour from Opus, Fable or an OpenAI model "
        "without the product depending on any one of them.",
        "Any turn in which the model needs information or an action.",
        "Add one method to the model port that sends the conversation and the "
        "tool definitions and returns either tool calls or a final answer. "
        "Implemented for the OpenAI adapter, an Anthropic adapter and the "
        "scripted test model. Provider differences in tool-call format stay "
        "inside the adapters. Tool definitions are declared once, in the core, "
        "and translated by each adapter.",
        "A tool call from any supported provider reaches the same tool with "
        "the same arguments.",
        "A malformed tool call is returned to the model as an error it can "
        "correct, never executed on a guess. A provider that cannot call tools "
        "is reported as unsupported for the loop, never silently downgraded.",
        "Never let a provider-specific tool format reach the core. Never "
        "execute a tool the harness did not declare.",
        "LB-127-AC1: one scripted conversation produces identical tool "
        "invocations through the OpenAI and Anthropic adapters.\n"
        "LB-127-AC2 (planted): a call to an undeclared tool is refused and "
        "returned to the model as an error.\n"
        "LB-127-AC3 (planted): a malformed argument set is returned for "
        "correction and never executed.",
        "backend/nm/ports/model.py (complete, structured and embed today; no "
        "tool calling); adapters/model/openai_adapter.py; scripted.py. Model "
        "identifiers stay pinned, never floating aliases (the port's existing "
        "rule).",
    ]),
    "LB-128": ("L.2", [
        "The reasoning loop: decide, act, verify, decide again",
        "The model works a matter the way an advocate does, choosing what to do "
        "next, within declared budgets, with every step logged.",
        "Receive help from a model that looks things up, checks and revises, "
        "rather than filling a fixed sequence of forms.",
        "Every turn once a matter is open.",
        "A loop runner gives the model the principles, the matter context and "
        "the message; runs the tool calls it asks for; returns the results; and "
        "repeats until the model submits an answer or a budget is reached. "
        "Budgets for steps, tokens, cost and elapsed time are declared, not "
        "tuned per call. Every step -- the model's choice, the tool, its "
        "arguments, the result and the check outcomes -- is logged against the "
        "turn so the owner can review what the model did and why. The model "
        "decides the order of work; no stage order is imposed. Built beside "
        "the existing turn engine behind a switch (LB-138).",
        "The turn ends with an answer that passed the harness, a question to "
        "the advocate, or a budget stop that says what was not finished.",
        "A budget stop is disclosed with what remains and resumes on the next "
        "turn. A tool failure is returned to the model as a result, never "
        "swallowed. An interrupted turn resumes from its log, never from "
        "scratch with duplicated effects.",
        "Never impose a fixed order of analysis. Never loop without a budget. "
        "Never present a budget stop as a finished answer. Never lose the step "
        "log.",
        "LB-128-AC1: on a golden conversation the loop reaches an answer with a "
        "complete step log.\n"
        "LB-128-AC2 (planted): a model that never submits is stopped at the "
        "step budget, and the stop is disclosed with what remains.\n"
        "LB-128-AC3: the order of work differs across two different matters -- "
        "no hidden fixed sequence.",
        "LB-01, LB-61, LB-64, LB-69, LB-70; MAX_EVIDENCE_ROUNDS and TracedModel "
        "today. OPEN: budget values, set from golden-set runs.",
    ]),
    "LB-129": ("L.2", [
        "What NM can already do becomes tools; closed lists only at the tool's door",
        "The retrieval, arithmetic and curated tables that work today become "
        "tools the model calls when it judges they are needed.",
        "Have the model use NM's retrieval, calculations and curated tables when "
        "it needs them, and reason in its own words everywhere else.",
        "The model decides it needs a provision, an authority, a computation, "
        "the matter file, a curated table, or an answer from the advocate.",
        "Expose what already works as declared tools: read a provision; search "
        "authorities; compute limitation; read the matter file; record a fact "
        "with its source; consult a curated table (elements, pre-institution "
        "conditions, governing code, interim tests, procedural periods, filing "
        "requirements); ask the advocate, which ends the turn with a question; "
        "and the check tools in LB-131. Each tool states what it does and what "
        "it cannot do, and returns held, not held or not assessed. Closed "
        "vocabularies (cause of action, relief, role) are required only as a "
        "tool's input key: the model reasons about any cause in its own words, "
        "and a cause a table does not cover returns 'no curated table' so the "
        "model continues with retrieval.",
        "Every capability the turn engine uses today is reachable as a tool "
        "with a declared contract.",
        "A tool that cannot run returns not assessed with the reason; the model "
        "decides what to do next.",
        "Never let a closed list stop the analysis of a matter it does not "
        "cover. Never identify an Act by fuzzy match inside a tool (CLAUDE.md "
        "section 5).",
        "LB-129-AC1: a partition suit -- a cause outside the closed list -- is "
        "analysed through retrieval rather than stopping at 'not established'.\n"
        "LB-129-AC2: every tool returns one of the three states, and a not "
        "assessed result is visible in the step log.\n"
        "LB-129-AC3 (planted): an Act named with a shared word resolves to no "
        "Act rather than the wrong one.",
        "backend/nm/knowledge/*; backend/nm/ports/*; LB-120-125. OPEN: the "
        "first tool set, chosen for the first pull request.",
    ]),
    "LB-130": ("L.3", [
        "Every answer is a set of claims, each tied to its source",
        "Zero invention made checkable: every statement the advocate reads "
        "points to the retrieved text or the advocate's own words it rests on.",
        "Rely on the fact that nothing NM says is invented.",
        "The model submits an answer.",
        "The answer is submitted as claims. Each claim carries its kind -- fact "
        "from the file, law from a retrieved provision, authority from a "
        "retrieved judgment, inference drawn from named claims, question, or "
        "limit -- and its support: the retrieved span with its locator, or the "
        "advocate's own words and where they said them. An inference names the "
        "claims it is drawn from. The advocate reads natural prose; the claim "
        "structure is what the harness checks and what opens when the advocate "
        "inspects a basis.",
        "Every sentence the advocate reads maps to a supported claim, a "
        "question or a stated limit.",
        "An unsupported claim goes to the repair step (LB-133). If it stays "
        "unsupported it is removed, or the answer is withheld with the reason.",
        "ZERO INVENTION, NON-NEGOTIABLE: never show a claim whose support was "
        "not retrieved on this matter or said by the advocate. Never cite from "
        "model memory. Never let an inference present itself as a fact.",
        "LB-130-AC1 (planted): an answer citing a provision the loop never "
        "retrieved is refused.\n"
        "LB-130-AC2: every sentence of a golden answer maps to a supported "
        "claim.\n"
        "LB-130-AC3: an inference names its premises and changes when a premise "
        "is corrected.",
        "G-GROUND, G-ATTRIB, G-QUOTE; LB-25, LB-57, LB-59. An extension of how "
        "grounding already works, not a new mechanism.",
    ]),
    "LB-131": ("L.4", [
        "The harness checks every answer: the eighteen output checks",
        "Whatever path the model took, what reaches the advocate has been "
        "checked for support, currency, consistency and honest absence.",
        "Know that whatever the model decided to do, what reaches me has been "
        "checked.",
        "Before any answer reaches the advocate, every time.",
        "The eighteen output checks run on every answer the loop produces: "
        "support (G-GROUND, G-QUOTE, G-ATTRIB); law in force and binding "
        "(G-INFORCE, G-BINDING, G-DATE); consistency with itself and the file "
        "(G-CONSISTENT, G-CONSERVE, G-CASCADE, G-CURRENCY, G-STALE); honest "
        "absence (G-NOTHELD, G-HELDNOTFOUND, G-NOTASSESSED, G-READ, G-MODEL); "
        "coverage (G-COVERAGE, G-COMPETENCE). The same checks are offered to "
        "the model as tools so it can check its own work before submitting; "
        "the final run is always the harness's, never the model's.",
        "An answer is served only when the checks pass or their disclosures "
        "are attached.",
        "A failed check goes to the repair step (LB-133).",
        "Never let the model skip, disable or waive a check. Never report a "
        "check as passed that could not run: not assessed is its own state "
        "(CLAUDE.md section 9).",
        "LB-131-AC1 (planted): for each of the eighteen checks, an answer that "
        "violates it is caught on the loop path.\n"
        "LB-131-AC2: a check that cannot run is shown as not assessed, never as "
        "passed.",
        "backend/nm/domain/gates.py (the matrix) and metrics.fire. The checks "
        "exist; this row puts them after the loop.",
    ]),
    "LB-132": ("L.4", [
        "Professional boundaries stay fixed",
        "Six boundaries are enforced by the harness on tool calls and answers "
        "alike, and cannot be reasoned around.",
        "Be protected by the same professional limits whatever the model "
        "decides.",
        "Any step or answer that touches danger, conflict, scope, capacity, a "
        "professional duty or unscreened material.",
        "The six boundaries are enforced by the harness: G-EMERGENCY, "
        "G-CONFLICT, G-SCOPE, G-CAPACITY, G-DUTY, G-UNSCREENED. They apply to "
        "tool calls as well as answers: a tool call outside scope, or on a "
        "conflicted matter, is refused before it runs. External effects -- "
        "sending, filing, paying, contacting anyone -- are never available as "
        "tools without the advocate's explicit approval of that specific act.",
        "The loop runs only within the boundaries; a boundary hit is shown to "
        "the advocate with its reason.",
        "A boundary that cannot be evaluated blocks, and says so (the third "
        "state).",
        "Never let the principles, a tool result or an instruction inside "
        "retrieved or uploaded material relax a boundary. Never take an "
        "external act without the advocate's approval of that act.",
        "LB-132-AC1 (planted): a tool call on a conflicted matter is refused "
        "before it executes.\n"
        "LB-132-AC2 (planted): an instruction inside an uploaded document to "
        "ignore scope has no effect.\n"
        "LB-132-AC3: an unevaluable boundary blocks with a visible reason.",
        "LB-31, LB-32, LB-33, LB-62, LB-119; gates.py.",
    ]),
    "LB-133": ("L.4", [
        "Repair: the model is told why a check failed and revises",
        "A failed check becomes feedback the model acts on, within a bound, "
        "rather than a refusal the advocate meets.",
        "Get a corrected answer rather than a refusal when the first draft "
        "fails a check.",
        "The harness finds a failed check on a submitted answer.",
        "The failing checks and their exact reasons go back to the model; it "
        "revises -- retrieving more, removing an unsupported claim, correcting "
        "an inconsistency -- and resubmits. Attempts are bounded. After the "
        "last attempt, anything still failing is withheld and the advocate is "
        "told what was withheld and why. Every attempt, its feedback and its "
        "revision are in the step log.",
        "A revised answer that passes, or a withheld part with its reason.",
        "Repair attempts count against the turn budget; exhausting them is "
        "disclosed, never hidden.",
        "Never let a repair relax the check it is repairing. Never show an "
        "intermediate draft. Never repair without a bound.",
        "LB-133-AC1 (planted): an answer with one fabricated quotation is "
        "repaired into one without it, and both attempts are logged.\n"
        "LB-133-AC2 (planted): an answer that cannot be repaired within the "
        "bound is withheld with the reason shown.\n"
        "LB-133-AC3: the served answer is the one the harness checked, byte "
        "for byte (LB-67).",
        "LB-66, LB-67. Today G-GROUND, G-QUOTE and G-ATTRIB withhold with no "
        "feedback to the model. OPEN: the attempt bound, set from golden-set "
        "runs.",
    ]),
    "LB-134": ("L.4", [
        "Thirteen judgment gates: principle or floor -- the owner decides",
        "Gates that encode advocate judgment are sorted, one by one, into "
        "floors the harness keeps and principles the model follows.",
        "Keep the safeguards that must always hold, and let the model exercise "
        "judgment where judgment is what is wanted.",
        "Owner review, before the loop replaces the pipeline for the behaviour "
        "a gate covers.",
        "Thirteen gates encode advocate judgment. For each, the owner decides "
        "whether it stays a FLOOR the harness enforces, or becomes a PRINCIPLE "
        "the model follows, with a harness check that the answer dealt with "
        "it. RECOMMENDATION FOR THE OWNER TO CORRECT. Floor: G-POSTURE (never "
        "advise without knowing the side), G-THREAD (an account bound to "
        "exactly one dispute), G-LIMITATION (no directive on an unresolved "
        "limitation), G-PREMISE (no limitation computed on an unestablished "
        "premise), G-CORRECTION (a correction replaces, never adds). Principle, "
        "with a check that it was addressed: G-SPLIT, G-GAP, G-SALVAGE, "
        "G-EXPOSURE, G-ADVERSE, G-PRESERVE, G-PROOF, G-REMEDY.",
        "Each of the thirteen has an owner decision recorded here.",
        "Until decided, a gate stays exactly as it is today.",
        "Never convert a floor into a principle without the owner's recorded "
        "decision.",
        "LB-134-AC1 (planted): for every gate converted to a principle, an "
        "answer that ignores what it covers is flagged by the 'addressed' "
        "check.",
        "gates.py; the archived advocate tenets. OPEN: all thirteen decisions.",
    ]),
    "LB-135": ("L.5", [
        "Context: the file on demand, the advocate's words verbatim, summaries checked",
        "The model keeps the whole matter within reach on long matters without "
        "losing or distorting what the advocate said.",
        "Have NM keep the whole matter in mind without losing or distorting "
        "what I said.",
        "Every turn; especially long matters and resumed sessions.",
        "The model's working context holds the principles, the current task "
        "and a compact matter summary. The matter file, documents and history "
        "are read through tools when needed rather than loaded wholesale. The "
        "advocate's own words are kept verbatim and always retrievable. When a "
        "conversation outgrows the budget, older turns are summarised; the "
        "summary names its sources, is re-checked against the file, and is "
        "never the only copy or the thing analysed (LB-58).",
        "The model can reach every material fact on the file within the "
        "context budget.",
        "A summary that disagrees with the file is discarded and rebuilt from "
        "the file.",
        "Never paraphrase the advocate's words where the exact words matter. "
        "Never analyse a summary in place of the file. Never let a summary "
        "become a source a claim can cite.",
        "LB-135-AC1: on a matter longer than the context budget, a fact stated "
        "in the first turn is used correctly in the last.\n"
        "LB-135-AC2 (planted): a summary that altered a date is detected "
        "against the file and rebuilt.\n"
        "LB-135-AC3 (planted): a claim citing a summary rather than the file or "
        "a source is refused by the harness.",
        "LB-10, LB-12, LB-58; memory.advocate_words; domain/summary.py "
        "as_context; ModelPort.context_budget.",
    ]),
    "LB-136": ("L.5", [
        "Research sub-loop with its own clean context",
        "Long reading is done by a nested loop that returns only findings with "
        "their spans, so the main conversation keeps its focus.",
        "Have NM dig through long material thoroughly without the main "
        "conversation losing focus.",
        "The model decides a question needs extensive reading -- a long "
        "judgment, a bundle of documents, many candidate authorities.",
        "A research tool starts a nested loop with a fresh context and a stated "
        "question. It reads what it needs and returns only its findings, each "
        "with the span and locator it rests on. The findings enter the main "
        "loop as tool results and are checked by the harness like everything "
        "else. A nested loop has its own budget, which counts against the "
        "turn.",
        "Findings with supporting spans, or a stated 'not found' naming what "
        "was searched.",
        "A research loop that exhausts its budget returns what it found and "
        "what it did not cover.",
        "Never return a finding without its span. Never let a research loop "
        "take an external act or reach material outside the matter's "
        "authorisation.",
        "LB-136-AC1: a question over a long judgment returns the relevant "
        "paragraph with its locator and nothing else.\n"
        "LB-136-AC2 (planted): a finding without a span is refused when it "
        "reaches the main loop.",
        "LB-105, LB-106; LB-128 budgets.",
    ]),
    "LB-137": ("L.3", [
        "Retrieval chosen by the model, identification kept exact",
        "The model decides what to look up and how often; which Act is read is "
        "still decided by exact match only.",
        "Have NM look up what the matter actually needs, and know that every "
        "Act and section it reads is the one it names.",
        "The model decides it needs law or authority.",
        "Retrieval is a set of tools the model calls as often as the budget "
        "allows, with its own queries, rather than a fixed number of rounds "
        "decided by code. Fuzzy search may RANK paragraphs and authorities; "
        "which Act is read is decided by exact match only (CLAUDE.md section "
        "5). Every result names the index it came from, and a zero names its "
        "index (B-163). Retrieved text is what claims cite (LB-130).",
        "The model has read what it relies on, and each reading is in the "
        "step log.",
        "A store that cannot be consulted returns not assessed; the model "
        "continues or discloses the gap.",
        "Never identify an Act by overlap. Never report absence from one store "
        "as absence from the corpus. Never let retrieved text carry an "
        "instruction the model follows.",
        "LB-137-AC1: a question naming s.53A of the Transfer of Property Act "
        "retrieves from that Act and no other.\n"
        "LB-137-AC2: a zero result names its index.\n"
        "LB-137-AC3 (planted): a retrieved passage containing an instruction "
        "does not change the model's behaviour.",
        "LB-100-108; backend/nm/domain/citation.py; MAX_EVIDENCE_ROUNDS, to be "
        "replaced by the loop budget.",
    ]),
    "LB-138": ("L.8", [
        "Build beside the pipeline, compare, then switch",
        "The loop is built alongside what works today, compared on the golden "
        "set, and switched on only when the evidence says it is better.",
        "Gain the new way of working without losing what works today.",
        "Each build slice of the loop.",
        "The loop is built beside the existing turn engine behind a switch, "
        "reading and writing the same matters. The golden set is run through "
        "both, with the owner's approval of each run, and compared on "
        "grounding, correctness and the advocate-facing outcome. The switch "
        "moves only when the loop matches or beats the pipeline; the pipeline "
        "parts it replaces are then removed so no second path survives "
        "(LB-73). Slices, one pull request each, reviewed by the owner: (1) "
        "tool calling, loop runner, step log; (2) principles and tools; (3) "
        "output checks and repair; (4) context and research; (5) switch, "
        "comparison and live progress in the interface.",
        "Each slice is merged only after owner review; the switch moves on a "
        "recorded comparison.",
        "A slice that regresses the golden set is not merged.",
        "Never run the golden set without the owner's approval of that run. "
        "Never leave two live paths after the switch.",
        "LB-138-AC1: each slice's pull request carries its test results and, "
        "from slice 3, a golden comparison.\n"
        "LB-138-AC2: after the switch, no code path reaches a retired pipeline "
        "stage.",
        "LB-40, LB-43, LB-44, LB-73, LB-74; docs/GOLDEN_SET.md; "
        "assurance/journeys/run_goldens.py.",
    ]),
    "LB-139": ("U.1", [
        "Show the work as it happens, and let the advocate correct it",
        "Each step of the loop is shown in plain words while it runs, and the "
        "advocate can stop it or correct a fact mid-turn.",
        "See what NM is doing while it works, and correct it before it goes "
        "down the wrong path.",
        "Any turn in which the loop takes more than one step.",
        "Stream each step to the interface in plain words -- 'reading s.138 of "
        "the Negotiable Instruments Act', 'checking limitation', 'asking you "
        "about service' -- never in internal identifiers. The advocate can stop "
        "the turn or correct a fact mid-turn; a correction is used by the next "
        "step and recorded like any other correction. The full step log stays "
        "inspectable after the turn.",
        "The advocate sees progress and can intervene; the finished answer is "
        "unchanged by the display.",
        "A lost connection keeps the turn's log; reopening shows where it "
        "reached.",
        "Never show an unchecked draft as an answer. Never show internal "
        "identifiers. Never let a progress line imply a check passed that has "
        "not run.",
        "LB-139-AC1: a multi-step turn shows each step in plain words.\n"
        "LB-139-AC2: a correction made mid-turn is used by the next step and "
        "recorded.",
        "LB-70, LB-82, LB-86; test_no_enum_value_reaches_the_advocate.",
    ]),
}

#: THE STRUCTURE. (code, level, title, what it is, objective, members). A member
#: is a row key: an ID, or one of the four preserved section notes.
GROUPS = [
    ("L", 1, "Legal brain \u2014 the model works freely within guiding principles; the harness checks everything it produces",
     f"Owner direction, {TODAY}. The models act autonomously in a loop -- decide, act, verify, decide next -- guided by "
     "principles rather than fixed rules that work against their capabilities. The harness checks every output. Zero "
     "invention and zero hallucination are non-negotiable: nothing that is not supported by retrieved material or the "
     "advocate's own words reaches the advocate. Retrieval is the ground everything here stands on.\n\n"
     "Every existing row keeps its own wording and readiness; this grouping changes where a row sits, not what it says. "
     "Where a preserved row prescribes a fixed step that the model should now decide, the L.2 loop rows govern and the "
     "row is read as a principle to follow, not a pipeline stage.",
     "Receive expert, fully grounded help from a model that reasons freely and is checked on everything it says.",
     ["NOTE:LB"]),
    ("L.0", 2, "Entry \u2014 open a matter and take the brief",
     "Opening a durable file and receiving the advocate's material in any medium. The F-B and F-C rows are preserved "
     "descriptions: column A keeps the prior wording and B-J supersede it where they conflict.",
     "Open a file and give NM the brief without being forced through a script.",
     ["OM-Q01", "OM-Q02", "OM-Q03", "OM-Q04", "OM-Q05", "OM-Q06", "OM-I01", "OM-I02", "OM-I03",
      *[f"F-B-{n:02d}" for n in range(1, 17)], *[f"F-C-{n:02d}" for n in range(1, 13)],
      "LB-02", "LB-37"]),
    ("L.1", 2, "Guiding principles \u2014 how NM works with the advocate",
     "What the model is told about how a careful advocate works. These guide; they do not script. A rule that must hold "
     "whatever the model does belongs in L.4, not here.",
     "Work with an assistant that behaves like a careful, candid colleague.",
     ["NOTE:OM", *[f"OM-P{n:02d}" for n in range(1, 15)], "LB-04", "LB-17", "LB-39", "LB-113", "LB-126"]),
    ("L.2", 2, "The reasoning loop and its tools",
     "The model decides what to do next, does it through declared tools, reads the result and decides again, within "
     "declared budgets and with every step logged.",
     "Receive help from a model that works the matter rather than filling a fixed sequence of forms.",
     ["LB-127", "LB-128", "LB-129", "LB-01", "LB-03", "LB-05", "LB-35", "LB-36", "LB-61", "LB-64", "LB-69",
      "LB-70", "LB-110", "LB-116", "LB-118"]),
    ("L.3", 2, "Retrieval and grounding \u2014 everything rests on retrieved text",
     "What the model reads, how it is found, and how every claim is tied to it. Fuzzy search may rank; only exact "
     "match identifies an Act.",
     "Rely on every statement being traceable to the text it rests on.",
     ["LB-130", "LB-137", "LB-11", "LB-13", "LB-14", "LB-25", "LB-26", "LB-27", "LB-57", "LB-59",
      *[f"LB-{n}" for n in range(100, 109)]]),
    ("L.4", 2, "The harness \u2014 output checks, professional boundaries and repair",
     "What holds whatever the model decides: the eighteen output checks, the six professional boundaries, and the repair "
     "step that turns a failed check into a correction.",
     "Know that nothing reaches me unchecked, and that professional limits cannot be reasoned around.",
     ["LB-131", "LB-132", "LB-133", "LB-134", "LB-31", "LB-32", "LB-33", "LB-38", "LB-60", "LB-62", "LB-63",
      "LB-66", "LB-67", "LB-71", "LB-114", "LB-119"]),
    ("L.5", 2, "Context and memory",
     "Keeping the whole matter within reach across long conversations, corrections and resumed sessions, without the "
     "file being replaced by a summary of it.",
     "Have NM remember the matter accurately and act on its current state.",
     ["LB-135", "LB-136", "LB-08", "LB-10", "LB-12", "LB-29", "LB-30", "LB-34", "LB-58", "LB-68"]),
    ("L.6", 2, "Legal reasoning and advice",
     "The substance: what is established, what must be proved, what could defeat the case, and what course to take.",
     "Receive candid, expert-quality analysis and advice on the matter in front of me.",
     ["LB-06", "LB-07", "LB-09", "LB-15", "LB-16", "LB-18", "LB-19", "LB-20", "LB-21", "LB-22", "LB-23", "LB-28",
      *[f"LB-{n}" for n in range(45, 56)], "LB-65", "LB-109", "LB-112", "LB-115", "LB-117"]),
    ("L.7", 2, "Indian practice layer \u2014 curated tables the model consults as tools",
     "The procedural craft that decides Indian matters before their merits are reached. Under the loop these tables are "
     "tools the model calls (LB-129), not pipeline stages.",
     "Never lose a matter to a procedural step nobody checked.",
     ["NOTE:PRACTICE", *[f"LB-{n}" for n in range(120, 126)]]),
    ("L.8", 2, "Models, evaluation and build discipline",
     "How quality is measured, how models are chosen and changed, and how the loop replaces the pipeline without two "
     "live paths.",
     "Trust that NM is measured before it is relied on, and changed without losing what works.",
     ["LB-138", "LB-40", "LB-43", "LB-44", "LB-56", "LB-72", "LB-73", "LB-74"]),
    ("U", 1, "Legal-brain interface \u2014 the conversation, the matter board and the sources",
     "How the advocate works with the legal brain on screen. Preserved rows keep their wording; LB-139 adds live progress "
     "for the loop.",
     "Work with NM comfortably and inspect the basis of anything it says.",
     []),
    ("U.1", 2, "Conversation, matter board and progress",
     "The working surface: one conversation beside the matter board, on every intended device.",
     "Move between conversation and the current position of the matter without losing context.",
     ["LB-139", "LB-24", "LB-41", "LB-75", "LB-76", "LB-80", "LB-81", "LB-82", "LB-83", "LB-84", "LB-85", "LB-86",
      "LB-88", "LB-89", "LB-90", "LB-93", "LB-111"]),
    ("U.2", 2, "Sources and the basis of claims",
     "Opening the exact source behind a claim, and seeing what was actually checked.",
     "Inspect the source of any statement with one action, without losing my place.",
     ["LB-42", "LB-77", "LB-78", "LB-79", "LB-87", "LB-91", "LB-92", "LB-94", "LB-95", "LB-96", "LB-97", "LB-98",
      "LB-99"]),
    ("P", 1, "Later journey phases \u2014 work the file, advise, act, carry, close and leave",
     "Preserved descriptions of the phases after the brief. Column A keeps the prior wording; B-J supersede it where "
     "they conflict. Under the loop, these phases are what the model does when the matter calls for it, not a sequence "
     "it must pass through.",
     "Carry a matter from analysis through action to a clean close.",
     ["NOTE:OTHER", *[f"F-D-{n:02d}" for n in range(1, 15)], *[f"F-E-{n:02d}" for n in range(1, 6)],
      *[f"F-F-{n:02d}" for n in range(1, 8)], *[f"F-G-{n:02d}" for n in range(1, 4)], "F-H-01", "F-H-02",
      "F-I-01"]),
    ("X", 1, "Diagnostics \u2014 measured findings and evidence, not feature acceptance",
     "Records of measured defects and the evidence behind their repair. They are not requirements to build.",
     "Keep what was measured visible without mistaking it for acceptance.",
     [f"DG-{n:02d}" for n in range(1, 19)]),
]

#: The four old section headers, preserved as notes and placed with their group.
NOTES = {
    "NOTE:OTHER": "Other phases — preserved descriptions and targeted refinements",
    "NOTE:OM": "Open a matter — approved interaction principles",
    "NOTE:LB": "Take the brief and the continuing legal brain",
    "NOTE:PRACTICE": "Indian practice layer",
}

KEY = re.compile(r"(LB-\d+|OM-[PIQ]\d+|DG-\d+|F-[A-Z]-\d+)\b")
LBOM = re.compile(r"(LB-\d+|OM-[PIQ]\d+)\b")
MIRROR = {8: 1, 9: 4, 10: 2, 12: 3, 15: 4, 16: 5, 17: 6, 19: 7, 28: 9, 30: 10, 33: 8}
NOTE_PATTERN = re.compile(r"All \d+ LB/OM requirements are linked into Implementation Plan")
ZEBRA = ("00F3F6F7", "00FFFFFF")
#: The heading separator the Arrive section already uses.
DASH = " \u2014 "


def sheet_features(sheet):
    """Native sheet controls, not only visible cell values."""
    from openpyxl.xml.functions import tostring
    return {
        "merged": tuple(str(area) for area in sheet.merged_cells.ranges),
        "freeze": sheet.freeze_panes,
        "state": sheet.sheet_state,
        "validations": tostring(sheet.data_validations.to_tree()),
        "conditional": repr(list(sheet.conditional_formatting)),
        "views": tostring(sheet.views.to_tree()),
        "protection": tostring(sheet.protection.to_tree()),
        "margins": tostring(sheet.page_margins.to_tree()),
        "print_options": tostring(sheet.print_options.to_tree()),
        "columns": {key: (dim.width, dim.hidden, dim.outlineLevel)
                    for key, dim in sheet.column_dimensions.items()},
        "hyperlinks": {cell.coordinate: (cell.hyperlink.target, cell.hyperlink.location)
                       for row in sheet for cell in row if cell.hyperlink},
    }


def key_of(value) -> str:
    text = str(value or "")
    for note, lead in NOTES.items():
        if text.startswith(lead):
            return note
    match = KEY.match(text)
    if not match:
        raise ValueError(f"unclassified Before Build row: {text[:80]!r}")
    return match[1]


def main() -> int:
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    book = load_workbook(SOURCE)
    sheet, plan = book["Before Build"], book["Implementation Plan"]
    features = {s.title: sheet_features(s) for s in book}
    before = {(s.title, c.coordinate): (c.value, copy.copy(c._style))
              for s in book for row in s for c in row}
    heights = {r: sheet.row_dimensions[r].height for r in range(1, sheet.max_row + 1)}

    # ---- snapshot every row being regrouped, by its key ----------------------
    old_last = sheet.max_row
    source: dict[str, dict] = {}
    for r in range(FIRST, old_last + 1):
        key = key_of(sheet.cell(r, 1).value)
        assert key not in source, f"duplicate row key {key}"
        source[key] = {
            "values": [sheet.cell(r, c).value for c in range(1, 11)],
            "styles": [copy.copy(sheet.cell(r, c)._style) for c in range(1, 11)],
            "height": sheet.row_dimensions[r].height,
            "row": r,
        }

    # ---- the new rows, in the existing row shape -----------------------------
    group_title = {code: title for code, _lvl, title, *_ in GROUPS}
    template = source["LB-121"]
    for ident, (group, cols) in NEW_ROWS.items():
        assert ident not in source, f"{ident} already exists"
        title, lead, *rest = cols
        area = "Legal-brain interface" if group.startswith("U") else "Legal brain"
        section = f"{area} › {group} {group_title[group].split(DASH)[0]}"
        column_a = f"{ident}\n{section}\n{title}\n\nOWNER-DIRECTED REQUIREMENT, DRAFTED FOR REVIEW:\n{lead}"
        values = [column_a, *rest, READINESS]
        assert len(values) == 10, ident
        source[ident] = {"values": values,
                         "styles": [copy.copy(s) for s in template["styles"]],
                         "height": 280, "row": None}

    # ---- every row placed exactly once --------------------------------------
    placed = [m for *_, members in GROUPS for m in members]
    assert len(placed) == len(set(placed)), "a row is placed twice"
    missing = set(source) - set(placed)
    extra = set(placed) - set(source)
    assert not missing, f"rows left unplaced: {sorted(missing)}"
    assert not extra, f"rows placed that do not exist: {sorted(extra)}"

    # ---- write the regrouped body -------------------------------------------
    level_style = {1: [copy.copy(sheet.cell(5, c)._style) for c in range(1, 11)],
                   2: [copy.copy(sheet.cell(6, c)._style) for c in range(1, 11)]}
    changed: set = set()
    layout: list[tuple[str, int]] = []

    def put(r, c, value, style):
        cell = sheet.cell(r, c)
        cell.value = value
        cell._style = style
        changed.add(("Before Build", cell.coordinate))

    r = FIRST
    for code, level, title, what, objective, members in GROUPS:
        header = [f"{code}  {title}\n\n{what}", objective, None, None, None, None, None, None, None,
                  f"{TODAY}: grouping adopted on owner direction. Rows keep their own content and readiness."]
        for c, value in enumerate(header, 1):
            put(r, c, value, copy.copy(level_style[level][c - 1]))
        sheet.row_dimensions[r].height = 118 if level == 1 else 74
        layout.append((f"HEADER:{code}", r))
        r += 1
        for index, member in enumerate(members):
            row = source[member]
            for c in range(1, 11):
                style = copy.copy(row["styles"][c - 1])
                put(r, c, row["values"][c - 1], style)
                sheet.cell(r, c).fill = PatternFill("solid", fgColor=ZEBRA[index % 2])
                font = copy.copy(sheet.cell(r, c).font)
                font.b = False
                sheet.cell(r, c).font = font
            sheet.row_dimensions[r].height = row["height"]
            layout.append((member, r))
            r += 1
    last = r - 1
    for extra_row in range(last + 1, old_last + 1):
        for c in range(1, 11):
            put(extra_row, c, None, copy.copy(sheet.cell(extra_row, c)._style))
    sheet.auto_filter.ref = f"A4:J{last}"

    # ---- the row-2 note, measured from the sheet ------------------------------
    population = sorted({LBOM.match(str(sheet.cell(x, 1).value or ""))[1]
                         for x in range(5, last + 1)
                         if LBOM.match(str(sheet.cell(x, 1).value or ""))})
    lb = [p for p in population if p.startswith("LB-")]
    note = (
        f"Grouped {TODAY} on owner direction: A Arrive; L Legal brain (L.0 entry, L.1 guiding principles, "
        "L.2 the reasoning loop and its tools, L.3 retrieval and grounding, L.4 the harness, L.5 context and memory, "
        "L.6 legal reasoning and advice, L.7 Indian practice layer, L.8 models, evaluation and build discipline); "
        "U legal-brain interface; P later journey phases; X diagnostics. Original descriptions remain in A; B–J "
        f"contains current requirements. The legal brain holds {len(lb)} LB requirements; LB-126–139 record the "
        "autonomous loop, the harness and context management and are drafted for owner review. Legal explanation and "
        "retrieved verbatim passages remain inline by default. "
        f"All {len(population)} LB/OM requirements are linked into Implementation Plan; release assignment and "
        "acceptance mapping remain OPEN. DIAGNOSTIC rows record measured findings and evidence, not feature "
        "acceptance. No application build or expert-quality judgment is established by reconciliation."
    )
    assert NOTE_PATTERN.search(note)
    put(2, 1, note, copy.copy(sheet.cell(2, 1)._style))

    # ---- the mirror for the new rows -----------------------------------------
    position = dict(layout)
    for ident, (group, cols) in NEW_ROWS.items():
        assert not any(plan.cell(x, 1).value == ident for x in range(2, plan.max_row + 1)), ident
        target = plan.max_row + 1
        for c in range(1, plan.max_column + 1):
            plan.cell(target, c)._style = copy.copy(plan.cell(target - 1, c)._style)
        src = position[ident]
        fixed = {1: ident, 3: "Requirement", 4: "Legal brain", 5: f"{group} {group_title[group].split(DASH)[0]}",
                 6: cols[0], 7: "Pilot", 38: "Draft", 39: "Not started", 40: "Not verified",
                 32: f"{TODAY}: owner-directed loop and harness row mirrored from Before Build. Not acceptance, "
                     "not an approved requirement."}
        for c, value in fixed.items():
            plan.cell(target, c, value)
            changed.add(("Implementation Plan", plan.cell(target, c).coordinate))
        for t, origin in MIRROR.items():
            plan.cell(target, t, sheet.cell(src, origin).value)
            changed.add(("Implementation Plan", plan.cell(target, t).coordinate))

    scratch = Path("/tmp/claude-0/-home-user-Nyaymalaw/82a504d3-1374-501f-9fe3-441ac9aa966e/scratchpad")
    scratch.mkdir(parents=True, exist_ok=True)
    out = scratch / SOURCE.name
    book.save(out)

    # ================= THE PROOF, on the saved bytes ==========================
    check = load_workbook(out)
    assert check.sheetnames == book.sheetnames
    for tab in check:
        assert sheet_features(tab) == features[tab.title], f"sheet features moved on {tab.title}"
    bb = check["Before Build"]

    # 1. Arrive and the header block are untouched, bar the row-2 note.
    for (title, coord), old in before.items():
        if title == "Before Build":
            row_number = int(re.sub(r"[A-Z]+", "", coord))
            if row_number < FIRST and coord != "A2":
                cell = bb[coord]
                assert (cell.value, cell._style) == old, coord
    for x in range(1, FIRST):
        assert bb.row_dimensions[x].height == heights.get(x), f"row {x} height moved"

    # 2. Every moved row keeps all ten values exactly. Only its zebra fill and
    #    bold weight may differ; border, alignment, number format and the rest
    #    of the font are compared against the ORIGINAL file, not memory.
    from openpyxl.xml.functions import tostring

    def xml(style_part):
        return tostring(style_part.to_tree())

    original = load_workbook(SOURCE)["Before Build"]
    for member, row_number in layout:
        if member.startswith("HEADER:"):
            continue
        want = source[member]
        got = [bb.cell(row_number, c).value for c in range(1, 11)]
        assert got == want["values"], f"{member}: content changed"
        if want["row"] is None:
            continue
        for c in range(1, 11):
            a, o = bb.cell(row_number, c), original.cell(want["row"], c)
            # SERIALISED, because openpyxl's `Border.__eq__` reported a
            # mismatch between two borders identical on every attribute. The
            # XML is what is written to the file, so it is what is compared.
            assert xml(a.border) == xml(o.border), f"{member}: border moved"
            assert xml(a.alignment) == xml(o.alignment), f"{member}: alignment moved"
            assert a.number_format == o.number_format, f"{member}: number format moved"
            assert (a.font.name, a.font.sz, a.font.color) == (o.font.name, o.font.sz, o.font.color), \
                f"{member}: font moved"

    # 3. The population: nothing lost, nothing doubled, only the new rows added.
    after_keys = [key_of(bb.cell(x, 1).value) for x in range(FIRST, bb.max_row + 1)
                  if bb.cell(x, 1).value and not str(bb.cell(x, 1).value).split("  ", 1)[0] in
                  {code for code, *_ in GROUPS}]
    assert sorted(after_keys) == sorted(source), "row population changed"
    assert len(after_keys) == len(set(after_keys))

    # 4. The Implementation Plan: every pre-existing cell unchanged.
    ip = check["Implementation Plan"]
    for (title, coord), old in before.items():
        if title == "Implementation Plan" and (title, coord) not in changed:
            cell = ip[coord]
            assert (cell.value, cell._style) == old, f"plan {coord} moved"

    # 5. The contract the reconciler enforces, and the feature population.
    from assurance.control_plane import plan_scenarios
    problems = plan_scenarios.requirement_problems(out)
    assert not problems, problems
    assert len(plan_scenarios.sheet_rows(out)) == 93, "the feature population moved"

    shutil.copy2(out, SOURCE)
    print(f"source digest before: {digest}")
    print(f"source digest after:  {hashlib.sha256(SOURCE.read_bytes()).hexdigest()}")
    print(f"rows regrouped: {len(source) - len(NEW_ROWS)}; rows added: {len(NEW_ROWS)}; "
          f"headers: {len(GROUPS)}; last row: {last}")
    print(f"LB/OM requirements: {len(population)} (LB {len(lb)})")
    for code, level, title, _w, _o, members in GROUPS:
        print(f"  {'  ' if level == 2 else ''}{code:<4} {len(members):>3}  {title.split(DASH)[0]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
