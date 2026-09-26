"""Give the legal brain its own Tools section, and add the tool catalogue.

OWNER DIRECTION, 26 September 2026: add the tool design and catalogue to the
plan, as part of the legal brain, placed properly.

WHERE IT GOES. Directly after the loop that uses the tools: L.3 Tools. The
sections after it move down by one -- retrieval becomes L.4, the harness L.5,
context L.6, legal reasoning L.7, the practice layer L.8, models and evaluation
L.9 -- and "The reasoning loop and its tools" becomes "The reasoning loop".
LB-129, which already said that NM's capabilities become tools, moves into L.3.

WHAT RENUMBERING TOUCHES, measured before writing: section numbers such as
"L.4" appear only in the section headers, the row-2 note, the owner-directed
rows LB-126..153 and those rows' Implementation Plan mirrors. No row the owner
authored carries one. The renumbering is applied to exactly that population
and the proof checks every other row is byte-identical.

Nine new rows, LB-154..162: the tool design rules and registry, then the
catalogue by family. Each is mirrored into the Implementation Plan sheet.
"""
import copy
import hashlib
import importlib.util
import re
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

_REPO = Path(__file__).resolve().parents[2]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))
from assurance.common._console import utf8_console  # noqa: E402

utf8_console()

_spec = importlib.util.spec_from_file_location(
    "regroup", Path(__file__).with_name("legal_brain_regroup_20260926.py"))
regroup = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(regroup)

SOURCE = regroup.SOURCE
TODAY = regroup.TODAY
FIRST = regroup.FIRST
READINESS = regroup.READINESS
HEADER = re.compile(r"^(L|L\.\d|U|U\.\d|P|X)  ")
SECTION_REF = re.compile(r"(?<![A-Za-z0-9-])L\.([0-9])(?![0-9])")
#: Old section number -> new. L.0-L.2 keep theirs; L.3 is the new Tools section.
RENUMBER = {3: 4, 4: 5, 5: 6, 6: 7, 7: 8, 8: 9}
OLD_LOOP = "The reasoning loop and its tools"
NEW_LOOP = "The reasoning loop"
TOOLS_TITLE = "Tools"

TOOLS_HEADER = (
    "L.3  Tools — what the model can do, and how every tool behaves\n\n"
    "The model acts only through declared tools. Most wrap code that already exists and is tested -- the "
    "search port, the evidence port, the limitation and deadline arithmetic, the curated tables -- so the "
    "work is the common envelope, the registry and a few new tools, not new legal logic. The professional "
    "boundaries are not tools: they run around every call. Actions that leave the product are proposals "
    "the advocate approves.",
    "Have NM use the right capability at the right moment, and know that every result it relied on is "
    "sourced, located and checkable.",
)


def renumber(text):
    if not isinstance(text, str):
        return text
    text = SECTION_REF.sub(lambda m: f"L.{RENUMBER.get(int(m[1]), int(m[1]))}", text)
    return text.replace(OLD_LOOP, NEW_LOOP)


NEW_ROWS = {
    "LB-154": [
        "How every tool is shaped: the envelope, the rules and the registry",
        "Every tool follows the same rules and returns the same envelope, and "
        "is declared once in a registry with its description, schema and tests.",
        "Rely on every tool result being sourced, located and honest about what "
        "it did not find.",
        "Any tool is added, changed or called.",
        "THE ENVELOPE: every result carries a status (found / not held / not "
        "found / not assessed / error), the source and its exact locator, the "
        "index or table consulted, and what was excluded and why; a zero never "
        "looks like absence (B-163). THE RULES: one job per tool, named in an "
        "advocate's words; the description says what the tool does AND what it "
        "cannot do, because the description is what the model chooses by; tools "
        "that identify never rank, and tools that rank never identify (CLAUDE.md "
        "section 5); large results come as a preview with a handle and are paged "
        "(LB-149); reads may run in parallel, writes run one at a time and carry "
        "the file version they rest on (LB-152); errors return to the model as "
        "results with a way forward, never swallowed; arithmetic lives in tools "
        "(LB-144); actions that leave the product are proposals the advocate "
        "approves (LB-132); the professional boundaries run around every call "
        "and are not tools the model can choose to skip. THE REGISTRY: each tool "
        "is declared once -- name, description, input schema, whether it is "
        "parallel-safe, whether it reads or writes or proposes, whether it is "
        "core or discovered on demand (LB-148) -- and each ships with its tests, "
        "including a planted failure its check must catch.",
        "Every tool the model can call is in the registry with its contract and "
        "tests.",
        "A tool whose result does not fit the envelope is refused before the "
        "model sees it (LB-143).",
        "Never declare a tool outside the registry. Never let a tool return a "
        "bare empty result. Never let a description promise what the tool does "
        "not do.",
        "LB-154-AC1: every registered tool returns the envelope on success, "
        "on 'not found' and on error.\n"
        "LB-154-AC2 (planted): a tool returning an empty list with no status is "
        "refused by the harness.\n"
        "LB-154-AC3: a repository check fails the build on a tool called by "
        "the loop that is not in the registry.",
        "LB-127, LB-129, LB-143, LB-144, LB-148, LB-149, LB-152. BUILD ORDER: "
        "core first (the envelope and registry; reading the matter; the statute "
        "and case-law reads and searches; compute_limitation; ask_advocate; "
        "submit_answer; check_claims), then the checked writes, the curated "
        "tables, date arithmetic, matter search, paged documents and read_turn, "
        "then verify_support, research, oppose, find_contrary_authority, "
        "playbooks and find_tool, propose_action and compute_interest.",
    ],
    "LB-155": [
        "Matter-file tools: read and write the file",
        "The model reads the matter through tools and writes to it only through "
        "checked tools.",
        "Have NM work from my actual file, and add to it only what it can show "
        "came from me or a document.",
        "The model needs the state of the matter, or has established, corrected "
        "or withdrawn something on it.",
        "READ: matter_brief (the current state, each line tagged with source and "
        "status); read_thread (one dispute -- posture, chronology, premises, "
        "issues, deadlines); read_facts (facts with the advocate's own words and "
        "where they said them); search_matter (full-text over the file and the "
        "uploads -- NEW); read_document (an upload, one page range at a time -- "
        "partly exists); read_turn (an earlier turn verbatim, used after a "
        "compaction -- NEW). WRITE, under the rules of LB-142: record_fact "
        "(must quote the advocate or a document), record_premise (stated with a "
        "quote, or inferred with its basis), record_issue and record_decision "
        "(each names what it rests on), correct_entry and withdraw_entry (the "
        "old value is kept in history and everything resting on it is marked "
        "for review). Reads wrap the existing matter model; writes wrap its "
        "existing update paths with the checks in front.",
        "The model can reach any material part of the file, and every write it "
        "makes carries its source.",
        "A read of something not on the file returns 'not on the file', naming "
        "what was searched; a refused write returns the reason.",
        "Never let a read tool return another matter's material. Never let a "
        "write land without its source or on a stale version.",
        "LB-155-AC1: search_matter finds a date the advocate mentioned only in "
        "an uploaded letter, with its page.\n"
        "LB-155-AC2 (planted): record_fact with a quotation absent from the "
        "file is refused.\n"
        "LB-155-AC3: read_turn returns an earlier turn word for word after a "
        "compaction.",
        "backend/nm/domain/matter.py (Matter, Thread, facts), "
        "backend/nm/domain/summary.py, backend/nm/edge/uploads.py; LB-142, "
        "LB-145, LB-150, LB-152.",
    ],
    "LB-156": [
        "Statute tools: identify the Act, read the provision as it stood",
        "Statutes are identified by exact match and read as in force on the "
        "date that matters; searching ranks provisions but never picks the Act.",
        "Know that every section NM reads is from the Act it names, in the form "
        "that applied on my date.",
        "The model needs a provision, or needs to find which provision governs.",
        "identify_act(name): exact title match only, returning 'no Act' rather "
        "than a near miss. read_provision(act, section, as_of): the text in "
        "force on that date, with its locator and the store it came from. "
        "search_provisions(query, act?): ranks provisions by relevance and "
        "never decides which Act applies. governing_code(dates): which of the "
        "replaced or replacing codes governs the conduct, the proceeding and the "
        "evidence (LB-120). All wrap existing code: the manifest's exact "
        "resolution, the evidence port's fetch, and the governing-law table.",
        "The provision the model relies on is the right one, in the right form.",
        "A provision not held returns 'not held', naming the store; a provision "
        "held but not retrieved returns 'held, not found' as the defect it is "
        "(G-HELDNOTFOUND).",
        "Never identify an Act by shared words. Never read a provision without "
        "its in-force date. Never report absence from one store as absence from "
        "the corpus.",
        "LB-156-AC1: 'section 53A of the Transfer of Property Act' reads the "
        "TPA and no other Act.\n"
        "LB-156-AC2 (planted): 'Indian Easements Act 1882' does not resolve to "
        "the Evidence Act or the TPA.\n"
        "LB-156-AC3: read_provision for a 2025 offence under a repealed code "
        "returns the replacing code's provision with the reason.",
        "backend/nm/knowledge/manifest.py (exact resolution), the evidence port "
        "(fetch), backend/nm/knowledge/governing_law.py, "
        "backend/nm/domain/citation.py; LB-13, LB-14, LB-107, LB-120, LB-137.",
    ],
    "LB-157": [
        "Case-law tools: find, read, resolve, weigh",
        "Judgments are searched by case, read in context, resolved by exact "
        "citation, and weighed for treatment, binding force and bench.",
        "Rely on authority that binds my court and has not been overruled, read "
        "in its own words and context.",
        "The model needs authority on a point, or needs to test an authority it "
        "has.",
        "search_authorities(query, court, years): results grouped by case, not "
        "scattered snippets. read_judgment(case, query): the judgment's index "
        "and the paragraphs that matched, paged by locator. read_paragraph"
        "(locator). resolve_citation(citation): exact key only -- a reporter "
        "citation resolves or it does not. treatment(case): later treatment -- "
        "overruled, doubted, followed. binds_this_court(case): binding status by "
        "court and date against the matter's forum. rank_authorities(cases): "
        "which supersedes which by court and bench (LB-122). "
        "find_contrary_authority(proposition): NEW -- searches for authority "
        "against or distinguishing a proposition, used by the model and by the "
        "opposing-counsel loop (LB-140). All but the last wrap the search port, "
        "the citator, the jurisdiction rule and the identity index.",
        "Every authority relied on has been read in context and weighed.",
        "A citation that does not resolve returns 'not resolved', never a "
        "similar case; an unbuilt index returns 'not assessed'.",
        "Never identify a case by its name alone (case names reached 0.83% of "
        "held judgments; reporter citations 90.9%). Never rely on an authority "
        "without its treatment and binding status.",
        "LB-157-AC1: resolve_citation on a held reporter citation returns that "
        "judgment and no other.\n"
        "LB-157-AC2: an authority later overruled carries that treatment when "
        "read.\n"
        "LB-157-AC3 (planted): find_contrary_authority on a proposition with a "
        "known contrary Supreme Court ruling returns it.",
        "The search port (search, discover, expand, passage, resolve, treatment, "
        "case_identity); backend/nm/knowledge/citator.py, jurisdiction.py, "
        "identity.py, authority_weight.py; LB-101, LB-106, LB-122.",
    ],
    "LB-158": [
        "Computation tools: every figure computed, never estimated",
        "Limitation, dates, deadlines, interest and fees are computed by tools "
        "that return the figure, the rule and the inputs.",
        "Rely on every date and figure NM gives without re-computing it.",
        "The model needs a date, a period, an amount or a deadline.",
        "compute_limitation: the figure, the Article or rule applied, the "
        "inputs, and whether it is conditional on an inferred premise. "
        "date_arithmetic: add, between, and -- once the court calendar is held "
        "-- holiday-aware dates (NEW). add_deadline and list_deadlines: the "
        "deadline register. compute_interest (NEW). court_fee: 'not assessed' "
        "until the Telangana schedule is held and versioned (LB-125). "
        "procedural_period: through the procedural-period table (LB-124). "
        "limitation, the register and the curated tables already exist; the "
        "tool is the envelope around them.",
        "Every figure in an answer traces to a computation (LB-144).",
        "A computation missing an input returns 'not assessed' with the input "
        "named, never a default.",
        "Never compute a figure in the model. Never present a conditional "
        "figure as definitive. Never estimate a fee from an unversioned "
        "schedule.",
        "LB-158-AC1: compute_limitation returns the Article, the accrual date "
        "used and the expiry, and labels an inferred accrual conditional.\n"
        "LB-158-AC2: date_arithmetic across a leap year and a month end matches "
        "a hand-checked table.\n"
        "LB-158-AC3: court_fee returns 'not assessed' and names the missing "
        "schedule.",
        "backend/nm/core/limitation.py and deadlines.py; "
        "backend/nm/knowledge/filing_requirement.py, procedural_period.py; "
        "LB-20, LB-114, LB-124, LB-125, LB-144.",
    ],
    "LB-159": [
        "Practice-table tools: the curated Indian practice layer, as tools",
        "The curated tables are consulted when the model judges they apply, "
        "and each answers with its source and one of the three states.",
        "Never lose a matter to a procedural step the model did not think to "
        "check.",
        "The model is working a matter of a kind a curated table covers.",
        "elements_of(cause): what must be proved. limitation_article_for(cause). "
        "pre_institution_steps(cause, against): notices, mediation, waiting "
        "periods (LB-121). interim_test(relief): the test for the interim "
        "order sought (LB-123). procedural_periods(role, track) (LB-124). "
        "filing_requirements: forum, valuation, fee (LB-125). "
        "playbook(practice_area): owner-edited guidance for a kind of matter "
        "(LB-148, NEW). Each table tool takes its closed-vocabulary key only at "
        "the tool's door; a matter the table does not cover returns 'no curated "
        "table' and the model continues with retrieval (LB-129).",
        "The curated craft reaches the model whenever the matter engages it.",
        "A key the table does not hold returns 'no curated table', never a "
        "neighbouring entry.",
        "Never let a table entry stand in for the provision it points to -- "
        "the provision is still read (LB-156). Never let a playbook assert law.",
        "LB-159-AC1: a cheque-dishonour matter reaches the demand-notice "
        "condition through pre_institution_steps.\n"
        "LB-159-AC2 (planted): interim_test for an uncurated relief returns 'no "
        "curated table', not the injunction test.\n"
        "LB-159-AC3: every table result names its curated_from source.",
        "backend/nm/knowledge/elements.py, resolution.py, institution.py, "
        "interim_relief.py, procedural_period.py, filing_requirement.py, "
        "governing_law.py; LB-120-125, LB-148.",
    ],
    "LB-160": [
        "Checking and delegation tools: check the work, research deeply, argue the other side",
        "The model can check its own draft, hand deep reading to a research "
        "loop, and send a draft to opposing counsel -- and the harness still "
        "runs its own final checks.",
        "Receive answers that were checked and attacked before they reached me.",
        "The model has a draft, or a question that needs extensive reading.",
        "check_claims(draft): the eighteen output checks run early, at the "
        "model's choice; the harness runs them again at the end regardless "
        "(LB-131). verify_support(claim, locator): the independent meaning "
        "check (LB-141, NEW). check_consistency(draft): the draft against the "
        "file (exists). research(question, budget): a fresh-context reading "
        "loop returning findings with their spans (LB-136, NEW). oppose(draft): "
        "the opposing-counsel loop returning objections with their spans "
        "(LB-140, NEW). Delegation tools spend from the turn's budget and their "
        "results come back through the envelope like any other tool.",
        "A draft reaches submission already checked, supported and tested "
        "against the other side.",
        "A check or delegation that cannot run says so, and the answer carries "
        "that limit.",
        "Never let the model's own check replace the harness's final one. Never "
        "let a delegated loop exceed its budget or reach outside the matter.",
        "LB-160-AC1: check_claims on a draft with an unsupported quotation "
        "names it before submission.\n"
        "LB-160-AC2: research over a long judgment returns only findings with "
        "locators.\n"
        "LB-160-AC3: oppose on a draft that ignores a limitation bar raises it.",
        "backend/nm/domain/gates.py, backend/nm/core/consistency.py; LB-131, "
        "LB-136, LB-140, LB-141.",
    ],
    "LB-161": [
        "Advocate and action tools, and the answer itself as a tool",
        "The model asks the advocate through a tool, proposes external acts "
        "for approval, and submits its answer as structured claims -- never as "
        "free prose the harness cannot check.",
        "Be asked only what matters, stay in control of every external act, "
        "and get answers whose every statement is checkable.",
        "The model needs a fact only the advocate has, wants an external act "
        "done, or is ready to answer.",
        "ask_advocate(question, why_it_matters, what_would_change): ends the "
        "turn with one decisive question and records it so it is not asked "
        "again (partly exists: the asked-question record). propose_action(kind, "
        "draft): a notice, filing, email or service -- held for the advocate's "
        "explicit approval of that act, then executed through the existing "
        "action and service paths (partly exists). submit_answer(claims): THE "
        "ANSWER IS A TOOL CALL. Each claim carries its kind -- fact, law, "
        "authority, inference, question, limit -- and its support; the harness "
        "checks the claims and renders them as natural prose (LB-130, NEW). "
        "Because the answer can only arrive this way, zero invention is "
        "structural rather than hoped for.",
        "The advocate sees one decisive question, an approval request, or a "
        "checked answer.",
        "An answer submitted with an unsupported claim goes to repair (LB-133); "
        "an unapproved action is never executed.",
        "Never let an external act run without approval of that act. Never "
        "accept an answer outside submit_answer. Never ask a question already "
        "answered on the file.",
        "LB-161-AC1 (planted): an answer returned as plain text instead of "
        "through submit_answer is refused.\n"
        "LB-161-AC2 (planted): propose_action for a notice does not send it "
        "before approval.\n"
        "LB-161-AC3: a question already answered on the file is not asked "
        "again.",
        "Matter.asked, action_proposals, service; backend/nm/core/service.py; "
        "LB-05, LB-31, LB-130, LB-132, LB-133.",
    ],
    "LB-162": [
        "Discovering tools on demand, and what is deliberately not a tool",
        "A small core is always loaded and the rest is found when needed; some "
        "capabilities are withheld on purpose, and the row says why.",
        "Have NM use the capability it needs without being given ones that "
        "would undermine grounding or control.",
        "The model needs a capability outside the core, or the catalogue is "
        "extended.",
        "find_tool(need) and load_playbook(practice_area) search the registry "
        "and the playbooks and append what is found (LB-148; NEW). NOT GIVEN, "
        "deliberately: open web search -- it is not the verified corpus, and "
        "law found there is not grounded (a later option only for named "
        "official sources, with the same envelope and checks); a general "
        "command line or code runner -- an opaque action the harness cannot "
        "gate or audit; direct database access; any tool that edits the "
        "principles, the registry or the harness itself.",
        "The model's tool set is the core plus what it found, and nothing that "
        "bypasses the harness.",
        "A need no tool meets is reported to the advocate as a limit, never "
        "met by improvisation.",
        "Never add a tool that bypasses the envelope or the boundaries. Never "
        "let the model extend its own tool set beyond the registry.",
        "LB-162-AC1: a matter needing the interim test finds and loads "
        "interim_test through find_tool.\n"
        "LB-162-AC2 (planted): a request to search the open web is met with "
        "the stated limit, not an improvised search.",
        "LB-148, LB-154; the registry.",
    ],
}


def main() -> int:
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    book = load_workbook(SOURCE)
    sheet, plan = book["Before Build"], book["Implementation Plan"]
    features = {s.title: regroup.sheet_features(s) for s in book}
    before = {(s.title, c.coordinate): (c.value, copy.copy(c._style))
              for s in book for row in s for c in row}
    heights = {r: sheet.row_dimensions[r].height for r in range(1, sheet.max_row + 1)}

    def owner_directed(key):
        m = re.match(r"LB-(\d+)$", key)
        return bool(m) and 126 <= int(m[1]) <= 153

    # ---- the current body ------------------------------------------------------
    old_last = sheet.max_row
    order = []
    for r in range(FIRST, old_last + 1):
        text = str(sheet.cell(r, 1).value or "")
        header = HEADER.match(text)
        key = f"HEADER:{header[1]}" if header else regroup.key_of(text)
        order.append((key, {
            "values": [sheet.cell(r, c).value for c in range(1, 11)],
            "styles": [copy.copy(sheet.cell(r, c)._style) for c in range(1, 11)],
            "height": sheet.row_dimensions[r].height, "row": r}))
    keys = [k for k, _ in order]
    for ident in NEW_ROWS:
        assert ident not in keys, ident

    # ---- which rows may change, and how: exactly the renumbering population ---
    def renumbered(key, values):
        if key.startswith("HEADER:") or owner_directed(key):
            values = [renumber(v) for v in values]
        if key == "LB-129":
            lines = values[0].split("\n")
            lines[1] = f"Legal brain › L.3 {TOOLS_TITLE}"
            values[0] = "\n".join(lines)
        return values

    # ---- rebuild the body: L.2 loses LB-129; the new L.3 follows L.2 ---------
    body = []
    group = None
    template = dict(order)["LB-153"]
    tools_rows = []
    for key, row in order:
        if key.startswith("HEADER:"):
            if group == "L.2":
                body.append(("HEADER:TOOLS", {
                    "values": [TOOLS_HEADER[0], TOOLS_HEADER[1], None, None, None, None, None, None, None,
                               f"{TODAY}: section added on owner direction; the sections after it moved "
                               "down by one."],
                    "styles": [copy.copy(s) for s in dict(order)["HEADER:L.2"]["styles"]],
                    "height": 118, "row": None}))
                body.extend(tools_rows)
            group = key.split(":", 1)[1]
        if key == "LB-129":
            tools_rows.insert(0, (key, row))
            continue
        body.append((key, row))
        if key == "HEADER:L.2":
            for ident, cols in NEW_ROWS.items():
                title, lead, *rest = cols
                column_a = (f"{ident}\nLegal brain › L.3 {TOOLS_TITLE}\n{title}\n\n"
                            f"OWNER-DIRECTED REQUIREMENT, DRAFTED FOR REVIEW:\n{lead}")
                values = [column_a, *rest, READINESS]
                assert len(values) == 10, ident
                tools_rows.append((ident, {"values": values,
                                           "styles": [copy.copy(s) for s in template["styles"]],
                                           "height": 280, "row": None}))
    assert len(body) == len(order) + len(NEW_ROWS) + 1

    # ---- write -----------------------------------------------------------------
    changed = set()
    layout = []
    zebra_groups = {"L.2", "L.3"}
    group, index = None, 0
    for offset, (key, row) in enumerate(body):
        r = FIRST + offset
        is_header = key.startswith("HEADER:")
        values = row["values"] if row["row"] is None else renumbered(key, list(row["values"]))
        if is_header:
            group = "L.3" if key == "HEADER:TOOLS" else key.split(":", 1)[1]
            if row["row"] is not None:
                old_code = group
                group = f"L.{RENUMBER[int(old_code[2])]}" if re.fullmatch(r"L\.[3-8]", old_code) else old_code
            index = 0
        for c in range(1, 11):
            cell = sheet.cell(r, c)
            if (cell.value, cell._style) != (values[c - 1], row["styles"][c - 1]):
                changed.add(("Before Build", cell.coordinate))
            cell.value = values[c - 1]
            cell._style = copy.copy(row["styles"][c - 1])
            if not is_header and group in zebra_groups:
                cell.fill = PatternFill("solid", fgColor=regroup.ZEBRA[index % 2])
                changed.add(("Before Build", cell.coordinate))
        if not is_header:
            index += 1
        sheet.row_dimensions[r].height = row["height"]
        layout.append((key, r, values))
    last = FIRST + len(body) - 1
    sheet.auto_filter.ref = f"A4:J{last}"

    # ---- the row-2 note --------------------------------------------------------
    population = sorted({regroup.LBOM.match(str(sheet.cell(x, 1).value or ""))[1]
                         for x in range(5, last + 1)
                         if regroup.LBOM.match(str(sheet.cell(x, 1).value or ""))})
    lb = [p for p in population if p.startswith("LB-")]
    note = str(sheet.cell(2, 1).value)
    old_structure = re.search(r"L Legal brain \(.*?\);", note)
    assert old_structure, "the row-2 note no longer lists the legal-brain structure"
    note = note.replace(old_structure[0], (
        "L Legal brain (L.0 entry, L.1 guiding principles, L.2 the reasoning loop, L.3 tools, "
        "L.4 retrieval and grounding, L.5 the harness, L.6 context and memory, L.7 legal reasoning and "
        "advice, L.8 Indian practice layer, L.9 models, evaluation and build discipline);"))
    note = note.replace("LB-126–153", "LB-126–162")
    note, n = re.subn(r"The legal brain holds \d+ LB requirements",
                      f"The legal brain holds {len(lb)} LB requirements", note)
    assert n == 1
    note, n = regroup.NOTE_PATTERN.subn(
        f"All {len(population)} LB/OM requirements are linked into Implementation Plan", note)
    assert n == 1
    note = note.replace("the opposing-counsel loop and context management",
                        "the opposing-counsel loop, context management and the tool catalogue")
    sheet.cell(2, 1).value = note
    changed.add(("Before Build", "A2"))

    # ---- the Implementation Plan: renumber the mirrors, add the new rows ------
    position = {k: r for k, r, _v in layout}
    for x in range(2, plan.max_row + 1):
        ident = plan.cell(x, 1).value
        if isinstance(ident, str) and owner_directed(ident):
            for c in range(1, plan.max_column + 1):
                v = plan.cell(x, c).value
                new = renumber(v)
                if ident == "LB-129" and c == 5:
                    new = f"L.3 {TOOLS_TITLE}"
                if new != v:
                    plan.cell(x, c).value = new
                    changed.add(("Implementation Plan", plan.cell(x, c).coordinate))
            for t, origin in regroup.MIRROR.items():
                want = sheet.cell(position[ident], origin).value
                if plan.cell(x, t).value != want:
                    plan.cell(x, t).value = want
                    changed.add(("Implementation Plan", plan.cell(x, t).coordinate))
    for ident, cols in NEW_ROWS.items():
        target = plan.max_row + 1
        for c in range(1, plan.max_column + 1):
            plan.cell(target, c)._style = copy.copy(plan.cell(target - 1, c)._style)
        fixed = {1: ident, 3: "Requirement", 4: "Legal brain", 5: f"L.3 {TOOLS_TITLE}", 6: cols[0],
                 7: "Pilot", 38: "Draft", 39: "Not started", 40: "Not verified",
                 32: f"{TODAY}: owner-directed tool row mirrored from Before Build. Not acceptance, "
                     "not an approved requirement."}
        for c, v in fixed.items():
            plan.cell(target, c, v)
            changed.add(("Implementation Plan", plan.cell(target, c).coordinate))
        for t, origin in regroup.MIRROR.items():
            plan.cell(target, t, sheet.cell(position[ident], origin).value)
            changed.add(("Implementation Plan", plan.cell(target, t).coordinate))

    out = Path("/tmp/claude-0/-home-user-Nyaymalaw/82a504d3-1374-501f-9fe3-441ac9aa966e/scratchpad") / SOURCE.name
    book.save(out)

    # ================= THE PROOF, on the saved bytes ==========================
    from openpyxl.xml.functions import tostring

    def xml(part):
        return tostring(part.to_tree())

    check = load_workbook(out)
    for tab in check:
        assert regroup.sheet_features(tab) == features[tab.title], tab.title
    bb = check["Before Build"]
    original = load_workbook(SOURCE)["Before Build"]
    for (title, coord), old in before.items():
        if title == "Before Build" and int(re.sub(r"[A-Z]+", "", coord)) < FIRST and coord != "A2":
            assert (bb[coord].value, bb[coord]._style) == old, coord
    for x in range(1, FIRST):
        assert bb.row_dimensions[x].height == heights.get(x)

    moved = dict(order)
    for key, r, _v in layout:
        if key in NEW_ROWS or key == "HEADER:TOOLS":
            continue
        # Header keys were renamed by renumbering; find the source row by its old key.
        source_key = key
        want = moved[source_key]
        expected = renumbered(source_key, list(want["values"]))
        got = [bb.cell(r, c).value for c in range(1, 11)]
        assert got == expected, f"{key}: content changed beyond the renumbering"
        if not (source_key.startswith("HEADER:") or owner_directed(source_key)):
            assert got == want["values"], f"{key}: an owner-authored row changed"
        for c in range(1, 11):
            a, o = bb.cell(r, c), original.cell(want["row"], c)
            assert xml(a.border) == xml(o.border) and xml(a.alignment) == xml(o.alignment), key
            assert (a.font.name, a.font.sz, a.font.color, a.font.b) == \
                (o.font.name, o.font.sz, o.font.color, o.font.b), key

    after = [k for k, _r, _v in layout]
    assert sorted(after) == sorted(keys + list(NEW_ROWS) + ["HEADER:TOOLS"]), "row population changed"
    heads = [(str(bb.cell(r, 1).value).split("  ", 1)[0], r) for k, r, _v in layout if k.startswith("HEADER:")]
    assert [h for h, _ in heads] == ["L", "L.0", "L.1", "L.2", "L.3", "L.4", "L.5", "L.6", "L.7", "L.8",
                                     "L.9", "U", "U.1", "U.2", "P", "X"], heads
    tools_at = dict(heads)["L.3"]
    next_at = dict(heads)["L.4"]
    for ident in list(NEW_ROWS) + ["LB-129"]:
        assert tools_at < position[ident] < next_at, f"{ident} is outside L.3"
    assert "L.3 Tools" in bb.cell(position["LB-129"], 1).value

    ip = check["Implementation Plan"]
    for (title, coord), old in before.items():
        if title == "Implementation Plan" and (title, coord) not in changed:
            assert (ip[coord].value, ip[coord]._style) == old, f"plan {coord} moved"
    from assurance.control_plane import plan_scenarios
    problems = plan_scenarios.requirement_problems(out)
    assert not problems, problems
    assert len(plan_scenarios.sheet_rows(out)) == 93

    shutil.copy2(out, SOURCE)
    print(f"source digest before: {digest}")
    print(f"source digest after:  {hashlib.sha256(SOURCE.read_bytes()).hexdigest()}")
    print(f"rows added: {len(NEW_ROWS)}; last row: {last}; LB/OM: {len(population)} (LB {len(lb)})")
    for h, r in heads:
        print(f"  {h:<4} row {r}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
