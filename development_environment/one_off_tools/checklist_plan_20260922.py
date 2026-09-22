"""Preserve the curated workbook; reconcile only authorised checklist/reader rows."""
import copy
import hashlib
import json
import shutil
from pathlib import Path

from openpyxl import load_workbook
from chat_reference_plan_20260922 import sheet_features
from reconcile_implementation_20260922 import preview

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'docs/Nyaymalaw_Implementation_Plan.xlsx'
OUT = ROOT / 'outputs/01a07b76-6b21-71f3-bb09-261f64617594'
ROWS = [
    ['LB-117\nDispute requirements\nA source-backed checklist within the conversation',
     'Understand what each dispute needs and what the file currently supports.',
     'A substantive conversation identifies a dispute or retrieves relevant law.',
     'Read applicable requirements from retrieved passages, with exact source anchors. Distinguish required from strengthening support. Do not create inapplicable rows. Derive four labelled states from attributed answers and current facts: held, outstanding, promised or unavailable. Preserve prior rows and answers across turns, failures and restarts. A changed source or corrected fact requires review, not silent deletion or continued green status.',
     'Asking is finished only with no outstanding items; follow-up is finished only with no outstanding or promised items. Neither means the dispute is legally won, the file is closed, or advice is authorised. An unestablished checklist cannot pass by being empty.',
     'No retrieved support, failed reading, incomplete output, changed text and invalid stored records remain distinguishable. Cache completed passage readings by content identity and dispute context, not locator alone. A failed read preserves existing work and remains retryable.',
     'No hard-coded legal lists, model-memory requirements, manual ticks, unsupported factual promotion, hidden inapplicable-item register or silent erasure. An advocate reporting a document exists does not mean NM has examined it. Handover uses the same validated state as the board.',
     'Prove source-span checks, empty and failed reads, duplicate identity, additive merge, changed-source invalidation, restart persistence, four states, stale/cross-dispute fact rejection and both completion conditions. Exercise the served conversation and source opening.',
     'F-B-17; LB-109–116; LB-92. Shared permission, grounding and maturity gates continue to apply.',
     '22 September 2026: PARTIAL. Domain, persistence, served-turn and source-reader checks pass. LB-100/101 remain unchanged. Classification/entailment and changed-dispute applicability still need independent live measurement; passage presence is not semantic proof. See conversational-checklist.md.'],
    ['LB-118\nConversational follow-up\nUpdate the checklist from ordinary replies',
     'Keep working naturally without filling a separate checklist or repeating facts.',
     'An advocate gives information, corrects it, says it is unknown/unavailable, promises it, or resumes a matter.',
     'Use the same conversation input read to propose source-bound answers for exact dispute and requirement IDs. The application validates quotations and scoped current facts before changing the record. Answer the immediate request first; ask at most one small group of decision-changing outstanding items. Keep required versus strengthening support distinct. For unavailable material explain its purpose and a supported course without it, or state that no alternative is established.',
     'Board, conversation and handover agree. Unknown stays outstanding with the answer retained; unavailable is not re-asked without a new reason; promised stays pending until fulfilled, withdrawn or corrected. No promise becomes evidence merely because time passed.',
     'Bring dated promises back when due and undated promises back on a later visit, without a repeated same-session interrogation. Preserve the promised expression when a date is ambiguous. Record reminders separately from statutory time limits. Failed parsing or wrong-target proposals leave previous answers intact.',
     'No separate mandatory form, extra per-reply classifier, automatic case closure, status inferred from silence, invented dates, accusation or canned sympathy. An unavailable item is not by itself a verdict; other supported proof routes and all existing gates remain relevant.',
     'Drive natural replies through the production turn: held/unknown/unavailable/promised, multiple disputes in one message, correction, replay, failed read, restart, due and undated return. Plant unknown IDs, fake quotes and conflicting updates. Verify accessible labels, no colour-only status and no manual tick endpoint.',
     'F-C-13; F-B-17; COMM-02/03; LB-117. Live legal reasoning and communication quality require separate evidence.',
     '22 September 2026: PARTIAL. Reply-read updates, strict attribution, preserved history and promise reminders implemented; controlled tests pass. Undated returns use authenticated session change or later day, not an arbitrary page refresh. Live checklist movement remains NOT ASSESSED: the new live run stopped at LB-119 before substance.'],
]
ROWS.append([
    'LB-119\nDiagnostic: capacity correction after opening\nA blocked assessment has an accessible correction path',
    'Record or revise the human capacity assessment without recreating a matter or bypassing a safeguard.',
    'Live Chrome run, 22 September: an unchecked opening assessment blocks substance; the opened matter has no visible way to amend it.',
    'Offer an explicit capacity assessment control in the matter cover after opening. Keep not assessed, in doubt and assessed not in doubt distinct. Require the advocate to choose and give a basis; authenticate the actor and record server time. Preserve prior assessments. Never infer clearance from age, instructions, login or a model interpretation.',
    'The owned matter records the assessment at the observed version, and the next substantive turn reads it through the existing screen. Recording it makes no model call and grants no settlement, filing or external-action authority.',
    'Wrong owner, stale versions, blank basis and forged actor/time are refused; revoking a favourable assessment restores the restriction. Changing accounts/matters cannot apply a late response to the wrong file.',
    'No automatic tick, clinical determination, privileged gate bypass, silent history rewrite or forced duplicate matter. The old blocked conversation remains visible.',
    'Exercise the actual route and browser control: explicit positive/negative/unassessed records, ownership, CSRF, version race, reload and historical preservation. A live follow-on requires a fresh bounded retest after the material diagnostic.',
    'B6; BK-91; LB-118. Diagnostic from Kavita Sen browser run; not a legal-quality pass.',
    'IMPLEMENTED / controlled tests passed (50-test capacity, cover and workspace run). Owner Chrome session also saved the explicit assessment through the new form and displayed confirmation; the model ledger remained at 25 calls. Live post-fix substance NOT ASSESSED. Two live GPT-4o mini calls, USD0.000469; paid messages stopped. Remaining original USD1 ledger balance USD0.990914.'
])


def main():
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    book = load_workbook(SOURCE)
    before = {(s.title, c.coordinate): (c.value, copy.copy(c._style))
              for s in book for row in s for c in row}
    features = {s.title: sheet_features(s) for s in book}
    changed = set()
    sheet = book['Before Build']
    indices = {}

    def write(tab, row, col, value):
        tab.cell(row, col, value)
        changed.add((tab.title, tab.cell(row, col).coordinate))

    for values in ROWS:
        ident = values[0].splitlines()[0]
        matches = [r for r in range(5, sheet.max_row + 1)
                   if str(sheet.cell(r, 1).value or '').startswith(ident + '\n')]
        assert len(matches) <= 1
        row = matches[0] if matches else sheet.max_row + 1
        indices[ident] = row
        for col, value in enumerate(values, 1):
            if not matches:
                sheet.cell(row, col)._style = copy.copy(sheet.cell(row - 1, col)._style)
            write(sheet, row, col, value)
        sheet.row_dimensions[row].height = 280
    reader = next(r for r in range(5, sheet.max_row + 1)
                  if str(sheet.cell(r, 1).value or '').startswith('LB-92\n'))
    indices['LB-92'] = reader
    additions = {
        4: '22 September source-reader decision: the Full document control is at the bottom of the passage view. Open the latest authorised local source at the exact cited passage when it still occurs. Label current text separately from the historical relied-on excerpt. No silent substitution of that excerpt.',
        6: 'When the current source no longer carries the relied-on passage, show a changed-source warning and retain the historical quotation with its identity. Do not fabricate a highlight or rewrite the saved answer. Full-document failure leaves the permitted excerpt readable.',
        8: 'LB-92-AC5: Verify the bottom Full document control, exact initial anchor, latest-source identification, changed/missing passage warning, denied access, reload and return focus. Prove no model call and no mutation of saved history.',
        10: '22 September 2026: source-reader decision now recorded here. Supersedes any earlier same-version-only reading for the explicit current full-document view, not for historical evidence. Code in 0090c6c is subject to tests; no full acceptance claimed.',
    }
    for col, addition in additions.items():
        old = str(sheet.cell(reader, col).value or '')
        if addition not in old:
            write(sheet, reader, col, old + '\n\n' + addition)
    # Reconcile the previously recorded requirements too: the older builder
    # appended Before Build LB-93–116 without updating their delivery mirrors.
    import re
    for r in range(5, sheet.max_row + 1):
        match = re.match(r'(LB-\d+|OM-[PIQ]\d+)\b', str(sheet.cell(r, 1).value or ''))
        if match:
            indices[match[1]] = r
    plan = book['Implementation Plan']
    mapping = {8: 1, 9: 4, 10: 2, 12: 3, 15: 4, 16: 5,
               17: 6, 19: 7, 28: 9, 30: 10, 33: 8}
    for ident, source in indices.items():
        matches = [r for r in range(2, plan.max_row + 1) if plan.cell(r, 1).value == ident]
        assert len(matches) <= 1
        row = matches[0] if matches else plan.max_row + 1
        if not matches:
            for col in range(1, plan.max_column + 1):
                plan.cell(row, col)._style = copy.copy(plan.cell(row - 1, col)._style)
            for col, value in {1: ident, 3: 'Requirement', 4: 'Legal brain',
                               6: sheet.cell(source, 1).value.splitlines()[-1],
                               7: 'Pilot', 38: 'Recorded', 39: 'In progress',
                               40: 'Not verified'}.items():
                write(plan, row, col, value)
        for target, origin in mapping.items():
            write(plan, row, target, sheet.cell(source, origin).value)
        if any(before.get((plan.title, plan.cell(row, col).coordinate), (None,))[0]
               != plan.cell(row, col).value for col in mapping):
            write(plan, row, 32, '22 September 2026: reconciled from Before Build; not acceptance.')
    for ident, ref in [('F-B-17', 'LB-117'), ('F-C-13', 'LB-118')]:
        row = next(r for r in range(2, plan.max_row + 1) if plan.cell(r, 1).value == ident)
        write(plan, row, 28, ref + '; LB-92. Before Build owns the current requirement.')
        write(plan, row, 41, '0090c6c foundations plus conversation-driven checklist updates and source-reader checks. Controlled tests passed; live substantive checklist movement NOT ASSESSED after LB-119 blocked the browser run. Before Build owns the requirement, not this delivery note.')
    OUT.mkdir(parents=True, exist_ok=True)
    target = OUT / SOURCE.name
    book.save(target)
    check = load_workbook(target)
    assert check.sheetnames == book.sheetnames
    for tab in check:
        assert sheet_features(tab) == features[tab.title]
    for key, old in before.items():
        if key not in changed:
            cell = check[key[0]][key[1]]
            assert (cell.value, cell._style) == old, key
    preview(check['Before Build'], list(indices.values())[:2], [1, 4, 10], OUT / 'checklist-plan.png')
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == digest, 'parallel workbook edit'
    shutil.copy2(target, SOURCE)
    print(json.dumps({'new_requirements': {k: indices[k] for k in ('LB-117', 'LB-118', 'LB-119')},
                      'reconciled_owner_rows': len(indices), 'unrelated_changes': 0}))


if __name__ == '__main__':
    main()
