"""Narrow append/update of the authorised whole-file dispute requirements."""
import copy
import hashlib
import json
import shutil
from pathlib import Path

from openpyxl import load_workbook
from chat_reference_plan_20260922 import sheet_features
from reconcile_implementation_20260922 import preview
import sys

_REPO = Path(__file__).resolve().parents[2]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))
from assurance.common._console import utf8_console  # noqa: E402

utf8_console()

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'docs/Nyaymalaw_Implementation_Plan.xlsx'
OUT = ROOT / 'outputs/01a07b76-6b21-71f3-bb09-261f64617594'
ROWS = [
    ('LB-109', 'Keep each dispute distinct',
     'Identify and retain every dispute in my brief, without treating each fact or legal issue as a new dispute.',
     'First or later substantive message, with or without case numbers.',
     'Read the current instructions against existing dispute IDs; require exact current-message spans. Register separate working disputes, scope their factual accounts, and preserve unallocated instructions. Ask only about unresolved binding.',
     'Each identified dispute is visible and traceable; no silent fact, posture or date transfer.',
     'Unavailable or invalid reads stay unassessed. Never merge old threads automatically. A clarification must not manufacture another duplicate.',
     'No scenario-specific routing or similarity-based identity; quotations are evidence of instructions, not proof of their truth.',
     'Test several disputes, shared parties, continuing updates, unknown IDs, fake spans, state preservation and mixed-message chronology isolation.',
     'BK-27/C4; diagnostics from the five-turn live browser matter.'),
    ('LB-110', 'Work through the whole matter',
     'By default NM should work through all identified disputes, obtain useful missing information and revisit its assessment as evidence changes. I can redirect it.',
     'Every matter work turn and board/readiness projection.',
     'Derive an agenda from persisted assessments, gaps, pauses, questions and currency. Prioritise real urgency and useful unanswered work. Work in bounded turns; offer the next pending dispute without forcing it over a specific request.',
     'Whole-file review is complete only when every identified dispute is currently assessed with no controlling gap. This never legally closes the matter.',
     'A paused or stale dispute remains outstanding. No progress stops with an explanation; no repeated unavailable-material loop. Empty or unassessed populations cannot pass.',
     'No fabricated completion, background infinite agent loop, automatic filing or silent authority expansion.',
     'Test unvisited, partial, stale, paused, repeated question, all-current and empty files; explicit navigation and acknowledgement-only behaviour.',
     'BK-54/BK-91/C1; Agentified NM queue and bounded board are design references, not copied production proof.'),
    ('LB-111', 'One compact dispute board',
     'See all disputes and their present status in one left-hand matter board, leaving the right pane for conversation.',
     'Open, reply, reload, return from History or choose another dispute.',
     'Render concise labelled dispute rows, current next need and an explicit focus control. Retain the matter cover only once; use safe text rendering and keyboard-accessible controls.',
     'Board and briefing read the same agenda; selecting a dispute cannot change facts or complete work.',
     'Failed board loading is visible. Focus cannot leak across matters, retries or sessions. Preserve the composer and committed history.',
     'No duplicate client cards, hidden intermediate reasoning, colour-only status or ambiguous persisted selection.',
     'Verify served projection, UI selection, send envelope, refresh, long labels and unchanged conversation.',
     'BK-27/BK-54; earlier one-board presentation requirement remains in force.'),
    ('LB-112', 'Pre-filing advice without a fictional proceeding',
     'Give bounded advice for our explicitly stated prospective claim or defence even though nothing has been filed.',
     'A known client seeks or resists identified relief before institution.',
     'Keep prospective substantive position separate from filed procedural roles. Require stated support; preserve role conflicts and ordinary evidence and limitation controls.',
     'No repeated demand to say who filed when the advocate has already said nothing is filed.',
     'Unknown intended position still asks a focused question; later filed roles require fresh stated evidence.',
     'Never convert no proceedings into plaintiff/defendant, bypass side-dependent controls or propagate one dispute position to another.',
     'Test prospective seeking/resisting, unknown intent, transition to filed role, conflict and multi-dispute separation.',
     'C3/C4; browser diagnostic, not a new grant of authority.'),
    ('LB-113', 'Converse with judgment, not a joined-up audit report',
     'NM should answer my immediate request in coherent, attentive paragraphs, using corrections and explaining their effect. It must not substitute stock sympathy for analysis.',
     'Every substantive response, re-entry and saved history.',
     'Separate internal issue/evidence inventories from the advocate-facing synthesis. Preserve material risks and retrieved passages visibly. Evaluate the actual served conversation, not just prompt wording.',
     'The response answers the request, distinguishes reported from verified material and poses only decision-changing questions.',
     'A failed synthesis must disclose its limitation; never conceal unsupported conclusions by polishing the text or collapsing warnings.',
     'No generic audit dump, invented emotional state, missing adverse caveats or repeated historical banner that dominates chat.',
     'Judge multi-dispute transcripts for relevance, calibration, attentiveness, question economy and UI readability, with negative controls.',
     'BK-37/C1; DIAGNOSTIC: Meera first turn produced extensive mechanical inventories instead of the requested candid assessment.'),
    ('LB-114', 'A selected event is not an established legal trigger',
     'A precise calculation must not make an unestablished limitation premise look certain.',
     'Limitation calculation, deadline register, board and recommendation.',
     'Keep retrieved rule identity separate from application to a dated event. Model-selected accrual remains inferred even with one dated event or a curated trigger; conditional calculations never populate the live deadline field.',
     'The same premise state appears in the response, persisted register and board. No definite deadline follows from source presence alone.',
     'Failed selection is unassessed. Explicit reviewed premises remain attributable and invalidate when their supporting facts change.',
     'No sale-specific date patch; apply this rule to every cause and chronology.',
     'Test one/many dates, correct rule plus unsupported event, reviewed premise, changed fact and conditional register/read-back.',
     'BK-35/BK-65; DIAGNOSTIC: live first turn promoted a model-selected sale date to a definite 2038 deadline.'),
    ('LB-115', 'Opposing arguments must not invent facts or law',
     'Prepare candid opposition and replies without reviving corrected facts, inventing rules or finding an answer merely to help our client.',
     'Adversarial reasoning, theory and final response release.',
     'Bind material assertions to current attributed facts and applicable retrieved passages; assess their semantic support, not citation presence alone. Preserve alternatives and missing support explicitly.',
     'An unsupported proposition cannot leave as advice merely because an authentic provision is cited nearby.',
     'Correction supersession, unavailable sources, failed support review and partial assessments remain visible and block dependent conclusions.',
     'No model-memory law, selective forgetting, fabricated opposition position or claim that a prompt alone enforces grounding.',
     'Plant contradicted facts and uncited legal propositions without identifiers; the served path must reject them. Independently review live transcripts.',
     'BK-91/BK-95/D7; DIAGNOSTIC: Meera reply revived recent-only knowledge despite the 2017 email and asserted unsupported evidential rules.'),
    ('LB-116', 'Distinguish my missing information from NM unfinished work',
     'NM should not ask me to obtain its own unfinished assessment or say a comparison found nothing when disputes were excluded.',
     'Readiness, queue, partial-read failures and whole-file comparison.',
     'Keep user-obtainable needs separate from pending internal work; derive review readiness over all disputes. Compare all source-bound sibling accounts. Honour explicit move-on without completing the skipped dispute.',
     'No cannot-get-this button for an unperformed review; no complete or none-found verdict over an omitted population.',
     'No-progress or no next dispute stops honestly. Provider truncation remains failure, not a shorter answer or a clean judgment.',
     'No automatic matter closure, unbounded retries, or hidden failed factors/theory.',
     'Test mixed gaps, no gaps but unreviewed disputes, single/empty/all-paused agenda, repeated urgent current dispute and full comparison population.',
     'BK-54/BK-91/BK-29; DIAGNOSTIC: first live reply exposed user controls for NM review work plus incomplete theory/factor reads.'),
]

STATUS = {
    'LB-109': 'PARTIAL: three source-bound disputes and pre-filing roles observed live; broad update/ambiguity quality remains open.',
    'LB-110': 'PARTIAL: derived queue, pauses and override implemented; live end-to-end iteration has not passed.',
    'LB-111': 'PARTIAL: compact board and responsive controlled browser checks implemented; live three-dispute board observed.',
    'LB-112': 'PARTIAL: prospective roles observed live without a fictional filing; substantive advice is not signed off.',
    'LB-113': 'OPEN: routine screening dump reduced; coherent final synthesis and independent live communication acceptance still fail.',
    'LB-114': 'PARTIAL: general inferred-accrual correction implemented and tested locally; fresh live validation pending.',
    'LB-115': 'OPEN: material live factual/legal support failure; prompts and citation-identity checks do not close it.',
    'LB-116': 'PARTIAL: internal-work/user-gap separation, navigation and source-bound sibling comparison repaired; live truncation/routing quality remains open.',
}


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    book = load_workbook(SOURCE)
    before = {(s.title, c.coordinate): (c.value, copy.copy(c._style))
              for s in book for cells in s for c in cells}
    features = {s.title: sheet_features(s) for s in book}
    sheet = book['Before Build']
    last = sheet.max_row
    changed, numbers = set(), []
    for ident, title, *values in ROWS:
        matches = [r for r in range(5, sheet.max_row + 1)
                   if str(sheet.cell(r, 1).value or '').startswith(ident + '\n')]
        assert len(matches) < 2
        row = matches[0] if matches else sheet.max_row + 1
        numbers.append(row)
        values = [f'{ident}\nDispute-loop diagnostic\n{title}', *values,
                  '22 September 2026: ' + STATUS[ident] + ' Evidence and current status: '
                  'dispute-loop.md in the 22 September legal-brain evidence folder. '
                  'No release or professional-quality sign-off.']
        assert len(values) == 10
        for col, value in enumerate(values, 1):
            cell = sheet.cell(row, col)
            if not matches:
                cell._style = copy.copy(sheet.cell(last, col)._style)
            cell.value = value
            changed.add((sheet.title, cell.coordinate))
        sheet.row_dimensions[row].height = 240
    for table in sheet.tables.values():
        if table.ref.startswith('A4:J'):
            table.ref = f'A4:J{sheet.max_row}'
            if table.autoFilter:
                table.autoFilter.ref = table.ref
    if sheet.auto_filter.ref:
        sheet.auto_filter.ref = f'A4:J{sheet.max_row}'
    target = OUT / SOURCE.name
    book.save(target)
    check = load_workbook(target)
    assert check.sheetnames == book.sheetnames
    for tab in check:
        assert sheet_features(tab) == features[tab.title]
    for key, old in before.items():
        if key not in changed:
            cell = check[key[0]][key[1]]
            assert (cell.value, cell._style) == old, key
    preview(check['Before Build'], numbers[4:6], [1, 4, 10], OUT / 'dispute-loop-plan.png')
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == digest, 'parallel workbook edit'
    shutil.copy2(target, SOURCE)
    print(json.dumps({'rows': numbers, 'unrelated_changes': 0, 'preserved_cells': len(before)}))


if __name__ == '__main__':
    main()
