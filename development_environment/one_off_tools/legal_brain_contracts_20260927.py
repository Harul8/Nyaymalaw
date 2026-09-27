"""Record the owner-approved amendments from the 27 September 2026 contract review.

OWNER DIRECTION, 27 September 2026: an external review of the legal-brain plan
found eleven places where two rows' contracts conflict or acceptance is not
precise enough to run. Each finding was checked against the sheet before
anything was written (all eleven hold, and two go further than the review:
LB-156-AC3 required the replacing code's provision -- the behaviour B-170
removed from the code the same day -- and LB-148 also cited BK-85-AC3).
The owner approved all eleven and decided two trade-offs:

  * the live scratch pad shows receipt-derived progress and checked
    explanations, never the model's own deliberation (item 2);
  * slices 1 and 2 run only on recorded and synthetic matters until slice 3's
    safeguards run on the loop path, and a first versioned principles file
    ships in slice 1 (item 11).

HOW IT IS WRITTEN. The discipline is the other workbook tools':

  * DRAFT rows are amended in place -- exact substring replacements, each
    asserted to occur exactly once;
  * OWNER-AGREED rows keep their wording; the amendment is added as a dated,
    owner-approved paragraph in the column it belongs to;
  * the Implementation Plan receives the mirrored columns and a dated line in
    'Requirement version / change summary';
  * the principles criterion (BK-101-AC1) moves from packet P53 to P49 in
    docs/blueprint/packets.json, and P54's switch rule takes the absolute bar.

Every cell is snapshotted, the book is written to a temporary file, reloaded,
and proved: nothing moved except what was meant to, native sheet features are
unchanged, and the reconciler reports no problem. Only then are the workbook
and the registry replaced, and the blueprint check must pass or the registry
is restored. A second run refuses rather than applying twice.
"""
import copy
import hashlib
import importlib.util
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

_REPO = Path(__file__).resolve().parents[2]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))
from assurance.common._console import utf8_console  # noqa: E402

utf8_console()

_spec = importlib.util.spec_from_file_location(
    "regroup", Path(__file__).with_name("legal_brain_regroup_20260926.py"))
regroup = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(regroup)

SOURCE, FIRST = regroup.SOURCE, regroup.FIRST
PACKETS = _REPO / "docs" / "blueprint" / "packets.json"
HEADER = re.compile(r"^(L|L\.\d|U|U\.\d|P|X)  ")
R = "27 September 2026 review (owner-approved)"
CHANGE_COLUMN = 32   # Implementation Plan: 'Requirement version / change summary'

# =============================================================================
# 1. EXACT REPLACEMENTS -- draft rows only.
# =============================================================================
REPLACE = [
    # ---- item 1: answer depth follows the request -------------------------
    ("LB-163", 4,
     "The harness requires the final answer to carry each section of LB-165 that applies, or to say why it does not.",
     "The seven parts of LB-165 are the harness's completeness check on the working record, not a template for "
     "every reply: the published answer carries what the advocate's request calls for. A needed analysis that is "
     "missing fails -- 'needed' meaning an item on the dispute's own needs list (LB-166) or a finding the step log "
     "holds that the request concerns -- and a heading left out because it does not bear on the request does not "
     "fail. An acknowledgement, a correction or a narrow question never produces a full brief (LB-24, LB-55)."),
    ("LB-165", 4,
     "Each item is a claim in submit_answer, rendered as readable prose; a section that does not apply says why.",
     "Each item is a claim in submit_answer, rendered as readable prose. This is the structure of a comprehensive "
     "dispute brief, published when the advocate asks for one; in ordinary conversation only the parts the request "
     "concerns are published, and the rest stays on the working record (LB-163) rather than being listed as 'does "
     "not apply'."),
    ("LB-161", 4,
     "ask_advocate(question, why_it_matters, what_would_change): ends the turn with one decisive question and "
     "records it so it is not asked again (partly exists: the asked-question record).",
     "ask_advocate(question, why_it_matters, what_would_change): asks one focused question and records it so it is "
     "not asked again (partly exists: the asked-question record). It does not force a choice between answering and "
     "asking: a turn gives the useful answer it can already support and asks the one question that would change it."),
    # ---- item 2: three records; only progress and checked findings shown ---
    ("LB-139", 4,
     "Everything in the panel is marked as working, not yet checked -- only the final answer has passed the "
     "harness -- and every model-written line in it has first passed the one citation detector every channel uses "
     "(LB-67), so an authority the loop never retrieved cannot appear even as working.",
     "THREE RECORDS, AND THE PANEL SHOWS TWO (owner decision, 27 September 2026). Execution diagnostics -- the "
     "model's own deliberation, raw tool traffic and anything a check removed -- are kept protected for diagnosis "
     "and never shown. Progress -- what was searched, read, completed or stopped -- is generated from the tool "
     "receipts, not written by the model. Explanations of findings and decisions appear only after they pass every "
     "check that applies to them on any channel (LB-67): the citation detector, and the checks on factual "
     "assertions and on the reading of a retrieved authority, because an authority genuinely retrieved can still be "
     "misstated. Labelling a line 'working' exempts it from no check."),
    ("LB-164", 4,
     "26 September 2026 review: an event carrying model-written text is released to the stream and the saved log "
     "only after LB-67's detector; what the detector removes is kept in encrypted diagnostics, never in the stream.",
     "26 September 2026 review, amended 27 September 2026 (owner-approved): every event is typed by the record it "
     "carries (LB-139). Receipt-derived progress is released as recorded; an explanation carrying model-written text "
     "reaches the stream and the saved log only after the checks LB-139 names -- the citation detector alone is not "
     "enough -- and diagnostics go only to encrypted storage, never to the stream."),
    # ---- item 3: the verifier's contract ------------------------------------
    ("LB-141", 4,
     "sees one claim and its cited passage and nothing else, and gives its reason first -- quoting the words it "
     "relies on -- then answers supports, does not support, or supports in part (the reason-first pattern recorded "
     "in LB-60).",
     "sees a typed evidence package and nothing else -- the claim, the minimum sources it rests on with their spans, "
     "the established premises it depends on (the dates, the jurisdiction, the governing law, the facts found) and "
     "any contrary material the step log holds -- and gives its reason first -- quoting the words it relies on -- "
     "then answers supports, does not support, or supports in part (the reason-first pattern recorded in LB-60). "
     "FOUR SEPARATE JUDGMENTS (27 September 2026, owner-approved): does the passage support the proposition; does "
     "the proposition apply here -- exception, jurisdiction, date, necessary premise; does the inference follow "
     "from its premises; is material contrary information unresolved. The last three run for claims a "
     "recommendation rests on, so the cost stays bounded; a purely textual claim takes the first. 'Supports in "
     "part' releases only the supported part; a disclaimer never makes an unsupported decisive recommendation "
     "releasable."),
    # ---- item 4: one tool-result contract -----------------------------------
    ("LB-154", 4,
     "THE ENVELOPE: every result carries a status (found / not held / not found / not assessed / error), the source "
     "and its exact locator, the index or table consulted, and what was excluded and why; a zero never looks like "
     "absence (B-163).",
     "THE ENVELOPE carries four separate things, never one status standing for all of them (27 September 2026, "
     "owner-approved): the EXECUTION OUTCOME (completed, refused, failed, or malformed input or output); the "
     "AVAILABILITY of the capability (it could run, or could not and why); the ASSESSMENT of what was asked (found, "
     "not held, held but not found, searched with no match, not assessed -- the existing Coverage, ResolutionState "
     "and Curation values, mapped rather than re-invented); and a RECEIPT shaped by the tool's kind -- a source "
     "read carries its store, version or corpus generation and exact span; a computation carries its inputs, the "
     "rule applied with where it was curated from, and the calculation; an action proposal carries its approval and "
     "effect identity. No tool fabricates a document locator to fill the envelope. Every result names the index or "
     "table consulted and what was excluded and why; a zero never looks like absence (B-163)."),
    ("LB-129", 8,
     "LB-129-AC2: every tool returns one of the three states, and a not assessed result is visible in the step log.",
     "LB-129-AC2: every tool returns the LB-154 envelope, whose assessment carries the three states -- found, not "
     "found or not held, not assessed -- at its core; a not assessed result is visible in the step log."),
    ("LB-143", 4,
     "After a tool returns: does the result carry its source, its locator and one of the three states?",
     "After a tool returns: does the result carry the LB-154 envelope -- outcome, availability, assessment and the "
     "receipt its kind requires?"),
    # ---- item 6: the stored record and the working context ------------------
    ("LB-147", 5,
     "Consecutive requests in a loop share their whole earlier history byte for byte.",
     "Consecutive requests within one context generation share their whole earlier history byte for byte. The "
     "stored record is immutable and complete; what the model is sent is a versioned working-context projection of "
     "it, and clearing (LB-149) or compaction starts a new generation -- recorded with its reason -- which is the "
     "only place the prefix, and so the cache expectation, may change."),
    ("LB-147", 8,
     "LB-147-AC1: across a ten-step loop, each request's history is a byte-identical prefix of the next.",
     "LB-147-AC1: across a ten-step loop within one context generation, each request's history is a byte-identical "
     "prefix of the next; a generation change is recorded, and is the only place the prefix changes."),
    ("LB-149", 4,
     "When results have been used and their findings recorded on the file (LB-142), older raw results are cleared "
     "and replaced by a stub naming the locator and the recorded finding; because retrieval can be repeated, "
     "clearing loses nothing.",
     "When results have been used and their findings recorded on the file (LB-142), older raw results are cleared "
     "and replaced by a stub naming the locator and the recorded finding; because retrieval can be repeated, "
     "clearing loses nothing. Clearing and compaction act on the working-context projection only, start a new "
     "context generation (LB-147) and never alter the stored record. A compaction carries more than dates: disputed "
     "facts, negations, the advocate's instructions, unresolved questions, pending actions and which assessments "
     "were superseded."),
    # ---- item 7: replay reproduces the environment --------------------------
    ("LB-146", 4,
     "A replay adapter feeds the recording back, so the loop, the tools and the harness run deterministically with "
     "no API key on every commit.",
     "A recording pins the environment as well as the model (27 September 2026, owner-approved): the initial matter "
     "snapshot, every tool receipt, the source identities and corpus generation, the clock and the permission "
     "state. STRICT REPLAY feeds all of it back and runs with no network and no API key on every commit; FRESH-WORLD "
     "REVALIDATION re-runs the tools against today's sources and reports every input that changed, invalidating the "
     "evidence that rested on it. A comparison between engines runs on isolated snapshots with external actions "
     "disabled, never twice on the same live matter."),
    # ---- item 8: a checklist reply updates state ----------------------------
    ("LB-166", 8,
     "LB-166-AC2: answering on the board updates the dispute and removes the item.",
     "LB-166-AC2 (amended 27 September 2026, owner-approved): a reply updates the item's state as LB-117 derives it "
     "-- held, outstanding, promised or unavailable -- with its attribution, basis and history kept inspectable. "
     "'I don't know', 'next week' and 'we cannot get it' are different outcomes, and only 'held' clears an item. "
     "One natural reply may update several correctly targeted items; an ambiguous reply turns nothing green; "
     "promised and unavailable items keep their consequences in the analysis. A fact several disputes need is asked "
     "once and linked to each, and a settled question reopens only on a material change or a stated reason."),
    # ---- item 9: source identity and Indian-practice details ----------------
    ("LB-156", 8,
     "LB-156-AC3: read_provision for a 2025 offence under a repealed code returns the replacing code's provision "
     "with the reason.",
     "LB-156-AC3 (reversed 27 September 2026, owner-approved): read_provision for a provision of a repealed code, "
     "asked for a 2025 date, returns THAT code's provision blocked by G-INFORCE with its in-force window -- never "
     "the replacing code's. Which code governs the date is a separate operation (governing_code, LB-120). The "
     "defect this refuses is B-170: a named CrPC section on a 2025 date was read from the BNSS."),
    ("LB-124", 7,
     "Never compute a period from a date the advocate has not given.",
     "Never compute a period from a date that is not established -- one the advocate gave, or one read from a "
     "document or order on the file, attributed and checked (LB-20 keeps asserted, computed and court-listed dates "
     "apart)."),
    ("LB-148", 9, "counsel review of each (BK-85-AC3).", "counsel review of each (BK-98-AC3)."),
    ("LB-168", 9, "counsel review of every entry (BK-85-AC3).", "counsel review of every entry (BK-98-AC3)."),
    ("LB-169", 9, "counsel review of every entry (BK-85-AC3).", "counsel review of every entry (BK-98-AC3)."),
    # ---- items 10 and 11: the bar, and foundations first --------------------
    ("LB-138", 4,
     "each only when the loop matches or beats the pipeline on that kind's golden subset;",
     "each only when the loop meets that kind's absolute acceptance bar (LB-39, LB-40, LB-72) and matches or beats "
     "the pipeline on its golden subset -- the comparison guards against regression and is never the bar, and no "
     "relative improvement offsets a critical error or an unassessed required capability;"),
    ("LB-138", 4,
     "(1) the loop foundation -- tool calling across providers, the loop runner with budgets, typed step events and "
     "the saved step log, the append-only conversation, and record-and-replay;",
     "(1) the loop foundation -- tool calling across providers, the loop runner with budgets, typed step events and "
     "the saved step log, the append-only conversation, record-and-replay, and the first versioned principles file "
     "the loop reads (LB-126);"),
    ("LB-126", 9,
     "Delivery owner (registered 26 September 2026, LB-43): BK-101-AC1, packet P53 (slice 7).",
     "Delivery owner (registered 26 September 2026, LB-43; moved 27 September 2026 on the owner's approval): "
     "BK-101-AC1, packet P49 (slice 1) -- the loop reads a versioned principles file from its first functioning "
     "slice; refining it with the owner continues in P53 (slice 7)."),
]

# =============================================================================
# 2. DATED ADDITIONS. `AC` in the text is numbered from the cell at run time.
# =============================================================================
APPEND = {
    # item 1
    "LB-163": [(8, "{AC} ({R}): across varied requests -- an acknowledgement, a date correction, a narrow statute "
                   "question, a request for the full dispute analysis -- the published answer's size follows the "
                   "request. Planted: a needed analysis removed from a full brief fails; a narrow answer without the "
                   "unrelated headings passes.")],
    "LB-24": [(10, "{R}: LB-163 and LB-165 are reconciled to this row -- their seven parts are an internal "
                   "completeness check and the shape of a full dispute brief, never a template for ordinary replies.")],
    "LB-55": [(10, "{R}: LB-163 and LB-165 are reconciled to this row -- a narrow advisory request publishes only "
                   "what it asks for, and the dispute structure is not forced on it.")],
    # item 2
    "LB-139": [(8, "{AC} ({R}, planted): a false factual assertion, and a false reading of a genuinely retrieved "
                   "authority, placed in progress text, a board item and a question -- none reaches the advocate.")],
    "LB-67": [(10, "{R}: the scratch pad and its step events are channels under this row. LB-139 and LB-164 now "
                   "keep diagnostics, receipt-derived progress and checked explanations apart, so nothing substantive "
                   "reaches the advocate unchecked through them; a 'working' label exempts nothing.")],
    # item 3
    "LB-141": [(8, "{AC} ({R}): changing an exception, the jurisdiction, the date or a necessary premise changes the "
                   "applicability judgment although the quoted words are identical. Planted: a decisive "
                   "recommendation supported only in part and carrying a disclaimer is not released.")],
    # item 4
    "LB-154": [(8, "{AC} ({R}): every registered tool is exercised on success, absence, an unavailable capability, "
                   "a failure and a malformed output, and each carries the same meaning on every provider.")],
    # item 5
    "LB-128": [(6, "{R}: THE RESUMPTION PROTOCOL. Every task, turn, tool call and effect has a stable identity; a "
                   "durable checkpoint is written after each accepted step, carrying the budget consumed; an effect "
                   "whose outcome is uncertain after an interruption is reconciled against its identity before "
                   "anything is retried; cancellation stops new work and deletes nothing accepted (LB-63); an event "
                   "is published only from the saved log, so a republished event keeps its sequence number."),
               (8, "{AC} ({R}): interrupt before an effect, after an effect but before its acknowledgement, and "
                   "after acceptance but before display -- each resumes with one logical effect, the correct state "
                   "and the budget already consumed.")],
    "LB-63": [(10, "{R}: the durable protocol connecting tool execution, accepted writes, budgets and event "
                   "publication is LB-128's; this row's invariants are its acceptance.")],
    "LB-69": [(10, "{R}: the consumed budget is carried by LB-128's durable checkpoints, so a resumed task cannot "
                   "reset it.")],
    # item 6
    "LB-149": [(8, "{AC} ({R}): a compaction neither revives a superseded conclusion, nor loses a qualification or "
                   "a negation, nor repeats a pending action.")],
    # item 7
    "LB-146": [(8, "{AC} ({R}): strict replay passes with the network disabled; a changed source invalidates "
                   "exactly the evidence resting on it; planted: a weakened guard makes a replay fail.")],
    # item 8
    "LB-117": [(10, "{R}: LB-166 (the matter board) is reconciled to this row's four states; an item is never "
                    "removed merely because a reply was given.")],
    "LB-118": [(10, "{R}: LB-166-AC2 now applies this row -- one natural reply may update several items, and an "
                    "ambiguous reply turns nothing green.")],
    # item 9
    "LB-168": [(4, "{R}: admissibility is assessed for a stated purpose and stage of the proceeding, not for the "
                   "document in the abstract. Unresolved admissibility never stops NM reporting accurately what an "
                   "examined document says; it limits only what the document may be relied on to prove.")],
    "LB-92": [(4, "{R}: the latest-source reader decision stands, and the drawer keeps two things apart -- the "
                  "excerpt the answer relied on, at its recorded version, and the current local document. Identical "
                  "wording in the current document does not establish an identical legal context: the surrounding "
                  "provisions, an amendment or the corpus generation may differ, and the drawer says which.")],
    "LB-91": [(10, "{R}: the historical relied-on excerpt and the current local document are distinguished as "
                   "LB-92 now records; identical wording is not identical legal context.")],
    "LB-95": [(10, "{R}: as LB-91 -- an exact saved identity binds the historical excerpt; the current document is "
                   "shown beside it, never in its place.")],
    # item 10
    "LB-40": [(8, "{R}: THE ABSOLUTE BAR, operationalised before the first switch -- pre-agreed severity "
                  "definitions, populations and thresholds covering legal errors and material omissions; "
                  "unsupported release and excessive refusal; independent judgment under pressure; communication, "
                  "question burden and comprehension; and cost and latency per satisfactorily completed task. The "
                  "severity definitions and thresholds are the owner's to approve. Development goldens "
                  "(docs/GOLDEN_SET.md) are kept apart from a held-out set, and a defect added to the regression "
                  "suite is no longer an unseen test.")],
    "LB-39": [(10, "{R}: the held-out set is kept apart from the development goldens, as LB-40 now records.")],
    "LB-72": [(10, "{R}: relative improvement over the pipeline cannot offset a critical error or an unassessed "
                   "required capability (LB-40, LB-138).")],
    # item 11
    "LB-138": [(10, "{R}, owner decision: slices 1 and 2 run only on recorded and synthetic matters -- no tool reads "
                    "or writes a live matter until slice 3's permission, budget, write-integrity and publication "
                    "safeguards run on the loop path. The first versioned principles file ships in slice 1 "
                    "(LB-126).")],
    "LB-43": [(8, "{R}: LB-43-AC4's population is every LB and OM requirement in the current sheet and every clause "
                  "of each, enumerated when the check runs -- 170 LB and 23 OM on 27 September 2026 -- not the 92 "
                  "and 324 recorded when it was written. Cell-to-cell reconciliation is necessary and is not this "
                  "check: a slice cannot start with an unowned clause among its selected rows or a missing mandatory "
                  "foundation.")],
}

ITEMS = {
    1: ("answer depth follows the request; the seven parts are an internal completeness check",
        ("LB-163", "LB-165", "LB-161", "LB-24", "LB-55")),
    2: ("diagnostics, receipt-derived progress and checked explanations are separate records; only the last two "
        "are shown", ("LB-139", "LB-164", "LB-67")),
    3: ("the verifier takes a typed evidence package and gives four separate judgments", ("LB-141",)),
    4: ("the tool envelope separates outcome, availability, assessment and a kind-specific receipt",
        ("LB-154", "LB-129", "LB-143")),
    5: ("a durable resumption protocol with stable task, turn, call and effect identities",
        ("LB-128", "LB-63", "LB-69")),
    6: ("the stored record is immutable; the model's context is a versioned projection with generations",
        ("LB-147", "LB-149")),
    7: ("replay pins the environment; strict replay and fresh-world revalidation are separate", ("LB-146",)),
    8: ("a checklist reply updates the item's state and never removes it by being given",
        ("LB-166", "LB-117", "LB-118")),
    9: ("exact provision reading is separate from governing law; attributed document dates; purpose-bound "
        "admissibility; BK-98-AC3; the relied-on excerpt apart from the current document",
        ("LB-156", "LB-124", "LB-148", "LB-168", "LB-169", "LB-91", "LB-92", "LB-95")),
    10: ("an absolute acceptance bar; comparison with the pipeline guards regression only",
         ("LB-138", "LB-39", "LB-40", "LB-72")),
    11: ("principles from slice 1; slices 1 and 2 on recorded and synthetic matters; ownership drawn from the "
         "current plan", ("LB-126", "LB-138", "LB-43")),
}
MARK = "27 September 2026: amended on the owner's approval of the contract review"


def next_ac(ident: str, text: str) -> str:
    numbers = [int(n) for n in re.findall(rf"{re.escape(ident)}-AC(\d+)", text or "")]
    return f"{ident}-AC{max(numbers, default=0) + 1}"


def packets_after(raw: str) -> str:
    """The registry with the principles criterion in slice 1 and the absolute bar at the switch."""
    doc = json.loads(raw)
    by = {p["id"]: p for p in doc["packets"]}
    p49, p50, p53, p54 = by["P49"], by["P50"], by["P53"], by["P54"]
    assert "BK-101-AC1" in p53["criteria"] and "BK-101-AC1" not in p49["criteria"], "already moved"
    for key in ("criteria", "final_criteria"):
        p53[key].remove("BK-101-AC1")
        p49[key].append("BK-101-AC1")
    p49["inputs"].append("The first distilled principles draft (LB-126).")
    p49["outputs"].append("A versioned principles file the loop reads at the start of every turn, its version "
                          "recorded on the turn; refined with the owner in P53.")
    live = ("No tool reads or writes a live matter before P51's safeguards run on the loop path: this slice runs "
            "on recorded and synthetic matters only.")
    p49["expected"].append(live)
    p50["expected"].append(live)
    old = "A principles document read every turn; a stable prefix"
    assert p53["outputs"][0].count(old) == 1
    p53["outputs"][0] = p53["outputs"][0].replace(old, "The principles refined with the owner; a stable prefix")
    old = "Read the versioned principles at the start of every turn and freeze the prefix for the conversation."
    assert p53["steps"].count(old) == 1
    p53["steps"][p53["steps"].index(old)] = ("Refine the principles with the owner, and freeze the prefix for "
                                             "the conversation.")
    old = ("Switch a kind of turn only when the loop matches or beats the pipeline, with the cost per turn "
           "measured as a release row.")
    assert p54["steps"].count(old) == 1
    p54["steps"][p54["steps"].index(old)] = (
        "Switch a kind of turn only when the loop meets that kind's absolute acceptance bar (LB-39, LB-40, LB-72) "
        "and matches or beats the pipeline, with the cost per turn measured as a release row; no relative "
        "improvement offsets a critical error or an unassessed required capability.")
    p54["expected"].append("Matching the pipeline is never the bar on its own.")
    return json.dumps(doc, indent=2, ensure_ascii=False) + "\n"


def main() -> int:
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    book = load_workbook(SOURCE)
    sheet, plan = book["Before Build"], book["Implementation Plan"]
    features = {s.title: regroup.sheet_features(s) for s in book}
    before = {(s.title, c.coordinate): (c.value, copy.copy(c._style)) for s in book for row in s for c in row}

    rows = {}
    for r in range(FIRST, sheet.max_row + 1):
        text = str(sheet.cell(r, 1).value or "")
        h = HEADER.match(text)
        key = f"HEADER:{h[1]}" if h else regroup.key_of(text)
        assert key not in rows, f"duplicate row key {key}"
        rows[key] = r

    touched = {k for k, *_ in REPLACE} | set(APPEND)
    for key in touched:
        assert key in rows, key
        assert MARK not in str(plan_cell(plan, key, CHANGE_COLUMN).value or ""), f"{key}: already applied"
    listed = {k for _item, (_s, keys) in ITEMS.items() for k in keys}
    assert listed == touched, f"items and edits disagree: {sorted(listed ^ touched)}"

    # ---- the expected values of every row that changes ---------------------
    expected = {}
    for key in touched:
        values = [sheet.cell(rows[key], c).value for c in range(1, 11)]
        for ident, col, old, new in REPLACE:
            if ident == key:
                assert str(values[col - 1]).count(old) == 1, (ident, col, old[:60])
                values[col - 1] = values[col - 1].replace(old, new)
        for col, text in APPEND.get(key, []):
            text = text.format(R=R, AC=next_ac(key, values[col - 1]))
            values[col - 1] = f"{values[col - 1]}\n\n{text}" if values[col - 1] else text
        expected[key] = values
        for c in range(1, 11):
            sheet.cell(rows[key], c).value = values[c - 1]

    # ---- the Implementation Plan: the mirror and the change summary --------
    plan_changed = set()
    summary = {key: [] for key in touched}
    for item, (what, keys) in ITEMS.items():
        for key in keys:
            summary[key].append(f"item {item}: {what}")
    for x in range(2, plan.max_row + 1):
        ident = plan.cell(x, 1).value
        if ident not in touched:
            continue
        for target, origin in regroup.MIRROR.items():
            v = sheet.cell(rows[ident], origin).value
            if plan.cell(x, target).value != v:
                plan.cell(x, target).value = v
                plan_changed.add(plan.cell(x, target).coordinate)
        cell = plan.cell(x, CHANGE_COLUMN)
        line = f"{MARK} -- {'; '.join(summary[ident])}."
        cell.value = f"{cell.value}\n\n{line}" if cell.value else line
        plan_changed.add(cell.coordinate)

    out_dir = Path(tempfile.gettempdir()) / "nm_legal_brain_contracts_20260927"
    out_dir.mkdir(exist_ok=True)
    out = out_dir / SOURCE.name
    book.save(out)

    # ================= THE PROOF, on the saved bytes ==========================
    check = load_workbook(out)
    for tab in check:
        assert regroup.sheet_features(tab) == features[tab.title], tab.title
    bb, ip = check["Before Build"], check["Implementation Plan"]
    changed_cells = {(k, c) for k in touched for c in range(1, 11)}
    for (title, coord), old in before.items():
        cell = check[title][coord]
        if title == "Before Build":
            key = next((k for k, r in rows.items() if r == cell.row), None)
            if key in touched and (key, cell.column) in changed_cells:
                assert cell._style == old[1], f"{coord}: style moved"
                continue
            assert (cell.value, cell._style) == old, f"Before Build {coord} moved"
        elif title == "Implementation Plan" and coord not in plan_changed:
            assert (cell.value, cell._style) == old, f"plan {coord} moved"
        elif title not in ("Before Build", "Implementation Plan"):
            assert (cell.value, cell._style) == old, f"{title} {coord} moved"
    for key in touched:
        got = [bb.cell(rows[key], c).value for c in range(1, 11)]
        assert got == expected[key], key
        assert got[0] == before[("Before Build", f"A{rows[key]}")][0], f"{key}: the owner's words moved"
    for x in range(2, ip.max_row + 1):
        if ip.cell(x, 1).value in touched:
            assert MARK in str(ip.cell(x, CHANGE_COLUMN).value)
    from assurance.control_plane import plan_scenarios
    problems = plan_scenarios.requirement_problems(out)
    assert problems == [], problems
    assert len(plan_scenarios.sheet_rows(out)) == len(plan_scenarios.sheet_rows(SOURCE))

    # ================= the registry, transactionally ==========================
    raw_packets = PACKETS.read_bytes()
    text = raw_packets.decode("utf-8")
    crlf = "\r\n" in text
    new = packets_after(text.replace("\r\n", "\n"))
    PACKETS.write_bytes((new.replace("\n", "\r\n") if crlf else new).encode("utf-8"))
    done = subprocess.run([sys.executable, str(_REPO / "assurance/control_plane/blueprint.py"), "check"],
                          cwd=_REPO, capture_output=True, text=True)
    if done.returncode != 0:
        PACKETS.write_bytes(raw_packets)
        print(done.stdout[-3000:], done.stderr[-3000:])
        raise SystemExit("blueprint check failed; the registry was restored and the workbook was not replaced")

    shutil.copy2(out, SOURCE)
    print(f"source digest before: {digest}")
    print(f"source digest after:  {hashlib.sha256(SOURCE.read_bytes()).hexdigest()}")
    print(f"rows amended: {len(touched)}; replacements: {len(REPLACE)}; additions: "
          f"{sum(len(v) for v in APPEND.values())}; plan cells written: {len(plan_changed)}")
    print("registry: BK-101-AC1 moved P53 -> P49; P49/P50 run on recorded and synthetic matters; P54 takes the "
          "absolute bar. blueprint check: " + (done.stdout.strip().splitlines() or ["ok"])[-1])
    return 0


def plan_cell(plan, key, column):
    for x in range(2, plan.max_row + 1):
        if plan.cell(x, 1).value == key:
            return plan.cell(x, column)
    raise KeyError(key)


if __name__ == "__main__":
    raise SystemExit(main())
