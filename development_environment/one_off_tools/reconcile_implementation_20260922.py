"""Bounded workbook migration and diagnostic updates; original/history stay intact.

Authoring fallback: the bundled artifact-tool package is unavailable. Run with
the bundled Python, not application dependencies. The workbook remains the owner.
"""
from __future__ import annotations

import argparse
import copy
import json
import re
import shutil
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.table import TableColumn
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image, ImageDraw, ImageFont
import sys

_REPO = Path(__file__).resolve().parents[2]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))
from assurance.common._console import utf8_console  # noqa: E402

utf8_console()

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs/Nyaymalaw_Implementation_Plan.xlsx"
OUTPUT = ROOT / "outputs/01a07b76-6b21-71f3-bb09-261f64617594"
HEADERS = ("Executable scenarios (Gherkin)", "Executable scenario state",
           "Executable scenario coverage notes")


def preview(sheet, rows, columns, target):
    """Render actual selected cells with wrapping and source styles for edit QA."""
    font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 15)
    widths = [min(700, max(160, int(sheet.column_dimensions[
        sheet.cell(1, c).column_letter].width * 7))) for c in columns]
    measuring = ImageDraw.Draw(Image.new("RGB", (1, 1)))
    wrapped, heights = [], []
    for r in rows:
        line_sets = []
        for c, width in zip(columns, widths):
            lines = []
            for paragraph in str(sheet.cell(r, c).value or "").split("\n"):
                current = ""
                for word in paragraph.split():
                    candidate = f"{current} {word}".strip()
                    if measuring.textlength(candidate, font=font) > width - 20 and current:
                        lines.append(current)
                        current = word
                    else:
                        current = candidate
                lines.append(current)
            line_sets.append(lines)
        wrapped.append(line_sets)
        heights.append(max(len(lines) for lines in line_sets) * 20 + 20)
    image = Image.new("RGB", (sum(widths), sum(heights)), "white")
    draw = ImageDraw.Draw(image)
    y = 0
    for r, line_sets, height in zip(rows, wrapped, heights):
        x = 0
        for c, lines, width in zip(columns, line_sets, widths):
            cell = sheet.cell(r, c)
            colour = cell.fill.fgColor
            fill = ("#" + colour.rgb[-6:]
                    if cell.fill.patternType and colour.type == "rgb" else "#FFFFFF")
            draw.rectangle((x, y, x + width, y + height), fill=fill, outline="#CBD5DA")
            ink = "#FFFFFF" if fill.lower() in {"#173e48", "#173e43", "#175448"} else "#182B35"
            draw.multiline_text((x + 10, y + 10), "\n".join(lines), font=font,
                                fill=ink, spacing=3)
            x += width
        y += height
    image.save(target)


def reconcile(diagnostics=None):
    OUTPUT.mkdir(parents=True, exist_ok=True)
    backup = OUTPUT / "implementation-before-20260922.xlsx"
    if not backup.exists():
        shutil.copy2(SOURCE, backup)
    w = load_workbook(SOURCE)
    before = {(s.title, c.coordinate): (c.value, copy.copy(c._style))
              for s in w for row in s for c in row}
    p, b = w["Implementation Plan"], w["Before Build"]
    preview(b, [248], [1, 2], OUTPUT / "before-build-before.png")
    changed = set()

    def put(sheet, row, col, value):
        sheet.cell(row, col, value)
        changed.add((sheet.title, sheet.cell(row, col).coordinate))

    def amend(sheet, address, addition, *, replace=False):
        cell = sheet[address]
        previous = str(cell.value or "")
        if replace or addition not in previous:
            put(sheet, cell.row, cell.column,
                addition if replace else f"{previous}\n\n{addition}".strip())

    # Apply the already-approved provider amendment to CURRENT requirements;
    # preserve column H/original descriptions and all historical scenario cells.
    assert p["A39"].value == "F-A-02" and p["A47"].value == "F-A-09"
    assert "A.2.6" in str(b["A17"].value)
    amend(p, "I39", "Register publicly with email and matching passwords, without professional-details collection or an enrolment invitation. Activation requires six-digit mailbox confirmation. Present the required registration notice/terms and separate adult acknowledgement. Offer a third, optional, initially unticked permission for disclosed global OpenAI API matter-text processing. Declining OpenAI never prevents registration or access to the advocate's own workspace or matters. Professional approval is separate. Owner decisions reflected 22 September 2026.", replace=True)
    plan_additions = {
        "N39": "21 September owner amendment: show the OpenAI notice and separate optional checkbox. Preserve the two required privacy/adult acknowledgements. Existing accounts review Profile > AI data sharing; never infer acceptance from an older notice.",
        "AG39": "F-A-02-AC04: retain affirmative OpenAI choice with the shown version, account and time through six-digit email confirmation. An unchecked choice grants no permission but permits owned-workspace access. Re-registration cannot overwrite an existing account's permission.",
        "AH39": "F-A-02-S04: register, confirm email and sign in with and without OpenAI permission. Both reach their own workspace; only explicit current permission allows external text dispatch. Verify three account fields, two required privacy/adult boxes and one optional OpenAI box, all initially unticked.",
        "AF39": "22 September 2026: reconciled owner-approved public registration and optional OpenAI text processing. Original descriptions and historical scenarios preserved; current requirement supersedes the obsolete invitation-only rollout wording.",
        "I47": "21 September owner amendment: separately disclose global OpenAI API processing of matter text, context and retrieved excerpts; default abuse-monitoring retention up to 30 days with exceptions; no India-only or zero-retention promise. This optional purpose-specific permission establishes neither client authority nor legal clearance.",
        "AG47": "F-A-09-AC04: check current account-bound permission and a live device-bound session before every OpenAI dispatch and retry. Withdrawal blocks later dispatch, cannot recall sent data, and cannot be undone by a stale tab. Refuse unreadable permission, alternative hosts, redirects, raw media and embeddings. Retain all matter-access, legal and professional-action controls.",
        "AH47": "F-A-09-S04: acceptance survives reload; withdrawal blocks future external calls while the owned matter remains accessible. A stale update, another account's record, revoked session or damaged permission cannot authorise a call. Provider settings or checkboxes never claim qualified legal approval.",
        "AF47": "22 September 2026: current notice openai-text-2026-09-21. BK-85-AC3 remains NOT_RUN; neither production acceptance nor counsel sign-off follows from this owner decision.",
    }
    for address, addition in plan_additions.items():
        amend(p, address, addition)
    before_additions = {
        "D12": "21 September owner amendment: separate optional, initially unticked OpenAI text-processing choice with global-processing and retention disclosure; existing required registration acknowledgements remain separate.",
        "E12": "Six-digit email confirmation preserves the exact provider-notice choice; no acceptance is inferred for existing accounts or unchecked registration.",
        "H12": "Also verify registration, confirmation and sign-in with and without OpenAI permission; exact six controls, unticked defaults and optional/required boundary.",
        "D17": "OpenAI choice is separate and optional: matter text, context and retrieved excerpts to the direct global OpenAI API only. Exclude raw files, images, audio, video, embeddings and other providers. Existing users accept or withdraw in Profile > AI data sharing.",
        "E17": "Versioned account-bound permission and live session are checked before every external dispatch and retry. No permission means no external AI response, not loss of workspace or matters. Withdrawal affects future dispatch and cannot recall requests already sent.",
        "H17": "Test persistence through confirmation/restart, missing/stale/forged/copied/corrupt records, CSRF, expired sessions, stale concurrent acceptance, retry after withdrawal, wrong endpoints, media exclusion and visible refusal. Failed storage writes cannot record acceptance.",
        "I17": "Owner approval 21 September permits disclosed global OpenAI text processing, not legal compliance. Disclosure: https://developers.openai.com/api/docs/guides/your-data . No API training by default does not mean zero retention. BK-85-AC3 remains NOT_RUN.",
        "J17": "22 September reconciliation: permission integration and controlled engineering checks are separate from live-model quality, qualified legal review and production acceptance; none is inferred from a checkbox.",
        "F18": "Provider disclosure must state default abuse-monitoring retention up to 30 days with exceptions. store=false disables stored completions, not all provider retention. Never promise India-only processing or unconditional no retention.",
    }
    for address, addition in before_additions.items():
        amend(b, address, addition)
    amend(b, "A2", "Original descriptions remain in A; B–J contains current requirements. Rows 156–248 cover 92 legal-brain requirements and 324 planned clauses. LB-90–92 carries the selected board, citation and source-view patterns. Legal explanation and retrieved verbatim passages remain inline by default. All 115 LB/OM requirements are linked into Implementation Plan; release assignment and acceptance mapping remain OPEN. DIAGNOSTIC rows below record measured findings and evidence, not feature acceptance. No application build or expert-quality judgment is established by reconciliation.", replace=True)

    if p.max_column == 56:
        for c, title in enumerate(HEADERS, 57):
            put(p, 1, c, title)
            p.cell(1, c)._style = copy.copy(p.cell(1, 34)._style)
            p.column_dimensions[p.cell(1, c).column_letter].width = (85, 22, 65)[c-57]
        count = 0
        for r in range(2, p.max_row + 1):
            if p.cell(r, 3).value != "Feature":
                continue
            body = p.cell(r, 49).value
            if p.cell(r, 1).value == "F-A-02":
                for old, new in (
                    ("a retype-password field and two privacy boxes", "a retype-password field, two required privacy boxes and one optional OpenAI permission box"),
                    ("carries no logo and no other text than the privacy notice and its boxes", "carries only account details, delivery status and confirmation navigation"),
                    ("registration is refused without saying the account exists", "the response does not say whether the account exists"),
                ):
                    assert old in body, old
                    body = body.replace(old, new)
            for c in range(57, 60):
                p.cell(r, c)._style = copy.copy(p.cell(r, 34)._style)
            put(p, r, 57, body)
            put(p, r, 58, "Executable" if body else "Planned")
            put(p, r, 59, (
                "Retained executable subset, migrated from Historical scenarios; original preserved. "
                "F-A-02 includes owner-approved registration amendments. Execution results are separate. "
                "Current prose obligations beyond this subset remain unverified."
                if body else "Current prose scenarios remain required. No executable Gherkin is registered yet; no coverage or verification is claimed."))
            count += bool(body)
        assert count == 19
        validation = DataValidation(type="list", formula1='"Executable,Planned"')
        p.add_data_validation(validation)
        validation.add(f"BF2:BF{p.max_row}")
    else:
        assert tuple(p.cell(1, c).value for c in range(57, 60)) == HEADERS

    # One traceable current row for each principle/brain requirement. Original
    # narrative stays in Before Build; status is not inferred from source presence.
    ids = {p.cell(r, 1).value: r for r in range(2, p.max_row + 1)}
    for r in range(5, 249):
        original = str(b.cell(r, 1).value or "")
        match = re.match(r"(LB-\d+|OM-[PIQ]\d+)\b", original)
        if not match or match[1] in ids:
            continue
        nr = p.max_row + 1
        rid = match[1]
        for c in range(1, 60):
            p.cell(nr, c)._style = copy.copy(p.cell(39, min(c, 59))._style)
        mapping = {1: rid, 3: "Requirement", 4: "Legal brain" if rid.startswith("LB") else "Open a matter",
                   6: original.split("\n")[2] if len(original.split("\n")) > 2 else rid,
                   7: "Release assignment pending", 8: original,
                   9: b.cell(r, 4).value, 10: b.cell(r, 2).value,
                   12: b.cell(r, 3).value, 15: b.cell(r, 4).value,
                   16: b.cell(r, 5).value, 17: b.cell(r, 6).value,
                   19: b.cell(r, 7).value, 28: b.cell(r, 9).value,
                   30: b.cell(r, 10).value, 32: f"22 September 2026: linked from Before Build row {r}; not acceptance.",
                   33: b.cell(r, 8).value, 38: "Recorded; engineering mapping pending",
                   39: "Not assessed", 40: "Not verified",
                   44: "Code mapping, release ownership and criterion-specific evidence remain open."}
        for c, value in mapping.items():
            put(p, nr, c, value)
        p.row_dimensions[nr].height = 240
        ids[rid] = nr

    if diagnostics:
        existing = {str(b.cell(r, 1).value or "").split("\n")[0]: r for r in range(5, b.max_row+1)}
        for diagnostic in diagnostics:
            rid = diagnostic["id"]
            assert rid.startswith("DG-") and len(diagnostic["cells"]) == 10
            r = existing.get(rid, b.max_row + 1)
            for c, value in enumerate(diagnostic["cells"], 1):
                put(b, r, c, value)
                b.cell(r, c)._style = copy.copy(b.cell(248, c)._style)
            b.row_dimensions[r].height = 250
            existing[rid] = r
        b.auto_filter.ref = f"A4:J{b.max_row}"

    if diagnostics and any(row["id"] == "DG-15" for row in diagnostics):
        # Current owner-approved discipline amendments. Preserve original
        # narratives and mirror only the same mapped current cells in the plan.
        additions = {
            "LB-04": {
                4: "Address the current purpose in natural professional paragraphs. Use the held record before asking a small group of neutral, consequential questions. Challenge a proposition and its support, not character. A criticism should identify a supported remedy or resolving information where available, never invent a fix.",
                7: "Do not force advice into acknowledgement-only messages or suppress new substantive work because it begins conversationally. Material qualifications survive rewriting. A user override cannot create evidence or authority.",
                8: "Check purpose, question economy and useful respectful challenge through COMM-01/02/03, each with a negative control. Test both acknowledgement and substantive follow-up inside a saved matter. Fixtures remain outside production instructions.",
            },
            "LB-08": {
                7: "Readable extraction, accurate understanding and factual truth are distinct. A confirmed reading is not corroboration. Keep the exact source and its limits.",
            },
            "LB-21": {
                4: "Keep extraction quality, faithful understanding, factual support, legal support, applicability, practical uncertainty and decision readiness distinct. State which unknown changes which conclusion and what would resolve it.",
                8: "REASON-01 must fail a readable-document-to-proven-fact leap, a quoted-rule-to-applicable-law leap, or one confidence score that conceals independent unknowns. Guidance and evaluator coverage are not proof of a persistent uncertainty model or measured legal accuracy.",
            },
            "LB-40": {
                8: "Each judged pass/fail needs exact excerpts from the tested material and a reason. Retain the complete input and its digest under contributing matter keys. Missing, invented or empty-population evidence is not assessed. Calibrate semantic judgments against qualified human review.",
                9: "Paid independent judge execution needs its own approved model and budget. Existing synthetic GPT-4o mini permission does not authorise a different judge. Qualified legal review remains separate.",
            },
            "LB-60": {
                4: "Compose the shared reasoning discipline once per model call. Add communication obligations only to advocate-facing outputs. Keep current instructions separate from historical context. Compare material alternatives before adopting a revisable preferred position.",
                7: "No instruction may force a single factual hypothesis, erase material caveats, invent a replacement remedy or turn every reference to the file into a new legal workup. Keep task schemas, evidence and permission controls effective.",
                8: "Test actual assembled route/repair prompts and schema-preserving composition. Retain exact-step dependency decisions, concise basis and input digest in encrypted diagnostics; never expose a refused candidate as advice. Separately measure semantic false-clear and false-block rates.",
            },
            "LB-70": {
                4: "Show relied-on legal explanation, retrieved passages and source links by default, with material risks stated in persistent text. Keep private internal deliberation out of the conversation. Optional audit detail is not a substitute for a clear answer.",
                8: "Browser-check paragraph order, visible sources and warnings, keyboard navigation, narrow layouts, composer recovery and saved History. A controlled renderer test must be labelled separately from a real provider turn.",
            },
        }
        column_map = {4: (9, 15), 7: (19,), 8: (33,), 9: (28,), 10: (30,)}
        for rid, cells in additions.items():
            matches = [r for r in range(5, b.max_row + 1)
                       if str(b.cell(r, 1).value or "").split("\n")[0] == rid]
            assert len(matches) == 1 and rid in ids, rid
            br, pr = matches[0], ids[rid]
            cells[10] = "DG-15 records this engineering increment and measured checks. DG-16 records the live routing defect and retest. These do not close the whole requirement, DG-13, independent evaluation or professional review."
            for col, addition in cells.items():
                note = "22 September 2026 amendment: " + addition
                amend(b, b.cell(br, col).coordinate, note)
                for pc in column_map[col]:
                    amend(p, p.cell(pr, pc).coordinate, note)
    if diagnostics and any(row["id"] == "DG-17" for row in diagnostics):
        rid = "LB-60"
        matches = [r for r in range(5, b.max_row + 1)
                   if str(b.cell(r, 1).value or "").split("\n")[0] == rid]
        assert len(matches) == 1 and rid in ids
        br, pr = matches[0], ids[rid]
        old_note = ("22 September 2026 DG-17/DG-18: review actual assembled prompts, nested "
                "transmitted schema descriptions and downstream consumers together. Inventory "
                "currently covers 24 constructors, 19 system constants and 21 schemas. "
                "Test empty populations, unknown-versus-negative states, context changes, "
                "purpose-sensitive answers, role progression and incomplete repairs. "
                "A source-only explanation is not full advisory assessment. "
                "Controlled checks and semantic evaluation remain separate.")
        note = ("DG-17/DG-18: check 24 prompt constructors, 19 system constants, 21 wire "
                "schemas and their consumers. See diagnostic rows for cases, results and limits.")
        for sheet, cell in ((b, b.cell(br, 8)), (p, p.cell(pr, 33))):
            if old_note in str(cell.value or ""):
                amend(sheet, cell.coordinate, str(cell.value).replace(old_note, note), replace=True)
            amend(sheet, cell.coordinate, note)
        b.row_dimensions[br].height = 409.5
        p.row_dimensions[pr].height = max(p.row_dimensions[pr].height or 0, 270)
    table = p.tables["ImplementationPlan"]
    table.ref = f"A1:BG{p.max_row}"
    if len(table.tableColumns) == 56:
        table.tableColumns.extend(TableColumn(id=c, name=p.cell(1,c).value) for c in range(57,60))
    if table.autoFilter:
        table.autoFilter.ref = table.ref
    p.auto_filter.ref = table.ref
    candidate = OUTPUT / SOURCE.name
    w.save(candidate)
    reopened = load_workbook(candidate)
    for (name, cell), (value, style) in before.items():
        if (name, cell) not in changed:
            assert reopened[name][cell].value == value, (name, cell, "value drift")
            assert reopened[name][cell]._style == style, (name, cell, "style drift")
    assert reopened.sheetnames == w.sheetnames
    assert [(s.title, s.freeze_panes) for s in reopened] == [(s.title,s.freeze_panes) for s in w]
    preview(reopened["Implementation Plan"], [39,40], [57,58,59], OUTPUT / "scenario-columns.png")
    if diagnostics:
        preview(reopened["Before Build"], [existing[diagnostics[-1]["id"]]], [1,2,10], OUTPUT / "diagnostic-row.png")
    print(json.dumps({"sheets": [(s.title,s.max_row,s.max_column) for s in reopened],
                      "preserved_original_cells": len(before)-len(set(before)&changed),
                      "linked_requirements": sum(str(i).startswith(("LB-","OM-")) for i in ids),
                      "diagnostics": len(diagnostics or [])}))
    reopened.close()
    w.close()
    shutil.copy2(candidate, SOURCE)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--diagnostics", type=Path)
    options = parser.parse_args()
    reconcile(json.loads(options.diagnostics.read_text(encoding="utf8")) if options.diagnostics else None)
