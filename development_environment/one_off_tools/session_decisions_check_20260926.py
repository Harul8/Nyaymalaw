"""Dry run: is every item discussed and settled this session present in the Before Build sheet?

Each item names the rows that should carry it and a phrase that must appear in
that row's text (any of its ten columns, or the section header). A phrase is
searched case-insensitively. Nothing is written.
"""
import re
import sys
from openpyxl import load_workbook

SOURCE = "docs/Nyaymalaw_Implementation_Plan.xlsx"
ws = load_workbook(SOURCE, read_only=True)["Before Build"]
HEADER = re.compile(r"^(L|L\.\d|U|U\.\d|P|X)  ")
ID = re.compile(r"^(LB-\d+|OM-[PIQ]\d+|F-[A-Z]-\d+)\b")
text = {}
for row in ws.iter_rows(min_row=1, values_only=True):
    a = str(row[0] or "")
    h, m = HEADER.match(a), ID.match(a)
    key = f"H:{h[1]}" if h else (m[1] if m else None)
    if key:
        text[key] = " ".join(str(v or "") for v in row).lower()

CHECKLIST = [
    # --- what makes Claude good, and the owner's intent -------------------------------------
    ("No model training; open models (Opus, Fable, OpenAI) used as they are", [("LB-127", "openai"), ("LB-127", "fable")]),
    ("Models work in a loop: decide, act, verify, decide next", [("LB-128", "decide, act, verify, decide again")]),
    ("Guiding principles, not fixed rules that work against the model", [("H:L", "rather than fixed rules"), ("LB-126", "not as steps to execute")]),
    ("Principles distilled from OM-P01-14, the D16 tenets and CLAUDE.md; owner edits without code", [("LB-126", "om-p01-14"), ("LB-126", "d16")]),
    ("Strong harness checks every output; the model acts on the result", [("LB-131", "every answer"), ("LB-133", "told why")]),
    ("Zero invention / hallucination is non-negotiable", [("H:L", "non-negotiable"), ("LB-130", "zero invention")]),
    ("Answers as claims tied to sources, so zero invention is checkable", [("LB-130", "set of claims"), ("LB-161", "submit_answer")]),
    ("18 output checks kept and strengthened", [("LB-131", "eighteen output checks")]),
    ("6 professional boundaries fixed", [("LB-132", "six boundaries")]),
    ("13 judgment gates: floor or principle, owner decides, with recommendation", [("LB-134", "thirteen"), ("LB-134", "recommendation")]),
    ("The pipeline (fixed reads, closed lists as a cage, fixed retrieval rounds) is what holds the model back", [("LB-129", "closed vocabularies"), ("LB-137", "fixed number of rounds")]),
    ("Closed lists only at the tool's door", [("LB-129", "only as a tool's input key")]),
    ("Build beside the pipeline behind a switch; switch on golden comparison", [("LB-138", "beside the existing turn engine"), ("LB-138", "golden set")]),
    ("Golden set is the improvement signal instead of training; every defect becomes a golden case", [("LB-146", "becomes a golden case")]),
    ("Show the work live; the advocate can correct mid-turn", [("LB-139", "correct a fact mid-turn")]),
    # --- structure ----------------------------------------------------------------------------
    ("Arrive first, then the legal brain holding loop, context and scaffolding; retrieval part of it", [("H:L", "legal brain"), ("H:L.4", "retrieved")]),
    ("Tools as their own section after the loop", [("H:L.3", "tools")]),
    # --- loops, tools, context, harness --------------------------------------------------------
    ("Five loops: turn, research, opposing counsel, matter, improvement", [("H:L.2", "five loops")]),
    ("Research sub-loop with fresh context returns findings with spans", [("LB-136", "fresh context")]),
    ("Opposing-counsel loop sees claims and sources, never the reasoning", [("LB-140", "never the main loop's reasoning")]),
    ("Independent verifier on a cheaper tier, never the author", [("LB-141", "cheaper tier"), ("LB-141", "never the instance that wrote")]),
    ("The matter written only through checked tools", [("LB-142", "checked tools")]),
    ("Checks before and after every tool call; step log as process evidence", [("LB-143", "before a tool runs"), ("LB-143", "process evidence")]),
    ("Every number and date from a tool", [("LB-144", "every number and date")]),
    ("Context tagged with source and status; retrieved text never an instruction", [("LB-145", "tagged"), ("LB-145", "never instructions")]),
    ("Record, replay, compare providers", [("LB-146", "replay adapter"), ("LB-146", "opus, fable and an openai")]),
    ("Every check keeps a planted-violation control", [("LB-131", "planted"), ("LB-154", "planted failure")]),
    ("Corpus is the ceiling: measured gaps close by acquisition, not by the model", [("H:L.4", "ceiling")]),
    ("Legal reasoning method as guidance: facts, issues, statute, authority, application, procedure, remedy, risk, next step", [("LB-163", "remedy")]),
    ("Cost: nested loops cost more; budgets; cheaper tier for checks", [("LB-128", "budgets"), ("LB-153", "cheaper models")]),
    # --- Claude's context management -----------------------------------------------------------
    ("Stable prefix, append-only history (cache and the model's own reasoning)", [("LB-147", "append-only")]),
    ("LB-145 'built fresh each turn' refined: assembled once, appended, rebuilt only at compaction", [("LB-147", "refines lb-145")]),
    ("Detail on demand: tool search and practice playbooks", [("LB-148", "playbooks"), ("LB-148", "tool-search")]),
    ("Spent results cleared, large results paged, re-fetchable by locator", [("LB-149", "locator")]),
    ("Compaction rebuilds the brief from the checked file, not a chat summary; native compaction must pass the same check", [("LB-150", "not from a model's summary"), ("LB-150", "not harness-checked")]),
    ("Advocate's own memory, apart from matters, with consent", [("LB-151", "approval")]),
    ("Freshness: versioned reads, stale writes refused, change notices", [("LB-152", "change notice")]),
    ("Harness owns context policy on every provider; budget visible; one model per loop", [("LB-153", "one model")]),
    # --- tools --------------------------------------------------------------------------------
    ("Ten tool rules, the envelope, the registry", [("LB-154", "the envelope"), ("LB-154", "registry")]),
    ("Tool families: matter read/write, statutes, case law, computation, practice tables, checking/delegation, advocate/action, discovery",
     [("LB-155", "search_matter"), ("LB-156", "identify_act"), ("LB-157", "find_contrary_authority"), ("LB-158", "date_arithmetic"),
      ("LB-159", "pre_institution_steps"), ("LB-160", "oppose"), ("LB-161", "ask_advocate"), ("LB-162", "find_tool")]),
    ("Deliberately not tools: open web search, command line, direct database access", [("LB-162", "open web search")]),
    ("External acts are proposals the advocate approves", [("LB-161", "propose_action")]),
    ("About 25 of ~40 tools wrap existing code; build order core first", [("LB-154", "build order")]),
    # --- scratch pad, answer, board, opposing counsel ------------------------------------------
    ("Scratch pad closed by default, openable, live, grouped by dispute", [("LB-139", "closed by default"), ("LB-139", "grouped under each dispute")]),
    ("Scratch pad shows understanding, disputes, Acts, sections considered/kept/set aside, judgments and relevant passages, needs", [("LB-139", "sections considered"), ("LB-139", "passages found relevant")]),
    ("Scratch pad shows the tools called", [("LB-164", "tool called")]),
    ("Interim marked as working, not advice; set-aside reasons labelled as NM's judgment", [("LB-139", "not yet checked"), ("LB-139", "nm's judgment")]),
    ("Decision: scratch pad saved with the matter", [("LB-139", "saved with the matter")]),
    ("Decision: order guided by principles; final sections required", [("LB-163", "not enforced step by step")]),
    ("Guided method per message", [("LB-163", "state the understanding")]),
    ("Final answer per dispute: Act passages, case-law passages, evidence, case, arguments, other side, strengthen", [("LB-165", "evidence to collect"), ("LB-165", "how to strengthen")]),
    ("Final answer shows the sections finalised in the scratch pad", [("LB-165", "finalised")]),
    ("Matter board: disputes listed; questions per dispute from passages, each saying why", [("LB-166", "each item names the passage")]),
    ("Decision: opposing counsel in three passes (early per dispute, full per dispute after details, across matter at end)", [("LB-140", "three passes"), ("LB-140", "pass 3")]),
    ("Pass 1 shapes the board's questions", [("LB-166", "anticipated defence")]),
    ("Guardrails: every defence backed by a passage; only questions that change the position; re-run only on change", [("LB-140", "never ask a pass-1 question"), ("LB-140", "re-runs only when")]),
    ("Open: what counts as a dispute's key details for pass 2", [("LB-140", "key details")]),
    # --- review and status ----------------------------------------------------------------------
    ("How the owner reviews what was built: behaviour first, legal tables, rules, mutation, code by risk, PR per slice", [("LB-167", "legal tables")]),
    ("Build status recorded per row and re-measured at each slice close", [("LB-167", "build status")]),
    ("Slices cover everything now planned (loop foundation, tools, harness, scratch pad, answer and board, opposing counsel, context, switch)", [("LB-138", "scratch pad")]),
]

missing = []
for topic, needs in CHECKLIST:
    absent = [(rid, phrase) for rid, phrase in needs if phrase.lower() not in text.get(rid, "")]
    mark = "OK " if not absent else "GAP"
    print(f"{mark} {topic}" + ("" if not absent else f"  <- missing {absent}"))
    if absent:
        missing.append(topic)
print(f"\n{len(CHECKLIST) - len(missing)} of {len(CHECKLIST)} present; {len(missing)} gaps")
sys.exit(0)
