"""Update semantic-hardening Before Build owners and their existing plan mirror.

Examples::

    python development_environment/one_off_tools/brain_semantic_build_record.py --initialize
    python development_environment/one_off_tools/brain_semantic_build_record.py \
        --updates /tmp/semantic-checkpoint.json --proof /tmp/checkpoint-proof.json

The JSON update shape is {"rows": {"LB-194": {"checkpoint": "...", ...}}}.
Supported row keys: checkpoint, current_position, evidence (list of paths),
tested_version, test_environment, remaining, build_status, verification_status,
and before_build (explicit column-letter/value edits in B:M). Original A values
are immutable after creation. Mirrored plan fields are derived from Before Build.

This is workbook authoring tooling, not application execution or semantic
acceptance. Unrelated values, styles and native workbook controls must survive
the saved-byte round trip. Existing mirror defects are reported and may not grow.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string
from openpyxl.xml.functions import tostring

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

WORKBOOK = ROOT / "docs/Nyaymalaw_Implementation_Plan.xlsx"
PLAN_NOTE = "docs/backlog/nm-brain-semantic-hardening-20261006.md"
MIRROR = {8: 1, 9: 4, 10: 2, 12: 3, 15: 4, 16: 5, 17: 6, 19: 7, 28: 9, 30: 10, 33: 8}
ROW_PATTERN = re.compile(r"(LB-\d+|OM-[PIQ]\d+)\b")
CELL_LIMIT = 32767

OWNERS = {
    "LB-194": {
        "title": "Constrained response and evidence rendering",
        "objective": "Prevent fabricated effect or factual prose from reaching the public reply after an incorrect semantic acceptance.",
        "requirement": "Fresh response expressions use a versioned constrained rendering contract across every block kind, including accounts, reasoning, questions and recovery. Code resolves original evidence expressions and confirmed execution results. A block label, valid citation or positive reviewer verdict cannot authorize arbitrary consequential text. Preserve genuinely requested independent work and saved historical replay.",
        "entry": "Writer, response reviewer, composition or saved-reply replay selects consequential text for release.",
        "exit": "Both demonstrated public false-prose releases are prevented through the saved-and-released boundary, with legitimate neighbouring replies preserved; real-model and browser acceptance remain unverified.",
        "recovery": "Reject only the unsafe dependent expression where the release contract permits. Use the existing bounded correction and truthful confirmed status; never release an unchecked caveat or ask the advocate to restart accepted work.",
        "exclusions": "No phrase lists, scenario branches, arbitrary free-text escape hatch or claim that exact evidence references prove semantic meaning. Keep versioned historical rendering applicable to its original durable reply.",
        "acceptance": "Inject a wrong accepting reviewer and operational/factual lies in every response kind. Inspect delivered wording, confirmed state, partial preservation and replay. Pair with legitimate explanations, questions, no-change outcomes, quotes and requested work products. Verify unchanged routine calls.",
        "dependencies": "LB-176/179/182/183/193 remain existing policy owners; LB-194 owns this concrete expression-contract slice. One production file and at most one prompt per verified checkpoint; surface inseparable handoffs before expanding scope.",
    },
    "LB-195": {
        "title": "Proposition grounding and substantive source portions",
        "objective": "Ground complete propositions against original sources without mistaking matching words or upstream labels for meaning.",
        "requirement": "Review original actor, event, attribution, negation, chronology, uncertainty and target association independently. Identify owned exact substantive portions in mixed passages while retaining their complete context. Code validates original text, bounds, ownership and references; semantic support remains an independently reviewed judgment. Preserve legacy durable source treatment replay explicitly.",
        "entry": "Material/dispute verification, source-purpose classification or reconsideration determines support for a candidate proposition.",
        "exit": "Independent reviewers receive original evidence without inherited classifier conclusions; supported propositions select substantive portions and invalid references cannot be admitted. Paired public-boundary checks preserve faithful paraphrases and original qualifications.",
        "recovery": "Use precise source-linked disagreements and existing candidate-free reconsideration. Anchor changes invalidate dependent decisions. Preserve supported peers and never silently resolve attribution, identity or competing factual interpretations.",
        "exclusions": "No keyword truth checks, isolated number matching, sentence-to-record quotas, confidence voting or assumption that an exact quote establishes an adopted fact or entails a paraphrase.",
        "acceptance": "Pair negated dates, attributed opponent accounts, several actors/events, mixed instructions/accounts and uncertain statements with their accurate alternatives. Exercise source identity, context-preserving overlapping portions, cache invalidation, saved bindings and replay.",
        "dependencies": "LB-178/179 source and grounding owners; LB-181 dependency validity; LB-183 call accounting. Strengthen existing independent reviewers and source owner rather than add a routine semantic stage.",
    },
    "LB-196": {
        "title": "Whole-account dispositions and targeted semantic recovery",
        "objective": "Expose omissions and recover source disagreements while preserving faithful existing records and valid empty extraction.",
        "requirement": "Existing coverage review records source-linked material proposition dispositions: accepted/currently represented, missing, unresolved, outside authorized scope or non-account. Rejected candidates cannot satisfy coverage. Coverage-backed source-purpose disagreements may trigger the existing turn-owned bounded reconsideration even without an extracted candidate.",
        "entry": "A material/dispute reader returns proposals or no proposals, coverage identifies an omission, or source-purpose and independent evidence review disagree.",
        "exit": "Missing substantive work remains explicitly unfinished; already represented account and legitimate empty extraction remain accepted. Recovery is scoped and bounded by one turn-owned ledger, with admitted input and independent work preserved.",
        "recovery": "Recheck original source purpose and the affected missing proposition within existing recovery bounds. No duplicated accepted input, repeated useful extraction or whole-turn discard for harmless metadata. Foundational ownership/history/commit failures remain whole-turn stops.",
        "exclusions": "No demand that every sentence creates a new record, invented coverage from an empty envelope, new routine review call or recovery bound reset in a nested activity.",
        "acceptance": "Exercise omission with no candidate, complete empty extraction, already faithful record, mixed source disagreement, changed anchor dependency, exhausted bound, unavailable reviewer and independently valid sibling. Inspect saved records and useful delivered work.",
        "dependencies": "LB-175/178/179 coverage/source owners; LB-182 recovery/save; LB-183 call inventory; LB-195 owned source portions. Handoffs must be verified through the existing public service.",
    },
    "LB-197": {
        "title": "End-to-end paired semantic qualification",
        "objective": "Separate unsafe admission, unnecessary rejection and useful delivery on the actual saved-and-released flow.",
        "requirement": "Turn eight remaining characterization probes into appropriately attributed public-flow regressions and paired legitimate neighbours. Keep standalone untrusted-proposal boundaries distinct from delivered failures. Include deliberately incorrect accepting reviewers to establish mechanical guarantees, then qualify semantic quality on the actual pinned model when available.",
        "entry": "Each semantic-hardening checkpoint changes a prompt, schema, grounding, recovery, rendering or replay boundary.",
        "exit": "Affected integrated offline behaviours and call counts have concrete evidence; residual semantic limits are explicitly recorded. No live semantic/browser qualification or zero-error-rate claim without those measurements.",
        "recovery": "Identify the failed generic invariant, retain independently valid outcomes, and return to the owning slice. New blocking checks need paired accurate counterexamples; gate counts alone cannot establish a false-positive rate.",
        "exclusions": "No real-model API/browser calls in this authorized offline phase, fabricated completion, repeated counts presented as population error rates, or production branches based on golden scenarios.",
        "acceptance": "Inspect exact original inputs, fabricated outputs, saved records, released expressions, partial delivery and durable replay. Measure false acceptance/rejection separately, retries, useful results and logical/provider call counts. Run unfamiliar compositions and paired meaningful variations.",
        "dependencies": "LB-183 measured call inventory; LB-194/195/196 concrete contracts; AGENTS standing rules and Before Build status owner. Real-model/browser evaluation is deferred by the user's explicit instruction.",
    },
}


def _xml(value):
    return tostring(value.to_tree()) if value is not None else None


def _cell(cell):
    comment = cell.comment
    return (
        copy.deepcopy(cell.value), cell.data_type, copy.copy(cell._style),
        _xml(cell.hyperlink),
        (comment.text, comment.author, comment.width, comment.height) if comment else None,
    )


def _row_dimension(value):
    return dict(value), copy.copy(value._style)


def _features(sheet):
    return {
        "merged": tuple(sorted(str(r) for r in sheet.merged_cells.ranges)),
        "freeze": sheet.freeze_panes, "state": sheet.sheet_state,
        "validations": _xml(sheet.data_validations),
        "conditional": tuple((_xml(k), tuple(_xml(r) for r in v))
                             for k, v in sheet.conditional_formatting._cf_rules.items()),
        "views": _xml(sheet.views), "protection": _xml(sheet.protection),
        "page_margins": _xml(sheet.page_margins), "print_options": _xml(sheet.print_options),
        "page_setup": _xml(sheet.page_setup), "sheet_properties": _xml(sheet.sheet_properties),
        "auto_filter": _xml(sheet.auto_filter), "format": _xml(sheet.sheet_format),
        "print_area": str(sheet.print_area), "print_titles": sheet.print_titles,
        "row_breaks": _xml(sheet.row_breaks), "col_breaks": _xml(sheet.col_breaks),
        "columns": {k: _xml(v) for k, v in sheet.column_dimensions.items()},
        "tables": {k: _xml(sheet.tables[k]) for k in sheet.tables},
        "drawings": (len(sheet._charts), len(sheet._images)),
    }


def _workbook_features(book):
    properties = copy.copy(book.properties)
    # openpyxl sets the normal last-modified timestamp on every save.
    properties.modified = properties.created
    return {
        "sheets": book.sheetnames, "properties": _xml(properties),
        "security": _xml(book.security), "calculation": _xml(book.calculation),
        "views": tuple(_xml(v) for v in book.views),
        "defined_names": tuple((k, _xml(v)) for k, v in book.defined_names.items()),
        "named_styles": tuple((v.name, v.builtinId, v.hidden, tuple(v.as_tuple())) for v in book._named_styles),
        "external_links": tuple(_xml(v) for v in book._external_links),
    }


def _indices(sheet, original):
    result = {}
    for row in range(5 if original else 2, sheet.max_row + 1):
        value = str(sheet.cell(row, 1).value or "")
        match = ROW_PATTERN.match(value)
        if match:
            ident = match[1]
            if ident in result:
                raise ValueError(f"duplicate {sheet.title} owner: {ident}")
            result[ident] = row
    return result


def _bounded(value):
    if isinstance(value, str) and len(value) > CELL_LIMIT:
        raise ValueError(f"cell exceeds Excel's {CELL_LIMIT}-character limit; use a new owned row or evidence file")
    return value


def update_workbook(path=WORKBOOK, *, initialize=False, updates=None, proof_path=None, dry_run=False):
    """Apply only named owner updates; publish after saved-byte preservation checks."""
    from assurance.control_plane.plan_scenarios import requirement_problems, sheet_rows

    path = Path(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    book = load_workbook(path)
    original_cells = {(s.title, cell.coordinate): _cell(cell)
                      for s in book for cell in s._cells.values()}
    original_features = {s.title: _features(s) for s in book}
    original_rows = {s.title: {k: _row_dimension(v) for k, v in s.row_dimensions.items()} for s in book}
    original_book = _workbook_features(book)
    before_problems = set(requirement_problems(path))
    feature_count = len(sheet_rows(path))
    before, plan = book["Before Build"], book["Implementation Plan"]
    indices, plan_indices = _indices(before, True), _indices(plan, False)
    touched, created, allowed = set(), [], set()

    def write(sheet, row, col, value):
        cell = sheet.cell(row, col)
        cell.value = _bounded(value)
        allowed.add((sheet.title, cell.coordinate))

    if initialize:
        for ident, owner in OWNERS.items():
            if (ident in indices) != (ident in plan_indices):
                raise ValueError(f"{ident}: existing owner/mirror presence disagrees")
            if ident in indices:
                continue
            row = before.max_row + 1
            indices[ident] = row
            for col in range(1, before.max_column + 1):
                before.cell(row, col)._style = copy.copy(before.cell(row - 1, col)._style)
                allowed.add((before.title, before.cell(row, col).coordinate))
            note = ("6 October 2026 planning checkpoint: user authorized the remaining semantic-hardening work. "
                    "Implementation in progress; this row records the plan, not completed code or semantic acceptance. "
                    f"Detailed queue and invariant/pass conditions: {PLAN_NOTE}. No real-model/API/browser calls.")
            values = [f"{ident}\n{owner['title']}\nOwner-authorized semantic-hardening plan recorded 6 October 2026",
                      owner["objective"], owner["entry"], owner["requirement"], owner["exit"], owner["recovery"],
                      owner["exclusions"], owner["acceptance"], owner["dependencies"], note, "Active",
                      "In progress; planning recorded, implementation not yet verified. Real-model/browser acceptance deferred.", PLAN_NOTE]
            for col, value in enumerate(values, 1):
                write(before, row, col, value)
            before.row_dimensions[row].height = before.row_dimensions[row - 1].height
            target = plan.max_row + 1
            plan_indices[ident] = target
            for col in range(1, plan.max_column + 1):
                plan.cell(target, col)._style = copy.copy(plan.cell(target - 1, col)._style)
                allowed.add((plan.title, plan.cell(target, col).coordinate))
            for col, value in {1: ident, 2: "L", 3: "Requirement", 4: "Legal brain",
                               5: "Semantic hardening", 6: owner["title"], 7: "Not a release decision",
                               31: "Implementation owner", 32: "6 October 2026 initial authorized plan; not acceptance.",
                               35: "Offline integrated/paired fixtures; live model/browser deferred",
                               38: "Review required", 39: "In progress", 40: "Not verified", 41: PLAN_NOTE,
                               42: "No implementation checkpoint asserted", 43: "2026-10-06; planning only",
                               44: before.cell(row, 12).value, 45: "No release or professional sign-off"}.items():
                write(plan, target, col, value)
            touched.add(ident)
            created.append(ident)

    for ident, spec in (updates or {}).items():
        if ident not in indices or ident not in plan_indices:
            raise ValueError(f"unknown owner/mirror: {ident}; initialize new authorized owners first")
        if not isinstance(spec, dict):
            raise ValueError(f"{ident}: update must be an object")
        expected = {"checkpoint", "current_position", "evidence", "tested_version", "test_environment",
                    "remaining", "build_status", "verification_status", "before_build"}
        if set(spec) - expected:
            raise ValueError(f"{ident}: unknown update fields {sorted(set(spec) - expected)}")
        row, target = indices[ident], plan_indices[ident]
        for letter, value in spec.get("before_build", {}).items():
            col = column_index_from_string(letter)
            if not 2 <= col <= 13:
                raise ValueError(f"{ident}: original A is immutable and only B:M can be updated")
            write(before, row, col, value)
        checkpoint = spec.get("checkpoint")
        if checkpoint:
            previous = str(before.cell(row, 10).value or "")
            if checkpoint not in previous:
                write(before, row, 10, previous + ("\n\n" if previous else "") + checkpoint)
            write(plan, target, 32, checkpoint)
        if "current_position" in spec:
            write(before, row, 12, spec["current_position"])
        if "evidence" in spec:
            evidence = spec["evidence"]
            if not isinstance(evidence, list) or not all(isinstance(x, str) for x in evidence):
                raise ValueError(f"{ident}: evidence must be a list of paths")
            previous = str(plan.cell(target, 41).value or "").splitlines()
            write(plan, target, 41, "\n".join(dict.fromkeys(previous + evidence)))
        for key, col in {"tested_version": 42, "test_environment": 43, "remaining": 44,
                         "build_status": 39, "verification_status": 40}.items():
            if key in spec:
                write(plan, target, col, spec[key])
        if "remaining" not in spec and "current_position" in spec:
            write(plan, target, 44, spec["current_position"])
        touched.add(ident)

    for ident in touched:
        for target_col, source_col in MIRROR.items():
            write(plan, plan_indices[ident], target_col, before.cell(indices[ident], source_col).value)

    wanted = {(s.title, cell.coordinate): _cell(cell) for s in book for cell in s._cells.values()
              if (s.title, cell.coordinate) in allowed}
    with tempfile.TemporaryDirectory(prefix="nm-semantic-workbook-", dir=path.parent) as temp:
        candidate = Path(temp) / path.name
        book.save(candidate)
        check = load_workbook(candidate)
        if _workbook_features(check) != original_book:
            raise ValueError("native workbook controls changed")
        for sheet in check:
            if _features(sheet) != original_features[sheet.title]:
                raise ValueError(f"native sheet controls changed: {sheet.title}")
            for number, old in original_rows[sheet.title].items():
                if _row_dimension(sheet.row_dimensions[number]) != old:
                    raise ValueError(f"existing row dimension changed: {sheet.title} {number}")
        for (title, coord), old in original_cells.items():
            got = _cell(check[title][coord])
            if got[2:] != old[2:]:
                raise ValueError(f"existing style/comment/link changed: {title} {coord}")
            if (title, coord) not in allowed and got != old:
                raise ValueError(f"unrelated cell changed: {title} {coord}")
        for (title, coord), expected_cell in wanted.items():
            if _cell(check[title][coord]) != expected_cell:
                raise ValueError(f"target cell failed saved-byte roundtrip: {title} {coord}")
        for sheet in check:
            for cell in sheet._cells.values():
                key = (sheet.title, cell.coordinate)
                if key not in original_cells and key not in allowed and cell.value is not None:
                    raise ValueError(f"unplanned cell appeared: {sheet.title} {cell.coordinate}")
                _bounded(cell.value)
        after_problems = set(requirement_problems(candidate))
        if after_problems - before_problems:
            raise ValueError(f"new mirror problems: {sorted(after_problems - before_problems)}")
        if len(sheet_rows(candidate)) != feature_count:
            raise ValueError("executable feature population changed")
        if _indices(check["Before Build"], True) != indices or _indices(check["Implementation Plan"], False) != plan_indices:
            raise ValueError("owner identity/order changed")
        changed = [f"{title}!{coord}" for (title, coord), value in wanted.items()
                   if original_cells.get((title, coord)) != value]
        proof = {"workbook": str(path.relative_to(ROOT)) if path.is_relative_to(ROOT) else str(path),
                 "source_sha256": digest, "output_sha256": hashlib.sha256(candidate.read_bytes()).hexdigest(),
                 "dry_run": dry_run, "created_owners": created, "touched_owners": sorted(touched),
                 "changed_cells": sorted(changed), "unrelated_cells_changed": 0,
                 "existing_styles_comments_links_changed": 0, "native_controls_changed": 0,
                 "normal_last_modified_timestamp_excluded_from_control_comparison": True,
                 "existing_row_dimensions_changed": 0, "executable_feature_population": feature_count,
                 "baseline_mirror_problems": sorted(before_problems), "remaining_mirror_problems": sorted(after_problems),
                 "new_mirror_problems": [], "max_cell_characters": max(len(str(c.value or "")) for s in check for c in s._cells.values()),
                 "no_completion_or_live_semantic_acceptance_asserted": True}
        check.close()
        if not dry_run and changed:
            if hashlib.sha256(path.read_bytes()).hexdigest() != digest:
                raise ValueError("workbook changed concurrently; rerun against current bytes")
            candidate.replace(path)
        book.close()
    if proof_path:
        Path(proof_path).write_text(json.dumps(proof, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return proof


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=WORKBOOK)
    parser.add_argument("--initialize", action="store_true")
    parser.add_argument("--updates", type=Path)
    parser.add_argument("--proof", type=Path)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    updates = json.loads(args.updates.read_text(encoding="utf-8")) if args.updates else {}
    if set(updates) - {"rows"}:
        parser.error("update document supports only the rows key")
    result = update_workbook(args.workbook, initialize=args.initialize,
                             updates=updates.get("rows", {}), proof_path=args.proof, dry_run=args.dry_run)
    print(json.dumps({k: v for k, v in result.items() if k != "changed_cells"}, indent=2))
    print(f"Saved-byte preservation checked; {len(result['changed_cells'])} cells changed.")


if __name__ == "__main__":
    main()
