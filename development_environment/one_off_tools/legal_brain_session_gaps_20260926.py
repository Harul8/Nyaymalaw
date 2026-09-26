"""Close the gaps found when the session's decisions were checked against the sheet.

OWNER DIRECTION, 26 September 2026: re-read everything discussed and settled in
this session and make sure all of it is in the sheet.

A checklist of 60 items was run against the Before Build text
(scratchpad session_checklist.py). 52 were present, one false alarm was a
search phrase, and seven were genuine gaps. This tool closes them:

  1. L.2 header      -- the five loops, named as a set
  2. L.4 header      -- the corpus is the ceiling
  3. LB-163          -- the legal reasoning order as guidance, including remedy
  4. LB-165          -- the answer shows the sections finalised, and the relief sought
  5-6. LB-167 (new)  -- how the owner reviews what is built, and a status that stays true
  7. LB-138          -- the slice list brought up to everything now planned

Edits to LB-163, LB-165 and LB-138 are exact substring replacements, each
asserted to match once, in rows that are owner-directed drafts. Every
owner-authored row is proved byte-identical.
"""
import copy
import hashlib
import importlib.util
import re
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

_REPO = Path(__file__).resolve().parents[2]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))
from assurance.common._console import utf8_console  # noqa: E402

utf8_console()

_spec = importlib.util.spec_from_file_location(
    "regroup", Path(__file__).with_name("legal_brain_regroup_20260926.py"))
regroup = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(regroup)

SOURCE, TODAY, FIRST = regroup.SOURCE, regroup.TODAY, regroup.FIRST
HEADER = re.compile(r"^(L|L\.\d|U|U\.\d|P|X)  ")

HEADER_ADDITIONS = {
    "L.2": (" Five loops run at different speeds: the turn loop; the research loop (LB-136); the "
            "opposing-counsel passes (LB-140); the matter loop, which records what each turn established "
            "into the file (LB-142); and the improvement loop, which turns every defect into a golden case "
            "(LB-146)."),
    "L.4": (" The corpus is the ceiling: retrieval cannot ground what is not held, so a measured gap -- today "
            "the Commercial Courts Act and the Telangana court-fee and suits-valuation schedules, measured "
            "against the manifest -- is closed by acquiring the source, never by the model filling it."),
}

#: (row, column, old, new). Each old string must occur exactly once.
EDITS = [
    ("LB-163", 4,
     "(5) from the passages, work out what the dispute needs -- facts, evidence, procedure;",
     "(5) from the passages, work out what the dispute needs -- facts, evidence, procedure, and the "
     "relief it can obtain;"),
    ("LB-163", 4,
     "The model may go back",
     "Within each dispute the reasoning follows an advocate's order as guidance -- facts, issues, statute, "
     "binding authority, application to the facts, procedure (limitation, forum, notice), remedy, risk and "
     "the next step -- taken in whatever order the matter calls for. The model may go back"),
    ("LB-165", 4,
     "RELEVANT ACT PASSAGES (retrieved text, with locators)",
     "RELEVANT ACT PASSAGES (the sections finalised for this dispute in the scratch pad, their retrieved "
     "text with locators)"),
    ("LB-165", 4,
     "THE CASE TO PREPARE (cause, forum and procedural steps --",
     "THE CASE TO PREPARE (cause, relief sought, forum and procedural steps --"),
    ("LB-138", 4,
     "Slices, one pull request each, reviewed by the owner: (1) tool calling, loop runner, step log; (2) "
     "principles and tools; (3) output checks and repair; (4) context and research; (5) switch, comparison "
     "and live progress in the interface.",
     "Slices, one pull request each, reviewed by the owner (LB-167) -- PROPOSED ORDER, for the owner to "
     "confirm: (1) the loop foundation -- tool calling across providers, the loop runner with budgets, typed "
     "step events and the saved step log, the append-only conversation, and record-and-replay; (2) the "
     "result envelope, the tool registry and the core tools wrapping existing code; (3) the harness after "
     "the loop -- the output checks, the professional boundaries on tool calls, repair and the independent "
     "verifier; (4) streaming and the scratch pad; (5) the per-dispute answer and the matter-board "
     "questions; (6) opposing counsel in three passes; (7) context -- compaction from the file, clearing, "
     "the advocate memory and playbooks; (8) document extraction, which unblocks the rows waiting on "
     "uploaded material; (9) the golden comparison and the switch."),
]

NEW_ID = "LB-167"
NEW_COLS = [
    "The owner's review of what is built, and a status that stays true",
    "The owner reviews the legal brain by risk -- behaviour first, then the legal tables, the rules, the "
    "controls and the riskiest code -- and every row's build status is re-measured from the code at each "
    "slice close.",
    "Know what NM does and how far to trust it, without reading 66,000 lines of code.",
    "The owner, before relying on a capability, and at every slice close.",
    "REVIEW BY RISK, in this order: (1) use it as an advocate -- put matters the owner knows well through "
    "the product and read the answers as a senior reads a junior's draft; every wrong answer becomes a "
    "golden case (LB-146); (2) the legal tables -- counsel review of every curated entry in "
    "backend/nm/knowledge (elements, limitation edges, pre-institution conditions, governing code, interim "
    "tests, procedural periods, filing requirements), each curated_from checked against the actual section "
    "or judgment, because no test can catch a rule that is legally wrong; (3) the rules -- the test names "
    "read as a rulebook, and the gate matrix; (4) that the controls bite -- the mutation tool breaks the code "
    "on purpose and the right test must fail; (5) code by risk -- the composition root, the edge's sign-in, "
    "CSRF and ownership checks, and the turn's run order. EVERY SLICE IS A PULL REQUEST the owner reviews -- "
    "its description, test names and a skim of the diff -- and nothing merges without the owner; an "
    "automated code review is a first pass, never the review. BUILD STATUS is recorded per row in the "
    "Implementation Plan -- Built, In progress (partially built) or Not started, with the evidence paths and "
    "what remains -- and re-measured from the code at each slice close; verification is separate and is set "
    "only when acceptance criteria are run as written.",
    "The owner has reviewed each capability before relying on it, and every row's status reflects the code.",
    "A status that cannot be supported by a path that exists is refused by the status tool.",
    "Never mark a row built on the strength of a plan or a commit message. Never treat an automated review "
    "as the owner's. Never let a slice merge unreviewed.",
    "LB-167-AC1: at each slice close the status tool re-runs and every cited evidence path exists.\n"
    "LB-167-AC2: each slice's pull request lists the rows it moves and their new status.\n"
    "LB-167-AC3: every curated legal table has a recorded counsel review before its rows are marked built.",
    "LB-43, LB-44, LB-72, LB-74, LB-138; assurance/gate/mutate.py; "
    "development_environment/one_off_tools/legal_brain_status_20260926.py; "
    "development_environment/reviews/LEGAL_BRAIN_STATUS_20260926.md.",
]


def main() -> int:
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    book = load_workbook(SOURCE)
    sheet, plan = book["Before Build"], book["Implementation Plan"]
    features = {s.title: regroup.sheet_features(s) for s in book}
    before = {(s.title, c.coordinate): (c.value, copy.copy(c._style)) for s in book for row in s for c in row}
    heights = {r: sheet.row_dimensions[r].height for r in range(1, sheet.max_row + 1)}

    order = []
    for r in range(FIRST, sheet.max_row + 1):
        text = str(sheet.cell(r, 1).value or "")
        h = HEADER.match(text)
        key = f"HEADER:{h[1]}" if h else regroup.key_of(text)
        order.append((key, {"values": [sheet.cell(r, c).value for c in range(1, 11)],
                            "styles": [copy.copy(sheet.cell(r, c)._style) for c in range(1, 11)],
                            "height": sheet.row_dimensions[r].height, "row": r}))
    source = dict(order)
    assert NEW_ID not in source

    # ---- the expected values of every row that changes ---------------------------
    expected = {}
    for key, row in order:
        values = list(row["values"])
        if key.startswith("HEADER:") and key[7:] in HEADER_ADDITIONS:
            values[0] += HEADER_ADDITIONS[key[7:]]
        for ident, col, old, new in EDITS:
            if key == ident:
                assert values[col - 1].count(old) == 1, (ident, old[:50])
                values[col - 1] = values[col - 1].replace(old, new)
        expected[key] = values
    new_values = [f"{NEW_ID}\nLegal brain › L.9 Models, evaluation and build discipline\n{NEW_COLS[0]}\n\n"
                  f"OWNER-DIRECTED REQUIREMENT, DRAFTED FOR REVIEW:\n{NEW_COLS[1]}",
                  *NEW_COLS[2:],
                  regroup.READINESS + " Owner direction, 26 September 2026: review what has been built, and keep "
                                      "its status recorded."]
    assert len(new_values) == 10
    expected[NEW_ID] = new_values

    body = []
    for key, row in order:
        body.append((key, row))
        if key == "LB-146":
            body.append((NEW_ID, {"values": None, "styles": [copy.copy(s) for s in source["LB-153"]["styles"]],
                                  "height": 280, "row": None}))
    assert len(body) == len(order) + 1

    group, index, layout = None, 0, []
    for offset, (key, row) in enumerate(body):
        r = FIRST + offset
        is_header = key.startswith("HEADER:")
        if is_header:
            group, index = key[7:], 0
        values = expected[key]
        for c in range(1, 11):
            cell = sheet.cell(r, c)
            cell.value = values[c - 1]
            cell._style = copy.copy(row["styles"][c - 1])
            if not is_header and group == "L.9":
                cell.fill = PatternFill("solid", fgColor=regroup.ZEBRA[index % 2])
        if not is_header:
            index += 1
        sheet.row_dimensions[r].height = row["height"]
        layout.append((key, r, group))
    last = FIRST + len(body) - 1
    sheet.auto_filter.ref = f"A4:J{last}"

    population = sorted({regroup.LBOM.match(str(sheet.cell(x, 1).value or ""))[1]
                         for x in range(5, last + 1) if regroup.LBOM.match(str(sheet.cell(x, 1).value or ""))})
    lb = [p for p in population if p.startswith("LB-")]
    note = str(sheet.cell(2, 1).value)
    assert "LB-126–166" in note
    note = note.replace("LB-126–166", "LB-126–167")
    note, n = re.subn(r"The legal brain holds \d+ LB requirements", f"The legal brain holds {len(lb)} LB requirements", note)
    assert n == 1
    note, n = regroup.NOTE_PATTERN.subn(f"All {len(population)} LB/OM requirements are linked into Implementation Plan", note)
    assert n == 1
    sheet.cell(2, 1).value = note

    # ---- the Implementation Plan ---------------------------------------------------
    position = {k: r for k, r, _g in layout}
    plan_changed = set()
    for x in range(2, plan.max_row + 1):
        ident = plan.cell(x, 1).value
        if ident in {e[0] for e in EDITS}:
            for t, origin in regroup.MIRROR.items():
                v = sheet.cell(position[ident], origin).value
                if plan.cell(x, t).value != v:
                    plan.cell(x, t).value = v
                    plan_changed.add(plan.cell(x, t).coordinate)
    target = plan.max_row + 1
    for c in range(1, plan.max_column + 1):
        plan.cell(target, c)._style = copy.copy(plan.cell(target - 1, c)._style)
    fixed = {1: NEW_ID, 3: "Requirement", 4: "Legal brain", 5: "L.9 Models, evaluation and build discipline",
             6: NEW_COLS[0], 7: "Pilot", 38: "Draft", 39: "In progress", 40: "Not verified",
             41: "development_environment/one_off_tools/legal_brain_status_20260926.py; "
                 "development_environment/reviews/LEGAL_BRAIN_STATUS_20260926.md",
             43: f"{TODAY}: the first status pass is recorded; no review step has begun.",
             44: "Owner review by risk, counsel review of the curated tables, and per-slice pull-request review "
                 "have not begun.",
             32: f"{TODAY}: owner-directed row mirrored from Before Build. Not acceptance, not an approved "
                 "requirement."}
    for c, v in fixed.items():
        plan.cell(target, c, v)
        plan_changed.add(plan.cell(target, c).coordinate)
    for t, origin in regroup.MIRROR.items():
        plan.cell(target, t, sheet.cell(position[NEW_ID], origin).value)
        plan_changed.add(plan.cell(target, t).coordinate)

    out = Path("/tmp/claude-0/-home-user-Nyaymalaw/82a504d3-1374-501f-9fe3-441ac9aa966e/scratchpad") / SOURCE.name
    book.save(out)

    # ================= THE PROOF, on the saved bytes ==========================
    from openpyxl.xml.functions import tostring

    def xml(part):
        return tostring(part.to_tree())

    check = load_workbook(out)
    for tab in check:
        assert regroup.sheet_features(tab) == features[tab.title], tab.title
    bb = check["Before Build"]
    original = load_workbook(SOURCE)["Before Build"]
    for (title, coord), old in before.items():
        if title == "Before Build" and int(re.sub(r"[A-Z]+", "", coord)) < FIRST and coord != "A2":
            assert (bb[coord].value, bb[coord]._style) == old, coord
    for x in range(1, FIRST):
        assert bb.row_dimensions[x].height == heights.get(x)
    touched = {e[0] for e in EDITS} | {f"HEADER:{g}" for g in HEADER_ADDITIONS} | {NEW_ID}
    for key, r, group in layout:
        got = [bb.cell(r, c).value for c in range(1, 11)]
        assert got == expected[key], key
        if key == NEW_ID:
            assert group == "L.9"
            continue
        if key not in touched:
            assert got == source[key]["values"], f"{key}: changed but not meant to"
        for c in range(1, 11):
            a, o = bb.cell(r, c), original.cell(source[key]["row"], c)
            assert xml(a.border) == xml(o.border) and xml(a.alignment) == xml(o.alignment), key
            assert (a.font.name, a.font.sz, a.font.color, a.font.b) == (o.font.name, o.font.sz, o.font.color, o.font.b), key
    ip = check["Implementation Plan"]
    for (title, coord), old in before.items():
        if title == "Implementation Plan" and coord not in plan_changed:
            assert (ip[coord].value, ip[coord]._style) == old, f"plan {coord} moved"
    from assurance.control_plane import plan_scenarios
    assert plan_scenarios.requirement_problems(out) == []
    assert len(plan_scenarios.sheet_rows(out)) == 93

    shutil.copy2(out, SOURCE)
    print(f"source digest before: {digest}")
    print(f"source digest after:  {hashlib.sha256(SOURCE.read_bytes()).hexdigest()}")
    print(f"{NEW_ID} at row {position[NEW_ID]}; edits: {len(EDITS)}; headers: {sorted(HEADER_ADDITIONS)}; "
          f"LB/OM {len(population)} (LB {len(lb)})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
