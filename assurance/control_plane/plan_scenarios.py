"""The Given/When/Then scenarios live in the implementation plan sheet.

    python assurance/control_plane/plan_scenarios.py --check   # exit 1 if a file differs
    python assurance/control_plane/plan_scenarios.py --write   # regenerate the files

WHY THE SHEET AND NOT THE FILES
---------------------------------
The product owner reads and decides from `docs/Nyaymalaw_Implementation_Plan.xlsx`.
Current prose requirements and historical evidence are retained separately from
the explicitly executable Gherkin subset. Only rows marked Executable generate
pytest-bdd files. Planned rows remain in the denominator: reconciliation proves
agreement with the sheet, not feature completion, coverage or a passing test run.

ONE OWNER, AND A CHECK THAT REFUSES THE SECOND. A generated file that someone edits by
hand would quietly test something the sheet does not say.
`tests/test_the_scenarios_come_from_the_plan.py` fails the build when any file differs
from the sheet, when a file has no row, and when a row has no file.
"""
from __future__ import annotations

import argparse
import re
import sys
from contextlib import closing
from pathlib import Path

_REPO = Path(__file__).resolve().parents[2]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))
from assurance.common._console import utf8_console  # noqa: E402

utf8_console()

ROOT = Path(__file__).resolve().parents[2]
SHEET = ROOT / "docs" / "Nyaymalaw_Implementation_Plan.xlsx"
FEATURES = ROOT / "tests" / "features"
COLUMN = "Executable scenarios (Gherkin)"
STATE_COLUMN = "Executable scenario state"
NOTES_COLUMN = "Executable scenario coverage notes"
WORKSHEET = "Implementation Plan"

#: The directory each journey phase's feature files are written to.
PHASE_DIRECTORIES = {
    "A": "arrive", "B": "open_a_matter", "C": "take_the_brief", "D": "work_the_file",
    "E": "advise", "F": "act", "G": "carry", "H": "close", "I": "leave",
}


def header(row_id: str) -> str:
    return (f"# GENERATED from docs/Nyaymalaw_Implementation_Plan.xlsx, row {row_id}, "
            f'column "{COLUMN}".\n# Edit the sheet, then run '
            f"`python assurance/control_plane/plan_scenarios.py --write`.\n")


def sheet_rows(sheet: Path = SHEET) -> list[dict]:
    """All planned features, including those without executable scenarios."""
    from openpyxl import load_workbook

    with closing(load_workbook(sheet, read_only=True, data_only=True)) as workbook:
        if WORKSHEET not in workbook.sheetnames:
            raise ValueError(f'the plan has no "{WORKSHEET}" worksheet')
        rows = workbook[WORKSHEET].iter_rows(values_only=True)
        headers = [str(cell or "").strip() for cell in next(rows)]
        required = ("ID", "Row type", "Feature / name", COLUMN, STATE_COLUMN, NOTES_COLUMN)
        for name in required:
            if headers.count(name) != 1:
                raise ValueError(f'the plan needs exactly one "{name}" column')
        found, seen = [], set()
        for values in rows:
            row = dict(zip(headers, values, strict=False))
            if str(row.get("Row type") or "").strip() != "Feature":
                continue
            row_id = str(row.get("ID") or "").strip()
            if not row_id or row_id in seen:
                raise ValueError(f"duplicate or missing feature ID: {row_id!r}")
            if not re.fullmatch(r"F-[A-I]-\d{2}", row_id):
                raise ValueError(f"invalid feature ID: {row_id!r}")
            seen.add(row_id)
            body = str(row.get(COLUMN) or "").strip()
            state = row.get(STATE_COLUMN)
            if state not in {"Executable", "Planned"}:
                raise ValueError(f"{row_id}: executable scenario state is missing or invalid")
            if (state == "Executable") != bool(body):
                raise ValueError(f"{row_id}: executable state and scenario body disagree")
            if not str(row.get(NOTES_COLUMN) or "").strip():
                raise ValueError(f"{row_id}: missing executable coverage limits")
            found.append(row)
        if not found or not any(row[STATE_COLUMN] == "Executable" for row in found):
            raise ValueError("the plan has an empty feature or executable population")
        return found


def requirement_problems(sheet: Path = SHEET) -> list[str]:
    """Both directions: every LB/OM requirement has one current, faithful row."""
    from openpyxl import load_workbook

    with closing(load_workbook(sheet, read_only=True, data_only=True)) as workbook:
        before, plan = workbook["Before Build"], workbook[WORKSHEET]
        authored = {}
        for row in before.iter_rows(min_row=5, values_only=True):
            match = re.match(r"(LB-\d+|OM-[PIQ]\d+)\b", str(row[0] or ""))
            if match:
                if match[1] in authored:
                    raise ValueError(f"duplicate Before Build requirement: {match[1]}")
                authored[match[1]] = row
        if not authored:
            return ["Before Build professional requirement population is empty"]
        current = {}
        found = []
        # Workbook coordinates are fixed source contracts, not substring matches.
        mapping = {8: 1, 9: 4, 10: 2, 12: 3, 15: 4, 16: 5,
                   17: 6, 19: 7, 28: 9, 30: 10, 33: 8}
        for row in plan.iter_rows(min_row=2, values_only=True):
            key = str(row[0] or "")
            if re.fullmatch(r"LB-\d+|OM-[PIQ]\d+", key):
                if key in current:
                    found.append(f"duplicate plan requirement: {key}")
                current[key] = row
        for key in sorted(authored.keys() ^ current.keys()):
            found.append(f"{key}: requirement missing from one sheet")
        for key in authored.keys() & current.keys():
            if current[key][2] != "Requirement":
                found.append(f"{key}: wrong row type")
            for target, source in mapping.items():
                if current[key][target-1] != authored[key][source-1]:
                    found.append(f"{key}: column {target} differs from Before Build")
        return found


def expected_files(rows: list[dict]) -> dict[Path, str]:
    """The exact bytes each generated file must hold, keyed by path under tests/features."""
    files: dict[Path, str] = {}
    for row in rows:
        if row[STATE_COLUMN] != "Executable":
            continue
        row_id = str(row.get("ID") or "").strip()
        parts = row_id.split("-")
        if len(parts) != 3 or parts[0] != "F" or parts[1] not in PHASE_DIRECTORIES:
            raise ValueError(f"row {row_id!r} carries scenarios but its ID is not F-<phase>-<nn>")
        scenarios = str(row[COLUMN]).replace("\r\n", "\n").strip("\n")
        if (not re.search(r"^\s*Scenario(?: Outline)?:\s+\S", scenarios, re.M)
                or re.search(r"^\s*Feature:", scenarios, re.M)):
            raise ValueError(f"{row_id}: executable scenarios are not Gherkin scenario bodies")
        title = str(row.get('Feature / name') or '').strip()
        text = (header(row_id) + f"Feature: {row_id} {title}\n\n"
                + scenarios + "\n")
        path = Path(PHASE_DIRECTORIES[parts[1]]) / f"{row_id}.feature"
        if path in files:
            raise ValueError(f"two rows claim the ID {row_id}")
        files[path] = text
    return files


def problems(expected: dict[Path, str], features: Path = FEATURES) -> list[str]:
    """Differences between the sheet and the files on disk. Empty means they agree."""
    found = []
    on_disk = {path.relative_to(features): path for path in features.rglob("*.feature")}
    for relative, text in sorted(expected.items()):
        path = on_disk.get(relative)
        if path is None:
            found.append(f"{relative}: in the sheet and not generated")
        elif path.read_text(encoding="utf8").replace("\r\n", "\n") != text:
            found.append(f"{relative}: differs from the sheet")
    for relative in sorted(set(on_disk) - set(expected)):
        found.append(f"{relative}: a feature file with no row in the sheet")
    return found


def write(expected: dict[Path, str], features: Path = FEATURES) -> None:
    for relative, text in expected.items():
        target = features / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(text, encoding="utf8", newline="\n")
    for path in features.rglob("*.feature"):
        if path.relative_to(features) not in expected:
            first = path.read_text(encoding="utf8").splitlines()[:1]
            if first and first[0].startswith("# GENERATED from docs/Nyaymalaw_Implementation_Plan"):
                path.unlink()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--write", action="store_true", help="regenerate the feature files")
    mode.add_argument("--check", action="store_true", help="check without writing (default)")
    args = parser.parse_args()
    rows = sheet_rows()
    expected = expected_files(rows)
    if args.write:
        write(expected)
    found = problems(expected) + requirement_problems()
    for line in found:
        print(line)
    print(f"{len(expected)} feature file(s) from the sheet; {len(found)} problem(s)")
    print(f"{len(rows)} planned features; {len(rows) - len(expected)} without executable "
          "Gherkin. Reconciliation is not acceptance or a coverage PASS.")
    return 1 if found else 0


if __name__ == "__main__":
    sys.exit(main())
