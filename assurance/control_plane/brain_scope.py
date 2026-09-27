"""Versioned clause contracts and prerequisite closure for a legal-brain slice.

Row ownership is necessary, not evidence coverage. This checker never turns a
planned check, a registered owner, or a build authorization into acceptance.
Recorded/synthetic engineering can be authorized while completion and client
cutover remain blocked. It does not grant permission to any runtime operation.
"""
from __future__ import annotations

import argparse
import ast
import hashlib
import json
import re
import sys
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
from assurance.common._console import utf8_console  # noqa: E402
from assurance.control_plane import requirement_owners as owners  # noqa: E402

utf8_console()

MAP = ROOT / "docs/blueprint/legal_brain_execution.json"
SELECTED_PACKETS = ("P49", "P50", "P51", "P52", "P53", "P54")
# These include every A-I Before Build promise mirrored into the delivery view,
# plus its explicit implementation/verification contracts. The J mirror mixes
# owner rationale with progress observations; it is not a promise fingerprint.
# Decisions remain separately blocked pending their canonical owner evidence.
PROMISE_COLUMNS = (
    "ID", "Original product-owner description", "Approved current requirement",
    "User objective", "Exclusions", "Actor and trigger", "Entry conditions",
    "Screen / layout reference", "Interaction and system response", "Exit conditions",
    "Failure / retry / resume", "Return triggers", "Must do", "Must never",
    "Inputs / outputs", "Data and provenance", "Permissions / privacy",
    "Interfaces and existing code to reuse", "AI / retrieval behaviour",
    "Applicable shared rules", "Quality targets", "Dependency IDs and required outputs",
    "Open questions", "Acceptance criteria", "Scenarios (Given / When / Then)",
    "Required verification methods", "Pass thresholds", "Negative controls",
)
REFERENCE = re.compile(r"\b(LB-\d+|OM-[PIQ]\d+|F-[A-I]-\d{2})\b")
RANGE = re.compile(r"\b(LB-|OM-[PIQ]|F-[A-I]-)(\d+)\s*[-–—�]\s*(\d+)\b")


@dataclass(frozen=True)
class SourceRow:
    ident: str
    clauses: tuple[str, ...]
    version: str
    dependencies: frozenset[str]
    ownership: owners.Requirement | None
    clause_text: tuple[tuple[str, str], ...] = ()


@dataclass(frozen=True)
class ScopeReport:
    packet: str
    selected_clauses: int
    contract_problems: tuple[str, ...]
    completion_blockers: tuple[str, ...]
    engineering_authorized: bool
    client_cutover_allowed: bool = False

    @property
    def contracts_ready(self) -> bool:
        return self.selected_clauses > 0 and not self.contract_problems

    @property
    def completion_ready(self) -> bool:
        return self.contracts_ready and not self.completion_blockers


def references(text: str) -> frozenset[str]:
    """Read row references, including the ranges used by the authored workbook."""
    found = set(REFERENCE.findall(text))
    for match in RANGE.finditer(text):
        prefix, first, last = match.groups()
        low, high = int(first), int(last)
        if high >= low:
            width = max(len(first), len(last))
            found.update(f"{prefix}{number:0{width}d}" for number in range(low, high + 1))
    return frozenset(found)


def clause_text(ident: str, text: str, clauses: tuple[str, ...]) -> tuple[tuple[str, str], ...]:
    """Keep definitions intact; a cited criterion is not a new definition."""
    starts = list(re.finditer(rf"(?m)^{re.escape(ident)}-AC\d+"
        r"(?:\s*\([^)]*\))?\s*:", text))
    defined = {}
    for index, match in enumerate(starts):
        clause = re.match(rf"{re.escape(ident)}-AC\d+", match.group()).group()
        defined.setdefault(clause, text[match.start():starts[index + 1].start()
            if index + 1 < len(starts) else len(text)].strip())
    # Unusually authored clauses keep their whole cell, not a guessed excerpt.
    return tuple((clause, defined.get(clause, text)) for clause in clauses)


def requirement_identity(contract: dict) -> str:
    """The promise being judged, never the verdict that judges it."""
    selected = {key: str(contract.get(key) or "") for key in PROMISE_COLUMNS}
    return hashlib.sha256(json.dumps(selected, ensure_ascii=False,
        sort_keys=True, separators=(",", ":")).encode("utf-8")).hexdigest()


def source_rows(sheet: Path = owners.SHEET,
                registry: owners.Registry | None = None) -> dict[str, SourceRow]:
    """The current complete row population; source text, not authored verdicts."""
    from openpyxl import load_workbook

    registry = registry or owners.Registry.load()
    population = {row.ident: row for row in owners.requirements(sheet, registry)}
    book = load_workbook(sheet, read_only=True)
    try:
        rows = book[owners.WORKSHEET].iter_rows(values_only=True)
        header = tuple(str(cell or "").strip() for cell in next(rows))
        at = {key: header.index(key) for key in PROMISE_COLUMNS}
        result = {}
        for values in rows:
            ident = str(values[at["ID"]] or "").strip()
            if not REFERENCE.fullmatch(ident):
                continue
            if ident in result:
                raise ValueError(f"duplicate requirement source: {ident}")
            contract = {key: str(values[index] or "") for key, index in at.items()}
            digest = requirement_identity(contract)
            owner = population.get(ident)
            text = contract["Acceptance criteria"]
            clauses = clause_text(ident, text, owner.clauses) if owner else ()
            result[ident] = SourceRow(ident, owner.clauses if owner else (), digest,
                references(contract["Dependency IDs and required outputs"]) - {ident}, owner,
                clauses)
        return result
    finally:
        book.close()


def closure(seeds: set[str], graph: dict[str, frozenset[str]]) -> frozenset[str]:
    """Transitive dependencies, cycle-safe; missing nodes remain visible."""
    seen = set()
    pending = list(seeds)
    while pending:
        ident = pending.pop()
        if ident not in seen:
            seen.add(ident)
            pending.extend(graph.get(ident, ()))
    return frozenset(seen)


def _nonempty(value: object) -> bool:
    return isinstance(value, str) and bool(value.strip())


def selected_rows(packet: str, rows: dict[str, SourceRow],
                  registry: owners.Registry) -> dict[str, SourceRow]:
    """The same owned-row selection used by the ownership check, never a verdict."""
    carried = registry.carries.get(packet, frozenset())
    return {ident: row for ident, row in rows.items() if row.ownership and (
        packet in row.ownership.packets or carried & set(row.ownership.criteria))}


def populations(rows: dict[str, SourceRow], registry: owners.Registry,
                packets: tuple[str, ...]) -> dict:
    """Exact whole-brain universe and selected subset; duplicates stay disclosed."""
    whole = {ident: row for ident, row in rows.items() if row.ownership is not None}
    selected = {packet: selected_rows(packet, rows, registry) for packet in packets}
    union = {ident: row for entries in selected.values() for ident, row in entries.items()}
    identity = [(ident, row.version, row.clauses) for ident, row in sorted(whole.items())]
    digest = hashlib.sha256(json.dumps(identity, separators=(",", ":"),
        ensure_ascii=False).encode("utf-8")).hexdigest()
    return {"whole_requirements": len(whole),
        "whole_clauses": sum(len(row.clauses) for row in whole.values()),
        "ownership_states": {state: sum(row.ownership.state == state for row in whole.values())
                             for state in ("owned", "problem", "declared", "unowned")},
        "selected_packets": list(packets), "selected_unique_requirements": len(union),
        "selected_unique_clauses": len({clause for row in union.values()
                                       for clause in row.clauses}),
        "selected_contracts_including_shared_clauses": sum(len(row.clauses)
            for entries in selected.values() for row in entries.values()),
        "requirements_identity": digest}


def execution_inventory(rows: dict[str, SourceRow], registry: owners.Registry,
                        document: dict) -> dict:
    """Expose the whole promise universe, not just a green selected subset.

    Known paths are candidate engineering owners. Neither their existence nor
    row ownership certifies a clause. Prose-only rows stay visible instead of
    vanishing behind a named-clause-only query.
    """
    whole = {ident: row for ident, row in rows.items() if row.ownership is not None}
    mapped = document.get("requirement_modules", {})
    entries = []
    for ident, row in sorted(whole.items()):
        claims = owners.registered_claims(row.ownership.line, registry)
        entries.append({
            "requirement": ident, "requirement_version": row.version,
            "row_ownership": row.ownership.state,
            "registered_row_owners": [{"packet": packet, "criterion": criterion}
                for packet, criteria in sorted(claims.items()) for criterion in sorted(criteria)],
            "engineering_candidates": list(mapped.get(ident, ())),
            "engineering_state": "partial_controls_not_clause_acceptance"
                if mapped.get(ident) else "not_mapped_by_this_execution_audit",
            "clauses": [{"clause": clause, "acceptance": "NOT_RUN"}
                        for clause in row.clauses],
            "named_clause_state": "defined" if row.clauses else "prose_only_not_clause_mapped",
            "complete": False,
        })
    return {"population": populations(rows, registry, SELECTED_PACKETS),
            "rows": entries, "client_cutover": False,
            "qualification": "An execution inventory is not version-pinned acceptance evidence."}


@lru_cache(maxsize=128)
def _test_names(path: str, digest: str) -> frozenset[str]:
    # The digest is in the cache key so a background edit cannot preserve a stale
    # symbol inventory. Reading a test is not execution or passing evidence.
    del digest
    tree = ast.parse(Path(path).read_text(encoding="utf-8"))
    return frozenset(node.name for node in ast.walk(tree)
                     if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)))


def _local_file(root: Path, value: object) -> Path | None:
    if not _nonempty(value):
        return None
    path = (root / value).resolve()
    return path if path.is_relative_to(root.resolve()) and path.is_file() else None


def module_problems(document: dict, root: Path) -> tuple[list[str], list[str]]:
    """Validate real module/test links separately from clause acceptance."""
    problems, blockers, seen = [], [], set()
    for module in document.get("modules", []):
        ident = module.get("id")
        if not _nonempty(ident) or ident in seen:
            problems.append(f"unknown or duplicate module: {ident}")
        seen.add(ident)
        paths = module.get("paths", [])
        if not paths:
            problems.append(f"{ident}: module has no implementation owner")
        for value in paths:
            if _local_file(root, value) is None:
                problems.append(f"{ident}: module path unavailable: {value}")
        for kind in ("checks", "rejecting_controls"):
            checks = module.get(kind, [])
            if not checks:
                problems.append(f"{ident}: module has no {kind}")
            for reference in checks:
                if not isinstance(reference, str) or reference.count("::") != 1:
                    problems.append(f"{ident}: invalid {kind} reference: {reference}")
                    continue
                name, symbol = reference.split("::")
                path = _local_file(root, name)
                try:
                    digest = hashlib.sha256(path.read_bytes()).hexdigest() if path else ""
                    if path is None or symbol not in _test_names(str(path), digest):
                        problems.append(f"{ident}: {kind} does not resolve: {reference}")
                except (SyntaxError, UnicodeError, OSError):
                    problems.append(f"{ident}: {kind} cannot be inspected: {reference}")
        measurement = module.get("measurement", {})
        if measurement.get("state") not in {"NOT_RUN", "CONTROLLED_TESTS_RECORDED"}:
            problems.append(f"{ident}: module measurement cannot declare clause PASS")
        if measurement.get("state") == "CONTROLLED_TESTS_RECORDED":
            subjects = measurement.get("subject_files", {})
            required = set(paths) | {ref.split("::")[0]
                for kind in ("checks", "rejecting_controls") for ref in module.get(kind, [])
                if isinstance(ref, str)}
            if set(subjects) != required:
                problems.append(f"{ident}: measured subject population differs "
                                "from module contract")
            if not _nonempty(measurement.get("execution")):
                problems.append(f"{ident}: controlled measurement lacks an execution record")
            for name, expected in subjects.items():
                path = _local_file(root, name)
                actual = hashlib.sha256(path.read_bytes()).hexdigest() if path else None
                if expected != actual:
                    blockers.append(f"{ident}: controlled measurement STALE: {name}")
        else:
            blockers.append(f"{ident}: module controls NOT_RUN")
    return problems, blockers


def refresh_contracts(document: dict, rows: dict[str, SourceRow],
                      registry: owners.Registry, packets: dict[str, dict]) -> dict:
    """Render derived source contracts, preserving authored limitations/decisions.

    This is a document projection, not promotion. Every rendered acceptance
    artifact remains NOT_RUN, even when a linked foundation has measured tests.
    """
    result = dict(document)
    packet_ids = tuple(document.get("selected_packets", ()))
    result["population"] = populations(rows, registry, packet_ids)
    modules = {module["id"]: module for module in document.get("modules", [])}
    row_graph = {ident: row.dependencies for ident, row in rows.items()}
    packet_graph = {ident: frozenset(p.get("prerequisites", []))
                    for ident, p in packets.items()}
    slices = []
    for packet in packet_ids:
        if packet not in packets:
            raise ValueError(f"unknown packet: {packet}")
        selected = selected_rows(packet, rows, registry)
        contracts = []
        for ident, row in sorted(selected.items()):
            linked = document.get("requirement_modules", {}).get(ident, [])
            if any(module not in modules for module in linked):
                raise ValueError(f"unknown module for {ident}")
            paths = list(dict.fromkeys(path for module in linked
                                       for path in modules[module]["paths"]))
            if not paths:
                paths = [boundary["path"] for boundary in packets[packet].get("boundaries", [])]
            owner_pairs = [{"criterion": criterion, "packet": owner_packet}
                for owner_packet, criteria in sorted(
                    owners.registered_claims(row.ownership.line, registry).items())
                for criterion in sorted(criteria)]
            remaining = document.get("requirement_remaining", {}).get(ident,
                "Exact clause acceptance, its rejecting control and qualified evidence "
                "remain unrecorded.")
            for clause in row.clauses:
                contracts.append({"clause": clause, "requirement_version": row.version,
                    "source_clause": dict(row.clause_text).get(clause), "owners": owner_pairs,
                    "implementation": {"state": "partial" if linked else "planned",
                        "paths": paths, "modules": linked, "remaining": remaining},
                    "verification": {"method": "controlled_foundations_then_clause_acceptance",
                        "coverage": "FOUNDATION_ONLY" if linked else "PLANNED",
                        "check": " | ".join(check for module in linked
                            for check in modules[module]["checks"]) or
                            f"Planned: execute the exact {clause} source criterion",
                        "negative_control": " | ".join(check for module in linked
                            for check in modules[module]["rejecting_controls"]) or
                            f"Planned: plant the contrary of {clause}; "
                            "a passing empty population is refused"},
                    "artifact": {"state": "NOT_RUN", "path": None, "tested_subject": None}})
        slices.append({"id": packet, "packet_prerequisite_closure": sorted(
            closure(set(packet_graph[packet]), packet_graph)),
            "requirement_prerequisite_closure": sorted(
                closure(set(selected), row_graph) - set(selected)), "clauses": contracts})
    result["slices"] = slices
    return result


def inspect_scope(packet: str, document: dict, rows: dict[str, SourceRow],
                  registry: owners.Registry, packets: dict[str, dict],
                  root: Path = ROOT) -> ScopeReport:
    """Reconcile scope and evidence contracts without executing any evidence."""
    problems: list[str] = []
    blockers: list[str] = []
    slices = [entry for entry in document.get("slices", []) if entry.get("id") == packet]
    if document.get("schema") != 1 or len(slices) != 1 or packet not in packets:
        return ScopeReport(packet, 0,
            ("unknown or duplicate slice, or invalid map schema",), (), False)
    spec = slices[0]
    selected = selected_rows(packet, rows, registry)
    if "selected_packets" in document and "population" not in document:
        problems.append("whole-brain population/source identity missing")
    if "population" in document:
        packet_ids = tuple(document["population"].get("selected_packets", []))
        if document["population"] != populations(rows, registry, packet_ids):
            problems.append("whole-brain or selected clause population/source identity differs")
        if set(packet_ids) != {entry.get("id") for entry in document.get("slices", [])}:
            problems.append("selected packet population differs from execution slices")
        if document.get("selected_packets") != list(packet_ids):
            problems.append("selected packet population differs from scope declaration")
    module_errors, module_blockers = module_problems(document, root)
    problems.extend(module_errors)
    blockers.extend(module_blockers)
    module_ids = {entry.get("id") for entry in document.get("modules", [])}
    expected = {clause: row for row in selected.values() for clause in row.clauses}
    mapped = spec.get("clauses", [])
    seen = set()
    for entry in mapped:
        clause = entry.get("clause")
        if clause in seen:
            problems.append(f"duplicate clause contract: {clause}")
        seen.add(clause)
        if clause not in expected:
            problems.append(f"unselected or unknown clause: {clause}")
            continue
        row = expected[clause]
        if entry.get("requirement_version") != row.version:
            problems.append(f"{clause}: requirement version is stale")
        if row.clause_text and entry.get("source_clause") != dict(row.clause_text).get(clause):
            problems.append(f"{clause}: acceptance clause text differs from workbook")
        if row.ownership.state != "owned":
            problems.append(f"{clause}: row ownership is {row.ownership.state}")
        claimed = owners.registered_claims(row.ownership.line, registry)
        links = entry.get("owners", [])
        if not links:
            problems.append(f"{clause}: no registered criterion contract")
        for link in links:
            owner_packet, criterion = link.get("packet"), link.get("criterion")
            if criterion not in claimed.get(owner_packet, ()):
                problems.append(f"{clause}: {criterion}/{owner_packet} "
                                "is not this row's registered owner")
        implementation = entry.get("implementation", {})
        verification = entry.get("verification", {})
        artifact = entry.get("artifact", {})
        paths = implementation.get("paths", [])
        if not paths or not all(_nonempty(path) for path in paths):
            problems.append(f"{clause}: missing implementation paths")
        if not all(_nonempty(verification.get(key))
                   for key in ("method", "check", "negative_control")):
            problems.append(f"{clause}: missing method, check or rejecting control")
        if artifact.get("state") not in {"NOT_RUN", "RECORDED"}:
            problems.append(f"{clause}: missing explicit artifact state")
        if implementation.get("state") not in {"planned", "partial", "implemented"}:
            problems.append(f"{clause}: invalid implementation state")
        linked = implementation.get("modules", [])
        if any(ident not in module_ids for ident in linked):
            problems.append(f"{clause}: unknown linked module")
        declared_modules = document.get("requirement_modules")
        if declared_modules is not None and linked != declared_modules.get(row.ident, []):
            problems.append(f"{clause}: linked modules differ from this requirement contract")
        if implementation.get("state") == "partial" and (
                not linked or not _nonempty(implementation.get("remaining"))):
            problems.append(f"{clause}: partial foundation lacks modules or remaining obligations")
        if implementation.get("state") != "implemented":
            blockers.append(f"{clause}: implementation {implementation.get('state')}")
        else:
            for path in paths:
                resolved = (root / path).resolve()
                if not resolved.is_relative_to(root.resolve()) or not resolved.is_file():
                    blockers.append(f"{clause}: implementation path unavailable: {path}")
        # Evidence evaluation has an existing owner. A map may point at it, but
        # cannot type PASS and thereby certify itself. No record is acceptance.
        if artifact.get("state") != "RECORDED":
            blockers.append(f"{clause}: executed versioned artifact NOT_RUN")
        else:
            blockers.append(f"{clause}: artifact requires canonical evidence evaluation")
    for clause in sorted(set(expected) - seen):
        problems.append(f"missing selected clause contract: {clause}")
    if not expected:
        problems.append("selected clause population is empty")
    packet_graph = {ident: frozenset(p.get("prerequisites", [])) for ident, p in packets.items()}
    packet_closure = closure(set(packet_graph.get(packet, ())), packet_graph)
    if set(spec.get("packet_prerequisite_closure", [])) != packet_closure:
        problems.append("packet prerequisite closure differs from current registry")
    row_graph = {ident: row.dependencies for ident, row in rows.items()}
    row_closure = closure(set(selected), row_graph) - set(selected)
    if set(spec.get("requirement_prerequisite_closure", [])) != row_closure:
        problems.append("requirement prerequisite closure differs from current workbook")
    for ident in sorted(row_closure):
        row = rows.get(ident)
        if row is None:
            blockers.append(f"{ident}: prerequisite unknown")
        elif row.ownership is None or row.ownership.state != "owned":
            blockers.append(f"{ident}: prerequisite ownership not established")
        else:
            blockers.append(f"{ident}: prerequisite acceptance not evaluated")
    for ident in sorted(packet_closure):
        if ident not in packets:
            blockers.append(f"{ident}: prerequisite packet unknown")
        else:
            blockers.append(f"{ident}: prerequisite packet acceptance not evaluated")
    for decision in document.get("decisions", []):
        if decision.get("state") != "RECORDED":
            blockers.append(f"{decision.get('id')}: owner decision NOT_RECORDED")
    if "LB-134" in rows and not any(decision.get("id") == "LB-134"
                                   for decision in document.get("decisions", [])):
        problems.append("LB-134: owner decision record missing")
    authorization = document.get("engineering_authorization", {})
    authorized = (authorization.get("environment") == "recorded_and_synthetic"
                  and _nonempty(authorization.get("source"))
                  and authorization.get("client_cutover") is False)
    return ScopeReport(packet, len(expected), tuple(dict.fromkeys(problems)),
        tuple(dict.fromkeys(blockers)), authorized)


def report(packet: str, path: Path = MAP) -> ScopeReport:
    document = json.loads(path.read_text(encoding="utf-8"))
    registry = owners.Registry.load()
    authored = json.loads(owners.PACKETS.read_text(encoding="utf-8"))
    packets = {p["id"]: p for p in authored["packets"]}
    result = inspect_scope(packet, document, source_rows(registry=registry), registry, packets)
    if document.get("selected_packets") != list(SELECTED_PACKETS):
        return ScopeReport(result.packet, result.selected_clauses,
            result.contract_problems + ("required P49-P54 scope declaration differs",),
            result.completion_blockers, result.engineering_authorized)
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("packet", nargs="?")
    parser.add_argument("--map", type=Path, default=MAP)
    parser.add_argument("--contracts-only", action="store_true",
                        help="check coverage contracts only, not completion readiness")
    parser.add_argument("--emit-refreshed-contracts", action="store_true",
                        help="print current derived contracts; never write or promote evidence")
    parser.add_argument("--whole-inventory", action="store_true",
                        help="print all legal-brain rows/clauses; never certify their acceptance")
    args = parser.parse_args(argv)
    if args.whole_inventory:
        document = json.loads(args.map.read_text(encoding="utf-8"))
        registry = owners.Registry.load()
        print(json.dumps(execution_inventory(source_rows(registry=registry), registry, document),
                         ensure_ascii=False, indent=2))
        return 0
    if args.emit_refreshed_contracts:
        document = json.loads(args.map.read_text(encoding="utf-8"))
        registry = owners.Registry.load()
        packets = {p["id"]: p for p in json.loads(
            owners.PACKETS.read_text(encoding="utf-8"))["packets"]}
        print(json.dumps(refresh_contracts(document, source_rows(registry=registry),
            registry, packets), ensure_ascii=False, indent=2))
        return 0
    if args.packet is None:
        parser.error("packet is required unless emitting contracts or whole inventory")
    result = report(args.packet, args.map)
    print(json.dumps({"packet": result.packet, "selected_clauses": result.selected_clauses,
        "contracts_ready": result.contracts_ready,
        "engineering_authorized": result.engineering_authorized,
        "completion_ready": result.completion_ready,
        "client_cutover_allowed": result.client_cutover_allowed,
        "contract_problems": result.contract_problems,
        "completion_blockers": result.completion_blockers}, indent=2))
    ready = result.contracts_ready if args.contracts_only else result.completion_ready
    return 0 if ready else 1


if __name__ == "__main__":
    raise SystemExit(main())
