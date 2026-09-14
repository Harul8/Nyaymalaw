"""Read-only verification of the saved current workbook against authored sources.

This does not regenerate a workbook, execute legal evaluations, or recompute
runtime readiness. Cached formulas are checked against independent populations.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from collections import Counter
from datetime import date, datetime, time, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from assurance.common._console import utf8_console  # noqa: E402
from assurance.common._documents import safe_load  # noqa: E402

utf8_console()
WORKBOOK = ROOT / "docs/Nyaymalaw_End_to_End_Project_Plan.xlsx"
SHEETS = {
    "Read Me",
    "Overview",
    "Delivery Plan",
    "Journey Steps",
    "Step Contracts",
    "Features",
    "Work Items",
    "Traceability",
    "Build Assurance",
    "Release Gates",
    "Release Profiles",
    "Readiness Plan",
    "Readiness Gaps",
    "Journey Scenarios",
    "Risks",
    "Advocate Standard",
    "Expert Workflow",
    "Advice Maturity",
    "Roles & Authority",
    "Gap Closure",
    "Modules",
    "Execution Packets",
    "Packet Guide",
    "Acceptance",
    "Decisions",
    "Command Contracts",
    "Evaluation Specs",
    "Evaluation Details",
    "Sources",
    "Reconciliation",
}
METADATA = {"__formulas__", "__load_errors__"}
ERROR_VALUES = {
    "#REF!",
    "#DIV/0!",
    "#VALUE!",
    "#NAME?",
    "#N/A",
    "#NUM!",
    "#NULL!",
    "#SPILL!",
    "#CALC!",
}
TEXT_SUFFIXES = {".md", ".json", ".yaml", ".js", ".mjs", ".py"}
# The snapshot protocol intentionally names the bounded non-blueprint inputs.
# The whole blueprint tree is discovered, including future nested contracts.
SNAPSHOT_BASE = {
    "docs/backlog/status.yaml",
    "docs/backlog/steps.yaml",
    "docs/backlog/plan.json",
    "docs/backlog/professional.json",
    "docs/backlog/build_rules.json",
    "docs/backlog/SCHEMA.md",
    "docs/BUILD_GUIDE.md",
    "docs/playbooks/START_A_CHANGE.md",
    "docs/playbooks/BUILD_A_CHANGE.md",
    "docs/playbooks/TEST_A_CHANGE.md",
    "docs/playbooks/SIGN_OFF_A_CHANGE.md",
    "assurance/specification/prd/part_a.js",
    "assurance/specification/prd/part_b.js",
    "assurance/specification/prd/part_c.js",
    "assurance/specification/prd/build.js",
    "assurance/specification/prd/helpers.js",
    "assurance/specification/prd/schemas.js",
    "assurance/specification/prd/gates.json",
    "assurance/specification/features.yaml",
    "docs/Nyaymalaw_PRD.docx",
    "assurance/specification/plan/view_content.json",
    "assurance/specification/plan/build_current_plan.mjs",
    "assurance/specification/plan/export_current_plan.py",
    "assurance/specification/plan/README.md",
}
CAPTURED_ARTIFACT = "docs/backlog/evidence/class_a.json"
CONTRACT_ASPECTS = {
    "actor": "Actor",
    "entry_conditions": "Entry conditions",
    "user_action": "User action",
    "expected_visible_result": "Visible result",
    "expected_domain_effect": "Domain effect",
    "failure_behaviour": "Failure / refusal",
    "recovery_behaviour": "Recovery / return",
    "exit_conditions": "Exit conditions",
    "notes": "Context",
}


def load_view(path: Path = WORKBOOK) -> dict[str, list[dict]]:
    """Read actual saved cells and their cached formula results; never save."""
    result: dict[str, list[dict]] = {"__formulas__": [], "__load_errors__": []}
    values = openpyxl.load_workbook(path, read_only=True, data_only=True)
    formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        for sheet in values:
            # Artifact exports need not declare worksheet dimensions.
            rows = list(sheet.iter_rows(values_only=True))
            if len(rows) < 3:
                result["__load_errors__"].append({"message": f"{sheet.title}: no header row"})
                result[sheet.title] = []
                continue
            headers = list(rows[2])
            if (
                not headers
                or any(not isinstance(v, str) or not v.strip() for v in headers)
                or len(headers) != len(set(headers))
            ):
                result["__load_errors__"].append({"message": f"{sheet.title}: malformed headers"})
                result[sheet.title] = []
                continue
            result[sheet.title] = []
            for excel_row, row in enumerate(rows[3:], 4):
                if not any(value is not None for value in row):
                    continue
                if len(row) > len(headers) and any(v is not None for v in row[len(headers) :]):
                    result["__load_errors__"].append(
                        {"message": f"{sheet.title}:{excel_row}: data outside headers"}
                    )
                result[sheet.title].append(
                    {h: row[i] if i < len(row) else None for i, h in enumerate(headers)}
                )
            for cells in formulas[sheet.title].iter_rows():
                for cell in cells:
                    if cell.data_type == "f":
                        cached = (
                            rows[cell.row - 1][cell.column - 1]
                            if cell.row <= len(rows) and cell.column <= len(rows[cell.row - 1])
                            else None
                        )
                        result["__formulas__"].append(
                            {
                                "sheet": sheet.title,
                                "cell": cell.coordinate,
                                "formula": cell.value,
                                "cached": cached,
                            }
                        )
    finally:
        values.close()
        formulas.close()
    return result


def _text(value) -> str:
    if isinstance(value, list):
        return "\n".join(_text(v) for v in value)
    if isinstance(value, dict):
        return "\n".join(f"{k}: {_text(v)}" for k, v in value.items())
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).lower()
    return str(value)


def _list(value) -> str:
    return ", ".join(value or [])


def _same_value(actual, expected) -> bool:
    if actual is None and expected == "":
        return True
    # A deliberately date-formatted Excel cell is returned as date/datetime.
    # An unformatted numeric serial is NOT accepted as equivalent to readable
    # ISO text: that export silently changes what the reader sees.
    if isinstance(actual, date) and isinstance(expected, str):
        try:
            instant = datetime.fromisoformat(expected.replace("Z", "+00:00"))
        except ValueError:
            return False
        if instant.tzinfo is not None:
            instant = instant.astimezone(timezone.utc).replace(tzinfo=None)
        observed = actual if isinstance(actual, datetime) else datetime.combine(actual, time())
        return observed == instant
    return actual == expected


def source_hash(path: Path) -> str:
    content = (
        path.read_text(encoding="utf-8").replace("\r\n", "\n").encode()
        if path.suffix in TEXT_SUFFIXES
        else path.read_bytes()
    )
    return hashlib.sha256(content).hexdigest()


def _read(root: Path, relative: str):
    text = (root / relative).read_text(encoding="utf-8")
    return safe_load(text) if relative.endswith(".yaml") else json.loads(text)


def check_view(tables: dict[str, list[dict]], root: Path = ROOT) -> list[str]:
    """Compare saved cell populations/content and captured source hashes.

    Effective evidence remains a dated snapshot. Do not recompute it against a
    fingerprint containing this running test or make a generated verdict its
    own freshness prerequisite.
    """
    errors: list[str] = []
    if not isinstance(tables, dict) or set(tables) != SHEETS | METADATA:
        return ["workbook: missing or unexpected sheet/metadata population"]
    if any(not isinstance(rows, list) for rows in tables.values()):
        return ["workbook: malformed table population"]
    errors.extend(
        row.get("message", "workbook: unreadable row") for row in tables["__load_errors__"]
    )
    for name in SHEETS:
        if not tables[name] or any(not isinstance(row, dict) for row in tables[name]):
            errors.append(f"{name}: empty or malformed table")
        for row in tables[name]:
            if isinstance(row, dict):
                for value in row.values():
                    if isinstance(value, str) and value in ERROR_VALUES:
                        errors.append(f"{name}: saved formula error {value}")
    if errors:
        return errors
    status = _read(root, "docs/backlog/status.yaml")
    steps = _read(root, "docs/backlog/steps.yaml")["steps"]
    plan = _read(root, "docs/backlog/plan.json")
    professional = _read(root, "docs/backlog/professional.json")
    rules = _read(root, "docs/backlog/build_rules.json")["rules"]
    context = _read(root, "assurance/specification/plan/view_content.json")
    modules = _read(root, "docs/blueprint/modules.json")["modules"]
    packets = _read(root, "docs/blueprint/packets.json")["packets"]
    decisions = _read(root, "docs/blueprint/decisions.json")["choices"]
    commands = _read(root, "docs/blueprint/contracts/commands.json")["x-commands"]
    evaluation_catalog = _read(root, "docs/blueprint/evaluations.json")
    # THE AUTHORED ROUTE TO RELEASE. Its measured columns are a dated derivation and
    # are not recomputed here, for the reason given in this function's docstring.
    stages = (plan.get("readiness_plan") or {}).get("stages") or []
    scenarios = evaluation_catalog["synthetic_cases"]
    items, features = status["items"], status["features"]
    wave = {row["id"]: row["wave"] or "Unscheduled" for row in plan["item_waves"]}
    acceptance = [
        {**ac, "item": item["id"]} for item in items for ac in item.get("acceptance") or []
    ]

    def identities(sheet: str, column: str, source: list[dict]) -> None:
        actual = [row.get(column) for row in tables[sheet]]
        expected = [row["id"] for row in source]
        if actual != expected:
            missing = sorted(set(expected) - set(v for v in actual if isinstance(v, str)))
            extra = sorted(set(v for v in actual if isinstance(v, str)) - set(expected))
            errors.append(
                f"{sheet}: exact ID population/order differs; missing={missing}, extra={extra}"
            )

    def compare(sheet: str, key: str, identifier: str, expected: dict) -> None:
        matches = [row for row in tables[sheet] if row.get(key) == identifier]
        if len(matches) != 1:
            errors.append(f"{sheet}: expected one {identifier}, found {len(matches)}")
            return
        row = matches[0]
        for column, value in expected.items():
            # Excel reads deliberately empty text cells as None.
            actual = row.get(column)
            if not _same_value(actual, value):
                errors.append(f"{sheet}/{identifier}: {column} differs from source")

    groups = [
        ("Items", "Work Items", "ID", items, "A"),
        ("Features", "Features", "Feature", features, "A"),
        ("Steps", "Journey Steps", "Step ID", steps, "C"),
        ("Advocate standards", "Advocate Standard", "ID", professional["advocate_standards"], "A"),
        ("Workflow states", "Expert Workflow", "ID", professional["workflow_states"], "A"),
        ("Advice levels", "Advice Maturity", "ID", professional["advice_maturity"], "A"),
        ("Roles", "Roles & Authority", "ID", professional["roles"], "A"),
        ("Gaps", "Gap Closure", "ID", professional["gap_closures"], "A"),
        ("Wave contracts", "Delivery Plan", "Wave", plan["wave_contracts"], "A"),
        ("Build rules", "Build Assurance", "Rule", rules, "A"),
        ("Modules", "Modules", "Module", modules, "A"),
        ("Execution packets", "Execution Packets", "Packet", packets, "A"),
        ("Acceptance criteria", "Acceptance", "Criterion", acceptance, "A"),
        ("Choices", "Decisions", "Choice", decisions, "A"),
        ("Command contracts", "Command Contracts", "Command", commands, "A"),
        ("Synthetic specifications", "Evaluation Specs", "Scenario", scenarios, "A"),
        ("Readiness stages", "Readiness Plan", "Stage", stages, "A"),
        ("Readiness step coverage", "Readiness Gaps", "Step ID", steps, "C"),
    ]
    for _, sheet, key, source, _ in groups:
        if not source:
            errors.append(f"{sheet}: source population is empty")
        identities(sheet, key, source)
    for row in items:
        next_action = row.get("next_action") or ""
        if row["delivery_status"] == "deferred":
            next_action = (
                (next_action + " " if next_action else "")
                + f"Review on {row['review_on']} (India date). Live status shows due/overdue; "
                "recorded reassessment is required before reactivation."
            )
        compare(
            "Work Items",
            "ID",
            row["id"],
            {
                "Next action": next_action,
                "Title": row["title"],
                "Kind": row["kind"],
                "Priority": row["priority"],
                "Delivery status": row["delivery_status"],
                "Implementation": row["implementation"],
                "Authored verification": row["verification"],
                "Wave": wave[row["id"]],
                "Hard dependencies": _list(row.get("depends_on")),
                "Acceptance count": len(row.get("acceptance") or []),
                "Record": row["record"],
            },
        )
    for row in features:
        compare(
            "Features",
            "Feature",
            row["id"],
            {
                "Phase": row["phase"],
                "Title": row["title"],
                "Build slice": row["slice"],
                "Implementation": row["implementation"],
                "Disposition": row["disposition"],
                "Delivery items": _list(row.get("delivery_items")),
            },
        )
    for row in steps:
        compare(
            "Journey Steps",
            "Step ID",
            row["id"],
            {
                "Phase": row["phase"],
                "Journey step": row["name"],
                "Basis": row["basis"],
                "Features": _list(row["features"]),
                "Registered work": _list(row["items"]),
                "Visible result": _text(row["expected_visible_result"]),
                "Failure behaviour": _text(row["failure_behaviour"]),
                "Recovery behaviour": _text(row["recovery_behaviour"]),
            },
        )
    for row in acceptance:
        contributors = [p for p in packets if row["id"] in p["criteria"]]
        control = row.get("negative_control") or {}
        compare(
            "Acceptance",
            "Criterion",
            row["id"],
            {
                "Work item": row["item"],
                "Required behaviour": row["requirement"],
                "Required evidence methods": _list(row["required_evidence"]),
                "Planted mutation": _text(control.get("mutation")),
                "Expected refusal / failure": _text(control.get("expected_failure")),
                "Contributing packets": _list([p["id"] for p in contributors]),
                "Final packet": _list(
                    [p["id"] for p in contributors if row["id"] in p["final_criteria"]]
                )
                or "Current planning delivery",
            },
        )
    for row in packets:
        compare(
            "Execution Packets",
            "Packet",
            row["id"],
            {
                "Module": row["module"],
                "Outcome": row["title"],
                "Kind": row["kind"],
                "Scoped predecessor outputs": _list(row["prerequisites"]),
                "Completed items required": _list(row["requires_completed_items"]),
                "Criteria": len(row["criteria"]),
                "Final criteria": len(row["final_criteria"]),
                "Required choices": _list(row["decisions"]),
                "Command contracts": _list(row["commands"]),
            },
        )
    for row in commands:
        compare(
            "Command Contracts",
            "Command",
            row["id"],
            {
                "Method": row["method"],
                "Target path": row["path"],
                "Owner criteria": _list(row["owner_ac"]),
                "Request schema": _text(row["request_schema"]),
                "Response schema": _text(row["response_schema"]),
            },
        )
    approval_records = _read(root, "docs/blueprint/approvals.json")["records"]
    for row in decisions:
        matching_adoptions = sorted(
            record["id"] for record in approval_records if record["choice"] == row["id"]
        )
        approval_label = (
            "approval not machine-resolved / manual verification required ("
            + ", ".join(matching_adoptions)
            + ")"
            if matching_adoptions
            else "no adoption record recorded"
        )
        compare(
            "Decisions",
            "Choice",
            row["id"],
            {
                "Recommendation": row["recommendation"],
                "Fallback": row["fallback"],
                "Accountable approver": row["approver"],
                "Approval required for": _list(row["approval_required_for"]),
                "Current approval": approval_label,
            },
        )
    for row in scenarios:
        compare(
            "Evaluation Specs",
            "Scenario",
            row["id"],
            {
                "Module": row["module"],
                "Purpose": row["title"],
                "Owner criteria": _list(row["owner_criteria"]),
                "Method": row["method"],
                "Execution": "NOT RUN — specification only",
            },
        )
    for row in modules:
        compare(
            "Modules",
            "Module",
            row["id"],
            {
                "Capability": row["title"],
                "Context modules": _list(row["requires"]),
                "Work items": _list(row["items"]),
                "Features": _list(row["features"]),
                "Steps": _list(row["steps"]),
                "Chapter": "docs/blueprint/" + row["guide"],
            },
        )
    for row in rules:
        compare(
            "Build Assurance",
            "Rule",
            row["id"],
            {
                "Stage": row["stage"],
                "Kind": row["kind"],
                "Requirement": row["statement"],
                "Enforcement": row["enforcement"],
                "Named check / review": row.get("check") or row.get("review") or "",
                "Coverage scope": row.get("coverage_scope") or "",
                "Remaining gap / reason": row.get("remaining_gap") or row.get("why_not") or "",
                "Stage guide": "docs/playbooks/" + row["card"],
            },
        )
    for row in plan["wave_contracts"]:
        compare(
            "Delivery Plan",
            "Wave",
            row["id"],
            {
                "Outcome": row["name"] + "\n" + row["goal"],
                "Scope": _text(row["scope"]),
                "Entry conditions": _text(row["entry_conditions"]),
                "Exit conditions": _text(row["exit_conditions"]),
                "Release claim limit": _text(row["release_claim"]),
                "Control": _text(row["control"]),
                "Assigned work": _list(
                    [w["id"] for w in plan["item_waves"] if w["wave"] == row["id"]]
                ),
            },
        )
    for row in stages:
        compare(
            "Readiness Plan",
            "Stage",
            row["id"],
            {
                "Stage name": row["name"],
                "Goal": row["goal"],
                "Journey steps": _list(row.get("journey_steps")),
                "Named work": _list(row.get("items")),
                "Release profile": row.get("release_profile") or "None",
                "Depends on": _list(row.get("depends_on")),
                "Authority needed": _list(row.get("authority")),
                "Actions": _text(row["actions"]),
                "Pilot exit": _text(row["pilot_exit"]),
                "Production exit": _text(row["production_exit"]),
                "Open decisions": (
                    _text(row["decisions"]) if row.get("decisions") else "None recorded"
                ),
            },
        )
    stage_of_step = {s: row["id"] for row in stages for s in row.get("journey_steps") or []}
    for row in steps:
        compare(
            "Readiness Gaps",
            "Step ID",
            row["id"],
            {
                "Phase": row["phase"],
                "Journey step": row["name"],
                "Readiness stage": stage_of_step.get(row["id"], "No stage"),
                "Registered work": _list(row.get("items")),
            },
        )
    for row in professional["advocate_standards"]:
        compare(
            "Advocate Standard",
            "ID",
            row["id"],
            {
                "Quality": row["quality"],
                "Observable behaviour": row["observable"],
                "Trust-destroying failure": row["failure"],
                "Registered work": _list(row["work_items"]),
                "Features / steps": _list(row["features"]) + "\n" + _list(row["steps"]),
            },
        )
    for row in professional["workflow_states"]:
        compare(
            "Expert Workflow",
            "ID",
            row["id"],
            {
                "Working state": row["state"],
                "Expert purpose": row["purpose"],
                "Required output": _text(row["required_output"]),
                "Registered work": _list(row["work_items"]),
                "Features / steps": _list(row["features"]) + "\n" + _list(row["steps"]),
            },
        )
    for row in professional["advice_maturity"]:
        compare(
            "Advice Maturity",
            "ID",
            row["id"],
            {
                "Level": row["level"],
                "When permitted": _text(row["when_permitted"]),
                "Must contain": _text(row["must_contain"]),
                "Must not imply": _text(row["must_not_imply"]),
                "Registered work": _list(row["work_items"]),
                "Features / steps": _list(row["features"]) + "\n" + _list(row["steps"]),
            },
        )
    for row in professional["roles"]:
        compare(
            "Roles & Authority",
            "ID",
            row["id"],
            {
                "Role": row["role"],
                "May do": _text(row["may"]),
                "Controlled / restricted": _text(row["controlled"]),
                "Registered work": _list(row["work_items"]),
            },
        )
    for row in professional["gap_closures"]:
        compare(
            "Gap Closure",
            "ID",
            row["id"],
            {
                "Gap": row["gap"],
                "Required plan change": _text(row["plan_change"]),
                "Minimum acceptance": _text(row["minimum_acceptance"]),
                "Registered stage links": _list(
                    [link["item"] + " (" + link["stage"] + ")" for link in row["links"]]
                ),
                "Foundation": row["foundation_wave"],
                "Feature complete": row["feature_complete_wave"],
                "Release gate": row["release_gate_wave"],
            },
        )

    # Check finite instruction populations as well as the one-row indexes.
    guide = [
        {
            "Packet": p["id"],
            "Outcome": p["title"],
            "Aspect": aspect,
            "Order": index + 1,
            "Instruction": instruction,
        }
        for p in packets
        for aspect in ("inputs", "outputs", "steps", "expected", "rollback")
        for index, instruction in enumerate(
            p[aspect] if isinstance(p[aspect], list) else [p[aspect]]
        )
    ]
    if tables["Packet Guide"] != guide:
        errors.append("Packet Guide: instruction population or content differs")
    step_contracts = [
        {
            "Step ID": s["id"],
            "Phase": s["phase"],
            "Step": s["name"],
            "Aspect": label,
            "Intended contract": _text(s[key]),
            "Registered work": _list(s["items"]),
        }
        for s in steps
        for key, label in CONTRACT_ASPECTS.items()
        if s.get(key)
    ]
    if tables["Step Contracts"] != step_contracts:
        errors.append("Step Contracts: clause population or content differs")
    evaluation_details = []
    for row in scenarios:
        evaluation_details.extend(
            {
                "Scenario": row["id"],
                "Aspect": "Input",
                "Field / order": key,
                "Specification": _text(value),
            }
            for key, value in row["inputs"].items()
        )
        evaluation_details.extend(
            {
                "Scenario": row["id"],
                "Aspect": "Sequence",
                "Field / order": i + 1,
                "Specification": value,
            }
            for i, value in enumerate(row["sequence"])
        )
        evaluation_details.extend(
            {
                "Scenario": row["id"],
                "Aspect": "Expected observation",
                "Field / order": value["path"],
                "Specification": value["operator"] + ": " + _text(value["value"]),
            }
            for value in row["expected"]
        )
        evaluation_details.extend(
            {
                "Scenario": row["id"],
                "Aspect": "Live observation",
                "Field / order": i + 1,
                "Specification": value,
            }
            for i, value in enumerate(row["live_observation"])
        )
    for key, label in (
        ("observation_contract", "Observation contract"),
        ("media_contract", "Media policy"),
    ):
        evaluation_details.extend(
            {
                "Scenario": "Shared contract",
                "Aspect": label,
                "Field / order": field,
                "Specification": _text(value),
            }
            for field, value in evaluation_catalog[key].items()
        )

    def autonomy_leaves(value, path=()):
        if isinstance(value, dict):
            for key, child in value.items():
                yield from autonomy_leaves(child, (*path, key))
        elif isinstance(value, list) and any(isinstance(child, (dict, list)) for child in value):
            for index, child in enumerate(value, 1):
                label = child.get("id", str(index)) if isinstance(child, dict) else str(index)
                yield from autonomy_leaves(child, (*path, label))
        else:
            yield {
                "Scenario": "Autonomy contract",
                "Aspect": path[0],
                "Field / order": ".".join(path[1:]) or path[0],
                "Specification": _text(value),
            }

    evaluation_details.extend(autonomy_leaves(_read(root, "docs/blueprint/autonomy.json")))
    saved_evaluation_details = tables["Evaluation Details"]
    differing = [
        i
        for i, (actual, expected) in enumerate(
            zip(saved_evaluation_details, evaluation_details, strict=False), 4
        )
        if set(actual) != set(expected)
        or any(not _same_value(actual.get(k), v) for k, v in expected.items())
    ]
    if len(saved_evaluation_details) != len(evaluation_details) or differing:
        mismatch = differing[0] if differing else "population"
        errors.append(
            f"Evaluation Details: specification population or content differs at row {mismatch}"
        )
    for sheet, key, source in (
        ("Release Gates", "Gate", context["review_gates"]),
        ("Journey Scenarios", "Scenario", context["scenarios"]),
        ("Risks", "Risk", context["risks"]),
    ):
        identities(sheet, key, source)

    formula_rows = tables["__formulas__"]
    formula_keys = [(r.get("sheet"), r.get("cell")) for r in formula_rows]
    if len(formula_keys) != len(set(formula_keys)):
        errors.append("formulas: duplicate saved formula cell")
    formulas = {(r.get("sheet"), r.get("cell")): r for r in formula_rows}

    def formula(sheet: str, address: str, expected_formula: str, expected_value: int):
        record = formulas.get((sheet, address), {})
        if record.get("formula") != expected_formula or record.get("cached") != expected_value:
            errors.append(f"{sheet}!{address}: formula or cached value differs")

    expected_reconciliation = []
    for index, (label, sheet, _, source, column) in enumerate(groups, 4):
        expected_reconciliation.append(label)
        compare(
            "Reconciliation",
            "Population",
            label,
            {
                "Source count": len(source),
                "Workbook count": len(source),
                "Difference": 0,
            },
        )
        formula(
            "Reconciliation",
            f"C{index}",
            f"=COUNTA('{sheet}'!{column}4:{column}{len(source) + 3})",
            len(source),
        )
        formula("Reconciliation", f"D{index}", f"=C{index}-B{index}", 0)
    profiles = plan["release_profiles"]
    profile_rows = tables["Release Profiles"]
    if [row.get("Profile") for row in profile_rows if row.get("Aspect") == "Scope"] != [
        p["id"] for p in profiles
    ]:
        errors.append("Release Profiles: exact profile population differs")
    for profile in profiles:
        expected = _list(profile.get("required_criteria")) or (
            "No unconditional criteria are listed. Activated conditional work "
            "contributes its complete criterion population."
        )
        rows = [
            row
            for row in profile_rows
            if row.get("Profile") == profile["id"] and row.get("Aspect") == "Required criteria"
        ]
        if len(rows) != 1 or rows[0].get("Requirement") != expected:
            errors.append(
                f"Release Profiles: {profile['id']} exact required criterion population differs"
            )
    profile_index = len(groups) + 4
    formula(
        "Reconciliation",
        f"C{profile_index}",
        f"=COUNTIF('Release Profiles'!B4:B{len(profile_rows) + 3},\"Scope\")",
        len(profiles),
    )
    formula("Reconciliation", f"D{profile_index}", f"=C{profile_index}-B{profile_index}", 0)
    full_steps = sum(all(s.get(k) for k in list(CONTRACT_ASPECTS)[:8]) for s in steps)
    for label, count in (
        ("Release profiles", len(profiles)),
        ("Wave assignments", len(plan["item_waves"])),
        ("Full intended step contracts", full_steps),
        ("Structural lint problems", 0),
    ):
        expected_reconciliation.append(label)
        compare(
            "Reconciliation",
            "Population",
            label,
            {"Source count": count, "Workbook count": count, "Difference": 0},
        )
    if [r.get("Population") for r in tables["Reconciliation"]] != expected_reconciliation:
        errors.append("Reconciliation: exact control-row population differs")
    for index, phase in enumerate("ABCDEFGHI", 4):
        fs = [f for f in features if f["phase"] == phase]
        ss = [s for s in steps if s["phase"] == phase]
        expected = [
            len(fs),
            *(
                sum(f["implementation"] == state for f in fs)
                for state in ("complete", "partial", "none")
            ),
            len(ss),
            sum(all(s.get(k) for k in list(CONTRACT_ASPECTS)[:8]) for s in ss),
        ]
        columns = [
            "Features",
            "Implemented complete",
            "Partial",
            "None",
            "Steps",
            "Full intended contracts",
        ]
        compare("Overview", "Phase", phase, dict(zip(columns, expected, strict=True)))
        for letter, value in zip("CDEFGH", expected, strict=True):
            record = formulas.get(("Overview", f"{letter}{index}"), {})
            if (
                not isinstance(record.get("formula"), str)
                or not record["formula"].startswith("=")
                or record.get("cached") != value
            ):
                errors.append(f"Overview!{letter}{index}: formula or cached value differs")
    for row in formula_rows:
        if row.get("cached") is None or (
            isinstance(row.get("cached"), str) and row["cached"] in ERROR_VALUES
        ):
            errors.append(f"{row.get('sheet')}!{row.get('cell')}: missing or errored formula cache")

    snapshot_rows = [r for r in tables["Sources"] if r.get("Role") == "Snapshot source"]
    expected_sources = SNAPSHOT_BASE | {
        p.relative_to(root).as_posix()
        for p in (root / "docs/blueprint").rglob("*")
        if p.is_file() and p.suffix in {".md", ".json"}
    }
    actual_sources = [r.get("Source") for r in snapshot_rows]
    if Counter(actual_sources) != Counter(expected_sources):
        errors.append("Sources: exact authored snapshot-source population differs")
    for row in snapshot_rows:
        relative = row.get("Source")
        if not isinstance(relative, str) or relative not in expected_sources:
            continue
        path = (root / relative).resolve()
        if not path.is_relative_to(root.resolve()) or not path.is_file():
            errors.append(f"Sources: missing or escaped source {relative}")
        elif row.get("Fingerprint / URL") != source_hash(path):
            errors.append(f"Sources: stale or incorrect hash {relative}")
    captured = [r for r in tables["Sources"] if r.get("Source") == CAPTURED_ARTIFACT]
    if (
        len(captured) != 1
        or captured[0].get("Role") != "Captured execution artifact"
        or not isinstance(captured[0].get("Fingerprint / URL"), str)
        or len(captured[0]["Fingerprint / URL"]) != 64
        or any(c not in "0123456789abcdef" for c in captured[0]["Fingerprint / URL"])
    ):
        errors.append(
            "Sources: captured execution artifact must remain historical, not current authority"
        )
    return errors


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("path", type=Path, nargs="?", default=WORKBOOK)
    args = parser.parse_args()
    try:
        tables = load_view(args.path)
        errors = check_view(tables)
    except (OSError, ValueError, KeyError) as exc:
        print(f"Current plan view: unreadable — {exc}")
        return 1
    for error in errors:
        print(f"  FAIL: {error}")
    print(
        f"Current plan view: {len(SHEETS)} sheets, {len(tables['Work Items'])} items, "
        f"{len(tables['Acceptance'])} criteria; {len(errors)} problems. "
        "Read-only; not release proof."
    )
    return int(bool(errors))


if __name__ == "__main__":
    sys.exit(main())
