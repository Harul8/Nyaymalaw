"""Every requirement clause resolves to a registered delivery owner -- or says it does not. LB-43-AC4.

    python assurance/control_plane/requirement_owners.py                # the whole population
    python assurance/control_plane/requirement_owners.py --slice P49    # may this slice start?

WHY THIS EXISTS
----------------
LB-43-AC4 asked that "all 92 LB requirements and 324 clauses" resolve to an
accountable owner. By 27 September 2026 the sheet held 170 LB and 23 OM
requirements and 497 clauses, and nothing computed either number: the count in
the criterion was a claim about the plan that the plan had outgrown. The
cell-to-cell reconciler (`plan_scenarios.requirement_problems`) proves the two
sheets agree; it cannot say whether anything owns a clause.

THE POPULATION IS READ FROM THE SHEET WHEN THIS RUNS. Every row of the
Implementation Plan whose ID is an LB or OM requirement, and every acceptance
clause (`<ID>-AC<n>`) its acceptance cell names. Nothing here counts them.

WHAT OWNS A ROW. The "Delivery owner" line LB-43 registers in the row's
dependency cell. Each criterion it names (`BK-n-ACm`, and `and ACk` after it)
must exist in docs/backlog/status.yaml and be carried by a packet in
docs/blueprint/packets.json; where a packet is written after it -- before the
next criterion or item the line names -- that packet must carry it. That last
rule is what catches a line
left behind when a criterion moves -- it is how LB-163 was found still pairing
BK-101-AC1 with P53 an hour after BK-101-AC1 moved to P49. A bare item
(`BK-98`) must exist. A clause is owned through its row: the sheet does not map
clauses to criteria one by one, and this says so rather than pretending to.

FOUR STATES, and the last two are not failures of the whole plan -- they block
the scope that selects them:

  owned       a registered, packaged criterion or item, and no problem
  problem     the line names something the registries do not carry
  declared    the line says there is no delivery owner (an owner decision)
  unowned     no delivery-owner line at all

A SLICE (`--slice PNN`) selects every row whose line names that packet or a
criterion it carries. It may start only if every selected row is owned and
names at least one clause, and every criterion of the packet is claimed by
some requirement row -- a criterion nothing requires is a foundation nobody
specified.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

_REPO = Path(__file__).resolve().parents[2]
if str(_REPO) not in sys.path:
    sys.path.insert(0, str(_REPO))
from assurance.common._console import utf8_console  # noqa: E402

utf8_console()

SHEET = _REPO / "docs" / "Nyaymalaw_Implementation_Plan.xlsx"
STATUS = _REPO / "docs" / "backlog" / "status.yaml"
PACKETS = _REPO / "docs" / "blueprint" / "packets.json"
WORKSHEET = "Implementation Plan"
ID_COLUMN, OWNER_COLUMN, CLAUSE_COLUMN = "ID", "Dependency IDs and required outputs", "Acceptance criteria"

REQUIREMENT = re.compile(r"^(LB-\d+|OM-[PIQ]\d+)$")
OWNER_LINE = re.compile(r"Delivery owner[^\n]*")
NO_OWNER = re.compile(r"No delivery owner", re.I)
CRITERION = re.compile(r"\bBK-(\d+)-AC(\d+)((?:\s*(?:,|and|&)\s*AC\d+)*)")
ITEM = re.compile(r"\bBK-(\d+)\b(?!-AC)")
PACKET = re.compile(r"\bP(\d{2})\b")


@dataclass
class Requirement:
    ident: str
    clauses: tuple[str, ...]
    line: str
    criteria: tuple[str, ...] = ()
    items: tuple[str, ...] = ()
    packets: tuple[str, ...] = ()
    state: str = "unowned"
    problems: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class Registry:
    criteria: frozenset[str]
    items: frozenset[str]
    carried_by: dict          # criterion -> packets that list it
    carries: dict             # packet -> criteria it lists

    @staticmethod
    def load(status: Path = STATUS, packets: Path = PACKETS) -> "Registry":
        import yaml
        doc = yaml.safe_load(status.read_text(encoding="utf-8"))
        items = frozenset(i["id"] for i in doc["items"])
        criteria = frozenset(a["id"] for i in doc["items"] for a in (i.get("acceptance") or []))
        carried_by: dict[str, set[str]] = {}
        carries: dict[str, set[str]] = {}
        for packet in json.loads(packets.read_text(encoding="utf-8"))["packets"]:
            listed = set(packet.get("criteria") or []) | set(packet.get("final_criteria") or [])
            carries[packet["id"]] = listed
            for criterion in listed:
                carried_by.setdefault(criterion, set()).add(packet["id"])
        return Registry(criteria, items, carried_by, carries)


def parse_owner(line: str) -> tuple[tuple[str, ...], tuple[str, ...], tuple[str, ...]]:
    """The criteria, bare items and packets a delivery-owner line names, in order."""
    criteria = [c for group, _ in references(line) if group[0].count("-") == 2 for c in group]
    items = [group[0] for group, _ in references(line) if group[0].count("-") == 1]
    packets = [f"P{m[1]}" for m in PACKET.finditer(line)]
    return tuple(dict.fromkeys(criteria)), tuple(dict.fromkeys(items)), tuple(dict.fromkeys(packets))


def references(line: str) -> list[tuple[tuple[str, ...], frozenset[str]]]:
    """Each criterion group or item the line names, with the packets written
    after it and before the next reference. `BK-91-AC1 and AC5 (P46)` is one
    group paired with P46; `reviewed under BK-67-AC3` is paired with nothing."""
    refs = []
    for m in CRITERION.finditer(line):
        group = (f"BK-{m[1]}-AC{m[2]}", *(f"BK-{m[1]}-AC{n}" for n in re.findall(r"AC(\d+)", m[3] or "")))
        refs.append((m.start(), m.end(), group))
    refs += [(m.start(), m.end(), (f"BK-{m[1]}",)) for m in ITEM.finditer(line)]
    refs.sort()
    out = []
    for i, (_start, end, group) in enumerate(refs):
        stop = refs[i + 1][0] if i + 1 < len(refs) else len(line)
        out.append((group, frozenset(f"P{m[1]}" for m in PACKET.finditer(line, end, stop))))
    return out


def resolve(requirement: Requirement, registry: Registry) -> Requirement:
    line = requirement.line
    if not line:
        requirement.state = "unowned"
        return requirement
    requirement.criteria, requirement.items, requirement.packets = parse_owner(line)
    if NO_OWNER.search(line) and not requirement.criteria:
        requirement.state = "declared"
        return requirement
    named = set(requirement.packets)
    for group, written in references(line):
        for ref in group:
            if ref.count("-") == 1:
                if ref not in registry.items:
                    requirement.problems.append(f"{ref} is not an item in status.yaml")
                elif written and not any(c.startswith(f"{ref}-AC") for p in written
                                         for c in registry.carries.get(p, ())):
                    requirement.problems.append(f"{ref} is paired with {', '.join(sorted(written))}, which "
                                                f"carries none of its criteria")
                continue
            if ref not in registry.criteria:
                requirement.problems.append(f"{ref} is not a criterion in status.yaml")
                continue
            carriers = registry.carried_by.get(ref, set())
            if not carriers:
                requirement.problems.append(f"{ref} is carried by no packet")
            elif written and not carriers & written:
                requirement.problems.append(
                    f"{ref} is paired with {', '.join(sorted(written))} but carried by "
                    f"{', '.join(sorted(carriers))}")
    for packet in requirement.packets:
        if packet not in registry.carries:
            requirement.problems.append(f"{packet} is not a packet in packets.json")
    owned = any(c in registry.criteria and registry.carried_by.get(c) for c in requirement.criteria) or any(
        i in registry.items and any(c.startswith(f"{i}-AC") for p in named for c in registry.carries.get(p, ()))
        for i in requirement.items)
    requirement.state = "problem" if requirement.problems else ("owned" if owned else "declared")
    if requirement.state == "declared" and not NO_OWNER.search(line):
        requirement.problems.append("the line names no registered criterion or packaged item")
        requirement.state = "problem"
    return requirement


def requirements(sheet: Path = SHEET, registry: Registry | None = None) -> list[Requirement]:
    """Every LB and OM requirement in the sheet, read now, with its clauses and owner."""
    from openpyxl import load_workbook
    registry = registry or Registry.load()
    book = load_workbook(sheet, read_only=True)
    try:
        rows = book[WORKSHEET].iter_rows(values_only=True)
        header = [str(c or "").strip() for c in next(rows)]
        at = {name: header.index(name) for name in (ID_COLUMN, OWNER_COLUMN, CLAUSE_COLUMN)}
        found = []
        for row in rows:
            ident = str(row[at[ID_COLUMN]] or "").strip()
            if not REQUIREMENT.match(ident):
                continue
            clauses = tuple(dict.fromkeys(re.findall(rf"{re.escape(ident)}-AC\d+", str(row[at[CLAUSE_COLUMN]] or ""))))
            lines = OWNER_LINE.findall(str(row[at[OWNER_COLUMN]] or ""))
            found.append(resolve(Requirement(ident, clauses, lines[-1] if lines else ""), registry))
    finally:
        book.close()
    return found


def slice_blockers(packet: str, population: list[Requirement], registry: Registry) -> list[str]:
    """Why the slice a packet delivers may not start. Empty means it may."""
    if packet not in registry.carries:
        return [f"{packet} is not a packet in packets.json"]
    carried = registry.carries[packet]
    selected = [r for r in population if packet in r.packets or carried & set(r.criteria)]
    blockers = []
    if not selected:
        blockers.append(f"no requirement row selects {packet}, so nothing says what it delivers")
    for r in selected:
        if r.state != "owned":
            blockers.append(f"{r.ident} is {r.state}: {'; '.join(r.problems) or r.line or 'no delivery owner line'}")
        if not r.clauses:
            blockers.append(f"{r.ident} names no acceptance clause, so nothing could show it done")
    claimed = {c for r in population for c in r.criteria}
    for criterion in sorted(carried - claimed):
        blockers.append(f"{criterion} ({packet}) is claimed by no requirement row -- work nobody specified")
    return blockers


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--slice", help="a packet id, e.g. P49: may that slice start?")
    args = parser.parse_args(argv)
    registry = Registry.load()
    population = requirements(registry=registry)
    by_state: dict[str, list[Requirement]] = {}
    for r in population:
        by_state.setdefault(r.state, []).append(r)
    lb = sum(r.ident.startswith("LB-") for r in population)
    print(f"{len(population)} requirements ({lb} LB, {len(population) - lb} OM) and "
          f"{sum(len(r.clauses) for r in population)} clauses, read from {SHEET.name} now")
    for state in ("owned", "problem", "declared", "unowned"):
        rows = by_state.get(state, [])
        clauses = sum(len(r.clauses) for r in rows)
        print(f"  {state:9} {len(rows):4} rows, {clauses:4} clauses")
    for r in by_state.get("problem", []):
        print(f"  PROBLEM {r.ident}: {'; '.join(r.problems)}")
    for r in by_state.get("declared", []):
        print(f"  declared {r.ident}: {r.line.split(':', 1)[-1].strip()[:140]}")
    print("A clause is owned through its row; the sheet does not map clauses to criteria one by one.")
    failed = bool(by_state.get("problem"))
    if args.slice:
        blockers = slice_blockers(args.slice, population, registry)
        print(f"\n{args.slice}: " + ("may start -- every selected row is owned and names a clause, and every "
                                     "criterion is claimed" if not blockers else f"{len(blockers)} blocker(s)"))
        for b in blockers:
            print(f"  - {b}")
        failed = failed or bool(blockers)
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
